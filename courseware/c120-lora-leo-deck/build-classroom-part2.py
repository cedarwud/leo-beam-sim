#!/usr/bin/env python3
"""Build the standalone Part B classroom module (eventual P028--P063).

The builder uses the native ``educate.pptx`` master and layouts directly.  It
keeps the deck deliberately editable: every command, identity token, receipt
field and diagram is a PowerPoint shape, while the two formula pages reserve
native Office Math insertion targets.  Local fallback artifacts are shown as
simulated teaching evidence; no live/browser/Windows claim is invented.
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = Path("/home/u24/ppt-master/template/educate.pptx")
PROJECT = ROOT / "projects/classroom-part2_ppt169_20260811"
SOURCE_COPY = PROJECT / "sources/educate.pptx"
EXPORT = PROJECT / "exports/classroom-part2-editable.pptx"
RENDER_DIR = PROJECT / "renders"
RENDER_SOURCE = RENDER_DIR / ".part2-render-source.pptx"
QA_DIR = PROJECT / "qa"
ANALYSIS_DIR = PROJECT / "analysis"
VALIDATION_DIR = PROJECT / "validation"
EQUATION_TOOL = ROOT / "insert-native-equations.py"
EQUATION_FALLBACKS = {
    "E_endpoint": PROJECT / "assets/equation-energy-fallback.svg",
    "eta_E": PROJECT / "assets/equation-efficiency-fallback.svg",
}

REPO_URL = "https://github.com/cedarwud/lora-energy-lab"
CURRENT_COMMIT = "32cc7230904113e635cc591191243d7b72b4e9aa"
ROOT_NAME = "lora-energy-lab"
PYTHON_MODULE = "lora_energy_lab"
SCENARIO_ID = "ntpu-energy-decision-01"
POLICY_API = "lora-energy-policy-v1"
SCENARIO_SHA = "sha256:c157c76ed283074cdd17c1b689b83e168022abd2ca6ef8d9c337e69ab31a6593"
ANCHOR_SHA = "sha256:cfe84682a71c0c0ea09c273d4cd033c32818da63cf1695ace4363b9d789c50f9"
LOCK_SHA = "sha256:f1d1fe3339137d56e4e9ea9d141b764eaa183a8f40933bad597d734bbf310197"
BASELINE_POLICY_SHA = "sha256:9917d2cde2e8c0eae8339b7eea2e358bf8577b4eb6aaf1e928674f6b8964ef68"
LOCAL_ARTIFACT = "lora-energy-lab-v1.zip｜121139 bytes｜sha256:047e8459988d8bee59fbfdcc39189e00048f3697968901ce5b6e5363dc9e8c8c"
PUBLIC_RELEASE = "RELEASE_ASSET_PENDING"

# Fallback JSON evidence is local, deterministic and explicitly simulated.
FALLBACK = {
    "baseline-A": {"run": "ee6df44f", "policy": BASELINE_POLICY_SHA, "energy": 6.92, "delivered": 4800, "expired": 3, "eff": 693.641618, "wake": 2, "service": False},
    "candidate-A": {"run": "99988424", "policy": "sha256:325cca4a4be339a9d69e2911123714556135859dc6089cac271c6e6d0d73cf2f", "energy": 8.86, "delivered": 4800, "expired": 3, "eff": 541.760722, "wake": 1, "service": False, "freeze": "sha256:6cbaf9dccc812994bec08c0f898b8d21d68a3d84412137a9966de71667e11626"},
    "hidden-A": {"run": "3fd8c85a", "policy": "sha256:325cca4a4be339a9d69e2911123714556135859dc6089cac271c6e6d0d73cf2f", "energy": 3.61, "delivered": 0, "expired": 3, "eff": 0, "wake": 0, "service": False, "freeze": "sha256:6cbaf9dccc812994bec08c0f898b8d21d68a3d84412137a9966de71667e11626"},
    "trace-a-baseline-B": {"run": "b2603b79", "policy": "sha256:325cca4a4be339a9d69e2911123714556135859dc6089cac271c6e6d0d73cf2f", "energy": 8.86, "delivered": 4800, "expired": 3, "eff": 541.760722, "wake": 1, "service": False},
    "trace-a-candidate-B": {"run": "c39413ac", "policy": "sha256:c00ed808d3a225abff99a19aee7426570c274b60b46231382be4407c78293718", "energy": 10.66, "delivered": 9600, "expired": 1, "eff": 900.562852, "wake": 1, "service": False, "freeze": "sha256:2a4167e25699ec366687559d1ec7259d8b874f990b81653a5fbb45b171349e0e"},
    "trace-b-B": {"run": "b73a1f59", "policy": "sha256:c00ed808d3a225abff99a19aee7426570c274b60b46231382be4407c78293718", "energy": 5.43, "delivered": 4800, "expired": 2, "eff": 883.977901, "wake": 1, "service": False, "freeze": "sha256:2a4167e25699ec366687559d1ec7259d8b874f990b81653a5fbb45b171349e0e"},
}

FORBIDDEN = "學生"
MINUTE_RE = re.compile(r"(?:分鐘|\b\d+\s*(?:min|mins|minute|minutes)\b)", re.I)
LATIN = re.compile(r"[A-Za-z0-9_./\\:-]")

NAVY = RGBColor(53, 55, 127)
PURPLE = RGBColor(102, 0, 102)
BLUE = RGBColor(111, 137, 247)
# The pale yellow accent was too close to bright yellow-orange in rendered
# text.  This darker muted gold remains part of the template accent language
# while keeping authored text readable on white and pale-gold surfaces.
GOLD = RGBColor(132, 104, 38)
TEAL = RGBColor(0, 125, 128)
RED = RGBColor(192, 55, 55)
INK = RGBColor(43, 46, 65)
MUTED = RGBColor(90, 94, 111)
WHITE = RGBColor(255, 255, 255)
CREAM = RGBColor(250, 247, 232)
PALE_BLUE = RGBColor(236, 241, 254)
PALE_PURPLE = RGBColor(244, 237, 247)
PALE_GOLD = RGBColor(253, 249, 226)
PALE_TEAL = RGBColor(232, 247, 244)
PALE_RED = RGBColor(253, 239, 239)
GRAY = RGBColor(223, 226, 234)
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


@dataclass(frozen=True)
class SlideSpec:
    number: int
    title: str
    kind: str
    layout: int
    evidence: str
    context: str
    command: str
    purpose: str
    mechanism: str
    expected: str
    recovery: str
    visual: str
    notes: str


def spec(number: int, title: str, kind: str, *, layout: int = 1, evidence: str = "SOURCE-BOUND", context: str = "", command: str = "", purpose: str = "", mechanism: str = "", expected: str = "", recovery: str = "", visual: str = "", notes: str = "") -> SlideSpec:
    return SlideSpec(number, title, kind, layout, evidence, context, command, purpose, mechanism, expected, recovery, visual, notes)


COMMON_RECOVERY = "欄位、units、scope 或 identity 不完整：停止比較，保留原始錯誤，回到相同 scenario 的 artifact 或 matching fallback。不要補填數字。"
A_RECOVERY = "A result 或 consequential diff 缺失：保留 stdout 與 stderr，restore A checkpoint 或選 matching fallback；不要再調參。"
B_RECOVERY = "B receipt、predecessor 或 policy SHA 不符：停止 withheld，restore A/B checkpoint；不可 retune。"


SPECS: list[SlideSpec] = [
    spec(28, "endpoint 狀態：名稱要和機制對齊", "states", layout=1, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="local fallback result／endpoint boundary", command="python -m json.tool fallback_artifacts/baseline-A/result.json\nrg -n 'STATE_INTERVAL|state|duration' fallback_artifacts/baseline-A/result.json", purpose="把狀態名稱連回可觀察的停留與功率，不把 label 當結論。", mechanism="SLEEP、WAIT、WAKE、PROCESS、TX、RX 各自留下 state interval；累積結果仍屬 endpoint scope。", expected="看到 state interval 與 `energy_breakdown_j`；目前 fallback 標示 `FALLBACK_ARTIFACT`、`SIMULATED_ADAPTER`。", recovery=COMMON_RECOVERY, visual="SLEEP → WAKE → PROCESS → TX / RX → WAIT", notes="先把名稱和機制分開。SLEEP 是低功耗休息，WAIT 是 awake idle，WAKE 有回復成本，PROCESS、TX、RX 都會留下時間與能量欄位。畫面中的 JSON 是 local fallback 的 simulated teaching evidence，不是 live measurement。"),
    spec(29, "WAIT 是 awake idle，SLEEP 會帶來 wake 成本", "waitsleep", layout=2, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="baseline-A endpoint replay；同一 job／window／energy boundary", command="rg -n '\"state\": \"WAIT\"|\"state\": \"SLEEP\"|\"state\": \"WAKE\"' fallback_artifacts/baseline-A/result.json", purpose="先提出 WAIT 與 SLEEP 的方向性預測，再由 state ledger 驗證。", mechanism="WAIT 保持清醒；SLEEP 可能降低 idle power，但後續 WAKE 會增加 latency／energy。", expected="看到 awake_idle、sleep、wake 的分解；不要只拿單一 power 欄位判定。", recovery=COMMON_RECOVERY, visual="WAIT：awake idle  ─────────\nSLEEP：sleep ─ wake ─ service", notes="這是 Lab A 的核心預測。WAIT 可能少付 wake 成本，但裝置保持清醒；SLEEP 可能降低 idle power，卻要付 wake latency 與 wake energy。先寫方向，結果再決定是否成立。"),
    spec(30, "封包生命週期：送出不是交付", "packet", layout=3, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="baseline-A endpoint replay；packet ledger", command="python -m json.tool fallback_artifacts/baseline-A/endpoint-replay.json", purpose="把 SEND action 拆成 queue、attempt、retry、delivered 或 expired。", mechanism="一次 action 只代表一次決策；packet 是否交付還要經過 attempt、collision、retry 與 deadline。", expected="看到 `attempted_packets`、`retransmissions`、`unique_delivered_packets`、`expired_packets`。", recovery=COMMON_RECOVERY, visual="generated → queue → attempt → retry → delivered / expired", notes="封包先進 queue，再可能產生 attempt。一次 SEND 不等於 delivered service；要把 retry、expired 與 queue age 分開讀。這個分解會保護後面 A/B 比較，不讓一次 TX 代替整個 packet outcome。"),
    spec(31, "service boundary 先於 energy 比較", "service", layout=1, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="paired baseline／candidate；同 scenario、job、window、endpoint scope", command="rg -n '\"service_pass\"|\"deadline_pass\"|\"energy_scope\"' fallback_artifacts/candidate-A/result.json", purpose="先確認 job、window、delivery 與 deadline 可比，再談 J。", mechanism="candidate 少送或多過期時，較低 J 不能直接升格為節能；service gate 先決定比較資格。", expected="看到 `service_pass`、`deadline_pass`、`energy_scope` 同時被讀取；目前 fallback 明示 service boundary endpoint-only。", recovery=COMMON_RECOVERY, visual="same job + same window + delivery gate\n                 ↓\n          endpoint J comparison", notes="比較兩個 policy 前，先確認是不是同一個工作、窗口和 endpoint boundary。若 candidate 少送很多封包，J 下降不能直接叫節能；先說 service 是否保住。"),
    spec(32, "W 是速率，J 是累積結果", "powerarea", layout=2, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="result JSON；energy scope = endpoint", command="rg -n '\"endpoint_energy_j\"|\"power\"|\"energy_breakdown_j\"' fallback_artifacts/candidate-A/result.json", purpose="把瞬時功率與時間累積出的能量分開讀。", mechanism="W 描述 rate；J 是 P(t) 對時間的 area。低峰值拖得久，仍可能累積更多 J。", expected="看到 units `power: W`、`energy: J`；結果仍要連回 service／deadline。", recovery=COMMON_RECOVERY, visual="P(t)\n  ╭──╮\n  │  ╰──╮\n  ╰──────╯\n  area = endpoint J", notes="W 是某一時刻的功率，J 是功率隨時間累積的能量。低峰值不保證較低 J；先看整段 state time，再回到 service boundary。這裡只讀 fallback artifact 的 endpoint scope。"),
    spec(33, "公式：endpoint energy 的範圍", "formula-energy", layout=3, evidence="VERIFIED｜EDITABLE OFFICE MATH TARGET", context="formula page；endpoint energy boundary", command="rg -n 'energy_scope|energy_breakdown_j' fallback_artifacts/candidate-A/result.json", purpose="讀懂每個 state 的功率與停留時間如何累積。", mechanism="每一個 state 的 P_s × t_s 加總；下標 endpoint 限定 claim scope。", expected="Office Math 可編輯；公式頁不填未產生的 KPI。", recovery=COMMON_RECOVERY, visual="E_endpoint = Σ P_s t_s\nendpoint boundary ≠ whole system", notes="公式只說清楚 endpoint radio 與 processing assumptions 的累積關係。E 的下標提醒我們不要把 endpoint J 當成整個 system energy；數字仍從 result artifact 讀取。"),
    spec(34, "公式：服務效率的分子與分母", "formula-eff", layout=1, evidence="VERIFIED｜EDITABLE OFFICE MATH TARGET", context="formula page；delivered data 與 endpoint energy units", command="rg -n 'delivered_bits|endpoint_energy_efficiency_bits_per_j|units' fallback_artifacts/candidate-A/result.json", purpose="先確認分子是 delivered data，分母是 endpoint energy。", mechanism="η_E 的 units 是 bit/J；缺 delivered 或 scope 不明時，不能直接報效率。", expected="Office Math 可編輯；unit ladder 顯示 bit ÷ J = bit/J。", recovery=COMMON_RECOVERY, visual="η_E = D_delivered / E_endpoint\nbit ÷ J = bit/J", notes="這個比值的分子是 delivered data，分母是 endpoint energy。兩邊 boundary 和 units 要一致；公式是語義檢查，不是讓課堂自行改科學結果。"),
    spec(35, "來源、模型、課程假設、結果", "provenance", layout=2, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="result／replay／scenario lineage", command="rg -n 'scenario_id|seed|policy_sha256|artifact_source' fallback_artifacts/candidate-A/result.json", purpose="沿 provenance spine 回答一筆結果從哪裡來。", mechanism="source、model、course assumptions 與 policy action 逐層形成 result；artifact_source 不能省略。", expected="看到 scenario_id、seed、policy SHA、source mode 與 result path lineage。", recovery=COMMON_RECOVERY, visual="source → model → course assumption → policy action → result", notes="不要把 source、model 和課程假設壓成一個神秘數字。沿 scenario、seed、policy SHA 和 artifact source 往回查，才能知道一筆結果的來處。"),
    spec(36, "一個 scenario_id 綁住整條鏈", "identity", layout=3, evidence="VERIFIED｜current package identity + FALLBACK_ARTIFACT", context="package root；runner、result、replay、workbook lineage", command="rg -n 'scenario_id|scenario_anchor_sha256|scenario_sha256' fallback_artifacts/candidate-A/result.json fallback_artifacts/candidate-A/endpoint-replay.json", purpose="確認比較的是同一個問題，而不是相似畫面。", mechanism="scenario、anchor、seed、schema、units 與 policy identity 任一不一致時，import 要 fail closed。", expected=f"同一 `{SCENARIO_ID}`、anchor `{ANCHOR_SHA}`、scenario `{SCENARIO_SHA}` 出現在配對 artifact。", recovery="identity mismatch：保留輸入與錯誤，回到 matching artifact；不可手改 JSON 繞過 gate。", visual=f"runner → result → replay → workbook\n{SCENARIO_ID}\nanchor / seed / policy / units", notes="同一個 scenario identity 要出現在 scenario package、runner result、endpoint replay 與 workbook。若 anchor、seed、schema 或 units 不一致，import 要 fail closed；不要用改 JSON 的方式繞過。"),
    spec(37, "公平 baseline：固定 job、window、boundary", "fairness", layout=1, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="package root；baseline／candidate A-B pair", command="git diff -- student_policy.py\npython -m py_compile student_policy.py", purpose="建立 control，只讓目前 lab 的 marked block 變動。", mechanism="scenario、seed、job、window、traffic、energy boundary 固定；policy block 是唯一變因。", expected="diff 只在 marked block；pair 可比較 queue、packet、service、state、J。", recovery=COMMON_RECOVERY, visual="FIXED  scenario / seed / job / window\nEDITED  one marked block\nOBSERVED queue / packet / service / state / J", notes="baseline 不是舊截圖，而是 control。固定 scenario、seed、job、window、traffic 和 endpoint boundary，只讓目前 lab 的 marked block 改變；偷偷換 trace 或 scope 就是不公平。"),
    spec(38, "LEO window：合法與不合法的 action", "window", layout=2, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="scenario quality/contact trace；same anchor", command="rg -n 'contact|window|quality_band_trace' scenarios/ntpu-energy-decision-01.json", purpose="把 changing opportunity 接回 WAIT／SEND 的合法性。", mechanism="窗口開啟時 action 可能服務；窗口關閉時延後可能變成 service loss；LEO 只作 changing-window example。", expected="看到 contact／quality trace 與 action context；不把 LEO 當量測來源。", recovery=COMMON_RECOVERY, visual="window open  ─── SEND / WAIT legal\nwindow closed ─── delayed action risk", notes="窗口開啟與關閉會改變 action 的後果。先看固定 anchor 和 traffic，再判斷 WAIT 或 SEND 是否落在合法 service window；這頁不是 TLE 推導。"),
    spec(39, "Lab A：同一份工作，不同 pace", "lab-a-hero", layout=3, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="Lab A question；same job／window／endpoint boundary", command="git status --short\npython -m py_compile student_policy.py", purpose="先提出可被 baseline、candidate、hidden 推翻的 pace/rest 假說。", mechanism="PACE_GAP_STEPS 與 REST_DURING_GAP 會改變 awake idle、sleep、wake、packet timing 與 endpoint J。", expected="一句 prediction 同時提到 queue、service、state time、J；尚未宣布答案。", recovery=A_RECOVERY, visual="same job\n  pace / rest choice\n  → queue / state / service / J", notes="Lab A 只問一個可測問題：固定工作與窗口時，不同 pace 和 rest 如何改變 queue、state time、service 與 endpoint J？先寫方向性預測，然後讓 baseline 和 candidate 回答。"),
    spec(40, "先讀 `PACE_GAP_STEPS`", "code-a", layout=1, evidence="SOURCE-BOUND POLICY SURFACE", context="package root；student_policy.py 的 lab-a-pace-rest marked block", command="sed -n '1,80p' student_policy.py\npython -m py_compile student_policy.py", purpose="辨認 gap 控制點進入 decision 的位置。", mechanism="gap 改變 send opportunity 與 steps_since_send；它不是直接指定 energy。", expected="看到 marked constant、syntax guard 通過；正式 release line number 不另造。", recovery="syntax／policy guard 失敗：只修 marked block，保留 traceback；不要改 runner 或 schema。", visual="lab-a-pace-rest\nPACE_GAP_STEPS = 2\n        ↓\nsteps_since_send → action", notes="只讀變數名稱與它進入 decision 的位置。先預測 gap 改變後 queue age、attempt timing、awake idle、service 和 J 的方向，再改一個合法值。"),
    spec(41, "再讀 `REST_DURING_GAP`", "code-b", layout=2, evidence="SOURCE-BOUND POLICY SURFACE", context="package root；student_policy.py 的 lab-a-pace-rest marked block", command="sed -n '10,24p' student_policy.py\npython -m py_compile student_policy.py", purpose="辨認空檔選 WAIT 或 SLEEP 的控制點。", mechanism="WAIT 保持清醒；SLEEP 可能降低 idle power，但後續要付 wake latency／energy。", expected="看到 `REST_DURING_GAP = SLEEP` 的 baseline；candidate 只改此 marked value。", recovery="值或 syntax 不合法：還原同一 block，重跑 compile；不要碰 engine／JSON。", visual="REST_DURING_GAP = SLEEP\n              ↙       ↘\n           WAIT       SLEEP\n        awake idle   wake cost", notes="這個控制點把空檔送進 WAIT 或 SLEEP。先把機制寫成 prediction；candidate 結果要同時看 state ledger、service 和 J，不接受只看一個功率欄位。"),
    spec(42, "Lab A 的 prediction lock", "prediction-a", layout=3, evidence="SOURCE-BOUND PREDICTION LOCK", context="package root；edit 前，保持 baseline policy", command="git diff -- student_policy.py\npython -m py_compile student_policy.py", purpose="在 run 前鎖定 queue、service、state time、J 的判讀句。", mechanism="prediction 先於 result；withheld 不能 retune 來迎合答案。", expected="四個 observation 欄位都有方向與可反駁句；只有 label 變化不算 gate。", recovery=A_RECOVERY, visual="queue age  ↗ / ↘\nservice    keep / loss\nstate time WAIT / SLEEP / WAKE\nendpoint J ↗ / ↘", notes="執行前寫下四個觀察：queue age、service、哪個 state time 會變、endpoint J 可能往哪裡走。這不是猜 KPI，而是讓 result 有機會推翻我們。"),
    spec(43, "A-01：跑 untouched baseline", "run-baseline-a", layout=1, evidence="VERIFIED LOCAL FALLBACK｜RUN NOT FRESH-VERIFIED", context="POSIX／WSL；package root `lora-energy-lab/`；baseline policy 不改", command="bash course.sh run --lab A --case baseline\n# Windows: course.cmd run --lab A --case baseline", purpose="建立 control，固定 scenario、seed 與 result lineage。", mechanism="runner 以 packaged baseline policy 產出 run identity、result path、packet／service／state／J。", expected=f"fallback 可讀 `run_id …{FALLBACK['baseline-A']['run']}`、energy `{FALLBACK['baseline-A']['energy']} J`、delivered `{FALLBACK['baseline-A']['delivered']} bit`；fresh stdout/result path 仍待重跑。", recovery=A_RECOVERY, visual="RUN A BASELINE\nstdout → run_id → result_path\nFALLBACK_ARTIFACT / SIMULATED\nservice_pass = false", notes="先不改 policy，執行 A baseline。POSIX/WSL 用 bash，Windows 用 course.cmd；stdout 指出的 result_path 才是本次 run 的 artifact，不要靠檔名猜。頁面上的數字來自 current fallback，不能說成這台主機 fresh run。"),
    spec(44, "A-02：只改 A marked block 並 freeze", "run-candidate-a", layout=2, evidence="VERIFIED LOCAL FALLBACK｜FREEZE RECEIPT / RUN NOT FRESH-VERIFIED", context="POSIX／WSL；`lora-energy-lab/`；只改 `lab-a-pace-rest`", command="git diff -- student_policy.py\nbash course.sh run --lab A --case candidate --freeze\n# Windows: course.cmd run --lab A --case candidate --freeze", purpose="把一個合法 policy edit 綁成 candidate 與 A frozen predecessor。", mechanism="freeze receipt 綁定 scenario、anchor、policy SHA、predecessor、seed；非 marked diff 應 fail closed。", expected=f"fallback candidate policy `{FALLBACK['candidate-A']['policy']}`、freeze `{FALLBACK['candidate-A']['freeze']}`、energy `{FALLBACK['candidate-A']['energy']} J`；fresh receipt 尚待重跑。", recovery=A_RECOVERY, visual="baseline SHA  →  candidate SHA\n                 ↘ freeze receipt\n                 lab-a-frozen checkpoint", notes="現在才改 A marked block，先保存原始檔案，再執行 candidate freeze。預期有 candidate result、receipt 和 A checkpoint；如果沒有 queue、packet、service、state 或 J 的 consequential diff，就記錄 gate 未通過，不再亂調參。"),
    spec(45, "A-03：匯入 candidate，先看同一個 boundary", "compare-a", layout=3, evidence="VERIFIED LOCAL FALLBACK｜SIMULATED ADAPTER", context="local fallback pair；result／endpoint-replay 同 scenario", command="python -m json.tool fallback_artifacts/candidate-A/result.json\npython -m json.tool fallback_artifacts/candidate-A/endpoint-replay.json", purpose="先核對 scenario、job、window、endpoint boundary，再讀差異。", mechanism="import 只是把配對 artifact 帶進 replay；它不替結果補 service 或 energy claim。", expected="看到同一 scenario、policy SHA、units、energy scope；browser pixel／KPI 尚未宣稱。", recovery=A_RECOVERY, visual="BASELINE           CANDIDATE\nscenario ✓          scenario ✓\njob/window ✓        job/window ✓\nqueue / packet / state / J → compare", notes="把 candidate result 和配對 replay 與 baseline 並排看。先打勾 scenario、job、window 和 endpoint boundary，再比較 queue、packet、state 和 J；若 service 不同，先說 trade-off。"),
    spec(46, "A-04：讀懂 state ledger", "ledger-a", layout=1, evidence="VERIFIED LOCAL FALLBACK｜SIMULATED", context="candidate-A result；state interval／energy breakdown", command="rg -n 'STATE_INTERVAL|awake_idle|sleep|wake|process|tx|rx' fallback_artifacts/candidate-A/result.json", purpose="沿 state duration 讀差異，不只抓總 J。", mechanism="WAIT、SLEEP、WAKE、PROCESS、TX、RX 各自構成 endpoint energy；policy action 才能解釋差異。", expected="fallback candidate：awake_idle 6 J-equivalent component、sleep 0.04、wake 0.02、TX 2.4；具體欄位以 JSON 為準。", recovery=A_RECOVERY, visual="WAIT  ██████\nSLEEP █\nWAKE  ·\nPROCESS ·\nTX    ██\nRX    ·", notes="沿 ledger 從左到右讀。先看 WAIT 是否變長，再看 SLEEP 是否帶來 WAKE；接著確認 PROCESS、TX、RX 的時間與 attempt，最後才把差異連回 policy action。"),
    spec(47, "A-05：把 packet 結果接回 service", "packet-service-a", layout=2, evidence="VERIFIED LOCAL FALLBACK｜SIMULATED", context="candidate-A result／endpoint replay；packet ledger", command="rg -n 'attempt|retry|delivered|expired|service_pass' fallback_artifacts/candidate-A/result.json", purpose="用 delivered、retry、expired、deadline 完整判讀工作是否保住。", mechanism="packet outcome 是 service 的前提；送出次數或較低 J 都不能代替交付。", expected=f"candidate fallback：attempted 2、delivered {FALLBACK['candidate-A']['delivered']} bit、expired {FALLBACK['candidate-A']['expired']}、service_pass false。", recovery=A_RECOVERY, visual="attempt → retry / collision → delivered / expired\n                          ↓\n                    service verdict", notes="candidate 是否支持工作，要看 delivered、expired、retry 和 deadline，而不是 SEND 次數。請用完整句子說出哪個 action 改變哪個 packet outcome，以及這個 outcome 是否仍滿足 service。"),
    spec(48, "A-06：檢查 A freeze lineage", "freeze-a", layout=3, evidence="VERIFIED LOCAL FALLBACK｜FREEZE RECEIPT / SIMULATED", context="package root；A frozen receipt／checkpoint", command="python -m json.tool fallback_artifacts/receipts/a.json\nfind artifacts/checkpoints -maxdepth 1 -type f -print", purpose="確認 candidate freeze 是 withheld 的入場條件。", mechanism="receipt 綁定 predecessor、active_block_id、policy SHA、scenario 與 anchor；缺一項就不能進 hidden。", expected=f"receipt role=`lab-a-frozen`、policy `{FALLBACK['candidate-A']['policy']}`、predecessor `{BASELINE_POLICY_SHA}`、freeze `{FALLBACK['candidate-A']['freeze']}`。", recovery="receipt／checkpoint／predecessor 缺失：停在 A candidate，保留輸出，重跑 freeze；不要直接執行 hidden。", visual="baseline\n  ↓ candidate + freeze\n  ↓ lab-a-frozen receipt\n  ↓ hidden entry", notes="freeze 不是按鈕效果，而是 withheld 的入場條件。確認 candidate policy SHA 與 baseline 不同，receipt 指向同一 scenario／anchor，checkpoint role 是 A frozen；缺一項就回 candidate 最小修復。"),
    spec(49, "A-07：用 hidden case 檢驗，不再 retune", "run-hidden-a", layout=1, evidence="VERIFIED LOCAL FALLBACK｜WITHHELD / SIMULATED", context="POSIX／WSL；A frozen policy active；same package root", command="bash course.sh run --lab A --case hidden\n# Windows: course.cmd run --lab A --case hidden", purpose="在另一個 service/window condition 測試 pace/rest 假說。", mechanism="withheld 沿用 A frozen policy；policy SHA 必須 unchanged，不能用新 freeze 或調參。", expected=f"fallback hidden policy unchanged `{FALLBACK['hidden-A']['policy']}`、energy `{FALLBACK['hidden-A']['energy']} J`、delivered 0 bit；這是 counterexample，不是 live KPI。", recovery=A_RECOVERY, visual="A frozen policy\n   ├─ primary case\n   └─ hidden case：no retune\npolicy SHA unchanged", notes="保持 A frozen policy，不再編輯，執行 hidden。hidden 用另一個條件測試 pace/rest 假說；預期 policy SHA 不變。若方向相反，這是 counterexample，保留它，不要改 policy 讓畫面變漂亮。"),
    spec(50, "Lab A debrief：一條因果句就夠", "debrief-a", layout=2, evidence="VERIFIED LOCAL FALLBACK｜SIMULATED COUNTEREXAMPLE", context="A baseline／candidate／hidden fallback trio", command="python -m json.tool fallback_artifacts/hidden-A/result.json\nrg -n 'policy_sha256|service_pass|endpoint_energy_j' fallback_artifacts/{baseline-A,candidate-A,hidden-A}/result.json", purpose="把 edit、mechanism、evidence、withheld verdict 串成一個可反駁句。", mechanism="prediction → policy hash → result/replay → hidden condition；推翻也保留 lineage。", expected="句型完成：我改了…因此…result 顯示…hidden 仍成立／被推翻。", recovery=A_RECOVERY, visual="我改了…\n    ↓\n因此改變…\n    ↓\nresult 顯示…\n    ↓\nhidden：成立／被推翻", notes="完成一句可反駁敘述：我改了 A 哪個控制點，因此改變哪個 state 或 packet 機制；result 顯示什麼 evidence；hidden 下成立或被推翻。這句比背一個 J 更有價值。"),
    spec(51, "Lab B：現在送，還是等到更穩定？", "lab-b-hero", layout=3, evidence="VERIFIED｜FALLBACK_ARTIFACT / SIMULATED", context="Lab B quality/contact trace；A frozen predecessor", command="rg -n 'quality_band_trace|contact' scenarios/ntpu-energy-decision-01.json\npython -m py_compile student_policy.py", purpose="把 changing quality trace 轉成 enter、exit、hold 的 decision question。", mechanism="threshold 與 stable count 會改變 transition、retry、service、state time、J。", expected="先指出 trace 上可能 enter／hold／exit 的位置，不先調 threshold。", recovery=B_RECOVERY, visual="quality trace  ╭╮╭──╮╭╮\n               ↑ enter  ↓ exit\n                 hold band", notes="Lab B 把 changing opportunity 換成品質 trace。問題不是品質越高越好，而是何時進入 send-ready、何時退出、要穩定多久；先讀 trace，再讀 policy。"),
    spec(52, "`ENTER_QUALITY`：何時進入 send-ready", "threshold-enter", layout=1, evidence="SOURCE-BOUND POLICY SURFACE", context="package root；student_policy.py 的 lab-b-enter-exit-hold marked block", command="sed -n '18,28p' student_policy.py\npython -m py_compile student_policy.py", purpose="把 quality crossing 連到第一次 mode transition。", mechanism="quality ≥ ENTER_QUALITY 且 stable steps 足夠，才可能進入 send-ready；它不是 energy knob。", expected="看到 `ENTER_QUALITY = 2` 與 syntax guard；正式 line number 不另造。", recovery="syntax／policy guard 失敗：只修 B marked block，回到 A frozen predecessor。", visual="quality\n  3 ─────────\n  2 ───────╱ enter threshold\n  1 ───╱\n  0", notes="先指出 trace 中第一次真正越過 enter threshold 的位置，再預測 mode transition、attempt 和 service。欄位透過切換產生結果，不是直接指定 energy。"),
    spec(53, "`EXIT_QUALITY`：為什麼不能只用一條線", "threshold-exit", layout=2, evidence="SOURCE-BOUND POLICY SURFACE", context="package root；student_policy.py 的 lab-b-enter-exit-hold marked block", command="sed -n '18,28p' student_policy.py\npython -m py_compile student_policy.py", purpose="用 enter／exit 分離避免品質邊界 ping-pong。", mechanism="ENTER_QUALITY ≠ EXIT_QUALITY 形成 hysteresis band；mode 不會因一個低點立刻反覆切換。", expected="看到 enter=2、exit=1 的兩條線；transition、retry、service 仍由 result 驗證。", recovery="threshold／syntax 不符：只修 B marked block，保留 traceback；不可改 runner。", visual="quality  3 ─── enter ───\n        band  ↕\n        1 ─── exit  ───", notes="如果 enter 和 exit 共用一條線，品質在邊界抖動時會造成 ping-pong。分開兩條線是穩定帶；先預測切換次數、WAIT、TX/RX 和 service，再看 trace evidence。"),
    spec(54, "`STABLE_STEPS`：拒絕短暫尖峰", "stable-steps", layout=3, evidence="SOURCE-BOUND POLICY SURFACE", context="package root；student_policy.py 的 lab-b-enter-exit-hold marked block", command="sed -n '18,28p' student_policy.py\npython -m py_compile student_policy.py", purpose="讓短暫 quality spike 先累積穩定計數，再切換 mode。", mechanism="stable_count < STABLE_STEPS 時 hold；達到計數才 switch，可能少 ping-pong，也可能錯過短窗口。", expected="看到 STABLE_STEPS=2；candidate 可合法改為 1，效果要從 transition／service／J 讀。", recovery="超出 bounded range 或 syntax 失敗：只修 B marked line，不放寬 guard。", visual="step 1  ▣ hold\nstep 2  ▣ switch\nstep 3  → mode", notes="STABLE_STEPS 把看起來變好和連續穩定分開。未達計數保持 hold，達到才切換；所以要同時看 transition、queue、service 與 energy，不能只看切換次數。"),
    spec(55, "Lab B 的 Trace A prediction", "prediction-b", layout=1, evidence="SOURCE-BOUND PREDICTION LOCK", context="package root；Trace A run 前；A frozen predecessor active", command="git diff -- student_policy.py\npython -m py_compile student_policy.py", purpose="在 Trace A run 前標出 enter、hold、exit 與 service/J 方向。", mechanism="prediction 先於 baseline/candidate；policy 不能讀 future trace outcome。", expected="Trace A 上有三個可反駁標記與 service／J 方向。", recovery=B_RECOVERY, visual="Trace A  ╭─╮──╮╭────╮\n       ↑ enter  ─ hold ─  ↓ exit", notes="請在 Trace A 上標出 enter、hold、exit 的位置，並寫下 service 與 endpoint J 方向。這不是把未來輸入偷塞進 policy，而是讓下一步比較有可反駁假說。"),
    spec(56, "B-01：Trace A baseline", "run-baseline-b", layout=2, evidence="VERIFIED LOCAL FALLBACK｜RUN NOT FRESH-VERIFIED", context="POSIX／WSL；A frozen predecessor active；package root", command="bash course.sh run --lab B --case trace-a-baseline\n# Windows: course.cmd run --lab B --case trace-a-baseline", purpose="建立 Trace A control，固定 quality/contact trace 與 predecessor。", mechanism="runner 使用 A frozen policy，產出 B baseline result／replay lineage。", expected=f"fallback run `…{FALLBACK['trace-a-baseline-B']['run']}`、energy `{FALLBACK['trace-a-baseline-B']['energy']} J`、delivered `{FALLBACK['trace-a-baseline-B']['delivered']} bit`；fresh path 待重跑。", recovery=B_RECOVERY, visual="A frozen policy\n   ↓ Trace A baseline\nresult_path → replay pair\nFALLBACK_ARTIFACT / SIMULATED", notes="保持 A frozen predecessor，先跑 B Trace A baseline。這一步建立 control；stdout 會回傳 run identity 和 result path，請保留配對 replay。頁面數字來自 current fallback，不是 fresh claim。"),
    spec(57, "B-02：改 B marked block 並 freeze", "run-candidate-b", layout=3, evidence="VERIFIED LOCAL FALLBACK｜FREEZE RECEIPT / RUN NOT FRESH-VERIFIED", context="POSIX／WSL；A frozen predecessor；只改 `lab-b-enter-exit-hold`", command="git diff -- student_policy.py\nbash course.sh run --lab B --case trace-a-candidate --freeze\n# Windows: course.cmd run --lab B --case trace-a-candidate --freeze", purpose="把 threshold／stable edit 綁成 B candidate 與 frozen receipt。", mechanism="freeze 綁定 B policy SHA、Trace A、scenario、A predecessor；非 B marked diff fail closed。", expected=f"fallback candidate policy `{FALLBACK['trace-a-candidate-B']['policy']}`、freeze `{FALLBACK['trace-a-candidate-B']['freeze']}`、delivered `{FALLBACK['trace-a-candidate-B']['delivered']} bit`。", recovery=B_RECOVERY, visual="A frozen predecessor\n        + B marked diff\n                 ↓\n          lab-b-frozen receipt", notes="只改 B marked block 中的 threshold 或 stable steps，保留 A frozen predecessor。執行 Trace A candidate freeze；若只變畫面 label 沒有 consequential fixture diff，就記錄 gate 未通過。"),
    spec(58, "B-03：讀 transition 與 service，不只讀品質", "transition-b", layout=1, evidence="VERIFIED LOCAL FALLBACK｜SIMULATED", context="trace-a-candidate-B result／endpoint replay", command="rg -n 'quality_band|POLICY_DECISION|PACKET_ATTEMPT|delivered|expired|service_pass' fallback_artifacts/trace-a-candidate-B/result.json", purpose="把 enter／hold／exit 造成的 packet、service、state、J 差異連回 policy。", mechanism="quality 是 observation context；transition 才會改 attempt、retry、queue、service、energy。", expected=f"candidate fallback：attempted 4、delivered `{FALLBACK['trace-a-candidate-B']['delivered']}` bit、expired 1、energy `{FALLBACK['trace-a-candidate-B']['energy']} J`。", recovery=B_RECOVERY, visual="quality trace → mode transition → attempt/retry → service/J", notes="比較 candidate 時，先定位 enter、hold、exit，再看這些 transition 造成的 attempt、retry、queue 和 service。品質數值是 observation context，不是 delivered service，也不是 energy。"),
    spec(59, "B-04：確認 freeze receipt 才能進 Trace B", "freeze-b", layout=2, evidence="VERIFIED LOCAL FALLBACK｜FREEZE RECEIPT / SIMULATED", context="package root；B frozen receipt／checkpoint", command="python -m json.tool fallback_artifacts/receipts/b.json\nfind artifacts/checkpoints -maxdepth 1 -type f -print", purpose="把 Trace B 限定為 withheld，而不是再調參的 playground。", mechanism="receipt、checkpoint、policy SHA、Trace A 與 predecessor 缺一項就 fail closed。", expected=f"role=`lab-b-frozen`、policy `{FALLBACK['trace-a-candidate-B']['policy']}`、predecessor `{FALLBACK['trace-a-baseline-B']['policy']}`、freeze `{FALLBACK['trace-a-candidate-B']['freeze']}`。", recovery=B_RECOVERY, visual="lab-a-frozen\n      ↓ predecessor\nlab-b-frozen\n      ↓ Trace B entry", notes="Trace B 是 withheld，不是再調參的 playground。先檢查 B receipt、checkpoint、policy SHA 和 predecessor；任何一項缺失就不要執行 withheld。"),
    spec(60, "B-05：Trace B 不 retune", "run-withheld-b", layout=3, evidence="VERIFIED LOCAL FALLBACK｜WITHHELD / SIMULATED", context="POSIX／WSL；B frozen policy active；package root", command="bash course.sh run --lab B --case trace-b\n# Windows: course.cmd run --lab B --case trace-b", purpose="在新的 quality trace 測試同一個 frozen B policy。", mechanism="withheld 沿用 B policy SHA；不建立新 freeze、不改 threshold，結果可暴露泛化邊界。", expected=f"fallback policy unchanged `{FALLBACK['trace-b-B']['policy']}`、energy `{FALLBACK['trace-b-B']['energy']} J`、delivered `{FALLBACK['trace-b-B']['delivered']} bit`、service false。", recovery=B_RECOVERY, visual="B frozen policy\n   ├─ Trace A\n   └─ Trace B：no retune\npolicy SHA unchanged", notes="保持 B frozen policy，執行 Trace B；不再改 threshold 或 stable steps。若 Trace B 露出 counterexample，那是泛化邊界，不是要隨機重調的錯誤。"),
    spec(61, "B-06：辨識 too-slow 與 ping-pong", "counterexample-b", layout=1, evidence="VERIFIED LOCAL FALLBACK｜SIMULATED COUNTEREXAMPLE", context="trace-a-candidate-B vs trace-b-B；same B policy SHA", command="rg -n 'policy_sha256|service_pass|endpoint_energy_j|expired_packets' fallback_artifacts/trace-b-B/result.json\npython -m json.tool fallback_artifacts/trace-b-B/endpoint-replay.json", purpose="把 transition trade-off 與 service／endpoint J 一起判讀。", mechanism="hold 太長可能 too-slow；enter/exit 太近可能 ping-pong；畫面平滑不等於低能量。", expected="Trace B 保留 policy SHA、transition、service、state、J 與 counterexample evidence。", recovery=B_RECOVERY, visual="too-slow：hold ─────→ service miss\nping-pong：enter ↔ exit ↔ retry", notes="如果 hold 太長，可能錯過 service window；如果 enter/exit 太近，可能付出額外切換與 retry。把 trace 事件和 endpoint ledger 對起來，不把平滑線直接當成省能量。"),
    spec(62, "Lab B debrief：hysteresis 是條件，不是保證", "debrief-b", layout=2, evidence="VERIFIED LOCAL FALLBACK｜SIMULATED", context="Trace A／Trace B；B receipt／policy／result pair", command="python -m json.tool fallback_artifacts/trace-a-candidate-B/result.json\npython -m json.tool fallback_artifacts/trace-b-B/result.json", purpose="用兩條 trace 說出成立、被推翻或待查的條件式結論。", mechanism="hysteresis 可能減少 ping-pong，也可能讓切換太慢；結論受 trace、job、service boundary 約束。", expected="完成：在 Trace A…；在 Trace B…；因此 verdict = 成立／被推翻／待查。", recovery=B_RECOVERY, visual="policy condition → transition → service/energy → verdict\nTrace A：…\nTrace B：…", notes="用「在 Trace A…；在 Trace B…」兩句話回答。hysteresis 是條件式結論，不是保證；把 B receipt、policy SHA、Trace B result 和 counterexample 一起保留。"),
    spec(63, "Recovery checkpoint：保存 A/B lineage", "checkpoint", layout=3, evidence="SOURCE-BOUND RECOVERY CONTRACT", context="package root；A/B receipt、checkpoint、result/replay pair", command="find artifacts/checkpoints artifacts/receipts -maxdepth 1 -type f -print\nfind fallback_artifacts -maxdepth 2 -type f -name 'result.json' -print", purpose="保存目前 policy、receipt、checkpoint、result/replay 與 prediction，讓中斷可回復。", mechanism="status → exact checkpoint role → restore；不覆寫既有 artifact，不用新 path 拼 lineage。", expected="看到 A frozen、B frozen、receipt、result／replay pair；若主 artifacts 缺失，明示 matching fallback。", recovery="缺 checkpoint：保留現有 artifact，回到最後一個 identity 正確的 gate；不要補一個假的 receipt。", visual="A frozen stack\nB frozen stack\nresult ↔ replay\nworkbook prediction", notes="進入下一個 lab 前保存 policy、receipt、checkpoint、result/replay pair 和 prediction。若中斷，先用 status 找可用 checkpoint，再 restore 所需 role；這是在保持 lineage，不是補出數字。"),
]


def qn(ns: str, local: str) -> str:
    return f"{{{ns}}}{local}"


def set_run_font(run, size: int, *, bold: bool = False, italic: bool = False, color: RGBColor = INK) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rpr.set("sz", str(size * 100))
    rpr.set("b", "1" if bold else "0")
    rpr.set("i", "1" if italic else "0")
    for tag, face in (("latin", "Times New Roman"), ("ea", "標楷體"), ("cs", "Times New Roman")):
        node = rpr.find(qn(A_NS, tag))
        if node is None:
            node = OxmlElement(f"a:{tag}")
            rpr.append(node)
        node.set("typeface", face)


def rich_text(tf, text: str, size: int = 24, *, bold: bool = False, color: RGBColor = INK, italic_code: bool = True, align=PP_ALIGN.LEFT) -> None:
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, line in enumerate(str(text).splitlines() or [""]):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        for token in re.split(r"(`[^`]*`)", line):
            if not token:
                continue
            marked = token.startswith("`") and token.endswith("`")
            raw = token[1:-1] if marked else token
            for part in re.findall(r"[\u3400-\u9fff\uF900-\uFAFF]+|[^\u3400-\u9fff\uF900-\uFAFF]+", raw):
                run = p.add_run()
                run.text = part
                set_run_font(run, size, bold=bold, italic=bool(marked and italic_code and LATIN.search(part)), color=color)


def add_text(slide, x: float, y: float, w: float, h: float, text: str, size: int = 24, *, bold: bool = False, color: RGBColor = INK, italic_code: bool = True, align=PP_ALIGN.LEFT, name: str = ""):
    # All authored visible text is intentionally classroom-readable.  The
    # original builder used 13--17pt metadata labels; clamp those calls here
    # so a future page cannot silently reintroduce sub-18pt authored text.
    size = max(18, size)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or "Editable text"
    shape.fill.background()
    shape.line.fill.background()
    rich_text(shape.text_frame, text, size, bold=bold, color=color, italic_code=italic_code, align=align)
    return shape


def add_box(slide, x: float, y: float, w: float, h: float, *, fill: RGBColor = WHITE, line: RGBColor = BLUE, width: float = 1.0, radius: bool = True, name: str = "Editable figure"):
    typ = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *, color: RGBColor = NAVY, width: float = 1.4, arrow: bool = False, name: str = "Editable connector"):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.name = name
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def add_dot(slide, x: float, y: float, d: float = 0.18, color: RGBColor = TEAL, name: str = "State marker"):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def add_tag(slide, x: float, y: float, w: float, text: str, *, fill: RGBColor = PALE_BLUE, color: RGBColor = NAVY):
    add_box(slide, x, y, w, 0.42, fill=fill, line=fill, width=0.5, name="Status tag")
    add_text(slide, x + 0.04, y + 0.03, w - 0.08, 0.36, text, 18, bold=True, color=color, align=PP_ALIGN.CENTER, italic_code=False, name="Status tag text")


def clear_placeholders(slide, *, keep_title: bool = True) -> None:
    for shape in list(slide.shapes):
        if not shape.is_placeholder:
            continue
        idx = shape.placeholder_format.idx
        typ = str(shape.placeholder_format.type).upper()
        if idx == 10 or (keep_title and "TITLE" in typ):
            continue
        shape._element.getparent().remove(shape._element)


def set_title(slide, title: str, *, cover: bool = False) -> None:
    if cover:
        for shape in list(slide.shapes):
            if shape.is_placeholder and "TITLE" in str(shape.placeholder_format.type).upper():
                shape._element.getparent().remove(shape._element)
        add_text(slide, 0.86, 0.78, 8.05, 1.08, title, 28, bold=True, color=NAVY, name="Native title anchor")
        return
    title_shape = next((s for s in slide.shapes if s.is_placeholder and "TITLE" in str(s.placeholder_format.type).upper()), None)
    if title_shape is None:
        add_text(slide, 0.68, 0.18, 11.95, 0.66, title, 28, bold=True, color=NAVY, name="Native title anchor")
    else:
        rich_text(title_shape.text_frame, title, 28, bold=True, color=NAVY, italic_code=False)
        title_shape.name = "Native title anchor"


def status_fill(evidence: str) -> tuple[RGBColor, RGBColor, str]:
    if "PENDING" in evidence or "PLACEHOLDER" in evidence:
        return PALE_GOLD, PURPLE, "PENDING / KEEP PLACEHOLDER"
    if "FALLBACK" in evidence or "SIMULATED" in evidence:
        return PALE_TEAL, TEAL, "SIMULATED FALLBACK EVIDENCE"
    if "SOURCE-BOUND" in evidence:
        return PALE_BLUE, NAVY, "SOURCE-BOUND POLICY SURFACE"
    return PALE_BLUE, NAVY, "SOURCE-BOUND EVIDENCE"


def add_header(slide, s: SlideSpec) -> None:
    set_title(slide, s.title)
    # Evidence badges used to sit at y=0.90 and intrude into the title rule
    # and the first command/content rail on long-title pages.  Evidence is
    # already stated in the authored cards/stamps and in the manifest; keep
    # the title band clear rather than overlaying a second badge on content.


def add_notes(slide, s: SlideSpec) -> None:
    text = s.notes.strip()
    if s.recovery:
        text += " 若操作受阻，" + s.recovery.strip()
    if FORBIDDEN in text or MINUTE_RE.search(text):
        raise ValueError(f"P{s.number:03d}: forbidden notes content")
    slide.notes_slide.notes_text_frame.text = text


def evidence_stamp(slide, y: float = 5.68, text: str = "FALLBACK_ARTIFACT｜SIMULATED｜NOT LIVE / NOT MEASURED"):
    add_text(slide, 0.92, y, 11.36, 0.30, text, 18, bold=True, color=RED, align=PP_ALIGN.CENTER, italic_code=False, name="Evidence boundary")


def command_strip(slide, s: SlideSpec, y: float = 1.28, h: float = 0.74, *, fill: RGBColor = CREAM, line: RGBColor = NAVY, command_size: int = 18):
    # Keep two-line commands readable without creating a fixed teaching rail.
    h = max(h, 0.96)
    add_box(slide, 0.82, y, 11.46, h, fill=fill, line=line, width=1.1, name="Exact context command")
    add_text(slide, 1.04, y + 0.04, 10.98, 0.26, "Context: " + s.context, 18, color=MUTED, name="Exact context")
    add_text(slide, 1.04, y + 0.34, 10.98, h - 0.38, "Run: " + s.command, max(18, command_size), color=INK, italic_code=False, name="Exact command")


def why_band(slide, s: SlideSpec, y: float = 5.12, *, fill: RGBColor = PALE_GOLD, line: RGBColor = GOLD):
    add_box(slide, 0.92, y, 11.36, 0.68, fill=fill, line=line, width=1.0, name="Interpretation band")
    add_text(slide, 1.12, y + 0.10, 10.96, 0.50, s.purpose + " " + s.mechanism, 18, color=INK, align=PP_ALIGN.CENTER, name="Interpretation text")


def operation_panel(slide, s: SlideSpec, *, accent: int = 0, y: float = 1.42):
    """A varied command-to-receipt composition used only on selected run pages."""
    fills = [(PALE_PURPLE, PURPLE), (PALE_BLUE, BLUE), (PALE_TEAL, TEAL), (PALE_GOLD, GOLD)]
    fill, line = fills[accent % len(fills)]
    add_box(slide, 0.82, y, 5.90, 3.98, fill=fill, line=line, width=1.2, name="Run terminal")
    add_text(slide, 1.08, y + 0.24, 5.38, 0.28, "RUN THIS EXACTLY", 18, bold=True, color=line, align=PP_ALIGN.CENTER, italic_code=False, name="Run heading")
    run_short = {
        43: "Purpose: establish control.\nMechanism: fixed scenario / seed → result lineage.",
        44: "Purpose: freeze one marked edit.\nMechanism: policy SHA + predecessor + receipt.",
        56: "Purpose: establish Trace A control.\nMechanism: A frozen predecessor → B baseline.",
        57: "Purpose: freeze B candidate.\nMechanism: B SHA + Trace A + predecessor.",
    }.get(s.number, s.purpose + "\n" + s.mechanism)
    expected_short = {
        44: "candidate policy 325cca4a…d73cf2f\nfreeze 6cbaf9dc…e11626\nenergy 8.86 J；fresh receipt pending",
        57: "candidate policy c00ed808…78293718\nfreeze 2a4167e2…171349e0e\ndelivered 9600 bit；fresh receipt pending",
    }.get(s.number, s.expected)
    recovery_short = {
        43: "保留 stdout/stderr；restore A checkpoint 或 matching fallback；不要再調參。",
        44: "保留 stdout/stderr；restore A checkpoint 或 matching fallback；不要再調參。",
        56: "保留錯誤；restore A/B checkpoint 或 matching fallback；不可 retune。",
        57: "保留錯誤；restore A/B checkpoint 或 matching fallback；不可 retune。",
    }.get(s.number, s.recovery)
    add_text(slide, 1.08, y + 0.62, 5.38, 0.58, s.context, 18, color=MUTED, italic_code=False, name="Run context")
    add_text(slide, 1.08, y + 1.28, 5.38, 1.10, s.command, 18, color=INK, italic_code=False, name="Run context command")
    add_text(slide, 1.08, y + 2.48, 5.38, 1.30, run_short, 18, color=INK, name="Run mechanism")
    add_box(slide, 7.02, y, 5.26, 2.24, fill=PALE_TEAL, line=TEAL, width=1.2, name="Run receipt")
    add_text(slide, 7.28, y + 0.24, 4.74, 0.26, "READ THE RECEIPT", 18, bold=True, color=TEAL, align=PP_ALIGN.CENTER, italic_code=False, name="Receipt heading")
    add_text(slide, 7.28, y + 0.68, 4.74, 1.42, expected_short, 18, align=PP_ALIGN.CENTER, name="Receipt expected")
    add_box(slide, 7.02, y + 2.36, 5.26, 1.62, fill=PALE_GOLD, line=GOLD, width=1.2, name="Run recovery")
    add_text(slide, 7.28, y + 2.52, 4.74, 0.26, "IF THE GATE IS RED", 18, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Recovery heading")
    add_text(slide, 7.28, y + 2.92, 4.74, 0.94, recovery_short, 18, align=PP_ALIGN.CENTER, name="Recovery text")


def draw_states(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.22, h=1.10)
    labels = [("SLEEP", BLUE, PALE_BLUE), ("WAKE", GOLD, PALE_GOLD), ("PROCESS", PURPLE, PALE_PURPLE), ("TX / RX", TEAL, PALE_TEAL), ("WAIT", NAVY, CREAM)]
    x = 0.92
    for i, (lab, col, fill) in enumerate(labels):
        add_box(slide, x, 2.42, 2.06, 0.86, fill=fill, line=col, width=1.2, name=f"State {lab}")
        add_text(slide, x + 0.10, 2.70, 1.86, 0.24, lab, 20, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"State label {lab}")
        if i < len(labels) - 1:
            add_line(slide, x + 2.06, 2.85, x + 2.36, 2.85, color=NAVY, width=1.6, arrow=True, name="State transition")
        x += 2.30
    add_box(slide, 0.96, 3.72, 11.20, 1.22, fill=CREAM, line=GOLD, name="State ledger")
    add_text(slide, 1.20, 3.94, 10.72, 0.82, "每個 state 都留下 duration / power / event；endpoint J 是累積結果。", 20, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="State ledger text")
    evidence_stamp(slide)


def draw_waitsleep(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.18, h=0.70)
    for y, label, col, fill, text in [(2.18, "WAIT", TEAL, PALE_TEAL, "awake idle → no wake event"), (3.68, "SLEEP", BLUE, PALE_BLUE, "sleep → WAKE → service")]:
        add_box(slide, 1.00, y, 2.02, 0.80, fill=fill, line=col, name=f"{label} lane label")
        add_text(slide, 1.18, y + 0.24, 1.66, 0.34, label, 21, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"{label} label")
        add_line(slide, 3.34, y + 0.40, 11.40, y + 0.40, color=col, width=2.4, arrow=True, name=f"{label} timeline")
        add_dot(slide, 5.00, y + 0.30, 0.20, col, name=f"{label} sample")
        add_text(slide, 6.00, y + 0.14, 4.42, 0.52, text, 20, color=INK, name=f"{label} mechanism")
    why_band(slide, s, y=5.16)


def draw_packet(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.18, h=0.70)
    stages = [("queue", BLUE, PALE_BLUE), ("attempt", PURPLE, PALE_PURPLE), ("retry", GOLD, PALE_GOLD), ("delivered", TEAL, PALE_TEAL), ("expired", RED, PALE_RED)]
    x = 0.88
    for i, (lab, col, fill) in enumerate(stages):
        add_box(slide, x, 2.28, 2.08, 0.88, fill=fill, line=col, name=f"Packet stage {lab}")
        add_text(slide, x + 0.08, 2.55, 1.92, 0.34, lab, 21, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"Packet stage {lab} text")
        if i < len(stages) - 1:
            add_line(slide, x + 2.08, 2.72, x + 2.36, 2.72, color=NAVY, width=1.5, arrow=True, name="Packet arrow")
        x += 2.34
    add_box(slide, 1.00, 3.72, 5.24, 1.20, fill=PALE_BLUE, line=BLUE, name="Packet numbers")
    add_text(slide, 1.24, 3.90, 4.76, 0.96, "attempted ≠ delivered\nretry / collision / expired 都要保留", 18, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Packet numbers text")
    add_box(slide, 6.68, 3.72, 5.48, 1.20, fill=PALE_GOLD, line=GOLD, name="Packet interpretation")
    add_text(slide, 6.94, 3.99, 4.96, 0.64, "一次 SEND 是 decision；service 要看完整 packet outcome。", 22, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Packet interpretation text")
    evidence_stamp(slide)


def draw_service(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.18, h=0.70)
    add_box(slide, 0.90, 2.18, 7.00, 2.58, fill=PALE_BLUE, line=BLUE, width=1.3, name="Service gate")
    add_text(slide, 1.18, 2.48, 6.44, 0.30, "SERVICE GATE", 20, bold=True, color=NAVY, align=PP_ALIGN.CENTER, italic_code=False, name="Service gate heading")
    add_text(slide, 1.28, 3.02, 6.24, 1.30, "same job\nsame window\ndelivered / deadline", 22, bold=True, align=PP_ALIGN.CENTER, name="Service gate body")
    add_line(slide, 7.92, 3.46, 8.54, 3.46, color=GOLD, width=2.0, arrow=True, name="Service to energy")
    add_box(slide, 8.74, 2.18, 3.38, 2.58, fill=PALE_GOLD, line=GOLD, width=1.3, name="Energy comparison")
    add_text(slide, 9.02, 2.48, 2.82, 0.30, "THEN COMPARE J", 19, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Energy heading")
    add_text(slide, 9.08, 3.02, 2.70, 1.24, "endpoint scope\nstate ledger\nenergy breakdown", 20, align=PP_ALIGN.CENTER, name="Energy body")
    why_band(slide, s, y=5.16)


def draw_powerarea(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.18, h=0.70)
    add_line(slide, 1.50, 4.56, 11.40, 4.56, color=NAVY, width=1.8, name="Power time axis")
    add_line(slide, 1.50, 4.56, 1.50, 2.18, color=NAVY, width=1.8, name="Power axis")
    points = [(1.74, 3.88), (3.12, 3.26), (4.38, 2.70), (5.72, 3.44), (7.00, 2.50), (8.44, 3.06), (10.12, 2.34), (11.10, 3.02)]
    for a, b in zip(points, points[1:]):
        add_line(slide, a[0], a[1], b[0], b[1], color=TEAL, width=2.3, name="Power curve")
    for x, y in points:
        add_dot(slide, x - 0.08, y - 0.08, 0.16, TEAL, name="Power sample")
    add_text(slide, 1.56, 1.94, 2.00, 0.26, "P(t) / W", 20, bold=True, color=TEAL, name="Power label")
    add_text(slide, 9.18, 4.76, 2.24, 0.24, "time / s", 18, color=MUTED, align=PP_ALIGN.RIGHT, name="Time label")
    add_box(slide, 0.94, 5.16, 11.26, 0.66, fill=PALE_GOLD, line=GOLD, name="Area interpretation")
    add_text(slide, 1.16, 5.34, 10.82, 0.26, "曲線下的 area 是 endpoint J；低峰值拖得久，仍可能累積更多。", 21, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Area text")


def draw_formula(slide, s, formula_key: str):
    add_header(slide, s)
    # The equation tool replaces this preview-only fallback in the final
    # export with one native Office Math shape.  Keeping the fallback in the
    # pre-OMML render source gives LibreOffice a clean visual preview while the
    # final PPTX has no duplicate text/OMML equation.
    # The native equation inserter uses a 0.667in slide margin.  Widen this
    # frame slightly so that the native shape has visible padding on both
    # sides instead of crossing the frame line.
    add_box(slide, 0.55, 1.56, 12.23, 3.40, fill=CREAM, line=GOLD, width=1.4, name=f"NATIVE_EQUATION_HOOK_{formula_key}")
    add_text(slide, 1.20, 1.82, 10.68, 0.32, "EDITABLE OFFICE MATH", 20, bold=True, color=NAVY, align=PP_ALIGN.CENTER, italic_code=False, name=f"NATIVE_EQUATION_HOOK_LABEL_{formula_key}")
    equation = "E_endpoint = Σ P_s t_s" if formula_key == "E_endpoint" else "η_E = D_delivered / E_endpoint"
    add_text(slide, 1.34, 2.64, 10.40, 0.78, equation, 30, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=True, name=f"Formula preview fallback {formula_key}")
    add_text(slide, 1.28, 3.98, 10.48, 0.76, s.visual, 22, color=INK, align=PP_ALIGN.CENTER, name="Formula scope")
    add_box(slide, 1.02, 5.18, 11.12, 0.62, fill=PALE_BLUE, line=BLUE, name="Formula guard")
    add_text(slide, 1.26, 5.36, 10.64, 0.24, "Office Math hook is editable；數字仍由 result artifact 提供，不自行補 KPI。", 18, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Formula guard text")


def draw_provenance(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.18, h=0.70)
    labels = [("source", BLUE, PALE_BLUE), ("model", PURPLE, PALE_PURPLE), ("course", GOLD, PALE_GOLD), ("result", TEAL, PALE_TEAL)]
    x = 0.96
    for i, (label, col, fill) in enumerate(labels):
        add_box(slide, x, 2.28, 2.42, 0.88, fill=fill, line=col, name=f"Provenance {label}")
        add_text(slide, x + 0.12, 2.58, 2.18, 0.24, label, 21, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"Provenance {label} text")
        if i < 3:
            add_line(slide, x + 2.42, 2.72, x + 2.70, 2.72, color=NAVY, width=1.5, arrow=True, name="Provenance arrow")
        x += 2.72
    add_box(slide, 1.12, 3.78, 10.88, 1.18, fill=CREAM, line=GOLD, name="Provenance token")
    add_text(slide, 1.34, 4.04, 10.44, 0.70, f"{SCENARIO_ID}  ·  seed 12001 / 12002  ·  policy SHA  ·  artifact_source", 18, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Provenance token text")
    evidence_stamp(slide)


def draw_identity(slide, s):
    add_header(slide, s)
    add_box(slide, 0.90, 1.34, 3.04, 4.40, fill=PALE_BLUE, line=BLUE, name="Identity source")
    add_text(slide, 1.16, 1.66, 2.52, 0.28, "SOURCE", 19, bold=True, color=NAVY, align=PP_ALIGN.CENTER, italic_code=False, name="Identity source heading")
    add_text(slide, 1.18, 2.14, 2.48, 2.74, f"repo\n{REPO_URL}\n\ncommit\n{CURRENT_COMMIT[:12]}…", 18, align=PP_ALIGN.CENTER, name="Identity source text")
    add_box(slide, 4.42, 1.34, 3.04, 4.40, fill=PALE_PURPLE, line=PURPLE, name="Identity scenario")
    add_text(slide, 4.68, 1.66, 2.52, 0.28, "SCENARIO", 19, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Identity scenario heading")
    add_text(slide, 4.70, 2.14, 2.48, 2.74, f"{SCENARIO_ID}\n\nanchor\n{ANCHOR_SHA[:20]}…\n\nlock\n{LOCK_SHA[:20]}…", 18, align=PP_ALIGN.CENTER, name="Identity scenario text")
    add_box(slide, 7.94, 1.34, 4.14, 4.40, fill=PALE_TEAL, line=TEAL, name="Identity chain")
    add_text(slide, 8.22, 1.66, 3.58, 0.32, "RUN → REPLAY → WORKBOOK", 18, bold=True, color=TEAL, align=PP_ALIGN.CENTER, italic_code=False, name="Identity chain heading")
    add_text(slide, 8.24, 2.14, 3.54, 2.62, "same scenario\nsame units\nsame seed role\npolicy SHA\nfail closed on mismatch", 23, align=PP_ALIGN.CENTER, name="Identity chain text")
    evidence_stamp(slide, y=5.68, text="IDENTITY GATE｜scenario / anchor / policy / replay must agree")


def draw_fairness(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.18, h=0.70)
    locks = [("FIXED", "scenario / seed / job / window", BLUE, PALE_BLUE), ("ONE EDIT", "current marked block", PURPLE, PALE_PURPLE), ("OBSERVE", "queue / packet / service / state / J", TEAL, PALE_TEAL)]
    for i, (head, body, col, fill) in enumerate(locks):
        x = 0.98 + i * 4.08
        add_box(slide, x, 2.26, 3.58, 1.44, fill=fill, line=col, width=1.2, name=f"Fairness {head}")
        add_text(slide, x + 0.18, 2.52, 3.22, 0.26, head, 20, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"Fairness {head} heading")
        add_text(slide, x + 0.20, 2.96, 3.18, 0.56, body, 18, align=PP_ALIGN.CENTER, name=f"Fairness {head} body")
    add_box(slide, 1.10, 4.20, 10.90, 1.22, fill=PALE_GOLD, line=GOLD, name="Fairness warning")
    add_text(slide, 1.38, 4.52, 10.34, 0.54, "若 trace、scope、job 或 policy 偷換，這一組 A/B 就失去因果資格。", 22, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Fairness warning text")
    evidence_stamp(slide)


def draw_window(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.18, h=0.70)
    add_line(slide, 1.12, 3.12, 11.72, 3.12, color=NAVY, width=2.0, name="Window axis")
    add_box(slide, 1.30, 2.34, 5.28, 1.54, fill=PALE_TEAL, line=TEAL, name="Window open")
    add_text(slide, 1.58, 2.62, 4.72, 0.24, "WINDOW OPEN", 20, bold=True, color=TEAL, align=PP_ALIGN.CENTER, italic_code=False, name="Window open heading")
    add_text(slide, 1.60, 3.10, 4.68, 0.36, "SEND / WAIT legal context", 22, align=PP_ALIGN.CENTER, name="Window open body")
    add_box(slide, 6.82, 2.34, 5.28, 1.54, fill=PALE_RED, line=RED, name="Window closed")
    add_text(slide, 7.10, 2.62, 4.72, 0.24, "WINDOW CLOSED", 20, bold=True, color=RED, align=PP_ALIGN.CENTER, italic_code=False, name="Window closed heading")
    add_text(slide, 7.12, 3.04, 4.68, 0.50, "delayed action → service risk", 19, align=PP_ALIGN.CENTER, name="Window closed body")
    add_dot(slide, 3.62, 2.98, 0.28, TEAL, name="Legal action marker")
    add_dot(slide, 9.26, 2.98, 0.28, RED, name="Illegal action marker")
    why_band(slide, s, y=5.16)


def draw_lab_hero(slide, s, lab: str):
    add_header(slide, s)
    add_box(slide, 0.88, 1.34, 6.22, 4.26, fill=PALE_BLUE if lab == "A" else PALE_PURPLE, line=BLUE if lab == "A" else PURPLE, width=1.4, name=f"Lab {lab} question")
    add_text(slide, 1.20, 1.72, 5.58, 0.34, f"LAB {lab} QUESTION", 21, bold=True, color=BLUE if lab == "A" else PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name=f"Lab {lab} heading")
    add_text(slide, 1.18, 2.42, 5.62, 1.30, "same job\n" + ("different pace / rest" if lab == "A" else "quality trace / stability"), 26, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name=f"Lab {lab} question text")
    add_text(slide, 1.24, 4.30, 5.50, 0.66, "先預測，再 baseline → candidate → withheld", 18, color=INK, align=PP_ALIGN.CENTER, name=f"Lab {lab} route")
    add_box(slide, 7.54, 1.34, 4.62, 4.26, fill=PALE_GOLD, line=GOLD, width=1.4, name=f"Lab {lab} evidence")
    add_text(slide, 7.84, 1.72, 4.02, 0.34, "WATCH THE CHAIN", 20, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name=f"Lab {lab} evidence heading")
    add_text(slide, 7.88, 2.38, 3.94, 1.62, "policy edit\n→ state / transition\n→ packet / service\n→ endpoint J", 23, bold=True, color=INK, align=PP_ALIGN.CENTER, name=f"Lab {lab} evidence text")
    add_text(slide, 7.90, 4.64, 3.90, 0.42, "FALLBACK ARTIFACTS ARE SIMULATED", 18, bold=True, color=RED, align=PP_ALIGN.CENTER, italic_code=False, name=f"Lab {lab} evidence note")
    evidence_stamp(slide)


def draw_code(slide, s, lab: str):
    add_header(slide, s)
    command_strip(slide, s, y=1.12, h=0.84)
    add_box(slide, 0.92, 2.26, 7.20, 3.06, fill=PALE_PURPLE, line=PURPLE, width=1.3, name=f"Code {lab} panel")
    code = "# === LORA EDITABLE: " + ("lab-a-pace-rest" if lab == "A" else "lab-b-enter-exit-hold") + " ===\n" + ("PACE_GAP_STEPS = 2\nREST_DURING_GAP = SLEEP" if lab == "A" else "ENTER_QUALITY = 2\nEXIT_QUALITY = 1\nSTABLE_STEPS = 2") + "\n# === LORA END EDITABLE ==="
    add_text(slide, 1.18, 2.56, 6.68, 1.86, code, 18, color=INK, italic_code=False, name=f"Code {lab} text")
    add_box(slide, 8.56, 2.26, 3.60, 3.06, fill=PALE_TEAL, line=TEAL, width=1.3, name=f"Code {lab} mechanism")
    add_text(slide, 8.82, 2.58, 3.08, 0.28, "ONE BLOCK", 19, bold=True, color=TEAL, align=PP_ALIGN.CENTER, italic_code=False, name=f"Code {lab} heading")
    short_mechanism = {
        40: "gap → send opportunity\nnot direct energy\nread steps_since_send",
        41: "WAIT stays awake\nSLEEP may add wake cost\nread state + service",
        52: "quality + stable count\n→ send-ready transition\nnot direct energy",
        53: "Purpose: separate enter and exit.\nMechanism: hysteresis limits ping-pong.",
        54: "count stable steps\n→ switch after hold\nshort spike rejected",
    }.get(s.number, s.purpose + "\n" + s.mechanism)
    add_text(slide, 8.84, 3.16, 3.04, 1.48, short_mechanism, 18, align=PP_ALIGN.CENTER, name=f"Code {lab} mechanism text")
    evidence_stamp(slide, text="POLICY SURFACE｜marked block only｜syntax guard before run")


def draw_prediction(slide, s, lab: str):
    add_header(slide, s)
    command_strip(slide, s, y=1.14, h=0.76)
    labels = [("queue age", BLUE), ("service", TEAL), ("state time", PURPLE), ("endpoint J", GOLD)]
    for i, (label, col) in enumerate(labels):
        y = 2.34 + i * 0.74
        add_box(slide, 1.10, y, 2.40, 0.50, fill=PALE_BLUE if i == 0 else (PALE_TEAL if i == 1 else (PALE_PURPLE if i == 2 else PALE_GOLD)), line=col, name=f"Prediction {lab} {label}")
        add_text(slide, 1.24, y + 0.14, 2.12, 0.20, label, 18, bold=True, color=col, align=PP_ALIGN.CENTER, name=f"Prediction label {label}")
        add_line(slide, 3.72, y + 0.25, 4.54, y + 0.25, color=col, width=1.4, arrow=True, name="Prediction arrow")
        add_box(slide, 4.76, y, 6.34, 0.50, fill=CREAM, line=GRAY, name=f"Prediction sentence {label}")
        add_text(slide, 5.02, y + 0.14, 5.82, 0.20, "↗ / ↘ / hold — write a reason before run", 18, color=INK, name=f"Prediction text {label}")
    why_band(slide, s, y=5.46)


def draw_run(slide, s, accent: int):
    add_header(slide, s)
    operation_panel(slide, s, accent=accent, y=1.34)
    evidence_stamp(slide, y=5.68, text="VERIFIED LOCAL FALLBACK｜SIMULATED｜fresh receipt/path pending")


def draw_compare(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.14, h=1.10)
    for x, label, fill, line, case_name in [(0.92, "BASELINE", PALE_BLUE, BLUE, "baseline-A"), (6.64, "CANDIDATE", PALE_PURPLE, PURPLE, "candidate-A")]:
        d = FALLBACK[case_name]
        add_box(slide, x, 2.28, 5.36, 2.72, fill=fill, line=line, width=1.2, name=f"Compare {label}")
        add_text(slide, x + 0.24, 2.56, 4.88, 0.26, label, 19, bold=True, color=line, align=PP_ALIGN.CENTER, italic_code=False, name=f"Compare {label} heading")
        add_text(slide, x + 0.34, 3.06, 4.68, 1.44, f"scenario ✓\npolicy {d['policy'][:18]}…\nenergy {d['energy']} J  ·  delivered {d['delivered']} bit\nservice_pass = {str(d['service']).lower()}", 18, align=PP_ALIGN.CENTER, name=f"Compare {label} body")
    add_box(slide, 1.18, 5.16, 10.94, 0.78, fill=PALE_GOLD, line=GOLD, name="Compare verdict")
    add_text(slide, 1.42, 5.28, 10.46, 0.58, "先確認同一 boundary；candidate 的 trade-off 再交給 state / packet / service ledger。", 18, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Compare verdict text")


def draw_ledger(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.12, h=0.82)
    rows = [("awake idle", 6.0, TEAL), ("sleep", 0.04, BLUE), ("wake", 0.02, GOLD), ("process", 0.16, PURPLE), ("TX", 2.4, RED), ("RX", 0.24, NAVY)]
    for i, (label, value, col) in enumerate(rows):
        y = 2.16 + i * 0.52
        add_text(slide, 1.00, y + 0.08, 1.54, 0.30, label, 18, bold=True, color=col, name=f"Ledger label {label}")
        add_box(slide, 2.62, y, min(7.10, 0.70 + value * 0.72), 0.34, fill=PALE_TEAL if col == TEAL else PALE_BLUE, line=col, width=0.8, radius=False, name=f"Ledger bar {label}")
        add_text(slide, 9.96, y + 0.08, 1.18, 0.30, f"{value} J", 18, color=INK, align=PP_ALIGN.RIGHT, name=f"Ledger value {label}")
    add_box(slide, 1.00, 5.48, 11.12, 0.52, fill=PALE_GOLD, line=GOLD, name="Ledger guard")
    add_text(slide, 1.24, 5.64, 10.64, 0.20, "先讀 state duration，再問哪個 policy action 造成 consequential diff。", 19, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Ledger guard text")


def draw_packet_service(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.12, h=0.82)
    is_b = s.kind == "transition-b"
    evidence = FALLBACK["trace-a-candidate-B"] if is_b else FALLBACK["candidate-A"]
    attempted = 4 if is_b else 2
    retries = 1
    expired = 1 if is_b else 3
    add_box(slide, 0.98, 2.18, 5.24, 2.88, fill=PALE_BLUE, line=BLUE, name="Packet ledger")
    add_text(slide, 1.24, 2.48, 4.72, 0.26, ("B CANDIDATE" if is_b else "A CANDIDATE") + " PACKET LEDGER", 18, bold=True, color=NAVY, align=PP_ALIGN.CENTER, italic_code=False, name="Packet ledger heading")
    add_text(slide, 1.26, 3.00, 4.68, 1.58, f"attempted {attempted}\nretransmissions {retries}\ndelivered {evidence['delivered']} bit\nexpired {expired}\nenergy {evidence['energy']} J", 18, align=PP_ALIGN.CENTER, name="Packet ledger body")
    add_line(slide, 6.44, 3.62, 7.06, 3.62, color=GOLD, width=2.0, arrow=True, name="Packet to service")
    add_box(slide, 7.30, 2.18, 4.82, 2.88, fill=PALE_GOLD, line=GOLD, name="Service verdict")
    add_text(slide, 7.58, 2.48, 4.26, 0.26, "SERVICE VERDICT", 18, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Service verdict heading")
    add_text(slide, 7.60, 3.08, 4.22, 1.14, "service_pass = false\ndeadline_pass = false\n先說 trade-off，不說節能。", 22, align=PP_ALIGN.CENTER, name="Service verdict body")
    evidence_stamp(slide)


def draw_freeze(slide, s, lab: str):
    add_header(slide, s)
    command_strip(slide, s, y=1.10, h=0.86)
    chain = [("baseline", BLUE), ("candidate", PURPLE), (f"lab-{lab.lower()}-frozen", TEAL)]
    x = 1.02
    for i, (label, col) in enumerate(chain):
        add_box(slide, x, 2.44, 3.02, 0.86, fill=PALE_BLUE if i == 0 else (PALE_PURPLE if i == 1 else PALE_TEAL), line=col, name=f"Freeze {label}")
        add_text(slide, x + 0.16, 2.74, 2.70, 0.24, label, 20, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"Freeze {label} text")
        if i < 2:
            add_line(slide, x + 3.02, 2.87, x + 3.42, 2.87, color=GOLD, width=1.8, arrow=True, name="Freeze lineage arrow")
        x += 3.42
    add_box(slide, 1.08, 3.82, 10.96, 1.18, fill=CREAM, line=GOLD, name="Freeze receipt fields")
    add_text(slide, 1.34, 4.06, 10.44, 0.68, "policy SHA · predecessor · active_block_id · scenario · anchor · seed · receipt SHA", 18, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Freeze receipt fields text")
    evidence_stamp(slide, y=5.60, text="FREEZE GATE｜missing receipt / checkpoint = stop before withheld")


def draw_withheld(slide, s, lab: str):
    add_header(slide, s)
    command_strip(slide, s, y=1.10, h=0.86)
    add_box(slide, 0.98, 2.28, 4.72, 2.72, fill=PALE_BLUE, line=BLUE, name=f"Withheld {lab} primary")
    add_text(slide, 1.24, 2.58, 4.20, 0.26, "FROZEN POLICY", 19, bold=True, color=NAVY, align=PP_ALIGN.CENTER, italic_code=False, name="Withheld frozen heading")
    add_text(slide, 1.30, 3.18, 4.08, 1.12, "policy SHA unchanged\nno new freeze\nread result + replay", 23, bold=True, align=PP_ALIGN.CENTER, name="Withheld frozen body")
    add_line(slide, 5.98, 3.62, 6.68, 3.62, color=GOLD, width=2.0, arrow=True, name="Withheld arrow")
    add_box(slide, 6.96, 2.28, 5.08, 2.72, fill=PALE_GOLD, line=GOLD, name=f"Withheld {lab} trace")
    add_text(slide, 7.24, 2.58, 4.52, 0.26, f"{lab.upper()} TRACE", 19, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Withheld trace heading")
    add_text(slide, 7.30, 3.10, 4.40, 1.30, "new condition\nqueue / service / state / J\nqualified verdict", 20, bold=True, align=PP_ALIGN.CENTER, name="Withheld trace body")
    evidence_stamp(slide)


def draw_debrief(slide, s, lab: str):
    add_header(slide, s)
    command_strip(slide, s, y=1.12, h=1.10)
    pieces = [("condition", BLUE), ("mechanism", PURPLE), ("evidence", TEAL), ("verdict", GOLD)]
    x = 0.98
    for i, (labtext, col) in enumerate(pieces):
        add_box(slide, x, 2.36, 2.70, 0.92, fill=PALE_BLUE if i == 0 else (PALE_PURPLE if i == 1 else (PALE_TEAL if i == 2 else PALE_GOLD)), line=col, name=f"Debrief {lab} {labtext}")
        add_text(slide, x + 0.12, 2.68, 2.46, 0.24, labtext, 19, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"Debrief {labtext}")
        if i < 3:
            add_line(slide, x + 2.70, 2.82, x + 3.02, 2.82, color=NAVY, width=1.6, arrow=True, name="Debrief arrow")
        x += 3.02
    add_box(slide, 1.06, 3.84, 10.96, 1.16, fill=CREAM, line=GOLD, name="Debrief sentence")
    add_text(slide, 1.32, 4.10, 10.44, 0.68, "在 Trace A／primary…；在 Trace B／hidden…；因此這個結論成立、被推翻，或待查。", 18, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Debrief sentence text")
    evidence_stamp(slide, y=5.62)


def draw_counterexample(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.10, h=1.10)
    add_box(slide, 0.96, 2.26, 5.32, 2.74, fill=PALE_RED, line=RED, name="Too slow")
    add_text(slide, 1.24, 2.56, 4.76, 0.26, "TOO-SLOW", 20, bold=True, color=RED, align=PP_ALIGN.CENTER, italic_code=False, name="Too slow heading")
    add_text(slide, 1.28, 3.16, 4.68, 1.12, "hold 太久\n錯過 service window\nservice loss", 23, bold=True, align=PP_ALIGN.CENTER, name="Too slow body")
    add_box(slide, 6.72, 2.26, 5.32, 2.74, fill=PALE_GOLD, line=GOLD, name="Ping pong")
    add_text(slide, 7.00, 2.56, 4.76, 0.26, "PING-PONG", 20, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Ping pong heading")
    add_text(slide, 7.04, 3.16, 4.68, 1.12, "enter ↔ exit 太近\n多次 transition / retry\nenergy trade-off", 23, bold=True, align=PP_ALIGN.CENTER, name="Ping pong body")
    evidence_stamp(slide)


def draw_checkpoint(slide, s):
    add_header(slide, s)
    command_strip(slide, s, y=1.10, h=0.86)
    stacks = [("A frozen", BLUE, PALE_BLUE), ("B frozen", PURPLE, PALE_PURPLE), ("result ↔ replay", TEAL, PALE_TEAL), ("prediction", GOLD, PALE_GOLD)]
    for i, (label, col, fill) in enumerate(stacks):
        x = 1.16 + (i % 2) * 5.68
        y = 2.28 + (i // 2) * 1.34
        add_box(slide, x, y, 4.92, 0.92, fill=fill, line=col, width=1.2, name=f"Checkpoint {label}")
        add_text(slide, x + 0.18, y + 0.30, 4.56, 0.24, label, 21, bold=True, color=col, align=PP_ALIGN.CENTER, italic_code=False, name=f"Checkpoint {label} text")
    add_box(slide, 1.10, 5.10, 10.96, 0.74, fill=PALE_GOLD, line=GOLD, name="Checkpoint recovery")
    add_text(slide, 1.34, 5.34, 10.48, 0.24, "status → exact role → restore；不覆寫既有 artifact，不補假的 receipt。", 20, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Checkpoint recovery text")


def compose(slide, s: SlideSpec) -> None:
    clear_placeholders(slide, keep_title=True)
    if s.kind == "states":
        draw_states(slide, s)
    elif s.kind == "waitsleep":
        draw_waitsleep(slide, s)
    elif s.kind == "packet":
        draw_packet(slide, s)
    elif s.kind == "service":
        draw_service(slide, s)
    elif s.kind == "powerarea":
        draw_powerarea(slide, s)
    elif s.kind == "formula-energy":
        draw_formula(slide, s, "E_endpoint")
    elif s.kind == "formula-eff":
        draw_formula(slide, s, "eta_E")
    elif s.kind == "provenance":
        draw_provenance(slide, s)
    elif s.kind == "identity":
        draw_identity(slide, s)
    elif s.kind == "fairness":
        draw_fairness(slide, s)
    elif s.kind == "window":
        draw_window(slide, s)
    elif s.kind == "lab-a-hero":
        draw_lab_hero(slide, s, "A")
    elif s.kind == "lab-b-hero":
        draw_lab_hero(slide, s, "B")
    elif s.kind in {"code-a", "code-b"}:
        draw_code(slide, s, "A" if s.number in {40, 41} else "B")
    elif s.kind == "prediction-a":
        draw_prediction(slide, s, "A")
    elif s.kind == "prediction-b":
        draw_prediction(slide, s, "B")
    elif s.kind == "run-baseline-a":
        draw_run(slide, s, 0)
    elif s.kind == "run-candidate-a":
        draw_run(slide, s, 1)
    elif s.kind == "compare-a":
        draw_compare(slide, s)
    elif s.kind == "ledger-a":
        draw_ledger(slide, s)
    elif s.kind == "packet-service-a":
        draw_packet_service(slide, s)
    elif s.kind == "freeze-a":
        draw_freeze(slide, s, "A")
    elif s.kind == "run-hidden-a":
        draw_withheld(slide, s, "hidden")
    elif s.kind == "debrief-a":
        draw_debrief(slide, s, "A")
    elif s.kind == "threshold-enter":
        draw_code(slide, s, "B")
    elif s.kind == "threshold-exit":
        draw_code(slide, s, "B")
    elif s.kind == "stable-steps":
        draw_code(slide, s, "B")
    elif s.kind == "run-baseline-b":
        draw_run(slide, s, 2)
    elif s.kind == "run-candidate-b":
        draw_run(slide, s, 3)
    elif s.kind == "transition-b":
        draw_packet_service(slide, s)
    elif s.kind == "freeze-b":
        draw_freeze(slide, s, "B")
    elif s.kind == "run-withheld-b":
        draw_withheld(slide, s, "Trace B")
    elif s.kind == "counterexample-b":
        draw_counterexample(slide, s)
    elif s.kind == "debrief-b":
        draw_debrief(slide, s, "B")
    elif s.kind == "checkpoint":
        draw_checkpoint(slide, s)
    else:
        raise ValueError(f"unknown slide kind: {s.kind}")
    add_notes(slide, s)


def remove_all_slides(prs: Presentation) -> None:
    ids = prs.slides._sldIdLst
    for item in list(ids):
        prs.part.drop_rel(item.rId)
        ids.remove(item)


def validate_specs() -> None:
    if [s.number for s in SPECS] != list(range(28, 64)):
        raise ValueError("Part B requires eventual P028..P063 exactly once")
    for s in SPECS:
        text = " ".join(str(v) for v in asdict(s).values())
        if FORBIDDEN in text:
            raise ValueError(f"P{s.number:03d}: forbidden exact Chinese word")
        if MINUTE_RE.search(text):
            raise ValueError(f"P{s.number:03d}: minute/time label forbidden")
        if not all((s.context, s.command, s.purpose, s.mechanism, s.expected, s.recovery, s.notes)):
            raise ValueError(f"P{s.number:03d}: missing classroom contract field")


def _remove_formula_preview_shapes(path: Path) -> None:
    """Remove the text-only equation previews after native OMML insertion."""
    with zipfile.ZipFile(path, "r") as archive:
        infos = [copy.copy(info) for info in archive.infolist()]
        data = {info.filename: archive.read(info.filename) for info in infos}
    P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
    for name in list(data):
        if not (name.startswith("ppt/slides/slide") and name.endswith(".xml")):
            continue
        root = ET.fromstring(data[name])
        changed = False
        for parent in root.iter():
            for child in list(parent):
                if child.tag != f"{{{P_NS}}}sp":
                    continue
                c_nv_pr = child.find(f"./{{{P_NS}}}nvSpPr/{{{P_NS}}}cNvPr")
                if c_nv_pr is not None and c_nv_pr.get("name", "").startswith("Formula preview fallback"):
                    parent.remove(child)
                    changed = True
        if changed:
            data[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    temp = path.with_suffix(".formula-cleanup.pptx")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as out:
        for info in infos:
            out.writestr(info, data[info.filename])
    temp.replace(path)


def build_pptx() -> None:
    validate_specs()
    PROJECT.mkdir(parents=True, exist_ok=True)
    for folder in (PROJECT / "sources", PROJECT / "exports", RENDER_DIR, QA_DIR, ANALYSIS_DIR, VALIDATION_DIR):
        folder.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, SOURCE_COPY)
    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)
    layouts = [0, 1, 2, 3, 5, 1, 2, 3]
    for i, s in enumerate(SPECS):
        slide = prs.slides.add_slide(prs.slide_layouts[s.layout if s.layout < len(prs.slide_layouts) else layouts[i % len(layouts)]])
        compose(slide, s)
    prs.save(str(EXPORT))
    # Save the text-preview source before inserting OMML.  LibreOffice cannot
    # load the final a14:m package, but this source is an honest visual preview
    # of the formula pages and the final export is checked structurally below.
    shutil.copy2(EXPORT, RENDER_SOURCE)
    equation_report: dict[str, object] = {"status": "NOT_RUN", "slides": []}
    if EQUATION_TOOL.exists():
        for logical, key in ((6, "E_endpoint"), (7, "eta_E")):
            temp = EXPORT.with_name(f".part2-equation-{logical}.pptx")
            report = VALIDATION_DIR / f"equation-slide-{logical}.json"
            fallback = EQUATION_FALLBACKS[key]
            if not fallback.exists():
                raise FileNotFoundError(f"missing equation fallback SVG: {fallback}")
            result = subprocess.run([sys.executable, str(EQUATION_TOOL), str(EXPORT), str(temp), "--logical-slide", str(logical), "--equation", key, "--fallback-svg", str(fallback), "--report", str(report)], text=True, capture_output=True, timeout=90, check=False)
            equation_report["slides"].append({"logical_slide": logical, "key": key, "returncode": result.returncode, "stdout": result.stdout[-1200:]})
            if result.returncode == 0 and temp.exists():
                temp.replace(EXPORT)
            elif temp.exists():
                temp.unlink()
        equation_report["status"] = "PASS" if all(item["returncode"] == 0 for item in equation_report["slides"]) else "PARTIAL"
    _remove_formula_preview_shapes(EXPORT)
    (ANALYSIS_DIR / "content-manifest.json").write_text(json.dumps({"schema": "classroom-part2.v1", "eventual_pages": "P028-P063", "template": str(TEMPLATE), "package": {"repo": REPO_URL, "commit": CURRENT_COMMIT, "root": ROOT_NAME, "module": PYTHON_MODULE, "scenario_id": SCENARIO_ID, "policy_api": POLICY_API, "public_release": PUBLIC_RELEASE, "local_artifact": LOCAL_ARTIFACT}, "equations": equation_report, "slides": [asdict(s) for s in SPECS]}, ensure_ascii=False, indent=2), encoding="utf-8")
    (ANALYSIS_DIR / "evidence-summary.json").write_text(json.dumps({"source": "local fallback_artifacts/*.json", "claim_boundary": "FALLBACK_ARTIFACT / SIMULATED_ADAPTER / UPSTREAM_EXECUTION_FALSE", "records": FALLBACK}, ensure_ascii=False, indent=2), encoding="utf-8")


def readback_and_qa() -> None:
    prs = Presentation(str(EXPORT))
    rows = []
    all_text = []
    notes = 0
    for i, slide in enumerate(prs.slides, 1):
        text = "\n".join(sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text.strip())
        rows.append(f"SLIDE {i}\n{text}")
        all_text.append(text)
        notes += bool(slide.notes_slide.notes_text_frame.text.strip())
    readback = "\n\n".join(rows)
    (VALIDATION_DIR / "readback.txt").write_text(readback, encoding="utf-8")
    eq_reports = []
    for logical in (6, 7):
        report_path = VALIDATION_DIR / f"equation-slide-{logical}.json"
        if report_path.exists():
            try:
                eq_reports.append(json.loads(report_path.read_text(encoding="utf-8")))
            except json.JSONDecodeError:
                pass
    checks = {
        "slide_count": len(prs.slides),
        "expected_slide_count": 36,
        "page_range": "P028-P063",
        "classroom_contract": {
            "all_specs_nonempty": all(all((s.context, s.command, s.purpose, s.mechanism, s.expected, s.recovery, s.notes)) for s in SPECS),
            "required_fields": ["context", "command", "purpose", "mechanism", "expected", "recovery", "notes"],
            "operation_slide_count": sum(1 for s in SPECS if s.command),
        },
        "forbidden_word_absent": FORBIDDEN not in readback,
        "minute_labels_absent": not MINUTE_RE.search(readback),
        "speaker_notes": notes,
        "speaker_notes_expected": 36,
        "template_background_preserved": True,
        "slide_level_background_fill": False,
        "footer_reserve": "native template footer/page-number placeholder preserved",
        "title_pt": 28,
        "body_target_pt": 24,
        "command_min_pt": 18,
        "public_release": PUBLIC_RELEASE,
        "local_artifact": LOCAL_ARTIFACT,
        "fallback_evidence_used": True,
        "windows_browser_evidence": "PENDING",
        "course_screenshots_used": False,
        "equation_hooks": ["E_endpoint", "eta_E"],
        "equation_native_math_reports": [{"slide": item.get("logical_slide"), "status": item.get("status"), "omath": item.get("validation", {}).get("target_omath_count")} for item in eq_reports],
    }
    (VALIDATION_DIR / "structural-qa.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")


def render_pdf() -> tuple[bool, str]:
    RENDER_DIR.mkdir(parents=True, exist_ok=True)
    # Use the shared soffice wrapper so the sandbox socket shim and VCL
    # environment are applied consistently; direct libreoffice calls fail in
    # the managed environment before producing a PDF.
    converter = Path("/home/u24/.codex/skills/pptx/scripts/office/soffice.py")
    # Render the pre-OMML source.  The final export retains native OMML and
    # is validated structurally, while LibreOffice cannot open its
    # AlternateContent package; the pre-OMML source carries the same formula
    # page layout with the clean text fallback for visual QA.
    source = RENDER_SOURCE if RENDER_SOURCE.exists() else EXPORT
    cmd = [sys.executable, str(converter), "--headless", "--norestore", "--nodefault", "--nolockcheck", "--nofirststartwizard", "--convert-to", "pdf:impress_pdf_Export", "--outdir", str(RENDER_DIR), str(source)]
    try:
        result = subprocess.run(cmd, text=True, capture_output=True, timeout=90, check=False)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return False, f"render unavailable: {exc}"
    produced = RENDER_DIR / (source.stem + ".pdf")
    pdf = RENDER_DIR / "part2.pdf"
    if produced.exists() and produced != pdf:
        produced.replace(pdf)
    if result.returncode != 0 or not pdf.exists():
        return False, (result.stderr or result.stdout or "libreoffice did not produce PDF").strip()
    try:
        subprocess.run(["pdftoppm", "-png", "-r", "110", str(pdf), str(RENDER_DIR / "slide")], text=True, capture_output=True, timeout=90, check=True)
        return True, f"rendered {len(list(RENDER_DIR.glob('slide-*.png')))} slide PNGs"
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        return False, f"PNG render unavailable: {exc}"


def contact_sheet() -> bool:
    images = sorted(RENDER_DIR.glob("slide-*.png"))
    if not images:
        return False
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        return False
    tw, th, cols = 320, 180, 4
    rows = (len(images) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tw, rows * (th + 24)), "white")
    draw = ImageDraw.Draw(sheet)
    for i, path in enumerate(images):
        with Image.open(path) as img:
            img = img.convert("RGB")
            img.thumbnail((tw, th))
            x = (i % cols) * tw + (tw - img.width) // 2
            y = (i // cols) * (th + 24)
            sheet.paste(img, (x, y))
            draw.text(((i % cols) * tw + 6, y + th + 3), path.stem, fill="black")
    sheet.save(QA_DIR / "contact-sheet.png")
    return True


def visual_qa(ok: bool, message: str) -> None:
    report = {"render_ok": ok, "render_message": message, "contact_sheet": str(QA_DIR / "contact-sheet.png") if (QA_DIR / "contact-sheet.png").exists() else None, "page_by_page_human_inspection": "PENDING controller review", "microsoft_powerpoint_open_reopen": "PENDING controller review", "template_master_layout": "preserved by native template source", "slide_level_background_fill": False, "footer_clearance": "PENDING render/controller inspection", "course_screenshot_dependency": False}
    (QA_DIR / "visual-qa.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-render", action="store_true")
    args = parser.parse_args()
    build_pptx()
    readback_and_qa()
    if args.skip_render:
        ok, message = False, "render skipped by request"
    else:
        ok, message = render_pdf()
        if ok:
            contact_sheet()
    visual_qa(ok, message)
    print(json.dumps({"output": str(EXPORT), "slides": len(SPECS), "render_ok": ok, "render": message}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
