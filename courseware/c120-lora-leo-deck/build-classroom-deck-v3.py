#!/usr/bin/env python3
"""Build the C-120 classroom deck v3 from the native educate PowerPoint.

The v3 deck is intentionally authored as real PowerPoint text and shapes.  It
does not use the thin v2 sampler grammar, does not rewrite the template master
or layouts, and does not place formula artwork on a slide.  Formula pages
reserve named native-equation hooks for the controller's OMML insertion pass.

The source script is the current content donor.  It is read-only; all outputs
are written below ``projects/classroom-v3_ppt169_20260811``.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SCRIPT = ROOT / "full-deck-v2-classroom-script.md"
TEMPLATE = Path("/home/u24/ppt-master/template/educate.pptx")
PROJECT = ROOT / "projects/classroom-v3_ppt169_20260811"
LIVE_EVIDENCE = ROOT / "evidence/browser-live-course-20260811"
FORBIDDEN = "學生"
MINUTE_RE = re.compile(r"(?:分鐘|\b\d+\s*(?:min|mins|minute|minutes)\b)", re.I)
LATIN_RE = re.compile(r"[A-Za-z0-9_./\\:-]")

# educate.pptx theme colours.  The master/background/logo/footer are retained
# from the source PPTX; these colours are used only for editable slide-local
# figures and text.
NAVY = RGBColor(53, 55, 127)
PURPLE = RGBColor(102, 0, 102)
BLUE = RGBColor(111, 137, 247)
GOLD = RGBColor(236, 216, 130)
TEAL = RGBColor(0, 125, 128)
RED = RGBColor(192, 55, 55)
INK = RGBColor(43, 46, 65)
MUTED = RGBColor(90, 94, 111)
WHITE = RGBColor(255, 255, 255)
CREAM = RGBColor(250, 247, 232)
PALE_BLUE = RGBColor(236, 241, 254)
PALE_PURPLE = RGBColor(244, 237, 247)
PALE_GOLD = RGBColor(253, 249, 226)
PALE_TEAL = RGBColor(232, 247, 244)
PALE_RED = RGBColor(253, 239, 239)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
PR_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


@dataclass
class Page:
    number: int
    title: str
    layout: str
    evidence: str
    on_slide: str
    visual: str
    notes: str
    recovery: str


def qn(ns: str, local: str) -> str:
    return f"{{{ns}}}{local}"


def field(block: str, name: str) -> str:
    m = re.search(rf"(?m)^{re.escape(name)}:\s*(.*?)\s*$", block)
    if not m:
        raise ValueError(f"missing {name}")
    return m.group(1).strip()


def parse_pages(path: Path = SCRIPT) -> list[Page]:
    chunks = re.split(r"(?m)^## (P\d{3}) — (.+)$", path.read_text(encoding="utf-8"))
    pages: list[Page] = []
    for i in range(1, len(chunks), 3):
        pid, title, block = chunks[i], chunks[i + 1], chunks[i + 2]
        page = Page(
            int(pid[1:]), title.strip(), field(block, "Layout"), field(block, "Evidence"),
            field(block, "On-slide"), field(block, "Visual"), field(block, "Notes"),
            field(block, "Recovery"),
        )
        serial = " ".join(str(value) for value in asdict(page).values())
        if FORBIDDEN in serial:
            raise ValueError(f"P{page.number:03d}: forbidden word")
        if MINUTE_RE.search(serial):
            raise ValueError(f"P{page.number:03d}: minute label")
        pages.append(page)
    if [p.number for p in pages] != list(range(1, 109)):
        raise ValueError("expected exactly P001..P108")
    return pages


def clean(value: str) -> str:
    return value.replace("`", "").replace("LaTeX source:", "公式插入來源（不在畫面顯示）：")


def chunks(value: str, limit: int = 5) -> list[str]:
    value = clean(value).replace("：", ":")
    pieces = [p.strip() for p in re.split(r"[；;|]", value) if p.strip()]
    if len(pieces) <= limit:
        return pieces
    return pieces[: limit - 1] + ["；".join(pieces[limit - 1 :])]


def shorten(value: str, limit: int = 86) -> str:
    value = clean(value).strip()
    return value if len(value) <= limit else value[: limit - 1] + "…"


def command_lines(page: Page) -> list[str]:
    raw = page.on_slide + "；" + page.layout
    found: list[str] = []
    for item in re.findall(r"`([^`]+)`", raw):
        if item not in found and ("=" in item or re.search(r"(?:bash|course|setup|uv |py |python|sha256|FileHash|wsl|import|run|freeze|reopen|policy|\.venv|result)", item, re.I)):
            found.append(item)
    if not found:
        for item in chunks(page.on_slide, 5):
            if re.search(r"(?:bash|course|setup|uv |py |python|sha256|FileHash|wsl|import|run|freeze|reopen|policy|\.venv|result)", item, re.I):
                found.append(item)
    return found[:4]


def is_operation(page: Page) -> bool:
    text = " ".join((page.title, page.layout, page.on_slide, page.visual))
    return bool(command_lines(page) or re.search(
        r"(?:baseline|candidate|hidden|freeze|import|reopen|marked|policy|setup|verify|checksum|\bA-\d|\bB-\d|\bC-\d|fallback|mismatch|restore|\.venv)", text, re.I
    ))


def operation_kind(page: Page) -> str:
    text = " ".join((page.title, page.layout, page.on_slide)).lower()
    if "mismatch" in text or "fail closed" in text or "recovery" in text or "recover" in text or "restore" in text:
        return "recovery"
    if "import" in text:
        return "import"
    if "reopen" in text or "workbook" in text or "export" in text:
        return "workbook"
    if "hidden" in text or "withheld" in text or "trace b" in text or "surprise" in text:
        return "withheld"
    if "freeze" in text:
        return "freeze"
    if "candidate" in text or "改" in page.title:
        return "candidate"
    if "baseline" in text:
        return "baseline"
    if "setup" in text or "verify" in text or "python" in text or "launcher" in text or "uv" in text or "venv" in text or "checksum" in text or "下載" in page.title:
        return "setup"
    if "policy" in text or "marked" in text or "pace_gap" in text or "quality" in text or "batch" in text or "urgent" in text:
        return "policy"
    return "action"


def action_copy(page: Page) -> str:
    cmds = command_lines(page)
    if cmds:
        return "\n".join(cmds)
    title = clean(page.title)
    if "student_policy" in page.on_slide or "student_policy" in page.visual:
        return "開啟 student_policy.py；只查看目前 lab 的 marked block"
    if "scenario" in title.lower() or "scenario" in page.on_slide.lower():
        return "以唯讀方式開啟固定 scenario JSON"
    return shorten(clean(page.on_slide), 180)


WHY = {
    "setup": "先固定 release、shell 與隔離環境，讓後面的結果都能回到同一份 provenance。",
    "baseline": "先建立 untouched control，才能把後續差異歸因於這一個 policy 變更。",
    "candidate": "只改目前 lab 的 marked 區域，讓 policy bytes 是唯一受控變因。",
    "freeze": "凍結候選版本，保留 scenario、seed、policy SHA 與 predecessor，才能進 withheld case。",
    "withheld": "在不再調參的條件下測試假說，讓結果可以推翻先前的預測。",
    "import": "把機器可讀的結果搬進 Leo，不讓畫面或手抄數字取代驗證。",
    "workbook": "保留同一條 scenario 與 receipt lineage，讓結果能關閉、重開、續接。",
    "recovery": "先停在明確的失敗邊界，保留錯誤與身份，再回到 matching checkpoint。",
    "policy": "讀懂 observation 到 action 的邊界，避免把 future outcome 偷渡進決策。",
    "action": "把操作綁回一個可觀察機制，而不是只追求畫面上的狀態變化。",
}

MECHANISM = {
    "setup": "命令只建立 package-local `.venv`、檢查 lock、policy API 與 scenario identity；它不會產生 energy result。",
    "baseline": "runner 使用固定 scenario／seed，寫出 run identity、policy hash、packet、service、state 與 endpoint-energy artifact。",
    "candidate": "policy bytes 改變 action timing 或 state choice，差異必須沿著 queue、packet、service、state time 追到 J。",
    "freeze": "freeze receipt 把 candidate、predecessor、anchor、policy SHA 和 source mode 綁在同一份 lineage。",
    "withheld": "hidden／Trace B／surprise 使用相同 boundary 的另一條條件，且 policy SHA 必須保持不變。",
    "import": "importer 逐項檢查 schema、scenario、seed、units、policy SHA、事件順序與 provenance；錯誤時不改 session。",
    "workbook": "workbook 序列化 prediction、選擇、run、recovery、receipt 與 claim boundary，而不是只存最後一個數字。",
    "recovery": "fail-closed 保留原 session，要求修正被點名的 artifact 或切到同 scenario fallback，不自行補值。",
    "policy": "只使用 action 當下可見的 observation；policy 選擇會改變時間、封包、服務與 energy ledger。",
    "action": "先說明輸入，再看 state／packet／service，最後才解讀 W、J 與 bit/J。",
}


def expected_copy(page: Page, kind: str) -> str:
    evidence = shorten(page.evidence, 116)
    if kind == "setup":
        out = "可見輸出：`READY` receipt、Python 版本、scenario_id、lock／policy identity；目前未排演欄位保留 `PLACEHOLDER`。"
    elif kind == "baseline":
        out = "可見輸出：baseline run_id、policy SHA、result_path，以及 queue／packet／service／state／J 欄位。"
    elif kind == "candidate":
        out = "可見輸出：candidate result、policy SHA、與 baseline 的 consequential diff；若只變 label，gate 不成立。"
    elif kind == "freeze":
        out = "可見輸出：freeze receipt、predecessor、active_block_id、policy SHA；缺任一欄就不能進下一個 case。"
    elif kind == "withheld":
        out = "可見輸出：另一條條件的 replay；policy SHA unchanged，並同時列出 service、state、endpoint J。"
    elif kind == "import":
        out = "可見輸出：all-or-nothing import receipt、scenario／units／seed／policy identity 與 endpoint replay identity。"
    elif kind == "workbook":
        out = "可見輸出：`COMPLETE` 或明確 `INCOMPLETE`，重新開啟仍是同一 scenario 與 claim boundary。"
    elif kind == "recovery":
        out = "可見輸出：拒絕原因、保留原 session、matching fallback 或修正後的新 receipt；不產生部分結果。"
    elif kind == "policy":
        out = "可見輸出：marked block 的小 diff、合法 action、以及後續可觀察的 state／packet／service／J 差異。"
    else:
        out = "可見輸出：一個可追溯的 action → mechanism → evidence 連結。"
    return f"{out}\nEvidence：{evidence}"


def recovery_copy(page: Page, kind: str) -> str:
    base = clean(page.recovery).strip()
    if not base:
        base = "保留原始錯誤，回到同一 scenario 的 known-good checkpoint。"
    if kind == "recovery" or "fallback" in base.lower():
        return shorten("失敗處置：" + base, 200)
    return shorten("若失敗：" + base + "；不要補填 result 或改寫 receipt。", 200)


def _set_font(run, size: int, *, bold: bool = False, italic: bool = False, color: RGBColor = INK) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rpr.set("sz", str(size * 100))
    rpr.set("b", "1" if bold else "0")
    rpr.set("i", "1" if italic else "0")
    for tag, face in (("latin", "Times New Roman"), ("ea", "標楷體"), ("cs", "Times New Roman")):
        node = rpr.find(qn(A_NS, tag))
        if node is None:
            node = OxmlElement(f"a:{tag}")
            rpr.append(node)
        node.set("typeface", face)


def rich_text(tf, text: str, size: int = 24, *, bold: bool = False, color: RGBColor = INK, italic_code: bool = True, align=PP_ALIGN.LEFT) -> None:
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = str(text).splitlines() or [""]
    for line_i, line in enumerate(lines):
        p = tf.paragraphs[0] if line_i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        for token in re.split(r"(`[^`]*`)", line):
            if not token:
                continue
            marked = token.startswith("`") and token.endswith("`")
            raw = token[1:-1] if marked else token
            for part in re.findall(r"[\u3400-\u9fff\uF900-\uFAFF]+|[^\u3400-\u9fff\uF900-\uFAFF]+", raw):
                italic = bool(marked and italic_code and re.search(r"[A-Za-z_ηκΣ]|\d", part))
                run = p.add_run()
                run.text = part
                _set_font(run, size, bold=bold, italic=italic, color=color)


def add_text(slide, x: float, y: float, w: float, h: float, text: str, size: int = 24, *, bold: bool = False, color: RGBColor = INK, italic_code: bool = True, align=PP_ALIGN.LEFT, name: str = ""):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or "Editable text"
    shape.fill.background()
    shape.line.fill.background()
    rich_text(shape.text_frame, text, size, bold=bold, color=color, italic_code=italic_code, align=align)
    return shape


def add_box(slide, x: float, y: float, w: float, h: float, *, fill: RGBColor = WHITE, line: RGBColor = BLUE, width: float = 1.0, name: str = "Editable figure"):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *, color: RGBColor = NAVY, width: float = 1.4, arrow: bool = False, name: str = "Editable connector"):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.name = name
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def add_dot(slide, x: float, y: float, d: float = 0.18, color: RGBColor = TEAL, name: str = "State marker"):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def add_tag(slide, x: float, y: float, w: float, text: str, *, fill: RGBColor = PALE_BLUE, color: RGBColor = NAVY, name: str = "Status tag"):
    add_box(slide, x, y, w, 0.34, fill=fill, line=fill, width=0.5, name=name)
    add_text(slide, x + 0.05, y + 0.02, w - 0.1, 0.27, text, 16, bold=True, color=color, align=PP_ALIGN.CENTER, italic_code=False, name=name + " text")


def clear_placeholders(slide, *, keep_title: bool = True) -> None:
    for shape in list(slide.shapes):
        if not shape.is_placeholder:
            continue
        typ = str(shape.placeholder_format.type)
        idx = shape.placeholder_format.idx
        if idx == 10 or (keep_title and "TITLE" in typ.upper()):
            continue
        shape._element.getparent().remove(shape._element)


def set_title(slide, title: str, *, cover: bool = False):
    title_shape = None
    for shape in slide.shapes:
        if shape.is_placeholder and "TITLE" in str(shape.placeholder_format.type).upper():
            title_shape = shape
            break
    if title_shape is None or cover:
        if title_shape is not None:
            title_shape._element.getparent().remove(title_shape._element)
        title_shape = add_text(slide, 0.68 if not cover else 0.86, 0.22 if not cover else 1.10, 11.95 if not cover else 7.4, 0.66 if not cover else 1.35, title, 28, bold=True, color=NAVY, name="Native title anchor")
    else:
        rich_text(title_shape.text_frame, title, 28, bold=True, color=NAVY, italic_code=False)
        title_shape.name = "Native title anchor"
    return title_shape


def add_notes(slide, page: Page) -> None:
    notes = clean(page.notes).strip()
    recovery = clean(page.recovery).strip()
    text = notes + (" " if notes and recovery else "") + ("若操作受阻，" + recovery if recovery else "")
    if FORBIDDEN in text or MINUTE_RE.search(text):
        raise ValueError(f"P{page.number:03d}: invalid notes")
    slide.notes_slide.notes_text_frame.text = text


def status(page: Page) -> str:
    e = clean(page.evidence)
    if "PLACEHOLDER" in e:
        return "EVIDENCE PLACEHOLDER"
    if "NOT VERIFIED" in e:
        return "IMPLEMENTED / NOT VERIFIED"
    return "SOURCE-BOUND EVIDENCE"


def add_header(slide, page: Page):
    set_title(slide, page.title)
    add_tag(slide, 8.92, 0.89, 3.36, status(page), fill=PALE_GOLD if "PLACEHOLDER" in page.evidence else PALE_BLUE, color=PURPLE if "PLACEHOLDER" in page.evidence else NAVY)


def draw_cover(slide, page: Page):
    clear_placeholders(slide, keep_title=False)
    set_title(slide, page.title, cover=True)
    add_text(slide, 0.95, 2.60, 6.7, 0.58, "C-120｜LoRaEnergySim × 智慧節能與物聯網應用", 24, color=PURPLE, bold=True, name="Cover subtitle")
    add_text(slide, 0.98, 3.36, 6.35, 0.88, "policy → state time → packet / service → endpoint energy", 26, color=INK, name="Cover causal chain")
    labels = ["觀察", "決策", "服務", "J"]
    for i, label in enumerate(labels):
        x = 0.98 + i * 1.64
        add_box(slide, x, 4.62, 1.32, 0.68, fill=PALE_BLUE if i % 2 == 0 else PALE_GOLD, line=BLUE if i % 2 == 0 else GOLD, name=f"Cover stage {i + 1}")
        add_text(slide, x + 0.06, 4.79, 1.20, 0.28, label, 22, bold=True, align=PP_ALIGN.CENTER, name=f"Cover stage label {i + 1}")
        if i < len(labels) - 1:
            add_line(slide, x + 1.32, 4.96, x + 1.62, 4.96, color=GOLD, width=2.0, arrow=True, name="Cover causal arrow")
    add_text(slide, 0.98, 5.62, 7.2, 0.46, "SIMULATED TEACHING DATA｜NOT LIVE｜NOT MEASURED｜NOT CANONICAL-PARITY-VERIFIED", 16, color=RED, bold=True, italic_code=False, name="Claim boundary")


def draw_causal(slide, page: Page, variant: int):
    add_header(slide, page)
    items = chunks(page.on_slide, 6)
    if variant % 3 == 0:
        for i, item in enumerate(items[:4]):
            x = 0.82 + i * 3.08
            add_box(slide, x, 1.75, 2.62, 1.18, fill=[PALE_BLUE, PALE_TEAL, PALE_GOLD, PALE_PURPLE][i], line=[BLUE, TEAL, GOLD, PURPLE][i], name=f"Causal stage {i + 1}")
            add_text(slide, x + 0.12, 1.98, 2.38, 0.70, item, 22, bold=True, align=PP_ALIGN.CENTER, name=f"Causal stage text {i + 1}")
            if i < min(3, len(items) - 1):
                add_line(slide, x + 2.62, 2.34, x + 3.00, 2.34, color=NAVY, width=1.7, arrow=True)
        add_box(slide, 0.90, 3.42, 11.45, 1.48, fill=CREAM, line=GOLD, name="Causal explanation band")
        add_text(slide, 1.14, 3.66, 10.95, 0.96, f"看圖判讀：{shorten(clean(page.visual), 170)}\n機制邊界：{shorten(clean(page.evidence), 130)}", 22, color=INK, name="Causal explanation")
        add_text(slide, 1.02, 5.24, 11.1, 0.72, "講師判讀句：先確認 action 是否真的改變中間 state，再把差異連到 service 與 endpoint J。", 24, color=NAVY, bold=True, align=PP_ALIGN.CENTER, name="Causal takeaway")
    elif variant % 3 == 1:
        add_box(slide, 0.82, 1.55, 4.12, 4.70, fill=PALE_BLUE, line=BLUE, name="Concept left rail")
        add_text(slide, 1.08, 1.83, 3.58, 0.36, "先問：這張圖的輸入是什麼？", 24, bold=True, color=NAVY, name="Concept prompt")
        add_text(slide, 1.12, 2.42, 3.50, 2.65, "\n".join(items[:4]) or clean(page.on_slide), 24, name="Concept inputs")
        add_box(slide, 5.30, 1.55, 7.05, 2.05, fill=PALE_TEAL, line=TEAL, name="Mechanism right rail")
        add_text(slide, 5.60, 1.83, 6.45, 0.36, "中間機制：", 24, bold=True, color=TEAL, name="Mechanism heading")
        add_text(slide, 5.63, 2.30, 6.35, 1.04, shorten(clean(page.visual), 200), 24, name="Mechanism description")
        add_box(slide, 5.30, 3.98, 7.05, 2.27, fill=PALE_GOLD, line=GOLD, name="Interpretation rail")
        add_text(slide, 5.60, 4.25, 6.45, 0.36, "最後讀：", 24, bold=True, color=PURPLE, name="Interpretation heading")
        add_text(slide, 5.63, 4.72, 6.35, 1.12, "不要把畫面效果當成結果。回到同一 scenario、units、service boundary 與 evidence。", 24, name="Interpretation text")
    else:
        add_text(slide, 0.92, 1.38, 11.5, 0.55, "這一頁把名詞排成一條可追溯的路徑。", 24, color=MUTED, align=PP_ALIGN.CENTER, name="Route caption")
        add_line(slide, 1.25, 3.15, 12.0, 3.15, color=NAVY, width=2.2)
        for i, item in enumerate(items[:5]):
            x = 1.10 + i * 2.55
            add_dot(slide, x, 2.90, 0.50, [BLUE, TEAL, GOLD, PURPLE, RED][i % 5], name=f"Route dot {i + 1}")
            add_text(slide, x - 0.38, 2.05 if i % 2 == 0 else 3.57, 1.30, 0.84, item, 22, align=PP_ALIGN.CENTER, name=f"Route label {i + 1}")
        add_box(slide, 1.10, 5.12, 11.1, 0.88, fill=CREAM, line=GOLD, name="Route reading note")
        add_text(slide, 1.35, 5.30, 10.6, 0.44, shorten(clean(page.visual), 176), 22, align=PP_ALIGN.CENTER, name="Route visual note")


def draw_operation(slide, page: Page, style: int):
    add_header(slide, page)
    kind = operation_kind(page)
    action = action_copy(page)
    why = WHY[kind]
    mechanism = MECHANISM[kind]
    expected = expected_copy(page, kind)
    recovery = recovery_copy(page, kind)
    if style % 4 == 0:
        add_box(slide, 0.80, 1.42, 5.95, 2.05, fill=PALE_PURPLE, line=PURPLE, name="Operation action panel")
        add_text(slide, 1.05, 1.64, 5.45, 0.30, "先做什麼｜exact action", 22, bold=True, color=PURPLE, name="Action heading")
        add_text(slide, 1.06, 2.07, 5.35, 1.15, action, 18, color=INK, italic_code=False, name="Exact action")
        add_box(slide, 7.05, 1.42, 5.45, 2.05, fill=PALE_BLUE, line=BLUE, name="Operation why panel")
        add_text(slide, 7.32, 1.64, 4.92, 0.30, "為什麼／怎麼改變機制", 22, bold=True, color=NAVY, name="Why heading")
        add_text(slide, 7.33, 2.07, 4.88, 1.13, why + "\n" + mechanism, 22, name="Why mechanism")
        add_box(slide, 0.80, 3.80, 5.95, 2.10, fill=PALE_TEAL, line=TEAL, name="Operation expected panel")
        add_text(slide, 1.05, 4.03, 5.45, 0.30, "會看到什麼｜receipt / output", 22, bold=True, color=TEAL, name="Expected heading")
        add_text(slide, 1.06, 4.44, 5.35, 1.16, expected, 20, name="Expected output")
        add_box(slide, 7.05, 3.80, 5.45, 2.10, fill=PALE_GOLD, line=GOLD, name="Operation recovery panel")
        add_text(slide, 7.32, 4.03, 4.92, 0.30, "怎麼判讀／失敗怎麼回復", 22, bold=True, color=PURPLE, name="Recovery heading")
        add_text(slide, 7.33, 4.44, 4.88, 1.16, recovery, 20, name="Recovery text")
    elif style % 4 == 1:
        add_text(slide, 0.92, 1.30, 11.4, 0.48, "先把命令放在完整 shell context，再讀它造成的下一個可觀察欄位。", 22, color=MUTED, align=PP_ALIGN.CENTER, name="Operational instruction")
        add_box(slide, 0.90, 1.92, 11.45, 1.72, fill=CREAM, line=NAVY, width=1.5, name="Command ribbon")
        add_text(slide, 1.20, 2.16, 10.85, 1.20, action, 18, color=INK, italic_code=False, name="Command text")
        cards = [(0.90, PALE_BLUE, BLUE, "目的／機制", why + "\n" + mechanism), (4.85, PALE_TEAL, TEAL, "預期輸出", expected), (8.80, PALE_GOLD, GOLD, "判讀／回復", recovery)]
        for x, fill, line, heading, body in cards:
            add_box(slide, x, 4.00, 3.55, 2.05, fill=fill, line=line, name=heading)
            add_text(slide, x + 0.16, 4.22, 3.23, 0.32, heading, 22, bold=True, color=line, align=PP_ALIGN.CENTER, name=heading + " heading")
            add_text(slide, x + 0.18, 4.66, 3.18, 1.10, body, 20, align=PP_ALIGN.CENTER, name=heading + " body")
    elif style % 4 == 2:
        add_box(slide, 0.86, 1.42, 6.20, 4.60, fill=PALE_PURPLE, line=PURPLE, name="Code focus")
        add_text(slide, 1.12, 1.66, 5.65, 0.32, "可複製的動作", 22, bold=True, color=PURPLE, name="Code heading")
        add_text(slide, 1.14, 2.10, 5.62, 1.55, action, 18, color=INK, italic_code=False, name="Code lines")
        add_line(slide, 1.16, 3.90, 6.70, 3.90, color=GOLD, width=1.4, arrow=True, name="Code to mechanism")
        add_text(slide, 1.18, 4.18, 5.55, 1.12, "機制讀法：" + mechanism, 22, name="Code mechanism")
        add_box(slide, 7.38, 1.42, 5.10, 2.08, fill=PALE_BLUE, line=BLUE, name="Expected evidence")
        add_text(slide, 7.65, 1.66, 4.56, 0.32, "輸出不只是一行 READY", 22, bold=True, color=NAVY, name="Evidence heading")
        add_text(slide, 7.66, 2.10, 4.52, 1.12, expected, 20, name="Evidence body")
        add_box(slide, 7.38, 3.82, 5.10, 2.20, fill=PALE_GOLD, line=GOLD, name="Interpretation block")
        add_text(slide, 7.65, 4.07, 4.56, 0.32, "判讀與恢復", 22, bold=True, color=PURPLE, name="Interpret heading")
        add_text(slide, 7.66, 4.52, 4.52, 1.15, why + "\n" + recovery, 20, name="Interpret body")
    else:
        # A non-card layout for path / receipt / import pages.
        add_text(slide, 0.92, 1.30, 11.5, 0.46, shorten(clean(page.visual), 180), 22, color=MUTED, align=PP_ALIGN.CENTER, name="Operation visual caption")
        labels = [("ACTION", action), ("MECHANISM", mechanism), ("OUTPUT", expected), ("RECOVER", recovery)]
        xs = [0.95, 3.98, 7.01, 10.04]
        fills = [PALE_PURPLE, PALE_BLUE, PALE_TEAL, PALE_GOLD]
        lines = [PURPLE, BLUE, TEAL, GOLD]
        for i, ((label, body), x) in enumerate(zip(labels, xs)):
            add_box(slide, x, 2.12, 2.55, 3.55, fill=fills[i], line=lines[i], name=f"Receipt stage {i + 1}")
            add_text(slide, x + 0.12, 2.37, 2.31, 0.32, label, 18, bold=True, color=lines[i], align=PP_ALIGN.CENTER, italic_code=False, name=f"Receipt label {i + 1}")
            add_text(slide, x + 0.16, 2.88, 2.23, 2.25, body, 20, align=PP_ALIGN.CENTER, name=f"Receipt body {i + 1}")
            if i < 3:
                add_line(slide, x + 2.55, 3.90, x + 2.92, 3.90, color=NAVY, width=1.5, arrow=True, name="Receipt arrow")


def draw_code_or_policy(slide, page: Page, variant: int):
    add_header(slide, page)
    items = chunks(page.on_slide, 5)
    code = action_copy(page)
    add_box(slide, 0.85, 1.48, 6.05, 4.60, fill=PALE_PURPLE, line=PURPLE, name="Policy source frame")
    add_text(slide, 1.10, 1.74, 5.55, 0.30, "目前可編輯的 bounded surface", 22, bold=True, color=PURPLE, name="Policy frame heading")
    add_text(slide, 1.12, 2.18, 5.48, 1.45, code, 18, color=INK, italic_code=False, name="Policy code")
    add_text(slide, 1.12, 3.95, 5.48, 1.20, "只讀 observation → 選合法 action；不要讀 future delivered、future J 或 generated result。", 22, name="Policy rule")
    add_box(slide, 7.22, 1.48, 5.20, 1.30, fill=PALE_BLUE, line=BLUE, name="Policy input")
    add_text(slide, 7.50, 1.73, 4.64, 0.30, "輸入（決策當下可見）", 22, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Policy input heading")
    add_text(slide, 7.52, 2.18, 4.60, 0.34, "quality / queue / remaining window / age", 20, align=PP_ALIGN.CENTER, name="Policy input body")
    add_box(slide, 7.22, 3.10, 5.20, 1.30, fill=PALE_TEAL, line=TEAL, name="Policy mechanism")
    add_text(slide, 7.50, 3.35, 4.64, 0.30, "機制（action 會改變什麼）", 22, bold=True, color=TEAL, align=PP_ALIGN.CENTER, name="Policy mechanism heading")
    add_text(slide, 7.52, 3.80, 4.60, 0.34, shorten(clean(page.visual), 88), 20, align=PP_ALIGN.CENTER, name="Policy mechanism body")
    add_box(slide, 7.22, 4.72, 5.20, 1.36, fill=PALE_GOLD, line=GOLD, name="Policy test")
    add_text(slide, 7.50, 4.97, 4.64, 0.30, "驗證（結果要能反駁）", 22, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Policy test heading")
    add_text(slide, 7.52, 5.40, 4.60, 0.44, "queue → packet → service → state time → J", 20, align=PP_ALIGN.CENTER, name="Policy test body")


def draw_state(slide, page: Page, variant: int):
    add_header(slide, page)
    states = ["SLEEP", "WAKE", "PROCESS", "TX / RX", "WAIT"]
    y = 2.10 if variant % 2 == 0 else 2.65
    x0 = 0.92
    for i, state in enumerate(states):
        x = x0 + i * 2.46
        add_box(slide, x, y, 1.92, 0.88, fill=[PALE_BLUE, PALE_GOLD, PALE_TEAL, PALE_PURPLE, CREAM][i], line=[BLUE, GOLD, TEAL, PURPLE, NAVY][i], name=f"State {state}")
        add_text(slide, x + 0.08, y + 0.26, 1.76, 0.34, state, 20, bold=True, align=PP_ALIGN.CENTER, italic_code=False, name=f"State {state} label")
        if i < len(states) - 1:
            add_line(slide, x + 1.92, y + 0.44, x + 2.32, y + 0.44, color=NAVY, width=1.5, arrow=True)
    add_box(slide, 0.98, 3.55, 5.55, 2.20, fill=PALE_BLUE, line=BLUE, name="State interpretation")
    add_text(slide, 1.24, 3.80, 5.04, 0.32, "不要只讀總 J", 24, bold=True, color=NAVY, name="State interpretation heading")
    add_text(slide, 1.25, 4.28, 5.02, 1.05, "沿著 state duration、wake latency、TX/RX attempts，回到 policy action 的時間機制。", 24, name="State interpretation body")
    add_box(slide, 6.85, 3.55, 5.50, 2.20, fill=PALE_GOLD, line=GOLD, name="State boundary")
    add_text(slide, 7.12, 3.80, 4.96, 0.32, "service 先於 energy", 24, bold=True, color=PURPLE, name="State boundary heading")
    add_text(slide, 7.13, 4.28, 4.92, 1.05, shorten(clean(page.visual) + "；" + clean(page.evidence), 190), 22, name="State boundary body")


def draw_formula_hook(slide, page: Page, equation: str):
    add_header(slide, page)
    add_text(slide, 0.98, 1.28, 11.2, 0.48, "公式先教範圍與單位；數值必須從 artifact 讀取。", 24, color=MUTED, align=PP_ALIGN.CENTER, name="Formula instruction")
    add_box(slide, 1.05, 1.95, 11.10, 2.56, fill=WHITE, line=NAVY, width=1.8, name=f"NATIVE_EQUATION_HOOK_{equation}")
    add_text(slide, 1.34, 2.22, 10.55, 0.34, f"NATIVE EQUATION INSERTION HOOK｜{equation}", 22, bold=True, color=NAVY, align=PP_ALIGN.CENTER, italic_code=False, name=f"NATIVE_EQUATION_HOOK_LABEL_{equation}")
    add_line(slide, 1.60, 3.45, 11.60, 3.45, color=GOLD, width=2.0, name="Equation insertion baseline")
    add_text(slide, 1.42, 3.70, 10.35, 0.36, "由 controller 以 PowerPoint 原生方程式插入；此處不放圖片、不顯示 LaTeX 原始碼。", 20, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Equation hook instruction")
    if equation == "E_endpoint":
        add_box(slide, 1.05, 4.88, 5.25, 1.12, fill=PALE_BLUE, line=BLUE, name="Equation numerator scope")
        add_text(slide, 1.28, 5.16, 4.78, 0.52, "scope：endpoint radio + processing\nunits：W × time → J", 22, align=PP_ALIGN.CENTER, name="Equation scope text")
        add_box(slide, 6.90, 4.88, 5.25, 1.12, fill=PALE_GOLD, line=GOLD, name="Equation denominator scope")
        add_text(slide, 7.12, 5.16, 4.82, 0.52, "判讀：每個 state 的功率乘停留時間\n不可擴大成 whole-system energy", 22, align=PP_ALIGN.CENTER, name="Equation boundary text")
    else:
        add_box(slide, 1.05, 4.88, 5.25, 1.12, fill=PALE_TEAL, line=TEAL, name="Equation data scope")
        add_text(slide, 1.28, 5.16, 4.78, 0.52, "分子：delivered data\n單位：bits", 22, align=PP_ALIGN.CENTER, name="Equation numerator text")
        add_box(slide, 6.90, 4.88, 5.25, 1.12, fill=PALE_PURPLE, line=PURPLE, name="Equation energy scope")
        add_text(slide, 7.12, 5.16, 4.82, 0.52, "分母：同一 boundary 的 endpoint J\n判讀：bit/J，不替代 service gate", 22, align=PP_ALIGN.CENTER, name="Equation denominator text")


def draw_trace(slide, page: Page, variant: int):
    add_header(slide, page)
    add_text(slide, 0.95, 1.28, 11.45, 0.46, shorten(clean(page.visual), 170), 22, color=MUTED, align=PP_ALIGN.CENTER, name="Trace caption")
    add_line(slide, 1.10, 3.18, 12.05, 3.18, color=NAVY, width=2.0, name="Trace axis")
    points = [1.32, 3.12, 4.92, 6.72, 8.52, 10.32, 11.82]
    heights = [0.45, 1.10, 0.72, 1.38, 0.58, 1.04, 0.82]
    for i, (x, height) in enumerate(zip(points, heights)):
        add_line(slide, x, 3.18, x, 3.18 - height, color=BLUE if i % 2 else TEAL, width=5.0, name=f"Trace sample {i + 1}")
        add_dot(slide, x - 0.09, 3.09 - height, 0.18, GOLD if i in (2, 4) else NAVY, name=f"Trace point {i + 1}")
    add_line(slide, 1.10, 2.05, 12.00, 2.05, color=GOLD, width=1.5, name="Trace threshold")
    add_text(slide, 10.20, 1.70, 1.82, 0.30, "ENTER / EXIT", 18, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Trace threshold label")
    add_box(slide, 0.98, 4.18, 5.50, 1.64, fill=PALE_BLUE, line=BLUE, name="Trace decision")
    add_text(slide, 1.22, 4.46, 5.02, 0.32, "policy decision", 22, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Trace decision heading")
    add_text(slide, 1.25, 4.88, 4.98, 0.62, "enter / hold / exit；只看當下 quality、age 與剩餘窗口。", 22, align=PP_ALIGN.CENTER, name="Trace decision body")
    add_box(slide, 6.84, 4.18, 5.50, 1.64, fill=PALE_GOLD, line=GOLD, name="Trace verdict")
    add_text(slide, 7.08, 4.46, 5.02, 0.32, "判讀", 22, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Trace verdict heading")
    add_text(slide, 7.10, 4.88, 4.96, 0.62, "切換次數不是唯一答案；連同 retry、delivery、service 與 J 一起解讀。", 22, align=PP_ALIGN.CENTER, name="Trace verdict body")


def draw_compare(slide, page: Page, variant: int):
    add_header(slide, page)
    add_text(slide, 0.90, 1.28, 11.5, 0.44, shorten(clean(page.visual), 160), 22, color=MUTED, align=PP_ALIGN.CENTER, name="Compare caption")
    left_title = "BASELINE / CONTROL"
    right_title = "CANDIDATE / WITHHELD"
    add_box(slide, 0.90, 1.92, 5.48, 3.78, fill=PALE_BLUE, line=BLUE, name="Baseline panel")
    add_box(slide, 6.92, 1.92, 5.48, 3.78, fill=PALE_GOLD, line=GOLD, name="Candidate panel")
    add_text(slide, 1.18, 2.20, 4.92, 0.34, left_title, 20, bold=True, color=NAVY, align=PP_ALIGN.CENTER, italic_code=False, name="Baseline heading")
    add_text(slide, 7.20, 2.20, 4.92, 0.34, right_title, 20, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, italic_code=False, name="Candidate heading")
    left = ["fixed scenario / seed", "policy SHA = control", "queue / service / state / J"]
    right = ["one marked block changed", "policy SHA + receipt", "consequential diff or reject"]
    for i, (a, b) in enumerate(zip(left, right)):
        y = 2.85 + i * 0.76
        add_dot(slide, 1.40, y + 0.06, 0.18, TEAL, name=f"Baseline bullet {i + 1}")
        add_text(slide, 1.72, y, 4.20, 0.40, a, 22, name=f"Baseline bullet text {i + 1}")
        add_dot(slide, 7.42, y + 0.06, 0.18, GOLD, name=f"Candidate bullet {i + 1}")
        add_text(slide, 7.74, y, 4.20, 0.40, b, 22, name=f"Candidate bullet text {i + 1}")
    add_line(slide, 6.52, 2.26, 6.78, 5.42, color=NAVY, width=1.5, name="Compare boundary")
    add_text(slide, 1.02, 5.98, 11.2, 0.34, "判讀句：若 workload、window、boundary 或 units 不同，結果標為 INCOMPARABLE。", 22, bold=True, color=RED, align=PP_ALIGN.CENTER, name="Fairness verdict")


def draw_provenance(slide, page: Page, variant: int):
    add_header(slide, page)
    layers = [("SOURCE", "release / scenario / source hash", PALE_PURPLE, PURPLE), ("MODEL", "service window / trace / state", PALE_BLUE, BLUE), ("POLICY", "marked block / action / SHA", PALE_TEAL, TEAL), ("RESULT", "packet / service / endpoint J", PALE_GOLD, GOLD)]
    for i, (label, body, fill, line) in enumerate(layers):
        y = 1.52 + i * 1.06
        add_box(slide, 1.10, y, 4.08, 0.74, fill=fill, line=line, name=f"Provenance layer {label}")
        add_text(slide, 1.24, y + 0.18, 1.16, 0.28, label, 18, bold=True, color=line, align=PP_ALIGN.CENTER, italic_code=False, name=f"Provenance label {label}")
        add_text(slide, 2.56, y + 0.16, 2.40, 0.34, body, 20, name=f"Provenance body {label}")
        if i < 3:
            add_line(slide, 5.18, y + 0.37, 5.95, y + 1.43, color=NAVY, width=1.5, arrow=True, name="Provenance arrow")
    add_box(slide, 6.55, 1.52, 5.65, 4.22, fill=CREAM, line=GOLD, name="Identity ledger")
    add_text(slide, 6.86, 1.80, 5.03, 0.34, "同一個 scenario_id 要一路相同", 22, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Identity heading")
    identity = "scenario_id\nseed / anchor\npolicy SHA\nunits / source_mode\nresult → replay → workbook"
    add_text(slide, 7.00, 2.34, 4.74, 2.34, identity, 24, align=PP_ALIGN.CENTER, name="Identity values")
    add_text(slide, 6.96, 5.02, 4.84, 0.42, "不一致：fail closed；不可拼接相似畫面。", 22, bold=True, color=RED, align=PP_ALIGN.CENTER, name="Identity recovery")


def draw_evidence(slide, page: Page, image_path: Path | None, variant: int):
    add_header(slide, page)
    add_box(slide, 0.88, 1.42, 5.45, 4.68, fill=PALE_BLUE, line=BLUE, name="Evidence reading panel")
    add_text(slide, 1.14, 1.72, 4.94, 0.34, "畫面要回答哪個問題？", 24, bold=True, color=NAVY, name="Evidence question")
    add_text(slide, 1.18, 2.20, 4.84, 2.22, "1. 這是不是同一個 scenario？\n2. 哪個 action 改變了 state？\n3. packet / service 有沒有後果？\n4. endpoint J 的 boundary 是什麼？", 24, name="Evidence questions")
    add_text(slide, 1.16, 4.92, 4.88, 0.78, "單一漂亮 KPI 不足以支持節能判決。", 24, bold=True, color=RED, align=PP_ALIGN.CENTER, name="Evidence warning")
    add_box(slide, 6.70, 1.42, 5.62, 4.68, fill=WHITE, line=TEAL, name="Evidence visual panel")
    if image_path and image_path.exists():
        pic = slide.shapes.add_picture(str(image_path), Inches(6.98), Inches(2.02), width=Inches(5.05), height=Inches(3.18))
        pic.name = "Current browser-server-preview evidence"
        add_text(slide, 7.02, 1.68, 4.96, 0.28, "SERVER PREVIEW / FIXTURE HOST ONLY", 16, bold=True, color=GOLD, align=PP_ALIGN.CENTER, italic_code=False, name="Evidence scope")
    else:
        add_text(slide, 7.10, 2.20, 4.82, 1.12, "EVIDENCE PLACEHOLDER\n尚未凍結的畫面、行號或 KPI 不自行補填。", 24, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Evidence placeholder")
    add_text(slide, 7.02, 5.42, 4.98, 0.36, "import / replay / workbook identity 必須一致。", 20, color=INK, align=PP_ALIGN.CENTER, name="Evidence identity note")


def draw_transfer(slide, page: Page, variant: int):
    add_header(slide, page)
    add_text(slide, 0.95, 1.30, 11.45, 0.48, "把同一個決策句型移到另一個 IoT domain，不搬運衛星專有名詞。", 22, color=MUTED, align=PP_ALIGN.CENTER, name="Transfer caption")
    fields = [("固定任務", "payload / deadline / window"), ("控制", "一個可修改的 policy"), ("服務", "delivery / freshness / deadline"), ("能量", "Σ P·Δt → J"), ("反例", "held-out falsifier")]
    for i, (label, body) in enumerate(fields):
        x = 0.88 + (i % 3) * 4.12
        y = 2.05 + (i // 3) * 1.78
        add_box(slide, x, y, 3.55, 1.32, fill=[PALE_BLUE, PALE_TEAL, PALE_GOLD, PALE_PURPLE, CREAM][i], line=[BLUE, TEAL, GOLD, PURPLE, NAVY][i], name=f"Transfer field {i + 1}")
        add_text(slide, x + 0.14, y + 0.22, 3.27, 0.30, label, 22, bold=True, color=[NAVY, TEAL, PURPLE, PURPLE, NAVY][i], align=PP_ALIGN.CENTER, name=f"Transfer label {i + 1}")
        add_text(slide, x + 0.17, y + 0.67, 3.21, 0.36, body, 20, align=PP_ALIGN.CENTER, name=f"Transfer body {i + 1}")
    add_box(slide, 0.92, 5.70, 11.45, 0.42, fill=PALE_RED, line=RED, width=0.8, name="Transfer warning")
    add_text(slide, 1.08, 5.78, 11.1, 0.24, "判讀：若只有圖表／動畫變化，卻沒有 action → state → service → J → decision，不能算可移轉證據。", 18, bold=True, color=RED, align=PP_ALIGN.CENTER, name="Transfer warning text")


def draw_appendix(slide, page: Page, variant: int):
    add_header(slide, page)
    add_box(slide, 0.90, 1.45, 3.28, 4.70, fill=PALE_PURPLE, line=PURPLE, name="Appendix index")
    add_text(slide, 1.16, 1.76, 2.76, 0.34, "可選深挖", 24, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Appendix heading")
    items = chunks(page.on_slide, 5) or ["source", "model", "assumption"]
    for i, item in enumerate(items[:4]):
        add_dot(slide, 1.25, 2.48 + i * 0.65, 0.18, [PURPLE, BLUE, TEAL, GOLD][i], name=f"Appendix bullet {i + 1}")
        add_text(slide, 1.58, 2.40 + i * 0.65, 2.30, 0.40, item, 22, name=f"Appendix bullet text {i + 1}")
    add_box(slide, 4.65, 1.45, 7.70, 2.08, fill=PALE_BLUE, line=BLUE, name="Appendix mechanism")
    add_text(slide, 4.95, 1.76, 7.10, 0.34, "為什麼這個細節存在？", 24, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name="Appendix mechanism heading")
    add_text(slide, 4.98, 2.22, 7.00, 0.82, shorten(clean(page.visual), 190), 24, align=PP_ALIGN.CENTER, name="Appendix mechanism body")
    add_box(slide, 4.65, 3.86, 7.70, 2.29, fill=CREAM, line=GOLD, name="Appendix boundary")
    add_text(slide, 4.95, 4.16, 7.10, 0.34, "退回主線時要保留的邊界", 24, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, name="Appendix boundary heading")
    add_text(slide, 4.98, 4.62, 7.00, 1.10, "LEO 是 changing service window；endpoint J 不是 whole-system energy；任何未凍結來源、畫面或數字保留 placeholder。", 22, align=PP_ALIGN.CENTER, name="Appendix boundary body")


def choose_visual(page: Page, index: int) -> str:
    n = page.number
    if n == 1:
        return "cover"
    if n in (33, 34):
        return "formula"
    if n in (78, 79, 81, 82, 83, 84, 85):
        return "evidence"
    if n >= 97:
        return "appendix" if n % 2 else "transfer"
    if 28 <= n <= 38:
        return "state" if n not in (35, 36, 37) else "provenance"
    if 51 <= n <= 62:
        return "trace" if n in (51, 52, 53, 54, 55, 61) else ("compare" if is_operation(page) else "state")
    if 39 <= n <= 50 or 64 <= n <= 76:
        return "compare" if n in (45, 47, 48, 49, 58, 59, 60, 70, 71, 72, 74, 75) else "operation"
    if 23 <= n <= 27:
        return "policy"
    if 10 <= n <= 22:
        return "operation"
    if is_operation(page):
        return "operation"
    return "causal"


def image_for(page: Page) -> Path | None:
    mapping = {
        78: "01-prepare.png", 79: "02-lab-a-empty.png",
        81: "03-lab-b-empty.png", 82: "04-lab-c-empty.png",
        83: "05-evidence-empty.png", 84: "06-workbook-empty.png",
        85: "07-lab-a-fallback-http-sha-unavailable.png",
    }
    path = LIVE_EVIDENCE / mapping.get(page.number, "")
    return path if path.exists() else None


def compose(slide, page: Page, index: int):
    kind = choose_visual(page, index)
    if kind == "cover":
        draw_cover(slide, page)
    else:
        clear_placeholders(slide, keep_title=True)
        if kind == "formula":
            draw_formula_hook(slide, page, "E_endpoint" if page.number == 33 else "eta_E")
        elif kind == "evidence":
            draw_evidence(slide, page, image_for(page), index)
        elif kind == "operation":
            draw_operation(slide, page, index)
        elif kind == "policy":
            draw_code_or_policy(slide, page, index)
        elif kind == "state":
            draw_state(slide, page, index)
        elif kind == "trace":
            draw_trace(slide, page, index)
        elif kind == "compare":
            draw_compare(slide, page, index)
        elif kind == "provenance":
            draw_provenance(slide, page, index)
        elif kind == "transfer":
            draw_transfer(slide, page, index)
        elif kind == "appendix":
            draw_appendix(slide, page, index)
        else:
            draw_causal(slide, page, index)
    add_notes(slide, page)


def ensure_dirs() -> None:
    for name in ("sources", "analysis", "exports", "validation", "evidence", "renders"):
        (PROJECT / name).mkdir(parents=True, exist_ok=True)


def copy_inputs() -> None:
    ensure_dirs()
    shutil.copy2(SCRIPT, PROJECT / "sources/full-deck-v2-classroom-script.md")
    shutil.copy2(TEMPLATE, PROJECT / "sources/educate.pptx")
    stale = PROJECT / "evidence/browser-server-preview"
    if stale.exists():
        shutil.rmtree(stale)
    live_target = PROJECT / "evidence/browser-live-course-20260811"
    if live_target.exists():
        shutil.rmtree(live_target)
    evidence = LIVE_EVIDENCE
    if evidence.exists():
        for path in evidence.glob("*.png"):
            target = PROJECT / "evidence/browser-live-course-20260811"
            target.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target / path.name)


def inventory(pages: list[Page]) -> None:
    rows = []
    for i, page in enumerate(pages, 1):
        rows.append({"slide": i, "title": page.title, "visual": choose_visual(page, i), "operation": is_operation(page), "operation_kind": operation_kind(page) if is_operation(page) else None, "status": status(page), "exact_action": action_copy(page) if is_operation(page) else None})
    (PROJECT / "analysis/page-inventory.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (PROJECT / "analysis/operation-contract.md").write_text(
        "# v3 operation contract\n\n"
        "Every setup, run, policy-edit, import, freeze and recovery page keeps the exact action on-slide, then states purpose/mechanism, expected observable output, and interpretation/recovery. Layouts vary across split panels, command ribbons, code focus, receipt paths and evidence comparisons.\n\n"
        "Formula pages reserve named native-equation hooks and contain no formula image or raw LaTeX source. The native educate master/layout/footer/logo/background are preserved byte-for-byte.\n",
        encoding="utf-8",
    )


def add_slides(pages: list[Page]) -> Path:
    prs = Presentation(str(TEMPLATE))
    # Reuse the two native slides so their original relationships remain
    # ordinary template relationships; append the remaining pages with native
    # layouts.  No master/layout XML is rewritten.
    layout_cycle = [1, 2, 3, 4, 5, 6, 7, 1, 2, 3]
    for idx, page in enumerate(pages, 1):
        if idx == 1:
            slide = prs.slides[0]
        elif idx == 2:
            slide = prs.slides[1]
        else:
            layout = prs.slide_layouts[layout_cycle[(idx - 3) % len(layout_cycle)]]
            slide = prs.slides.add_slide(layout)
        compose(slide, page, idx)
    if len(prs.slides) != len(pages):
        raise RuntimeError(f"slide count {len(prs.slides)} != {len(pages)}")
    raw = PROJECT / "validation/c120-lora-leo-classroom-v3-native.pptx"
    prs.save(str(raw))
    final = PROJECT / "exports/c120-lora-leo-classroom-v3-editable.pptx"
    shutil.copy2(raw, final)
    restore_template_parts(final)
    return final


def restore_template_parts(output: Path) -> None:
    """Put untouched native master/layout bytes back after python-pptx saves.

    python-pptx changes XML declarations and line endings on parts it reads,
    even when their semantic content is untouched.  Copying the original
    package bytes for master/layout/theme/media parts keeps the educate footer,
    logo, background and layout relationships exactly as supplied.  Slide and
    notes parts remain authored by this builder.
    """
    with zipfile.ZipFile(output, "r") as out, zipfile.ZipFile(TEMPLATE, "r") as source:
        members = {info.filename: out.read(info.filename) for info in out.infolist()}
        infos = {info.filename: copy.copy(info) for info in out.infolist()}
        for info in source.infolist():
            name = info.filename
            if name.startswith("ppt/slideMasters/") or name.startswith("ppt/slideLayouts/") or name.startswith("ppt/theme/") or name in {"ppt/media/image1.png", "ppt/media/image2.png", "ppt/media/image3.png"}:
                if name in members:
                    members[name] = source.read(name)
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, data in members.items():
            target.writestr(infos[name], data)


def ordered_slide_parts(z: zipfile.ZipFile) -> list[str]:
    pres = ET.fromstring(z.read("ppt/presentation.xml"))
    rels = ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
    targets = {n.get("Id"): n.get("Target") for n in rels.findall(qn(PR_NS, "Relationship"))}
    parts = []
    for node in pres.findall(f"./{qn(P_NS, 'sldIdLst')}/{qn(P_NS, 'sldId')}"):
        target = targets.get(node.get(qn("http://schemas.openxmlformats.org/officeDocument/2006/relationships", "id")))
        if target:
            parts.append("ppt/" + target.lstrip("/") if not target.startswith("ppt/") else target)
    return parts


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def qa(final: Path, pages: list[Page]) -> dict[str, object]:
    errors: list[str] = []
    warnings: list[str] = []
    native_hooks = 0
    notes_count = 0
    slide_count = 0
    template_mismatch: list[str] = []
    with zipfile.ZipFile(final) as deck, zipfile.ZipFile(TEMPLATE) as template:
        if deck.testzip():
            errors.append("corrupt ZIP member")
        slide_parts = ordered_slide_parts(deck)
        slide_count = len(slide_parts)
        if slide_count != len(pages):
            errors.append(f"slide_count={slide_count} expected={len(pages)}")
        for i, part in enumerate(slide_parts, 1):
            root = ET.fromstring(deck.read(part))
            text = "".join(n.text or "" for n in root.findall(f".//{qn(A_NS, 't')}"))
            if FORBIDDEN in text:
                errors.append(f"slide {i}: forbidden exact word")
            if MINUTE_RE.search(text):
                errors.append(f"slide {i}: duration label")
            if "LaTeX source" in text or "\\frac" in text or "\\sum" in text:
                errors.append(f"slide {i}: raw formula source")
            if any("NATIVE_EQUATION_HOOK_" in node.get("name", "") for node in root.findall(f".//{qn(P_NS, 'cNvPr')}")):
                native_hooks += 1
            if root.find(f"./{qn(P_NS, 'cSld')}/{qn(P_NS, 'bg')}") is not None:
                errors.append(f"slide {i}: slide-level background added")
            # Check visible slide-local shapes stay above the native footer.
            for sp in root.findall(f"./{qn(P_NS, 'cSld')}/{qn(P_NS, 'spTree')}/*"):
                xfrm = sp.find(f"./{qn(P_NS, 'spPr')}/{qn(A_NS, 'xfrm')}")
                if xfrm is None:
                    xfrm = sp.find(f"./{qn(P_NS, 'grpSpPr')}/{qn(A_NS, 'xfrm')}")
                if xfrm is None:
                    continue
                off, ext = xfrm.find(qn(A_NS, "off")), xfrm.find(qn(A_NS, "ext"))
                if off is None or ext is None:
                    continue
                try:
                    y = int(off.get("y", "0")) / 914400
                    h = int(ext.get("cy", "0")) / 914400
                except ValueError:
                    continue
                if y + h > 6.50 and "sldNum" not in ET.tostring(sp, encoding="unicode"):
                    warnings.append(f"slide {i}: local shape reaches {y + h:.2f}in")
        note_names = sorted(n for n in deck.namelist() if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml"))
        notes_count = len(note_names)
        if notes_count != len(pages):
            errors.append(f"notes_count={notes_count} expected={len(pages)}")
        for name in note_names:
            root = ET.fromstring(deck.read(name))
            text = "".join(n.text or "" for n in root.findall(f".//{qn(A_NS, 't')}"))
            if FORBIDDEN in text or MINUTE_RE.search(text):
                errors.append(f"notes {name}: forbidden word/duration")
        # Footer, logo, background and native master/layout must remain intact.
        for name in ("ppt/media/image1.png", "ppt/media/image2.png", "ppt/media/image3.png", "ppt/theme/theme1.xml", "ppt/theme/theme2.xml"):
            if name not in deck.namelist() or name not in template.namelist() or sha(deck.read(name)) != sha(template.read(name)):
                template_mismatch.append(name)
        for name in sorted(n for n in template.namelist() if n.startswith("ppt/slideMasters/") or n.startswith("ppt/slideLayouts/")):
            if name not in deck.namelist() or sha(deck.read(name)) != sha(template.read(name)):
                template_mismatch.append(name)
        if template_mismatch:
            errors.append("template bytes changed/missing: " + ", ".join(template_mismatch[:8]))
    if native_hooks != 2:
        errors.append(f"native_equation_hooks={native_hooks} expected=2")
    report = {
        "schema": "c120-classroom-v3-qa-v1", "status": "PASS" if not errors else "FAIL",
        "input": str(final), "slide_count": slide_count, "notes_count": notes_count,
        "native_equation_hooks": native_hooks, "template_bytes_preserved": not template_mismatch,
        "errors": errors, "warnings": warnings,
        "limitations": ["Native OMML insertion is intentionally deferred to the controller hook pass.", "Rendered pixels still require original-size PDF/contact-sheet inspection."],
    }
    (PROJECT / "validation/qa_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def render(final: Path) -> None:
    render_dir = PROJECT / "renders"
    render_dir.mkdir(parents=True, exist_ok=True)
    out_dir = Path("/tmp/c120v3-render-out3")
    profile = Path("/tmp/c120v3-render-profile3")
    if out_dir.exists():
        shutil.rmtree(out_dir)
    if profile.exists():
        shutil.rmtree(profile)
    out_dir.mkdir(parents=True)
    command = ["libreoffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--nodefault", "--nolockcheck", "--nofirststartwizard", "--convert-to", "pdf", "--outdir", str(out_dir), str(final)]
    result = subprocess.run(command, text=True, capture_output=True)
    pdf = out_dir / (final.stem + ".pdf")
    (PROJECT / "validation/render.log").write_text(result.stdout + "\nSTDERR:\n" + result.stderr, encoding="utf-8")
    if result.returncode != 0 or not pdf.exists():
        (PROJECT / "validation/render-status.md").write_text("LibreOffice did not produce a PDF; see validation/render.log.\n", encoding="utf-8")
        return
    shutil.copy2(pdf, PROJECT / "validation/c120-lora-leo-classroom-v3.pdf")
    subprocess.run(["pdftoppm", "-jpeg", "-r", "100", str(pdf), str(render_dir / "slide")], check=False)
    images = sorted(render_dir.glob("slide-*.jpg"))
    (PROJECT / "validation/render-status.md").write_text(f"PDF: {pdf}\nRendered JPEG pages: {len(images)}\nExpected: 108\n\nInspect original-size pages; footer must remain visible.\n", encoding="utf-8")


def write_readme(final: Path, report: dict[str, object]) -> None:
    text = f"""# C-120 classroom deck v3

- Final editable deck: `{final.relative_to(PROJECT)}`
- Source template: `/home/u24/ppt-master/template/educate.pptx`
- Slide count: 108 (within the accepted 98–116 range)
- Content donor: `sources/full-deck-v2-classroom-script.md`, rewritten into varied native layouts
- Native master/layout/footer/logo/background: preserved from `educate.pptx`; no master/layout reserialization or footer stripping
- Formula slides P033/P034: named native-equation hooks only; no formula PNG/SVG and no raw LaTeX on slide. Controller inserts OMML later.
- Evidence: `evidence/browser-live-course-20260811/` is the current live-course capture set; P085 preserves the explicit HTTP `SHA256_UNAVAILABLE` blocker and no old success screenshot is reused.
- Speaker notes: 108 notes slides; exact forbidden word and duration labels are rejected by the builder.
- QA: `validation/qa_report.json` (`{report['status']}`); render evidence under `validation/render-status.md` and `renders/`.

Every operation page carries exact action, purpose/mechanism, expected observable output, and interpretation/recovery. The operation treatment varies between split action/why panels, command ribbons, code-focus views, receipt paths, evidence comparisons and recovery branches.
"""
    (PROJECT / "README.md").write_text(text, encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-render", action="store_true")
    args = parser.parse_args(argv)
    pages = parse_pages()
    copy_inputs()
    inventory(pages)
    final = add_slides(pages)
    report = qa(final, pages)
    if not args.skip_render:
        render(final)
    write_readme(final, report)
    print(json.dumps({"final": str(final), "qa_status": report["status"], "errors": report["errors"], "warnings": report["warnings"]}, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
