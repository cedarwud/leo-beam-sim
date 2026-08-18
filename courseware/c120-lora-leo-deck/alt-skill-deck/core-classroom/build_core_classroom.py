#!/usr/bin/env python3
"""Build a compact classroom-first LoRaEnergySim core deck."""

from __future__ import annotations

import importlib.util
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent
ALT_ROOT = ROOT.parent
HELPER_PATH = ALT_ROOT / "part-a-v2-p019-p027b" / "build_section.py"
PRELAB_OUTPUT = ROOT / "build" / "LoRaEnergySim-LEO-ALT-CORE-PRELAB.pptx"
LABS_FALLBACK_OUTPUT = ROOT / "build" / "LoRaEnergySim-LEO-ALT-CORE-LABS-FALLBACK.pptx"
REPLAY_IMAGE = ALT_ROOT / "part-c-v2-p064-p080" / "assets" / "replay-frame.png"
LEDGER_IMAGE = ALT_ROOT / "part-c-v2-p064-p080" / "assets" / "ledger.png"

spec = importlib.util.spec_from_file_location("c120_p019_helpers", HELPER_PATH)
if spec is None or spec.loader is None:
    raise RuntimeError(f"cannot load helper: {HELPER_PATH}")
H = importlib.util.module_from_spec(spec)
spec.loader.exec_module(H)

# Keep IDs in source/build metadata while making every visible title splice-safe.
_prepare_slide_with_internal_id = H.prepare_slide


def _prepare_slide_without_visible_id(prs, title):
    semantic_title = re.sub(r"^[A-Z]+(?:-?\d+)+[A-Za-z]?\s*｜\s*", "", title)
    return _prepare_slide_with_internal_id(prs, semantic_title)


H.prepare_slide = _prepare_slide_without_visible_id


def title_lead(slide, text: str) -> None:
    H.add_text(slide, 0.94, 0.98, 11.45, 0.52, text, 24, H.NAVY, True,
               PP_ALIGN.CENTER, name="Plain-language lead")


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def add_three_cards(slide, cards, y=1.72, h=2.72) -> None:
    xs = [0.92, 4.77, 8.62]
    for index, (heading, body, fill, line) in enumerate(cards):
        H.add_visual_card(slide, xs[index], y, 3.54, h, heading, body, fill, line,
                          f"Teaching card {index + 1}", heading_size=21, body_size=19)


def add_four_cards(slide, cards, y=1.74, h=2.55) -> None:
    xs = [0.92, 3.97, 7.02, 10.07]
    for index, (heading, body, fill, line) in enumerate(cards):
        H.add_visual_card(slide, xs[index], y, 2.74, h, heading, body, fill, line,
                          f"Teaching card {index + 1}", heading_size=19, body_size=17)


def add_result_order(slide, y=3.92) -> None:
    H.add_label_box(slide, 1.00, y, 2.44, 0.68, "① 服務\n是否完成任務", H.BLUE_PALE, H.BLUE,
                    18, H.NAVY, True, name="Read service")
    H.add_chevron(slide, 3.53, y + 0.12, 0.38, 0.42, H.TEAL, "Read connector 1")
    H.add_label_box(slide, 4.00, y, 2.44, 0.68, "② 封包\n送達／重試／逾期", H.TEAL_PALE, H.TEAL,
                    17, H.NAVY, True, name="Read packet")
    H.add_chevron(slide, 6.53, y + 0.12, 0.38, 0.42, H.GOLD, "Read connector 2")
    H.add_label_box(slide, 7.00, y, 2.44, 0.68, "③ 狀態\n休眠／清醒／傳送", H.GOLD_PALE, H.GOLD,
                    17, H.NAVY, True, name="Read state")
    H.add_chevron(slide, 9.53, y + 0.12, 0.38, 0.42, H.PURPLE, "Read connector 3")
    H.add_label_box(slide, 10.00, y, 2.44, 0.68, "④ 端點能量\n最後才讀焦耳", H.PURPLE_PALE, H.PURPLE,
                    17, H.NAVY, True, name="Read energy")


def intro_problem(prs):
    slide = H.prepare_slide(prs, "C001｜這個實驗真正要解決什麼")
    title_lead(slide, "有限電池的物聯網端點，必須在變動服務窗口中決定何時休息、等待或傳送")
    add_three_cards(slide, [
        ("端點有工作", "感測資料持續產生\n封包進入佇列", H.BLUE_PALE, H.BLUE),
        ("服務窗口會變", "有時可傳、有時關閉\n品質與剩餘時間不同", H.TEAL_PALE, H.TEAL),
        ("每個動作都有代價", "休眠省電但要喚醒\n等待反應快但持續耗能", H.GOLD_PALE, H.GOLD),
    ])
    H.add_causal_strip(slide, "問題不是只追求最低焦耳，而是在完成服務的前提下解釋能量取捨",
                       "C001 causal", fill=H.PURPLE_PALE, line=H.PURPLE, color=H.NAVY)
    H.add_field_band(slide, "endpoint", "物聯網端點", "固定課程情境", "產生資料並執行策略動作",
                     "裝置角色", "只分析端點服務與端點能量，不等於整體 LEO 系統能量", H.PURPLE)
    notes(slide, "本課從物聯網端點的能源決策開始。端點有有限電池，資料封包持續產生，但服務窗口會開啟、關閉並改變品質。策略每次只能選擇休眠、清醒等待或送出。休眠可能降低待機能量，卻增加喚醒成本與延遲；等待保持反應能力，卻累積清醒閒置能量；傳送增加服務機會，也帶來處理與無線事件。因此判斷不能只看最低焦耳，必須先確認服務是否完成。")


def intro_simulator(prs):
    slide = H.prepare_slide(prs, "C002｜LoRaEnergySim 在做什麼")
    title_lead(slide, "它把策略動作轉成可重播的狀態、封包、服務與端點能量結果")
    add_four_cards(slide, [
        ("讀取情境", "窗口、品質、期限\n封包到達與固定種子", H.BLUE_PALE, H.BLUE),
        ("呼叫策略", "讀 `student_policy.py`\n取得一個合法動作", H.TEAL_PALE, H.TEAL),
        ("推進事件", "狀態停留、喚醒\n嘗試、重試、送達", H.GOLD_PALE, H.GOLD),
        ("輸出證據", "服務成敗\n端點焦耳與事件重播", H.PURPLE_PALE, H.PURPLE),
    ])
    H.add_causal_strip(slide, "同一情境重跑 → 只改一個策略值 → 比較事件路徑是否真的改變",
                       "C002 causal")
    H.add_field_band(slide, "simulator", "模擬器", "課程封裝的固定執行器", "重複推進端點與封包事件",
                     "一次執行", "產生一致的課程模擬證據，不是即時量測", H.BLUE)
    notes(slide, "LoRaEnergySim 的核心用途是把策略選擇放入一個可重複執行的事件流程。固定情境先提供服務窗口、品質、期限、封包到達與種子；執行器在每個決策步驟呼叫 student_policy.py，取得一個合法動作，再推進無線狀態、封包嘗試、重試、送達或逾期。最後輸出服務結果、端點能量與事件重播。這使一次程式修改能沿著動作、事件與結果被檢查，而不是只留下單一分數。")


def intro_lora_value(prs):
    slide = H.prepare_slide(prs, "C003｜LoRaEnergySim 提供的三項價值")
    title_lead(slide, "它提供成熟的 LoRa 端點問題框架，讓狀態、封包與能量能沿同一條路徑解讀")
    add_three_cards(slide, [
        ("端點狀態與能量", "把休眠、處理、TX、RX\n連到端點能量邊界", H.BLUE_PALE, H.BLUE),
        ("封包事件與可靠度", "用封包、碰撞、重試、送達\n解釋策略後果", H.TEAL_PALE, H.TEAL),
        ("可重複研究架構", "固定模型來源與版本\n讓案例與結果可以比較", H.GOLD_PALE, H.GOLD),
    ])
    H.add_causal_strip(slide, "LoRaEnergySim 讓策略動作可以連到狀態、封包、服務與端點能量證據",
                       "LoRaEnergySim value", fill=H.PURPLE_PALE, line=H.PURPLE, color=H.NAVY)
    H.add_field_band(slide, "model_reference", "模型參考", "固定的 LoRaEnergySim repository 與版本",
                     "提供端點狀態、封包事件與能量欄位的可追溯來源", "來源識別",
                     "以相同模型語言比較不同策略的事件與結果", H.PURPLE)
    notes(slide, "LoRaEnergySim 提供三項價值。第一，它建立 LoRa 端點的狀態與能量模型，使休眠、處理、發射、接收與端點能量邊界有可追溯來源。第二，它提供封包、碰撞、重試與送達的事件語言，讓策略結果能由中間事件解釋。第三，固定模型來源與版本，使不同案例與策略結果能以一致的欄位和判讀順序比較。課程操作因此可以從一行策略修改，一路追蹤到狀態、封包、服務與端點能量。")


def intro_wrapper(prs):
    slide = H.prepare_slide(prs, "C004｜上游模型與課程封裝分別完成什麼")
    title_lead(slide, "研究模型提供端點與封包概念；課程封裝把它整理成可教、可驗、可匯入的固定流程")
    H.add_card(slide, 0.94, 1.68, 5.55, 2.72, "上游 LoRaEnergySim 能力",
               "• 端點無線狀態與停留時間\n• 封包、碰撞、重傳與送達事件\n• 端點能量模型與統計\n\n本課把它視為能力來源，不宣稱本次直接執行上游程式。",
               H.BLUE_PALE, H.BLUE, title_size=22, body_size=19, name="Upstream capability")
    H.add_card(slide, 6.82, 1.68, 5.55, 2.72, "課程封裝另外完成",
               "• 固定 LEO 服務窗口情境與案例\n• 只可修改的 `student_policy.py` 區塊\n• 可重複 runner、schema 與兩個 JSON\n• Leo 網站驗證、重播與 Workbook 比較",
               H.TEAL_PALE, H.TEAL, title_size=22, body_size=19, name="Course wrapper")
    H.add_causal_strip(slide, "上游能力來源 ≠ 本次上游執行；課程結果維持 coherent simulated evidence",
                       "C003 boundary", fill=H.RED_PALE, line=H.RED, color=H.RED)
    H.add_field_band(slide, "engine_mode", "執行模式", "結果中繼資料", "標示資料如何產生",
                     "模式值", "不得把課程模擬結果改稱實測或上游直接執行", H.RED)
    notes(slide, "需要明確區分兩層。上游 LoRaEnergySim 提供端點狀態、封包事件、重傳與端點能量等研究能力；課程封裝則固定情境、案例順序、可修改區塊、runner、schema、結果檔與網站匯入流程。本課結果的 engine_mode 維持 coherent course simulated adapter，不能把它說成即時量測，也不能說成本次直接執行上游程式。")


def intro_competition(prs):
    slide = H.prepare_slide(prs, "C005｜為什麼競賽要使用這套工具")
    title_lead(slide, "競賽需要的不只是想法，而是能說明『改了什麼、為何有效、證據在哪裡』")
    add_three_cards(slide, [
        ("可重複", "固定情境、案例與種子\n相同條件可以重跑", H.BLUE_PALE, H.BLUE),
        ("可歸因", "一次只改一個值\n事件差異可連回修改", H.TEAL_PALE, H.TEAL),
        ("可反駁", "保留 withheld case\n結果不符也要保存", H.GOLD_PALE, H.GOLD),
    ])
    H.add_causal_strip(slide, "三個 Lab 練習的核心：提出預測 → 受控修改 → 執行 → 匯入 → 解釋取捨",
                       "C004 evidence")
    H.add_field_band(slide, "evidence", "證據紀錄", "result、replay 與 Workbook", "保存修改至結果的因果路徑",
                     "紀錄", "只能支持端點層級的受控比較", H.PURPLE)
    notes(slide, "在智慧節能與物聯網競賽中，單純說策略比較省電仍不足。需要交代修改位置、機制、預期方向、執行條件與可檢查結果。固定情境提供可重複性，一次只改一個值提供可歸因性，保留 withheld case 則讓機制可以被反駁。三個 Lab 因此不是追求漂亮分數，而是練習建立一條可重建的端點證據鏈。")


def intro_artifacts(prs):
    slide = H.prepare_slide(prs, "C006｜每次執行會產生什麼檔案")
    title_lead(slide, "只有 run case 會產生案例結果；setup 與 verify 只準備環境")
    H.add_label_box(slide, 0.94, 1.72, 3.40, 1.20, "執行一個 case\n`course.sh run ...`", H.BLUE_PALE, H.BLUE,
                    23, H.NAVY, True, name="Run one case")
    H.add_chevron(slide, 4.52, 2.05, 0.48, 0.54, H.TEAL, "Artifact connector")
    H.add_label_box(slide, 5.18, 1.72, 3.25, 1.20, "`result.json`\n摘要、服務、封包、端點能量", H.TEAL_PALE, H.TEAL,
                    19, H.NAVY, True, name="Result artifact")
    H.add_label_box(slide, 8.72, 1.72, 3.65, 1.20, "`endpoint-replay.json`\n動作、狀態、佇列、封包時序", H.GOLD_PALE, H.GOLD,
                    18, H.NAVY, True, name="Replay artifact")
    H.add_card(slide, 0.94, 3.36, 5.56, 1.38, "一般執行",
               "兩個 JSON 位於同一個 `artifacts/<run_id>/` 資料夾；以命令輸出的 `result_path` 找檔案。",
               H.PURPLE_PALE, H.PURPLE, title_size=20, body_size=18, name="Normal run")
    H.add_card(slide, 6.82, 3.36, 5.56, 1.38, "使用 `--freeze`",
               "除了兩個 JSON，還會保存 receipt 與 policy checkpoint；它們不是一般上傳檔。",
               H.RED_PALE, H.RED, title_size=20, body_size=18, name="Freeze run")
    H.add_field_band(slide, "result_path", "結果檔路徑", "runner 命令輸出", "定位該次 result.json",
                     "路徑文字", "不要自行猜資料夾，也不要混用另一個 run 的 replay", H.PURPLE)
    notes(slide, "setup 與 verify 只建立並核對環境，不會產生案例結果。每次成功的 run case 會建立一個新的 run_id 資料夾，其中有 result.json 與 endpoint-replay.json。result.json 保存摘要、服務、封包與端點能量；endpoint-replay.json 保存動作、狀態、佇列與封包時序。命令會印出 result_path，應依該路徑找檔案。使用 --freeze 時另外保存 receipt 與 policy checkpoint，但這些不是網站的一般上傳檔。")


def intro_website(prs):
    slide = H.prepare_slide(prs, "C007｜本機 runner 與網站各做什麼")
    title_lead(slide, "模擬在本機完成；網站只接收結果、驗證身分、呈現重播並保存比較")
    H.add_card(slide, 0.94, 1.66, 5.54, 2.06, "本機 LoRa runner",
               "讀固定情境與 `student_policy.py`\n執行 case\n產生 result／replay pair",
               H.BLUE_PALE, H.BLUE, title_size=22, body_size=20, name="Local runner")
    H.add_card(slide, 6.84, 1.66, 5.54, 2.06, "Leo `/course` 網站",
               "只選擇上傳 `result.json`\n驗證 schema／identity／units／provenance\nmaterialize replay 並寫入 Workbook",
               H.TEAL_PALE, H.TEAL, title_size=22, body_size=19, name="Course website")
    H.add_causal_strip(slide, "不要上傳 `student_policy.py`；`endpoint-replay.json` 留在同一 run 資料夾配對",
                       "C006 upload rule", fill=H.RED_PALE, line=H.RED, color=H.RED)
    H.add_field_band(slide, "result.json", "結果資料檔", "本機 runner", "作為網站匯入入口",
                     "JSON 檔案", "逐次匯入 baseline、candidate、withheld，再從服務開始判讀", H.TEAL)
    notes(slide, "真正的模擬發生在本機 runner。網站不執行 student_policy.py，也不重新計算策略；網站只接收 result.json，驗證 schema、scenario identity、run identity、units、policy 與 provenance，通過後才 materialize endpoint replay 並寫入 Workbook。endpoint-replay.json 要留在相同 run 資料夾做配對與稽核。案例應依順序逐次匯入，先看服務是否完成，再看封包、狀態與端點能量。")


def website_replay(prs):
    slide = H.prepare_slide(prs, "C008｜網站重播頁要看哪裡")
    title_lead(slide, "畫面用來解釋事件順序；數值仍以本次匯入的 result.json 為準")
    if REPLAY_IMAGE.exists():
        slide.shapes.add_picture(str(REPLAY_IMAGE), Inches(0.96), Inches(1.58), width=Inches(6.00), height=Inches(3.35))
    H.add_card(slide, 7.26, 1.58, 5.10, 3.35, "講解順序",
               "① Queue：當下有多少資料等待\n② Action：策略選擇什麼動作\n③ Radio／state：停留在哪種狀態\n④ Packet／service：是否嘗試、送達或逾期\n⑤ Energy：最後才解釋累積焦耳",
               H.BLUE_PALE, H.BLUE, title_size=22, body_size=19, name="Replay reading")
    H.add_field_band(slide, "event", "事件", "匯入結果的重播資料", "依時間呈現動作與後果",
                     "事件紀錄", "畫面只幫助閱讀，不取代 result identity 與原始數值", H.BLUE)
    notes(slide, "這張圖是目前網站重播區塊的位置參考。講解時先讀 Queue，確認資料是否等待；再讀 Action，確認策略選擇；接著看 Radio 或 state，確認休眠、清醒、喚醒、處理與傳送的停留；之後讀 Packet 與 service，判斷嘗試、重試、送達或逾期；最後才解釋端點能量。畫面上的數值必須以本次匯入的 result.json 為準，截圖本身不能冒充新的實驗結果。")


def website_ledger(prs):
    slide = H.prepare_slide(prs, "C009｜網站比較頁如何保存證據")
    title_lead(slide, "Workbook 把預測、修改、結果路徑與解釋放在同一次比較中")
    if LEDGER_IMAGE.exists():
        slide.shapes.add_picture(str(LEDGER_IMAGE), Inches(0.96), Inches(1.58), width=Inches(6.00), height=Inches(3.35))
    H.add_card(slide, 7.26, 1.58, 5.10, 3.35, "每個案例至少保存",
               "• case 與 run identity\n• 修改前的 prediction\n• result_path 與配對 replay\n• service／packet／state／endpoint J\n• 結論與仍存在的限制",
               H.GOLD_PALE, H.GOLD, title_size=22, body_size=19, name="Workbook reading")
    H.add_field_band(slide, "workbook", "比較紀錄", "瀏覽器課程工具", "保存多次案例的證據鏈",
                     "JSON 紀錄", "匯入失敗時原 Workbook 不應被改寫", H.GOLD)
    notes(slide, "Workbook 不是漂亮結果表，而是可重開的證據紀錄。每個案例至少要保存 case 與 run identity、執行前 prediction、result_path、配對 replay、服務、封包、狀態、端點能量、結論與限制。若 schema、identity、units、policy 或 provenance 不一致，匯入應 fail closed，原有 Workbook 不被悄悄改寫。")


def env_sequence(prs):
    slide = H.prepare_slide(prs, "從 Python 3.11 到 READY 的完整順序")
    title_lead(slide, "每一步都在 lora-energy-lab 根目錄完成；上一關成功才進入下一關")
    add_three_cards(slide, [
        ("① 找到正確版本", "先檢查 Python 3.11.x\n沒有才依平台安裝", H.BLUE_PALE, H.BLUE),
        ("② 建立隔離環境", "確認 venv → 建立 `.venv`\n啟用後再次核對版本", H.TEAL_PALE, H.TEAL),
        ("③ 安裝並驗證", "在 venv 內安裝 locked requirements\nverify 必須顯示 READY", H.GOLD_PALE, H.GOLD),
    ])
    H.add_causal_strip(slide, "Python 3.11.x → `.venv` 已啟用 → 安裝 `requirements-lock.txt` → `READY`",
                       "Environment sequence", fill=H.PURPLE_PALE, line=H.PURPLE, color=H.NAVY)
    H.add_field_band(slide, "requirements-lock.txt", "鎖定需求檔", "課程套件根目錄",
                     "固定安裝契約與可重現環境", "文字檔",
                     "實際檔名含 -lock；不是 requirements.txt", H.PURPLE)
    notes(slide, "環境準備必須依固定順序完成。先在 lora-energy-lab 根目錄檢查或安裝 Python 3.11.x，再確認標準函式庫的 venv 模組可用，建立並啟用 `.venv`。啟用後再次執行版本與執行器路徑檢查，確定命令已進入隔離環境。接著安裝 `requirements-lock.txt`，最後執行 verify；只有看到 READY 才能開始案例。實際檔名包含 `-lock`，不是一般的 requirements.txt。")


def env_windows_python(prs):
    slide = H.prepare_slide(prs, "Windows：先檢查或安裝 Python 3.11")
    title_lead(slide, "先使用 Python Launcher 檢查；找不到 3.11 才安裝 uv 與精確版本")
    H.add_code_box(slide, 0.94, 1.60, 5.54, 2.42,
                   "先檢查：\n`py -3.11 --version`\n`$python311 = (py -3.11 -c \"import sys; print(sys.executable)\").Trim()`",
                   fill=H.BLUE_PALE, line=H.BLUE, size=18, name="Windows Python check")
    H.add_code_box(slide, 6.84, 1.60, 5.54, 2.42,
                   "找不到 3.11 時：\n`winget install --id=astral-sh.uv -e`\n`uv python install 3.11`\n`$python311 = (uv python find 3.11).Trim()`\n`& $python311 --version`",
                   fill=H.TEAL_PALE, line=H.TEAL, size=17, name="Windows Python install")
    H.add_causal_strip(slide, "成功訊號：輸出 `Python 3.11.x`，並取得 `$python311` 的執行器路徑",
                       "Windows Python success", fill=H.GOLD_PALE, line=H.GOLD, color=H.NAVY)
    H.add_field_band(slide, "python311", "Python 3.11 執行器變數", "PowerShell 命令結果",
                     "保存建立 venv 時要使用的精確執行器", "檔案路徑",
                     "版本不是 3.11.x 就停止，不修改 runner 放寬版本", H.BLUE)
    notes(slide, "Windows 先用 Python Launcher 執行 `py -3.11 --version`。若輸出 Python 3.11.x，再把實際執行器路徑存入 `$python311`。若找不到版本，先以 winget 安裝 uv，重新開啟 PowerShell 後執行 `uv python install 3.11`，再用 `uv python find 3.11` 取得路徑。最後以該路徑再次顯示版本；版本正確才建立 venv，不可修改 runner 迴避版本關卡。")


def env_windows_venv(prs):
    slide = H.prepare_slide(prs, "Windows：建立、啟用 venv，再安裝與驗證")
    title_lead(slide, "所有 pip 與 verify 命令都要在已啟用的 `.venv` 中執行")
    H.add_code_box(slide, 0.94, 1.58, 5.54, 3.18,
                   "建立並啟用：\n`& $python311 -m venv --help`\n`& $python311 -m venv .venv`\n`.\\.venv\\Scripts\\Activate.ps1`\n`python --version`\n`python -c \"import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)\"`",
                   fill=H.BLUE_PALE, line=H.BLUE, size=16.5, name="Windows venv")
    H.add_code_box(slide, 6.84, 1.58, 5.54, 3.18,
                   "在 venv 內安裝與驗證：\n`python -m pip install --disable-pip-version-check`\n`  --require-hashes --no-deps`\n`  -r .\\requirements-lock.txt`\n`.\\course.cmd verify`\n\n預期：`status = READY`",
                   fill=H.TEAL_PALE, line=H.TEAL, size=16.5, name="Windows install verify")
    H.add_field_band(slide, ".venv", "專案虛擬環境", "Python venv 模組",
                     "隔離本課執行器與安裝契約", "資料夾",
                     "啟用後 Python 路徑應位於 .venv；READY 才繼續", H.TEAL)
    notes(slide, "Windows 先以 `$python311 -m venv --help` 確認 venv 模組存在，再建立 `.venv` 並執行 Activate.ps1。`python --version` 必須仍為 3.11.x，下一條命令則以 `sys.prefix` 與 `sys.base_prefix` 的差異確認已進入隔離環境。接著在此環境中安裝 `requirements-lock.txt`，再執行 `course.cmd verify`。最終成功訊號是 status 為 READY，而不是只有 pip 命令沒有報錯。")


def env_linux_python(prs):
    slide = H.prepare_slide(prs, "Linux／WSL：先檢查或安裝 Python 3.11")
    title_lead(slide, "WSL 使用其 Linux distribution；先檢查，沒有 3.11 才使用 uv 安裝")
    H.add_code_box(slide, 0.94, 1.56, 5.54, 3.24,
                   "先檢查：\n`command -v python3.11`\n`python3.11 --version`\n\n若成功：\n`PYTHON_BIN=\"$(command -v python3.11)\"`",
                   fill=H.BLUE_PALE, line=H.BLUE, size=18, name="Linux Python check")
    H.add_code_box(slide, 6.84, 1.56, 5.54, 3.24,
                   "找不到 3.11 時：\n`sudo apt update`\n`sudo apt install -y curl`\n`curl -LsSf https://astral.sh/uv/install.sh | sh`\n`uv python install 3.11`\n`PYTHON_BIN=\"$(uv python find 3.11)\"`\n`\"$PYTHON_BIN\" --version`",
                   fill=H.TEAL_PALE, line=H.TEAL, size=16, name="Linux Python install")
    H.add_field_band(slide, "PYTHON_BIN", "Python 執行器路徑變數", "Linux shell 命令結果",
                     "指定建立 venv 的精確 Python 3.11", "檔案路徑",
                     "輸出必須是 Python 3.11.x；WSL 也使用此 Linux 流程", H.BLUE)
    notes(slide, "Linux 與 WSL 先執行 `command -v python3.11` 與版本檢查。若版本正確，把路徑保存於 `PYTHON_BIN`。若找不到版本，Ubuntu 24.04 不應反覆嘗試不存在的 python3.11 apt 套件；先準備 curl，再依 uv 官方安裝流程安裝 uv 與 Python 3.11，最後以 `uv python find 3.11` 取得路徑並再次核對版本。WSL 使用同一套 Linux 命令，不使用 Windows Python。")


def env_linux_venv(prs):
    slide = H.prepare_slide(prs, "Linux／WSL：建立、啟用 venv，再安裝與驗證")
    title_lead(slide, "建立 `.venv` 後必須執行 source；終端機才會改用隔離環境")
    H.add_code_box(slide, 0.94, 1.56, 5.54, 3.24,
                   "建立並啟用：\n`\"$PYTHON_BIN\" -m venv --help`\n`\"$PYTHON_BIN\" -m venv .venv`\n`source .venv/bin/activate`\n`python --version`\n`python -c 'import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)'`",
                   fill=H.BLUE_PALE, line=H.BLUE, size=16.5, name="Linux venv")
    H.add_code_box(slide, 6.84, 1.56, 5.54, 3.24,
                   "在 venv 內安裝與驗證：\n`python -m pip install`\n`  --disable-pip-version-check`\n`  --require-hashes --no-deps`\n`  -r requirements-lock.txt`\n`bash course.sh verify`\n\n預期：`status = READY`",
                   fill=H.TEAL_PALE, line=H.TEAL, size=16.5, name="Linux install verify")
    H.add_field_band(slide, "source", "啟用 shell 環境的命令", "Linux／macOS shell",
                     "把目前終端機切換到 .venv 的 Python", "命令",
                     "啟用後執行器路徑應位於 .venv；READY 才開始 Lab", H.TEAL)
    notes(slide, "Linux 與 WSL 先以 `PYTHON_BIN -m venv --help` 確認支援，再建立 `.venv`。執行 `source .venv/bin/activate` 後，版本必須是 3.11.x，執行器路徑也必須落在 `.venv`。接著使用目前環境的 Python 執行 pip，依 hash 與 no-deps 契約安裝 `requirements-lock.txt`，最後執行 `bash course.sh verify`。只有 status 為 READY 才表示環境、情境與策略契約可以進入案例。")


def env_macos_python(prs):
    slide = H.prepare_slide(prs, "macOS：先檢查或安裝 Python 3.11")
    title_lead(slide, "不要假設 macOS 的 python3 版本；先檢查，再選 Homebrew 或 uv")
    H.add_code_box(slide, 0.94, 1.56, 3.48, 3.26,
                   "先檢查：\n`command -v python3.11`\n`python3.11 --version`\n\n若成功：\n`PYTHON_BIN=\"$(command -v python3.11)\"`",
                   fill=H.BLUE_PALE, line=H.BLUE, size=17, name="macOS check")
    H.add_code_box(slide, 4.68, 1.56, 3.48, 3.26,
                   "使用 Homebrew：\n`brew install python@3.11`\n`PYTHON_BIN=\"$(brew --prefix python@3.11)/bin/python3.11\"`\n`\"$PYTHON_BIN\" --version`",
                   fill=H.TEAL_PALE, line=H.TEAL, size=16.5, name="macOS brew")
    H.add_code_box(slide, 8.42, 1.56, 3.96, 3.26,
                   "使用 uv：\n`curl -LsSf https://astral.sh/uv/install.sh | sh`\n`uv python install 3.11`\n`PYTHON_BIN=\"$(uv python find 3.11)\"`\n`\"$PYTHON_BIN\" --version`",
                   fill=H.GOLD_PALE, line=H.GOLD, size=16, name="macOS uv")
    H.add_field_band(slide, "python@3.11", "Homebrew Python 3.11 套件", "Homebrew formula",
                     "提供 macOS 的精確 Python 版本", "套件版本",
                     "最終仍以 PYTHON_BIN --version 顯示 3.11.x 為準", H.GOLD)
    notes(slide, "macOS 不應假設系統的 `python3` 已是 3.11。先檢查 `python3.11`；若已存在，保存其路徑。若需要安裝，可以使用 Homebrew 的 `python@3.11`，或依 uv 官方流程安裝精確版本。兩條路徑只需選一條，最後都必須設定 `PYTHON_BIN` 並由該執行器輸出 Python 3.11.x。確認版本後才進入 venv 步驟，不混用另一個 Python 建立環境。")


def env_macos_venv(prs):
    slide = H.prepare_slide(prs, "macOS：建立、啟用 venv，再安裝與驗證")
    title_lead(slide, "Homebrew 與 uv 的 Python 都包含 venv；啟用後再進行 locked install")
    H.add_code_box(slide, 0.94, 1.56, 5.54, 3.24,
                   "建立並啟用：\n`\"$PYTHON_BIN\" -m venv --help`\n`\"$PYTHON_BIN\" -m venv .venv`\n`source .venv/bin/activate`\n`python --version`\n`python -c 'import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)'`",
                   fill=H.BLUE_PALE, line=H.BLUE, size=16.5, name="macOS venv")
    H.add_code_box(slide, 6.84, 1.56, 5.54, 3.24,
                   "在 venv 內安裝與驗證：\n`python -m pip install`\n`  --disable-pip-version-check`\n`  --require-hashes --no-deps`\n`  -r requirements-lock.txt`\n`bash course.sh verify`\n\n預期：`status = READY`",
                   fill=H.TEAL_PALE, line=H.TEAL, size=16.5, name="macOS install verify")
    H.add_field_band(slide, "verify", "環境驗證", "course.sh／verify_setup.py",
                     "核對 Python、情境、策略與安裝契約", "狀態",
                     "READY 只代表環境可用；尚未產生 Lab result", H.TEAL)
    notes(slide, "macOS 先以選定的 `PYTHON_BIN` 確認 venv 模組，建立 `.venv` 後以 source 啟用。版本與執行器路徑都正確，才在此環境中安裝 `requirements-lock.txt`。接著執行 `bash course.sh verify`，由同一個套件檢查 Python 版本、情境、baseline policy 與契約。READY 代表環境可以進入案例，並不表示已有服務或能量結果；案例 JSON 仍要由後續 run command 產生。")


def env_ready(prs):
    slide = H.prepare_slide(prs, "看到 READY 後，環境準備才算完成")
    title_lead(slide, "成功不是只有安裝指令沒有報錯；版本、venv 路徑與 verify 狀態都要成立")
    add_three_cards(slide, [
        ("版本正確", "`python --version`\n顯示 Python 3.11.x", H.BLUE_PALE, H.BLUE),
        ("venv 已啟用", "執行器路徑位於 `.venv`\n不是系統 Python", H.TEAL_PALE, H.TEAL),
        ("契約通過", "verify 輸出 `status = READY`\n寫入 verify receipt", H.GOLD_PALE, H.GOLD),
    ])
    H.add_causal_strip(slide, "READY 之後保持 `.venv` 啟用 → 只改目前 Lab 的 marked block → 執行 case",
                       "READY next step", fill=H.PURPLE_PALE, line=H.PURPLE, color=H.NAVY)
    H.add_field_band(slide, "status", "環境狀態", "verify receipt", "表示環境與課程契約是否可進入案例",
                     "狀態值", "READY 不等於 Lab 通過，也不等於已產生 result.json", H.PURPLE)
    notes(slide, "環境完成需要三個一致訊號。第一，啟用後的 Python 顯示 3.11.x；第二，執行器路徑位於本專案的 `.venv`，表示不是系統環境；第三，verify 輸出 status 為 READY，並建立環境收據。READY 只說明版本、情境與策略契約可以執行，不是 Lab 結果。接著保持 `.venv` 啟用，只修改目前 Lab 的 marked block，再執行指定 case 產生 result 與 replay。")


def lab_locate(prs, lab, title, block, variable, purpose, fixed):
    slide = H.prepare_slide(prs, f"{lab}-01｜先找到唯一要改的地方")
    title_lead(slide, purpose)
    add_three_cards(slide, [
        ("檔案", "`student_policy.py`\n只修改這一份策略檔", H.BLUE_PALE, H.BLUE),
        ("標記區塊", f"`{block}`\n只開目前 Lab", H.TEAL_PALE, H.TEAL),
        ("唯一變數", f"`{variable}`\n其他值與檔案不動", H.GOLD_PALE, H.GOLD),
    ])
    H.add_causal_strip(slide, f"保持不變：{fixed}、runner、scenario、schemas 與其他 Lab 區塊",
                       f"{lab} locate boundary", fill=H.PURPLE_PALE, line=H.PURPLE, color=H.NAVY)
    H.add_field_band(slide, variable, title, "student_policy.py marked block", "控制本 Lab 的單一策略差異",
                     "設定值", "只有此值可改，才能把結果差異歸因於本次修改", H.PURPLE)
    notes(slide, f"本頁只做定位。實際檔案是 student_policy.py，區塊是 {block}，唯一要比較的變數是 {variable}。{fixed}、runner、scenario、schemas 與其他 Lab 區塊保持不變。先建立這個可編輯邊界，再看下一頁的放大修改，不需要閱讀整個函式。")


def before_after(prs, lab, variable, before, after, before_meaning, after_meaning, takeaway):
    slide = H.prepare_slide(prs, f"{lab}-02｜只改這一行：{before} → {after}")
    title_lead(slide, "兩邊程式結構完全相同；只比較反白的一個值")
    H.add_card(slide, 0.94, 1.62, 5.54, 2.58, "修改前｜Before",
               f"`{variable} = {before}`\n\n{before_meaning}", H.BLUE_PALE, H.BLUE,
               title_size=23, body_size=27, name=f"{lab} before")
    H.add_card(slide, 6.84, 1.62, 5.54, 2.58, "修改後｜After",
               f"`{variable} = {after}`\n\n{after_meaning}", H.GOLD_PALE, H.GOLD,
               title_size=23, body_size=27, name=f"{lab} after")
    H.add_causal_strip(slide, takeaway, f"{lab} one line diff",
                       fill=H.RED_PALE, line=H.RED, color=H.RED)
    H.add_field_band(slide, variable, "本次修改值", "目前 Lab 的 marked block", "改變策略判斷",
                     "依變數定義", "修改後先做語法檢查，再執行指定 case", H.RED)
    notes(slide, f"這一頁故意不顯示整段函式。修改前是 {variable} = {before}，修改後是 {variable} = {after}。{before_meaning}；{after_meaning}。除這一行外，檔案、縮排、函式、其他常數與其他 Lab 區塊全部保持不變。完成後先儲存，再依後續頁面的 exact command 執行。")


def lab_a_mechanism_run(prs):
    slide = H.prepare_slide(prs, "A-03｜為什麼 SLEEP → WAIT，接著怎麼執行")
    title_lead(slide, "WAIT 可能減少再次喚醒，卻會累積清醒閒置能量；必須同時看服務與焦耳")
    add_four_cards(slide, [
        ("SLEEP", "空檔進入低功耗\n之後傳送要 WAKE", H.BLUE_PALE, H.BLUE),
        ("WAIT", "空檔保持清醒\nawake-idle 持續累積", H.TEAL_PALE, H.TEAL),
        ("封包後果", "動作時機改變\nattempt／retry／delivery", H.GOLD_PALE, H.GOLD),
        ("最後判斷", "先 service\n再比較 endpoint J", H.PURPLE_PALE, H.PURPLE),
    ])
    H.add_code_box(slide, 0.94, 4.55, 11.42, 0.78,
                   "`bash course.sh run --lab A --case baseline`   →   `bash course.sh run --lab A --case candidate --freeze`   →   `bash course.sh run --lab A --case hidden`",
                   fill="F6F7FB", line=H.BLUE, size=17, name="Lab A commands")
    H.add_field_band(slide, "REST_DURING_GAP", "空檔動作", "Lab A marked block", "決定送出間隔中的休息方式",
                     "action", "Windows 將 bash course.sh 改為 course.cmd，其餘參數不變", H.TEAL, y=5.47, h=0.78)
    notes(slide, "Lab A 比較空檔休息與清醒等待。SLEEP 讓端點進入低功耗狀態，但下一次傳送可能多一次 WAKE；WAIT 保持清醒，可能更快反應，卻持續累積 awake-idle 能量。依序執行 baseline、candidate --freeze 與 hidden。Linux、macOS 與 WSL 使用 bash course.sh；Windows 使用 course.cmd 並保留相同的 lab、case 與 freeze 參數。")


def lab_result(prs, lab, title, expected, field, chinese, mechanism):
    slide = H.prepare_slide(prs, f"{lab}-04｜上傳後怎麼解讀")
    title_lead(slide, "每次只上傳該 run 的 result.json；先確認服務，再解釋原因與能量")
    H.add_card(slide, 0.94, 1.62, 5.52, 1.98, "網站操作",
               "① 依 stdout 的 `result_path` 找檔案\n② `/course` 選擇 `result.json`\n③ 保留同目錄 `endpoint-replay.json`\n④ 核對 scenario ID 與 run ID",
               H.BLUE_PALE, H.BLUE, title_size=21, body_size=18, name=f"{lab} upload")
    H.add_card(slide, 6.84, 1.62, 5.52, 1.98, "如何回答結果",
               f"{expected}\n\n若沒有看到預期事件差異，就不能把最終數值歸因於這次修改。",
               H.GOLD_PALE, H.GOLD, title_size=21, body_size=18, name=f"{lab} interpret")
    add_result_order(slide, 3.92)
    H.add_field_band(slide, field, chinese, "result／replay artifact", "支持本 Lab 的機制判讀",
                     "依欄位定義", mechanism, H.PURPLE)
    notes(slide, f"每次案例完成後，依 stdout 的 result_path 找到該次 result.json，到 /course 逐次匯入，並保留同目錄的 endpoint-replay.json。核對 scenario ID 與 run ID 後，先判斷 service，再看封包的送達、重試或逾期，再看狀態分布，最後才讀 endpoint J。{expected}。若中間事件沒有預期差異，即使最後焦耳不同，也不能直接歸因於本次修改。")


def lab_b_mechanism_run(prs):
    slide = H.prepare_slide(prs, "B-03｜為什麼 2 → 1，接著怎麼執行")
    title_lead(slide, "需要的穩定觀察變少，可能更早進入可送模式，也可能對短暫品質變化過度反應")
    H.add_card(slide, 0.94, 1.62, 5.54, 2.16, "原本需要兩次穩定觀察",
               "第 1 步看到品質達標：先等待\n第 2 步仍達標：才讓 quality_ready 成立",
               H.BLUE_PALE, H.BLUE, title_size=22, body_size=20, name="B two steps")
    H.add_card(slide, 6.84, 1.62, 5.54, 2.16, "改成一次穩定觀察",
               "第 1 步看到品質達標：立刻讓 quality_ready 成立\nMODE_CHANGE 可能提早出現",
               H.TEAL_PALE, H.TEAL, title_size=22, body_size=20, name="B one step")
    H.add_code_box(slide, 0.94, 4.08, 11.42, 1.08,
                   "`bash course.sh run --lab B --case trace-a-baseline`\n`bash course.sh run --lab B --case trace-a-candidate --freeze`   →   `bash course.sh run --lab B --case trace-b`",
                   fill="F6F7FB", line=H.BLUE, size=17, name="Lab B commands")
    H.add_field_band(slide, "STABLE_STEPS", "穩定步數", "Lab B marked block", "決定品質達標要連續維持幾個 step",
                     "steps", "Windows 將 bash course.sh 改為 course.cmd；Trace B 保持 frozen policy", H.TEAL, y=5.47, h=0.78)
    notes(slide, "STABLE_STEPS 從 2 改為 1，表示品質達標後不必再等待第二個穩定觀察，quality_ready 與 MODE_CHANGE 可能提早出現。好處可能是更早掌握服務窗口，風險是對短暫品質變化過度反應。依序執行 Trace A baseline、Trace A candidate --freeze 與 withheld Trace B。Windows 只替換 launcher 為 course.cmd，其餘參數不變。")


def lab_c_value(prs, seq, value, label, meaning, color_pair):
    slide = H.prepare_slide(prs, f"C-{seq}｜{label}：URGENT_MARGIN_S = {value}")
    title_lead(slide, "這一頁只讀一個值；BATCH_SIZE = 3 與其他程式全部不變")
    fill, line = color_pair
    H.add_card(slide, 2.06, 1.62, 9.22, 2.52, label,
               f"`URGENT_MARGIN_S = {value}`\n\n{meaning}", fill, line,
               title_size=24, body_size=29, name=f"Lab C value {value}")
    H.add_causal_strip(slide, f"只改急件提前量；先預測 SEND_URGENT 時機，再執行對應 case",
                       f"Lab C {value} causal", fill=H.PURPLE_PALE, line=H.PURPLE, color=H.NAVY)
    H.add_field_band(slide, "URGENT_MARGIN_S", "急件提前量", "Lab C marked block", "決定剩餘期限多早進入急件分支",
                     "秒", "數值越大代表越早嘗試急件傳送；是否有利仍由服務與事件證據判斷", line)
    notes(slide, f"Lab C 的唯一比較值是 URGENT_MARGIN_S。此頁設定為 {value} 秒。{meaning}。BATCH_SIZE 固定為 3，Lab A、Lab B、runner、scenario、schemas 與其他程式行都不變。不要把三個值擠在同一段程式碼中；每次只設定一個值、儲存並執行對應 case。")


def lab_c_mechanism(prs):
    slide = H.prepare_slide(prs, "C-04｜為什麼 5 秒可能太晚，30 秒可能更早")
    title_lead(slide, "急件門檻控制的是『何時放棄等待批次、改成優先送出』")
    add_three_cards(slide, [
        ("20 秒｜基準", "在期限剩 20 秒時\n允許 SEND_URGENT", H.BLUE_PALE, H.BLUE),
        ("5 秒｜候選", "等待更久才急送\n可能錯過窗口或期限", H.RED_PALE, H.RED),
        ("30 秒｜修訂", "更早進入急件分支\n可能恢復服務，也可能增加能量", H.TEAL_PALE, H.TEAL),
    ])
    H.add_causal_strip(slide, "門檻 → SEND_URGENT 時機 → wake／TX／delivery／deadline → service → endpoint J",
                       "Lab C mechanism", fill=H.GOLD_PALE, line=H.GOLD, color=H.NAVY)
    H.add_field_band(slide, "deadline_pass", "期限是否通過", "result service summary", "確認必要封包是否在期限內完成",
                     "布林值", "先確認期限與服務，再比較端點焦耳；較低焦耳但失敗不是改善", H.RED)
    notes(slide, "URGENT_MARGIN_S 控制的是剩餘期限小於多少秒時，放棄等待批次並使用 SEND_URGENT。5 秒代表等待更久，可能錯過服務窗口或期限；30 秒代表更早急送，可能恢復服務，也可能增加 wake、TX 或處理能量。20 與 30 在固定主要案例中也可能產生相同結果，這種 null contrast 要如實保留，不能硬說 30 一定更好。")


def lab_c_run(prs):
    slide = H.prepare_slide(prs, "C-05｜四次執行的固定順序")
    title_lead(slide, "每次先設定正確值，再執行對應 case；revision 才使用 freeze")
    add_four_cards(slide, [
        ("① baseline", "值為 20\n`--lab C --case baseline`", H.BLUE_PALE, H.BLUE),
        ("② candidate", "改成 5\n`--lab C --case candidate`", H.RED_PALE, H.RED),
        ("③ revision", "改成 30 並凍結\n`--lab C --case revision --freeze`", H.TEAL_PALE, H.TEAL),
        ("④ surprise", "保持 frozen 30\n`--lab C --case surprise`", H.GOLD_PALE, H.GOLD),
    ])
    H.add_code_box(slide, 0.94, 4.55, 11.42, 0.78,
                   "Linux／macOS／WSL：`bash course.sh run ...`     Windows：`course.cmd run ...`；每次保存 stdout 的 `result_path`",
                   fill="F6F7FB", line=H.BLUE, size=18, name="Lab C launcher rule")
    H.add_field_band(slide, "freeze", "策略凍結", "revision 執行選項", "保存要帶入 surprise 的策略身分",
                     "狀態", "freeze 缺失或身分不一致時停止 surprise，不重新修改策略", H.PURPLE, y=5.47, h=0.78)
    notes(slide, "Lab C 依序執行 baseline、candidate、revision --freeze 與 surprise。baseline 使用 20，candidate 使用 5，revision 使用 30 並凍結，surprise 保持同一份 frozen 30，不再修改。Linux、macOS 與 WSL 使用 bash course.sh run；Windows 使用 course.cmd run。每次保存命令輸出的 result_path，不能用上一個案例的路徑。")


def new_presentation():
    prs = Presentation(str(H.TEMPLATE))
    H.remove_all_slides(prs)
    return prs


def save_presentation(prs, output, subject):
    prs.core_properties.title = "LoRaEnergySim + LEO｜核心課堂版"
    prs.core_properties.subject = subject
    prs.core_properties.author = "C-120 classroom core lane"
    prs.core_properties.comments = "Editable review deck built only from educate slideLayout2.xml"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    H.overlay_template_parts(output)


def build() -> None:
    prelab = new_presentation()

    intro_lora_value(prelab)

    for builder in [
        H.slide_019,
        env_sequence, env_windows_python, env_windows_venv,
        env_linux_python, env_linux_venv,
        env_macos_python, env_macos_venv, env_ready,
        H.slide_021c, H.slide_022, H.slide_023, H.slide_024a, H.slide_024b,
        H.slide_027a, H.slide_027b,
    ]:
        builder(prelab)
    save_presentation(prelab, PRELAB_OUTPUT, "LoRaEnergySim value, platform setup, policy boundary, and baseline handoff")

    labs = new_presentation()

    lab_locate(labs, "B", "穩定步數", "lab-b-enter-exit-hold", "STABLE_STEPS",
               "Lab B 比較品質需要連續穩定幾步，觀察進入可送模式的時機", "ENTER_QUALITY = 2、EXIT_QUALITY = 1")
    before_after(labs, "B", "STABLE_STEPS", "2", "1",
                 "需要連續兩個穩定觀察", "一個穩定觀察就成立",
                 "唯一修改：2 → 1；進入與退出品質門檻都不變")
    lab_b_mechanism_run(labs)
    lab_result(labs, "B", "穩定步數", "預期 quality_ready 或 MODE_CHANGE 可能提早；再看 attempt、retry、delivery、service 與端點焦耳。",
               "MODE_CHANGE", "模式變更事件", "Trace A 的改善必須再由 frozen policy 的 Trace B 檢查")

    lab_locate(labs, "C", "急件提前量", "lab-c-batch-urgent", "URGENT_MARGIN_S",
               "Lab C 比較急件要提前多久送出，觀察期限、服務與能量代價", "BATCH_SIZE = 3")
    lab_c_value(labs, "02", "20", "基準值｜Baseline", "期限剩 20 秒時進入急件分支", (H.BLUE_PALE, H.BLUE))
    lab_c_value(labs, "03", "5", "候選值｜Candidate", "等到只剩 5 秒才急送，可能太晚", (H.RED_PALE, H.RED))
    lab_c_value(labs, "04", "30", "修訂值｜Revision", "期限剩 30 秒就急送，較早但可能增加成本", (H.TEAL_PALE, H.TEAL))
    lab_c_mechanism(labs)
    lab_c_run(labs)
    lab_result(labs, "C", "急件提前量", "先確認 deadline 與 service；若 20 與 30 結果相同，保留『此案例無差異』，不要硬造改善。",
               "deadline_pass", "期限是否通過", "5 秒失敗或 30 秒恢復都要由本次 result／replay 證據支持")
    save_presentation(labs, LABS_FALLBACK_OUTPUT, "Lab B and Lab C classroom fallback pages")
    print(json.dumps({
        "prelab_output": str(PRELAB_OUTPUT),
        "prelab_slides": len(prelab.slides),
        "labs_fallback_output": str(LABS_FALLBACK_OUTPUT),
        "labs_fallback_slides": len(labs.slides),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
