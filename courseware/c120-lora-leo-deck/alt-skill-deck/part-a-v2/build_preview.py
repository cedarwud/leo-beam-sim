#!/usr/bin/env python3
"""Build a sparse, visual-first Part A preview from educate slide layout 2.

The preview is intentionally rebuilt from the native template with python-pptx.
It does not copy or repair the earlier Part A package.  Every authored slide is
created from the template's second slide layout, while the background, logo,
divider, and footer remain inherited from the template.
"""

from __future__ import annotations

import json
import copy
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ALT_ROOT = ROOT.parent
LATEST_ROOT = ALT_ROOT / "latest"
LATEST_ROOT.mkdir(parents=True, exist_ok=True)
TEMPLATE = Path("/home/u24/pptx-wrap/assets/templates/educate.pptx")
OUTPUT = LATEST_ROOT / "LoRaEnergySim-LEO-ALT-PART-A-V2-P001-P008-REVIEW.pptx"
QA_OUTPUT = ROOT / "qa" / "preview-qa.json"
MANIFEST_OUTPUT = ROOT / "slides.json"

NAVY = "25356B"
INK = "25324A"
MUTED = "65758B"
BLUE = "4967C2"
BLUE_PALE = "ECF1FE"
TEAL = "007374"
TEAL_PALE = "E8F7F4"
PURPLE = "660066"
PURPLE_PALE = "F4EDF7"
GOLD = "846826"
GOLD_PALE = "FDF9E2"
RED = "992A2A"
RED_PALE = "FDEFEF"
WHITE = "FFFFFF"
LINE = "D7DEEA"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_fonts(run) -> None:
    """Set script-specific typefaces on one DrawingML run."""
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("latin", "Times New Roman"), ("ea", "標楷體"), ("cs", "Times New Roman")):
        child = rpr.find(qn(f"a:{tag}"))
        if child is None:
            child = OxmlElement(f"a:{tag}")
            rpr.append(child)
        child.set("typeface", face)


def add_marked_runs(paragraph, text: str, size: float, color: str,
                    bold: bool = False) -> None:
    """Backtick spans are editable italic variables or field names."""
    pieces = re.split(r"(`[^`]+`)", text)
    for piece in pieces:
        if not piece:
            continue
        italic = piece.startswith("`") and piece.endswith("`")
        value = piece[1:-1] if italic else piece
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
        set_run_fonts(run)


def write_text(shape, text: str, size: float = 24, color: str = INK,
               bold: bool = False, align=PP_ALIGN.LEFT,
               valign=MSO_ANCHOR.MIDDLE, margins=(0.10, 0.05, 0.10, 0.05),
               line_spacing: float = 1.0) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        add_marked_runs(p, line, size, color, bold)


def add_text(slide, x, y, w, h, text, size=24, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margins=(0.02, 0.01, 0.02, 0.01), line_spacing=1.0,
             name="Text"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    write_text(shape, text, size, color, bold, align, valign, margins, line_spacing)
    return shape


def add_box(slide, x, y, w, h, fill, line=LINE, radius=True,
            name="Box", line_width=1.2):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_label_box(slide, x, y, w, h, text, fill, line, size=22,
                  color=INK, bold=False, align=PP_ALIGN.CENTER,
                  name="Label"):
    shape = add_box(slide, x, y, w, h, fill, line, name=name)
    write_text(shape, text, size, color, bold, align, MSO_ANCHOR.MIDDLE,
               margins=(0.12, 0.05, 0.12, 0.05), line_spacing=0.95)
    return shape


def add_chevron(slide, x, y, w=0.42, h=0.52, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_field_band(slide, field: str, chinese: str, source: str,
                   purpose: str, unit: str, interpretation: str) -> None:
    band = add_box(slide, 0.88, 5.42, 11.55, 0.90, "F7F8FC", BLUE,
                   name=f"Field contract {field}", line_width=1.0)
    text = (
        f"`{field}`（{chinese}）｜來源：{source}｜作用：{purpose}\n"
        f"單位：{unit}｜判讀：{interpretation}"
    )
    write_text(band, text, 17, NAVY, False, PP_ALIGN.LEFT,
               MSO_ANCHOR.MIDDLE, margins=(0.18, 0.05, 0.18, 0.05),
               line_spacing=0.92)


def delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def prepare_slide(prs: Presentation, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    if title_shape is None:
        raise RuntimeError("slideLayout2 title placeholder missing")
    for placeholder in list(slide.placeholders):
        # python-pptx may return a fresh wrapper for the same placeholder;
        # compare the underlying XML element rather than object identity so
        # the authored title survives the cleanup pass.
        if (placeholder._element is title_shape._element
                or placeholder.placeholder_format.type == PP_PLACEHOLDER.TITLE):
            continue
        if placeholder.placeholder_format.type not in {
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.SLIDE_NUMBER,
        }:
            delete_shape(placeholder)
    title_shape.left = Inches(0.718057)
    title_shape.top = Inches(0.204514)
    title_shape.width = Inches(10.34861)
    title_shape.height = Inches(0.525)
    title_shape.name = "Native layout2 title"
    write_text(title_shape, title, 28, NAVY, True, PP_ALIGN.LEFT,
               MSO_ANCHOR.MIDDLE, margins=(0.02, 0.0, 0.02, 0.0),
               line_spacing=1.0)
    return slide


def remove_all_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        rel_id = sld_id.rId
        prs.part.drop_rel(rel_id)
        sld_id_lst.remove(sld_id)


def overlay_template_parts(path: Path) -> None:
    """Restore template-owned inherited parts after python-pptx saves.

    python-pptx can rewrite an inherited master/layout/theme even when no
    authored background is present.  The content slides and their
    relationships remain from the build; the native educate parts are copied
    byte-for-byte from the requested template.
    """
    owned_prefixes = (
        "ppt/slideLayouts/",
        "ppt/slideMasters/",
        "ppt/theme/",
        "ppt/media/",
        "ppt/notesMasters/",
    )
    with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(path, "r") as old:
        data = {info.filename: old.read(info.filename) for info in old.infolist()}
        infos = {info.filename: copy.copy(info) for info in old.infolist()}
        for info in source.infolist():
            if info.filename.startswith(owned_prefixes):
                data[info.filename] = source.read(info.filename)
                infos[info.filename] = copy.copy(info)
    overlay = path.with_suffix(".overlay.pptx")
    with zipfile.ZipFile(overlay, "w", zipfile.ZIP_DEFLATED) as out:
        for name, payload in data.items():
            out.writestr(infos[name], payload)
    overlay.replace(path)


def emit_manifest(path: Path) -> None:
    """Emit an ordered, merge-friendly source record for the parent lane."""
    # The expanded opening redraws the current course narrative.  Donor entries
    # are topic-level references only; no donor bytes or evidence are embedded.
    donor_sources = [
        {
            "donor": "donor-insertion-map.md",
            "slides": "e2 2, 4-6",
            "mapping": "changing opportunity and visible service window; redraw reference only",
        },
        {
            "donor": "ADR-004 / C120 SDD / README.zh-TW.md",
            "slides": "authority",
            "mapping": "plain-language LoRaEnergySim, endpoint, gateway, packet, and simulator definition",
        },
        {
            "donor": "ADR-004 / THIRD_PARTY_NOTICES.md",
            "slides": "authority",
            "mapping": "source-backed upstream capabilities and explicit upstream-execution boundary",
        },
        {
            "donor": "README.zh-TW.md / lora_energy_lab source",
            "slides": "authority",
            "mapping": "course-owned scenario, bounded policy, runner, result/replay, and Leo import seam",
        },
        {
            "donor": "ADR-004 / C120 SDD",
            "slides": "authority",
            "mapping": "competition relevance, repeatable causal evidence, and claim ceiling",
        },
        {
            "donor": "README.zh-TW.md / INSTRUCTOR-REHEARSAL-RUNBOOK.md",
            "slides": "authority",
            "mapping": "run-only result/replay artifact creation; setup and verify are not run artifacts",
        },
        {
            "donor": "INSTRUCTOR-REHEARSAL-RUNBOOK.md / teaching-rewrite",
            "slides": "authority",
            "mapping": "local runner simulation versus /course result import and same-directory replay pairing",
        },
        {
            "donor": "INSTRUCTOR-REHEARSAL-RUNBOOK.md / C120 SDD",
            "slides": "authority",
            "mapping": "baseline/candidate/withheld import order, service-to-endpoint-J reading, and freeze boundary",
        },
        {
            "donor": "donor-insertion-map.md",
            "slides": "e2 14, 16/17, 19-21",
            "mapping": "state, packet, service, and endpoint-energy causal chain; redraw reference only",
        },
        None,
        None,
        {
            "donor": "donor-insertion-map.md",
            "slides": "e2 20, 36-37",
            "mapping": "same-observation A/B comparison and attribution gate; redraw reference only",
        },
        {
            "donor": "donor-insertion-map.md",
            "slides": "e2 14, 16/17, 19-21",
            "mapping": "endpoint versus system/canonical evidence boundary; redraw reference only",
        },
        {
            "donor": "donor-insertion-map.md",
            "slides": "e2 2, 4, 19",
            "mapping": "changing-window transfer to other IoT contexts; redraw reference only",
        },
        {
            "donor": "donor-insertion-map.md",
            "slides": "e2 96-97, 113",
            "mapping": "prediction, evidence, and claim classification; redraw reference only",
        },
        {
            "donor": "README.zh-TW.md / LAB-EXPLANATION-ACCEPTANCE.md",
            "slides": "authority",
            "mapping": "Lab A/B/C controlled edits and observable evidence targets",
        },
    ]
    source_ids = [
        "P001", "P001A", "P001B", "P001C", "P001D",
        "P008B", "P008C", "P008D",
        "P002", "P003", "P004", "P005", "P006", "P007", "P008", "P008A",
    ]
    prs = Presentation(str(path))
    pages = []
    for order, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        body = []
        for shape in slide.shapes:
            if shape.is_placeholder or not getattr(shape, "text", "").strip():
                continue
            if shape.text.strip() == title.strip():
                continue
            body.append(shape.text)
        pages.append({
            "order": order,
            "source-page": source_ids[order - 1],
            "title": title,
            "body": body,
            "layout": "slideLayout2.xml",
            "notes": slide.notes_slide.notes_text_frame.text,
            "evidence-status": "待補：fresh runner receipt/current evidence not embedded in this preview",
            "donor-source": donor_sources[order - 1],
        })
    manifest = {
        "deck": "LoRaEnergySim-LEO-ALT-PART-A-V2-P001-P008-EXPANDED",
        "template": str(TEMPLATE),
        "shell": "source slide 2 / slideLayout2.xml only",
        "background-author-fill": "unset",
        "font-contract": {"cjk": "標楷體", "latin": "Times New Roman"},
        "title-pt": 28,
        "body-baseline-pt": 24,
        "minimum-body-pt": 16,
        "current-evidence": "待補：deck contains no fabricated screenshot, KPI, or run result",
        "slides": pages,
    }
    MANIFEST_OUTPUT.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def slide_1(prs):
    slide = prepare_slide(prs, "有限電池 × 變動服務窗口")
    endpoint = add_box(slide, 0.92, 1.38, 2.75, 2.86, BLUE_PALE, BLUE,
                       name="Finite battery endpoint")
    write_text(endpoint, "有限電池\n端點", 30, NAVY, True, PP_ALIGN.CENTER,
               MSO_ANCHOR.MIDDLE, margins=(0.18, 0.18, 0.18, 0.18))
    for idx, color in enumerate((RED, GOLD, TEAL)):
        cell = add_box(slide, 1.34 + idx * 0.60, 3.55, 0.44, 0.22, color, None,
                       radius=False, name=f"Battery cell {idx + 1}")
        cell.line.fill.background()
    add_text(slide, 4.03, 1.62, 1.70, 0.42, "資料封包", 22, NAVY, True,
             PP_ALIGN.CENTER, name="Packet label")
    for idx in range(3):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.12 + idx * 0.52), Inches(2.33), Inches(0.32), Inches(0.32))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(TEAL)
        dot.line.fill.background()
    add_chevron(slide, 5.55, 2.23, 0.50, 0.50, TEAL)
    add_text(slide, 3.95, 2.94, 2.12, 0.58, "產生 → 佇列 → 送出", 18,
             MUTED, False, PP_ALIGN.CENTER, name="Packet path")
    add_box(slide, 6.42, 1.34, 5.82, 2.94, TEAL_PALE, TEAL,
            name="Changing service window")
    add_text(slide, 6.78, 1.58, 5.10, 0.52, "會開、會關、品質會變", 27,
             TEAL, True, PP_ALIGN.CENTER, name="Window headline")
    widths = [1.36, 1.04, 1.48, 1.16]
    fills = [TEAL, "CBD5E1", BLUE, GOLD]
    labels = ["開啟", "關閉", "開啟", "縮短"]
    cursor = 6.88
    for idx, (width, fill, label) in enumerate(zip(widths, fills, labels)):
        segment = add_box(slide, cursor, 2.52, width, 0.72, fill, None,
                          radius=False, name=f"Window segment {idx + 1}")
        write_text(segment, label, 18, WHITE if fill != "CBD5E1" else INK,
                   True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
                   margins=(0.02, 0.01, 0.02, 0.01))
        cursor += width + 0.07
    add_label_box(slide, 1.25, 4.56, 10.70, 0.58,
                  "先看窗口，再看動作；最後讀服務與端點能量",
                  GOLD_PALE, GOLD, 22, NAVY, True, name="Reading order")
    add_field_band(slide, "scenario", "情境定義", "course package",
                   "固定端點、窗口與資料規則", "identifier",
                   "相同輸入可重跑")
    slide.notes_slide.notes_text_frame.text = (
        "本頁建立課程的固定情境。環境感測端點以有限電池運作，服務窗口依固定軌跡開啟、關閉或縮短。"
        "情境定義 scenario 來自 course package，作用是固定端點、窗口與資料規則，單位為 identifier；"
        "相同輸入應產生可重跑的比較。後續判讀依序確認窗口、策略動作、服務結果與端點能量。"
    )


def slide_1a(prs):
    slide = prepare_slide(prs, "LoRaEnergySim 是什麼")
    add_text(slide, 1.00, 1.02, 11.30, 0.50,
             "把端點的一次選擇，放進可重複的封包與能量模擬",
             23, NAVY, True, PP_ALIGN.CENTER, name="Simulator lead")

    endpoint = add_box(slide, 0.92, 1.82, 2.70, 2.38, BLUE_PALE, BLUE,
                       name="Endpoint concept")
    add_text(slide, 1.14, 2.10, 2.26, 0.54, "端點", 27, NAVY, True,
             PP_ALIGN.CENTER, name="Endpoint heading")
    add_text(slide, 1.18, 2.86, 2.18, 0.98,
             "感測器＋有限電池\n產生資料、選擇動作",
             21, INK, False, PP_ALIGN.CENTER, name="Endpoint body")

    add_chevron(slide, 3.82, 2.70, 0.46, 0.54, TEAL)
    packet = add_box(slide, 4.42, 1.82, 3.78, 2.38, GOLD_PALE, GOLD,
                     name="Packet concept")
    add_text(slide, 4.72, 2.10, 3.18, 0.54, "資料封包", 27, GOLD, True,
             PP_ALIGN.CENTER, name="Packet heading")
    add_text(slide, 4.76, 2.86, 3.10, 0.98,
             "一份待傳資料\n產生 → 佇列 → 嘗試傳送",
             21, INK, False, PP_ALIGN.CENTER, name="Packet body")

    add_chevron(slide, 8.42, 2.70, 0.46, 0.54, TEAL)
    gateway = add_box(slide, 9.02, 1.82, 3.34, 2.38, TEAL_PALE, TEAL,
                      name="Gateway concept")
    add_text(slide, 9.32, 2.10, 2.74, 0.54, "閘道", 27, TEAL, True,
             PP_ALIGN.CENTER, name="Gateway heading")
    add_text(slide, 9.34, 2.86, 2.70, 0.98,
             "接收資料封包\n留下送達、碰撞與重傳事件",
             21, INK, False, PP_ALIGN.CENTER, name="Gateway body")

    add_label_box(slide, 1.24, 4.56, 10.70, 0.58,
                  "它回答：同一份工作，在等待、休眠或傳送的不同決策下，服務與端點能量（J）如何一起變？",
                  GOLD_PALE, GOLD, 20, NAVY, True, name="Simulator question")
    add_field_band(slide, "simulator", "模擬器", "course package 參照 LoRaEnergySim",
                   "在固定時鐘重播端點、封包與能量假設", "一次固定執行",
                   "產生可比較證據，不是實測")
    slide.notes_slide.notes_text_frame.text = (
        "本頁從零定義模擬器。端點是有感測器與有限電池的資料裝置；資料封包是一次要送出的資料；"
        "閘道是接收端點封包並留下事件的接收節點。LoRaEnergySim 在本課被當作能源決策的模型參照，"
        "把端點狀態、封包事件與能量後果放進可重複的固定執行。這裡的問題不是猜一個漂亮的總分，"
        "而是比較同一份工作在等待、休眠或傳送的不同動作下，服務與端點焦耳是否同時改變。"
        "模擬器 simulator 的來源是 course package 對 LoRaEnergySim 的參照，作用是在固定時鐘重播課程假設，"
        "單位是一次固定執行；判讀時只代表可比較證據，不代表實測。"
    )


def slide_1b(prs):
    slide = prepare_slide(prs, "原始上游套件已做什麼")
    add_text(slide, 1.00, 1.02, 11.30, 0.50,
             "上游研究模型提供「端點狀態 → 封包事件 → 能量」的可重用能力",
             23, NAVY, True, PP_ALIGN.CENTER, name="Upstream lead")

    left = add_box(slide, 0.92, 1.82, 5.56, 2.80, TEAL_PALE, TEAL,
                   name="Upstream capabilities")
    add_text(slide, 1.22, 2.08, 4.96, 0.48, "上游已有的研究能力", 24, TEAL,
             True, PP_ALIGN.CENTER, name="Upstream heading")
    add_label_box(slide, 1.32, 2.84, 4.74, 0.60,
                  "休眠／處理／發射／接收狀態",
                  BLUE_PALE, BLUE, 21, NAVY, True, name="Upstream states")
    add_label_box(slide, 1.32, 3.62, 4.74, 0.60,
                  "封包、碰撞、重傳、端點能量",
                  GOLD_PALE, GOLD, 21, NAVY, True, name="Upstream packet energy")

    right = add_box(slide, 6.85, 1.82, 5.56, 2.80, RED_PALE, RED,
                    name="Course gap")
    add_text(slide, 7.15, 2.08, 4.96, 0.48, "上游沒有替本課固定的部分", 24, RED,
             True, PP_ALIGN.CENTER, name="Course gap heading")
    add_label_box(slide, 7.25, 2.84, 4.74, 0.60,
                  "固定情境、LEO 服務窗口、案例順序",
                  BLUE_PALE, BLUE, 19, NAVY, True, name="Missing scenario")
    add_label_box(slide, 7.25, 3.62, 4.74, 0.60,
                  "受控策略、結構化匯入、Leo 重播",
                  GOLD_PALE, GOLD, 19, NAVY, True, name="Missing course seam")

    add_text(slide, 1.10, 4.84, 11.12, 0.40,
             "因此：上游能力是參照；課程流程、案例與證據邊界由本課封裝。",
             21, NAVY, True, PP_ALIGN.CENTER, name="Upstream boundary")
    add_field_band(slide, "upstream_model", "上游模型能力", "GillesC/LoRaEnergySim@f854462c",
                   "提供端點狀態、封包與能量概念", "能力集合",
                   "僅是上游能力，不是本課 runner 執行")
    slide.notes_slide.notes_text_frame.text = (
        "本頁只列出有來源支持的上游能力。ADR-004 與第三方 provenance 指出，"
        "GillesC/LoRaEnergySim 的研究模型可連接端點的休眠、處理、發射、接收，以及封包、碰撞、重傳與能源結果。"
        "這些是可重用的模型概念，不是本課已執行上游 repository 的證明。"
        "固定 scenario、LEO 服務窗口、bounded policy、JSON 匯入與 Leo replay 是本課需要另外封裝的路徑，"
        "不能回寫成上游原生功能。上游模型能力 upstream_model 的來源是指定 commit，作用是提供狀態、封包與能量研究概念，"
        "單位是能力集合；判讀時只表示上游能力，不表示本課 runner 已執行上游。"
    )


def slide_1c(prs):
    slide = prepare_slide(prs, "課程封裝補了什麼")
    add_text(slide, 1.00, 1.02, 11.30, 0.50,
             "把研究模型包成一條可教、可驗、可匯入的固定路徑",
             23, NAVY, True, PP_ALIGN.CENTER, name="Wrapper lead")
    steps = [
        (0.92, "固定情境", "scenario.json\n同一 scenario_id", BLUE_PALE, BLUE),
        (3.94, "受控決策", "student_policy.py\n課程策略插槽；只改 marked block", GOLD_PALE, GOLD),
        (6.96, "本機執行", "runner\nbaseline → candidate → withheld", TEAL_PALE, TEAL),
        (9.98, "匯入證據", "result.json（結果）\nreplay（端點重播）\nLeo 網站匯入／Workbook", PURPLE_PALE, PURPLE),
    ]
    for idx, (x, heading, body, fill, line) in enumerate(steps):
        add_box(slide, x, 1.84, 2.46, 2.62, fill, line,
                name=f"Course wrapper step {idx + 1}")
        add_text(slide, x + 0.16, 2.04, 2.14, 0.50, heading, 23, line, True,
                 PP_ALIGN.CENTER, name=f"Wrapper heading {idx + 1}")
        add_text(slide, x + 0.18, 2.82, 2.10, 1.04, body, 19, INK, False,
                 PP_ALIGN.CENTER, name=f"Wrapper body {idx + 1}")
        if idx < len(steps) - 1:
            add_chevron(slide, x + 2.58, 2.86, 0.32, 0.46, line)
    add_label_box(slide, 1.20, 4.68, 10.92, 0.62,
                  "student_policy.py 是課程策略插槽；runner 不直接執行上游 LoRaEnergySim，上游僅能力參照。",
                  RED_PALE, RED, 18, NAVY, True, name="Wrapper boundary")
    add_field_band(slide, "runner", "課程執行器", "course package",
                   "讀固定 scenario、載入 bounded policy、輸出 result／replay",
                   "run artifact", "把一次修改連到可重播證據")
    slide.notes_slide.notes_text_frame.text = (
        "本頁把上游能力與課程封裝拆開。student_policy.py 是課程策略插槽，只開放 marked block；"
        "目前 runner 不直接執行上游 LoRaEnergySim，上游僅作能力來源與參照。第一步固定 scenario.json 與 scenario_id，"
        "第二步只開放 student_policy.py 的 marked block，第三步由本機 runner 依序產生 baseline、candidate 與 withheld，"
        "第四步把配對的 result.json 與 endpoint replay 匯入 Leo 網站，再留在 Workbook 中。"
        "這條路徑讓證據可驗證、可重播，也讓伺服器不接收或執行 Python。"
        "課程執行器 runner 的來源是 course package，作用是讀固定情境、載入受控策略並輸出結果與重播，"
        "單位是 run artifact；判讀時表示一次可追溯執行，不是上游能力或整個系統指標。"
    )


def slide_1d(prs):
    slide = prepare_slide(prs, "為什麼競賽要用它")
    add_text(slide, 1.00, 1.02, 11.30, 0.50,
             "先得到可重複的因果證據，再回到 LEO 網站解讀",
             23, NAVY, True, PP_ALIGN.CENTER, name="Competition lead")
    cards = [
        (0.92, "可重複", "同一情境、相同策略檔與固定種子\n重跑同一條事件路徑", BLUE_PALE, BLUE),
        (4.77, "可檢驗", "改一個受控區塊\n檢查封包／服務／狀態／J 是否改變", TEAL_PALE, TEAL),
        (8.62, "可轉移", "智慧農業、暖通空調、物流\n重訂窗口、期限與能量邊界", GOLD_PALE, GOLD),
    ]
    for idx, (x, heading, body, fill, line) in enumerate(cards):
        add_box(slide, x, 1.82, 3.54, 2.76, fill, line,
                name=f"Competition reason {idx + 1}")
        add_text(slide, x + 0.18, 2.08, 3.18, 0.52, heading, 25, line, True,
                 PP_ALIGN.CENTER, name=f"Competition heading {idx + 1}")
        add_text(slide, x + 0.24, 2.88, 3.06, 1.16, body, 20, INK, False,
                 PP_ALIGN.CENTER, name=f"Competition body {idx + 1}")
    add_label_box(slide, 1.10, 4.82, 11.12, 0.46,
                  "本套件不是即時資料、典範系統指標或整體 LEO 能量；它先建立可反駁的端點因果證據。",
                  RED_PALE, RED, 19, NAVY, True, name="Competition claim boundary")
    add_field_band(slide, "evidence", "證據", "result.json + endpoint replay",
                   "保存 policy → event → service／J 的路徑", "record",
                   "支持 bounded comparison，不擴大 claim")
    slide.notes_slide.notes_text_frame.text = (
        "本頁回答競賽主題的關聯。智慧節能與物聯網應用需要的是：一次端點決策能否連到封包服務、"
        "無線狀態時間與端點焦耳，並能在固定情境下重跑與反駁。課程 package 讓同一 scenario、策略 bytes 與 seed"
        "留下可比的事件路徑，再把 result 與 replay 匯入 Leo 讀取。這不是即時資料、典範系統指標或整體 LEO 能量，"
        "而是先把端點層的 policy → event → service／J 因果證據做實，再轉移到智慧農業、暖通空調或物流情境，"
        "並重新定義它們的窗口、期限與能量邊界。證據 evidence 的來源是 result.json 與 endpoint replay，"
        "作用是保存因果路徑，單位是 record；判讀時只能支持 bounded comparison。"
    )


def slide_8b(prs):
    slide = prepare_slide(prs, "只有 run case 才產生結果檔")
    add_text(slide, 1.00, 1.02, 11.30, 0.50,
             "setup／verify 只檢查環境；run case 才會寫入一次執行的產物",
             23, NAVY, True, PP_ALIGN.CENTER, name="Artifact gate lead")
    steps = [
        (0.92, "setup／verify\n環境檢查", "套件 package\n情境 scenario\n策略 policy\n不產生此次 run 產物", "F7F8FC", MUTED),
        (4.37, "run case\n執行案例", "本機 runner 依固定輸入模擬\n唯一產生結果的動作", TEAL_PALE, TEAL),
        (7.82, "artifacts/<run_id>/\n執行產物目錄", "result.json（結果）\nendpoint-replay.json（端點重播）", GOLD_PALE, GOLD),
    ]
    for idx, (x, heading, body, fill, line) in enumerate(steps):
        add_box(slide, x, 1.82, 3.05, 2.82, fill, line,
                name=f"Artifact gate step {idx + 1}")
        add_text(slide, x + 0.16, 2.08, 2.73, 0.68, heading, 22, line, True,
                 PP_ALIGN.CENTER, name=f"Artifact gate heading {idx + 1}")
        add_text(slide, x + 0.22, 3.02, 2.61, 1.04, body, 18, INK, False,
                 PP_ALIGN.CENTER, name=f"Artifact gate body {idx + 1}")
        if idx < len(steps) - 1:
            add_chevron(slide, x + 3.16, 2.92, 0.34, 0.48, line)
    add_label_box(slide, 1.04, 4.86, 11.20, 0.48,
                  "只有 run case 會產生 artifacts/<run_id>/result.json 與 endpoint-replay.json；setup／verify 不算。",
                  RED_PALE, RED, 19, NAVY, True, name="Artifact gate rule")
    add_field_band(slide, "run artifact", "執行產物", "本機 runner",
                   "保存一次執行的 result 與 endpoint replay", "同一 run 目錄",
                   "只有 run 產生；setup／verify 不算")
    slide.notes_slide.notes_text_frame.text = (
        "本頁先釐清檔案何時出現。setup 與 verify 只檢查 package、scenario 與 policy 的環境條件，"
        "不代表已完成一次案例執行，也不會產生此次 run 的結果檔。只有 run case 由本機 runner 依固定輸入做模擬，"
        "才會在 artifacts/<run_id>/ 寫入配對的 result.json 與 endpoint-replay.json。"
        "在 Linux／macOS 終端機執行時使用 Linux 指令；WSL 也使用 Linux 指令。"
        "執行產物 run artifact 的來源是本機 runner，作用是保存一次執行的結果與端點事件重播，"
        "單位是同一 run 目錄；判讀時只有 run 產生，setup／verify 不算。"
    )


def slide_8c(prs):
    slide = prepare_slide(prs, "課程頁 /course 只匯入結果檔")
    add_text(slide, 1.00, 1.02, 11.30, 0.50,
             "本機 runner 做模擬；網站做驗證、具象化、重播與工作簿",
             23, NAVY, True, PP_ALIGN.CENTER, name="Course import lead")
    local = add_box(slide, 0.92, 1.82, 4.00, 2.82, BLUE_PALE, BLUE,
                    name="Local runner boundary")
    add_text(slide, 1.18, 2.06, 3.48, 0.52,
             "本機 runner（模擬）", 24, BLUE, True, PP_ALIGN.CENTER,
             name="Local runner heading")
    add_text(slide, 1.22, 2.82, 3.40, 1.30,
             "Linux／macOS 終端機\nWSL 也使用 Linux 指令\n讀 student_policy.py\n（課程策略插槽）",
             18, INK, False, PP_ALIGN.CENTER, name="Local runner body")
    add_chevron(slide, 5.05, 2.88, 0.46, 0.56, TEAL)
    web = add_box(slide, 5.70, 1.82, 6.54, 2.82, TEAL_PALE, TEAL,
                  name="Course import boundary")
    add_text(slide, 5.98, 2.06, 5.98, 0.52,
             "/course（課程頁）只選 result.json", 23, TEAL, True,
             PP_ALIGN.CENTER, name="Course import heading")
    add_label_box(slide, 6.10, 2.82, 1.76, 0.58,
                  "驗證", BLUE_PALE, BLUE, 20, NAVY, True,
                  name="Course validation")
    add_label_box(slide, 8.06, 2.82, 1.76, 0.58,
                  "具象化", GOLD_PALE, GOLD, 20, NAVY, True,
                  name="Course materialization")
    add_label_box(slide, 10.02, 2.82, 1.76, 0.58,
                  "重播／工作簿", PURPLE_PALE, PURPLE, 18, NAVY, True,
                  name="Course replay workbook")
    add_text(slide, 6.06, 3.68, 5.82, 0.62,
             "endpoint-replay.json 留在同一目錄配對\n不選、不上傳 student_policy.py",
             18, INK, False, PP_ALIGN.CENTER, name="Course pairing rule")
    add_label_box(slide, 1.06, 4.86, 11.18, 0.48,
                  "policy 原始檔只在本機 runner 讀取；/course 接收 result.json，不接收 Python。",
                  RED_PALE, RED, 19, NAVY, True, name="Course upload boundary")
    add_field_band(slide, "course import", "課程匯入", "/course 控制項",
                   "驗證 result.json，建立重播與工作簿", "JSON pair",
                   "選結果檔；replay 留同目錄")
    slide.notes_slide.notes_text_frame.text = (
        "本頁把本機與網站的責任拆開。student_policy.py 是課程策略插槽，不是上游 LoRaEnergySim 的直接執行入口。"
        "Linux／macOS 終端機或使用 Linux 指令的 WSL 先由本機 runner 讀取這個插槽並做固定情境模擬，"
        "目前不直接執行上游 LoRaEnergySim；上游僅提供能力來源與參照。產出 result.json 與同一 run 目錄的 endpoint-replay.json。"
        "/course 課程頁只選 result.json；它不要求、不上傳也不執行 student_policy.py。"
        "網站的工作是驗證資料身份與結構、具象化結果、建立端點事件重播並寫入工作簿。"
        "課程匯入 course import 的來源是 /course 控制項，作用是驗證 result.json 並建立重播與工作簿，"
        "單位是 JSON pair；判讀時選結果檔，並讓 replay 留在同一個 run 目錄配對。"
    )


def slide_8d(prs):
    slide = prepare_slide(prs, "三次匯入，沿同一解讀順序")
    add_text(slide, 1.00, 1.02, 11.30, 0.50,
             "baseline → candidate → withheld：每次匯入一對檔案，再按同一順序解讀",
             23, NAVY, True, PP_ALIGN.CENTER, name="Import order lead")
    cases = [
        (0.92, "baseline\n基準", "result.json +\n同目錄 replay", "先看 service\n服務判定", BLUE_PALE, BLUE),
        (4.36, "candidate\n候選", "result.json +\n同目錄 replay", "再看 packet\n封包事件", TEAL_PALE, TEAL),
        (7.80, "withheld\n保留反例", "result.json +\n同目錄 replay", "再看 state → J\n狀態 → 端點焦耳", GOLD_PALE, GOLD),
    ]
    for idx, (x, heading, files, observe, fill, line) in enumerate(cases):
        add_box(slide, x, 1.82, 3.05, 2.82, fill, line,
                name=f"Import order case {idx + 1}")
        add_text(slide, x + 0.16, 2.04, 2.73, 0.58, heading, 22, line, True,
                 PP_ALIGN.CENTER, name=f"Import order heading {idx + 1}")
        add_text(slide, x + 0.28, 2.86, 2.49, 0.76, files, 18, INK, False,
                 PP_ALIGN.CENTER, name=f"Import order files {idx + 1}")
        add_text(slide, x + 0.22, 3.88, 2.61, 0.48, observe, 18, line, True,
                 PP_ALIGN.CENTER, name=f"Import order observe {idx + 1}")
        if idx < len(cases) - 1:
            add_chevron(slide, x + 3.16, 2.92, 0.34, 0.48, line)
    add_label_box(slide, 0.98, 4.84, 8.00, 0.50,
                  "service → packet → state → endpoint J（端點焦耳）",
                  "F7F8FC", PURPLE, 21, NAVY, True, name="Import reading order")
    freeze = add_box(slide, 9.20, 4.70, 3.04, 0.72, RED_PALE, RED,
                     name="Freeze boundary")
    write_text(freeze, "--freeze（凍結）\n額外 receipt／checkpoint；非一般上傳", 16, NAVY, True,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
               margins=(0.05, 0.01, 0.05, 0.01), line_spacing=0.82)
    add_field_band(slide, "import order", "匯入順序", "/course + 配對 artifacts",
                   "依序比較 baseline、candidate、withheld", "三次 JSON pair",
                   "service→packet→state→endpoint J；freeze 另存")
    slide.notes_slide.notes_text_frame.text = (
        "本頁把網站讀取順序固定下來。baseline（基準）、candidate（候選）與 withheld（保留反例）依序匯入，"
        "每一次都選該次 result.json，並保留同一個 run 目錄的 endpoint-replay.json 配對。"
        "判讀順序固定為 service → packet → state → endpoint J：先確認服務判定，再讀封包事件、狀態時間，"
        "最後才讀端點焦耳，避免只用一個總數字宣稱因果。"
        "--freeze 是受控案例額外的 receipt／checkpoint，用來保存策略與案例 lineage；它不是一般上傳，也不改變 /course 只選 result.json 的邊界。"
        "匯入順序 import order 的來源是 /course 與配對 artifacts，作用是依序比較三次 JSON pair，"
        "單位是三次 JSON pair；判讀時沿 service、packet、state、endpoint J，freeze 資料另存。"
    )


def slide_2(prs):
    slide = prepare_slide(prs, "一次動作，同時改變服務與能量")
    add_text(slide, 1.05, 1.05, 11.15, 0.52,
             "同一個 `action` 會沿不同路徑留下可觀察後果", 24,
             NAVY, True, PP_ALIGN.CENTER, name="Action lead")
    cards = [
        (0.92, BLUE_PALE, BLUE, "`SLEEP`｜低功耗休息", "休息功率降低\n喚醒延遲與能量增加"),
        (4.77, GOLD_PALE, GOLD, "`WAIT`｜清醒等待", "反應能力保留\n清醒閒置能量累積"),
        (8.62, TEAL_PALE, TEAL, "`SEND`｜送出封包", "服務機會增加\n處理與無線事件增加"),
    ]
    for idx, (x, fill, line, heading, body) in enumerate(cards):
        card = add_box(slide, x, 1.82, 3.54, 2.70, fill, line,
                       name=f"Action card {idx + 1}")
        add_text(slide, x + 0.18, 2.03, 3.18, 0.56, heading, 23, line, True,
                 PP_ALIGN.CENTER, name=f"Action heading {idx + 1}")
        add_text(slide, x + 0.26, 2.85, 3.02, 1.18, body, 20, INK, False,
                 PP_ALIGN.CENTER, name=f"Action body {idx + 1}")
    add_label_box(slide, 2.02, 4.73, 4.15, 0.48, "服務結果", "F7F8FC", BLUE,
                  20, NAVY, True, name="Service outcome")
    add_label_box(slide, 7.17, 4.73, 4.15, 0.48, "端點能量（J）", "F7F8FC", TEAL,
                  20, NAVY, True, name="Energy outcome")
    add_field_band(slide, "action", "策略動作", "policy API",
                   "每次決策的唯一輸出", "enum",
                   "觸發狀態與封包轉移")
    slide.notes_slide.notes_text_frame.text = (
        "本頁說明策略動作的三條主要路徑。SLEEP 降低休息功率，但必須支付喚醒延遲與喚醒能量；"
        "WAIT 保留反應能力，但會累積清醒閒置能量；SEND 類動作建立服務機會，也會增加處理、發射、接收或重試事件。"
        "策略動作 action 來自 policy API，作用是提供每次決策的唯一輸出，單位為 enum；判讀時以其觸發的狀態與封包轉移為準。"
    )


def slide_3(prs):
    slide = prepare_slide(prs, "先讀中間事件，再讀最終結果")
    add_text(slide, 1.08, 1.03, 11.10, 0.48,
             "能量或服務差異，必須由中間事件支持", 24, NAVY, True,
             PP_ALIGN.CENTER, name="Causal lead")
    labels = [
        ("服務窗口", BLUE_PALE, BLUE),
        ("策略動作", GOLD_PALE, GOLD),
        ("無線狀態\n封包事件", PURPLE_PALE, PURPLE),
        ("服務判定", TEAL_PALE, TEAL),
        ("端點能量", RED_PALE, RED),
    ]
    x = 0.84
    for idx, (label, fill, line) in enumerate(labels):
        add_label_box(slide, x, 1.92, 2.07, 1.46, label, fill, line,
                      21, NAVY, True, name=f"Causal node {idx + 1}")
        if idx < len(labels) - 1:
            add_chevron(slide, x + 2.15, 2.37, 0.38, 0.52, line)
        x += 2.49
    add_label_box(slide, 1.12, 3.77, 11.05, 0.72,
                  "窗口關閉 → 傳輸類動作不執行 → 安全分支為 `SLEEP`",
                  "F7F8FC", PURPLE, 22, NAVY, True, name="Closed-window branch")
    add_text(slide, 1.24, 4.69, 10.82, 0.42,
             "若中間事件沒有差異，最後的數值變化不得歸因於本次修改。",
             20, RED, True, PP_ALIGN.CENTER, name="Attribution rule")
    add_field_band(slide, "radio_state", "無線狀態", "runner event ledger",
                   "表示各狀態區間", "state enum + duration",
                   "提供端點能量的時間分段")
    slide.notes_slide.notes_text_frame.text = (
        "本頁建立結果歸因的順序。先確認服務窗口與策略動作，再讀取無線狀態和封包事件，最後判讀服務與端點能量。"
        "窗口關閉時，傳輸類動作不進入執行器，安全分支為 SLEEP。無線狀態 radio_state 來自 runner event ledger，"
        "作用是表示各狀態區間，單位為 state enum 加 duration；其判讀用途是提供端點能量的時間分段。"
    )


def slide_4(prs):
    slide = prepare_slide(prs, "固定情境中的四步操作閉環")
    add_text(slide, 1.08, 1.02, 11.08, 0.48,
             "只改一個受控常數，再沿同一條證據路徑比較", 24,
             NAVY, True, PP_ALIGN.CENTER, name="Loop lead")
    steps = [
        ("01", "修改", "`student_policy.py`\nmarked block", BLUE_PALE, BLUE),
        ("02", "執行", "固定 case\n固定輸入", GOLD_PALE, GOLD),
        ("03", "讀取", "狀態、封包\n服務、能量", TEAL_PALE, TEAL),
        ("04", "解釋", "動作 → 事件\n→ 結果", PURPLE_PALE, PURPLE),
    ]
    x = 0.92
    for idx, (num, heading, body, fill, line) in enumerate(steps):
        card = add_box(slide, x, 1.82, 2.72, 2.63, fill, line,
                       name=f"Loop step {num}")
        add_text(slide, x + 0.18, 2.00, 0.55, 0.46, num, 24, line, True,
                 PP_ALIGN.CENTER, name=f"Loop number {num}")
        add_text(slide, x + 0.78, 1.98, 1.66, 0.50, heading, 24, NAVY, True,
                 PP_ALIGN.LEFT, name=f"Loop heading {num}")
        add_text(slide, x + 0.22, 2.78, 2.28, 1.05, body, 20, INK, False,
                 PP_ALIGN.CENTER, name=f"Loop body {num}")
        if idx < len(steps) - 1:
            add_chevron(slide, x + 2.83, 2.86, 0.36, 0.46, line)
        x += 3.06
    add_text(slide, 1.18, 4.72, 10.96, 0.42,
             "預測與結果同時保存；反例保留為可檢驗證據。",
             20, NAVY, True, PP_ALIGN.CENTER, name="Loop evidence rule")
    add_field_band(slide, "result.json", "結果資料檔", "package runner",
                   "保存 run identity、狀態、封包、服務與能量", "JSON artifact",
                   "作為 replay 與 import 的輸入")
    slide.notes_slide.notes_text_frame.text = (
        "本頁定義固定情境中的操作閉環。可編輯面限於 student_policy.py 的 marked block，且一次只修改一個受控常數。"
        "執行固定 case 後，依序讀取狀態、封包、服務與能量，再以動作、事件與結果的關係說明差異。"
        "結果資料檔 result.json 由 package runner 產生，作用是保存 run identity、狀態、封包、服務與能量，單位為 JSON artifact；"
        "其判讀用途是作為 replay 與 import 的輸入。"
    )


def slide_5(prs):
    slide = prepare_slide(prs, "中間事件是因果歸因的必要橋樑")
    add_text(slide, 0.98, 1.03, 11.38, 0.48,
             "同一觀測下，比較基準與候選策略的事件路徑", 24,
             NAVY, True, PP_ALIGN.CENTER, name="Comparison lead")
    add_label_box(slide, 0.92, 1.76, 1.72, 1.02, "同一觀測", "F7F8FC", NAVY,
                  21, NAVY, True, name="Shared observation top")
    add_label_box(slide, 0.92, 3.20, 1.72, 1.02, "同一觀測", "F7F8FC", NAVY,
                  21, NAVY, True, name="Shared observation bottom")
    lanes = [
        (1.72, "基準策略", "動作 A", "事件路徑 A", "服務／能量 A", BLUE_PALE, BLUE),
        (3.16, "候選策略", "動作 B", "事件路徑 B", "服務／能量 B", TEAL_PALE, TEAL),
    ]
    for row, (y, label, action, events, result, fill, line) in enumerate(lanes):
        add_chevron(slide, 2.73, y + 0.26, 0.38, 0.46, line)
        add_label_box(slide, 3.14, y, 2.12, 1.10, f"{label}\n{action}", fill, line,
                      20, NAVY, True, name=f"Policy lane {row + 1}")
        add_chevron(slide, 5.35, y + 0.26, 0.38, 0.46, line)
        add_label_box(slide, 5.78, y, 2.52, 1.10, events, "F7F8FC", line,
                      20, NAVY, True, name=f"Event lane {row + 1}")
        add_chevron(slide, 8.40, y + 0.26, 0.38, 0.46, line)
        add_label_box(slide, 8.83, y, 3.45, 1.10, result, fill, line,
                      20, NAVY, True, name=f"Result lane {row + 1}")
    add_text(slide, 1.05, 4.64, 11.18, 0.48,
             "先確認事件路徑不同，才比較服務與能量。",
             21, RED, True, PP_ALIGN.CENTER, name="Causal comparison rule")
    add_field_band(slide, "packet_progress", "封包進度", "result／replay artifact",
                   "記錄產生、佇列、嘗試、重試、送達與逾期", "event count + time",
                   "連接策略動作與服務結果")
    slide.notes_slide.notes_text_frame.text = (
        "本頁將因果歸因具體化。基準策略與候選策略必須接收相同觀測，差異由策略動作開始，並經由不同的事件路徑傳到服務與端點能量。"
        "若事件路徑沒有差異，最終數值差異不得歸因於策略修改。封包進度 packet_progress 來自 result 或 replay artifact，"
        "作用是記錄產生、佇列、嘗試、重試、送達與逾期，單位為 event count 加 time；其判讀用途是連接策略動作與服務結果。"
    )


def slide_6(prs):
    slide = prepare_slide(prs, "Runner 證據有明確邊界")
    add_text(slide, 1.03, 1.02, 11.22, 0.48,
             "可重跑，不等於可以外推到所有能量層級", 24,
             NAVY, True, PP_ALIGN.CENTER, name="Evidence lead")
    left = add_box(slide, 0.92, 1.78, 5.56, 3.18, TEAL_PALE, TEAL,
                   name="Supported evidence")
    right = add_box(slide, 6.85, 1.78, 5.56, 3.18, RED_PALE, RED,
                    name="Unsupported evidence")
    add_text(slide, 1.22, 2.00, 4.96, 0.52, "本套件可支持", 25, TEAL, True,
             PP_ALIGN.CENTER, name="Supported heading")
    add_text(slide, 1.36, 2.73, 4.68, 1.70,
             "固定 scenario／seed\n端點狀態與封包事件\n服務判定與端點 J",
             21, INK, False, PP_ALIGN.CENTER, name="Supported body")
    add_text(slide, 7.15, 2.00, 4.96, 0.52, "本套件不支持", 25, RED, True,
             PP_ALIGN.CENTER, name="Unsupported heading")
    add_text(slide, 7.34, 2.73, 4.58, 1.70,
             "LEO／system energy\ncanonical efficiency\n實測或即時 KPI",
             21, INK, False, PP_ALIGN.CENTER, name="Unsupported body")
    add_field_band(slide, "endpoint_energy_j", "端點能量", "endpoint radio／processing model",
                   "累積狀態能量", "J",
                   "只屬於 endpoint evidence layer")
    slide.notes_slide.notes_text_frame.text = (
        "本頁界定執行器可支持的證據範圍。固定 scenario、policy 與 seed 可使事件順序重跑，並支持端點狀態、封包、服務與端點能量的比較。"
        "這些結果不等同於 LEO 或 system energy，也不等同於 canonical efficiency、實測或即時 KPI。"
        "端點能量 endpoint_energy_j 來自 endpoint radio 與 processing model，作用是累積狀態能量，單位為 J；"
        "判讀時只屬於 endpoint evidence layer。"
    )


def slide_7(prs):
    slide = prepare_slide(prs, "LEO 是變動服務窗口的工作範例")
    add_text(slide, 1.04, 1.02, 11.20, 0.48,
             "LEO 改變可服務機會；主要控制仍在端點策略", 24,
             NAVY, True, PP_ALIGN.CENTER, name="LEO lead")
    add_label_box(slide, 0.94, 1.76, 2.10, 1.08, "低地球軌道\n服務機會", BLUE_PALE, BLUE,
                  21, NAVY, True, name="LEO opportunity")
    add_chevron(slide, 3.18, 2.04, 0.42, 0.52, BLUE)
    timeline = add_box(slide, 3.70, 1.58, 5.82, 1.44, "F7F8FC", LINE,
                       name="LEO window timeline")
    segments = [("窗口開啟", TEAL, 1.50), ("品質下降", GOLD, 1.32), ("窗口關閉", "CBD5E1", 1.26), ("再度開啟", BLUE, 1.36)]
    cursor = 3.95
    for idx, (label, fill, width) in enumerate(segments):
        segment = add_box(slide, cursor, 2.01, width, 0.56, fill, None,
                          radius=False, name=f"LEO timeline segment {idx + 1}")
        write_text(segment, label, 16, WHITE if fill != "CBD5E1" else INK,
                   True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
                   margins=(0.01, 0.01, 0.01, 0.01))
        cursor += width + 0.06
    add_chevron(slide, 9.66, 2.04, 0.42, 0.52, TEAL)
    add_label_box(slide, 10.20, 1.76, 2.16, 1.08, "等待／休息\n／送出", TEAL_PALE, TEAL,
                  21, NAVY, True, name="Endpoint choice")
    add_text(slide, 1.00, 3.32, 11.28, 0.42,
             "相同機制可以轉移；窗口、期限與能量邊界必須重新定義。",
             21, NAVY, True, PP_ALIGN.CENTER, name="Transfer rule")
    examples = [
        (1.16, "暖通空調", "低負載時段"),
        (4.70, "邊緣閘道", "可連線時段"),
        (8.24, "物流追蹤", "回報機會"),
    ]
    for idx, (x, heading, body) in enumerate(examples):
        add_label_box(slide, x, 3.92, 3.05, 0.86, f"{heading}｜{body}",
                      [GOLD_PALE, BLUE_PALE, PURPLE_PALE][idx],
                      [GOLD, BLUE, PURPLE][idx], 19, NAVY, True,
                      name=f"Transfer example {idx + 1}")
    add_field_band(slide, "service_window", "服務窗口", "fixed scenario trace",
                   "定義合法傳輸區間", "trace time",
                   "決定動作是否具有服務機會")
    slide.notes_slide.notes_text_frame.text = (
        "本頁將低地球軌道視為變動服務窗口的工作範例。LEO 的作用是改變服務機會；主要控制仍是端點策略在等待、休息與送出之間做選擇。"
        "暖通空調、邊緣閘道與物流追蹤可以沿用相同機制，但各情境必須重新定義窗口、期限與能量邊界。"
        "服務窗口 service_window 來自 fixed scenario trace，作用是定義合法傳輸區間，單位為 trace time；判讀時決定動作是否具有服務機會。"
    )


def slide_8(prs):
    slide = prepare_slide(prs, "從預測到結果，保留可反駁路徑")
    add_text(slide, 1.04, 1.02, 11.20, 0.48,
             "先寫下方向，再執行；結果不符時保留反例", 24,
             NAVY, True, PP_ALIGN.CENTER, name="Prediction lead")
    steps = [
        (0.92, "01", "記錄預測", "狀態／封包／服務／J", BLUE_PALE, BLUE),
        (3.90, "02", "受控修改", "一次一個常數", GOLD_PALE, GOLD),
        (6.88, "03", "精確執行", "固定 case 與輸入", TEAL_PALE, TEAL),
        (9.86, "04", "比較與保存", "結果、解釋、反例", PURPLE_PALE, PURPLE),
    ]
    for idx, (x, num, heading, body, fill, line) in enumerate(steps):
        card = add_box(slide, x, 1.80, 2.50, 2.55, fill, line,
                       name=f"Prediction step {num}")
        add_text(slide, x + 0.18, 2.02, 0.52, 0.44, num, 23, line, True,
                 PP_ALIGN.CENTER, name=f"Prediction number {num}")
        add_text(slide, x + 0.72, 2.00, 1.55, 0.48, heading, 21, NAVY, True,
                 PP_ALIGN.LEFT, name=f"Prediction heading {num}")
        add_text(slide, x + 0.22, 2.82, 2.06, 0.96, body, 19, INK, False,
                 PP_ALIGN.CENTER, name=f"Prediction body {num}")
        if idx < len(steps) - 1:
            add_chevron(slide, x + 2.58, 2.77, 0.32, 0.44, line)
    add_label_box(slide, 1.34, 4.62, 10.66, 0.56,
                  "保留條件案例維持 frozen policy，用來檢查機制是否可轉移。",
                  "F7F8FC", RED, 20, NAVY, True, name="Withheld case")
    add_field_band(slide, "prediction", "預測", "workbook record",
                   "指定可反駁的變化方向", "statement",
                   "作為執行前的比較基準")
    slide.notes_slide.notes_text_frame.text = (
        "本頁把預測、執行與反例放在同一個證據脈絡。執行前先記錄狀態、封包、服務或端點能量的預期方向，"
        "再進行一次一個常數的受控修改與固定 case 執行。結果不符合預測時，反例仍須保存。"
        "預測 prediction 來自 workbook record，作用是指定可反駁的變化方向，單位為 statement；判讀時作為執行前的比較基準。"
    )


def slide_8a(prs):
    slide = prepare_slide(prs, "三個實驗到底改什麼")
    add_text(slide, 0.98, 1.02, 11.34, 0.50,
             "三次都只改 student_policy.py 的一個 marked block；runner、scenario、schema 不動",
             22, NAVY, True, PP_ALIGN.CENTER, name="Lab map lead")
    labs = [
        (
            0.82,
            "Lab A｜休息方式",
            "SLEEP → WAIT\nPACE_GAP_STEPS = 2 不變",
            "觀察：SLEEP／AWAKE_IDLE／WAKE 的時間與 J；封包服務",
            BLUE_PALE,
            BLUE,
        ),
        (
            4.55,
            "Lab B｜穩定門檻",
            "STABLE_STEPS：2 → 1\nENTER_QUALITY=2／EXIT_QUALITY=1 不變",
            "觀察：進入時機、attempt／retry／delivery、service、J",
            TEAL_PALE,
            TEAL,
        ),
        (
            8.28,
            "Lab C｜急件提前量",
            "URGENT_MARGIN_S：20 → 5 → 30\nBATCH_SIZE = 3 不變",
            "觀察：urgent timing、deadline／expired、service、J",
            GOLD_PALE,
            GOLD,
        ),
    ]
    for idx, (x, heading, change, observe, fill, line) in enumerate(labs):
        add_box(slide, x, 1.78, 3.44, 2.92, fill, line,
                name=f"Lab map card {idx + 1}")
        add_text(slide, x + 0.16, 2.02, 3.12, 0.50, heading, 22, line, True,
                 PP_ALIGN.CENTER, name=f"Lab map heading {idx + 1}")
        add_text(slide, x + 0.20, 2.78, 3.04, 0.90, change, 18, INK, True,
                 PP_ALIGN.CENTER, name=f"Lab map change {idx + 1}")
        add_text(slide, x + 0.22, 3.86, 3.00, 0.62, observe, 16, INK, False,
                 PP_ALIGN.CENTER, name=f"Lab map observe {idx + 1}")
    add_label_box(slide, 1.06, 4.90, 11.20, 0.44,
                  "同一檔案／同一 scenario／固定 runner → prediction → baseline → candidate → withheld",
                  "F7F8FC", PURPLE, 19, NAVY, True, name="Lab map causal strip")
    add_field_band(slide, "marked block", "受控區塊", "student_policy.py",
                   "限制可修改的最小決策面", "one block",
                   "其他行與檔案不動，才能歸因")
    slide.notes_slide.notes_text_frame.text = (
        "本頁把三個實驗的變更面一次說清楚。Lab A 只把 REST_DURING_GAP 從 SLEEP 改成 WAIT，"
        "PACE_GAP_STEPS、Lab B/C 區塊、runner、scenario 與 schema 都不變；觀察休眠、清醒閒置、喚醒與端點 J，"
        "再連到封包服務。Lab B 只把 STABLE_STEPS 從 2 改成 1，ENTER_QUALITY=2 與 EXIT_QUALITY=1 不變；"
        "觀察進入時間、嘗試、重傳、交付、service 與 J。Lab C 只依序把 URGENT_MARGIN_S 從 20 改為 5 再改為 30，"
        "BATCH_SIZE=3 不變；觀察急件分支時間、deadline、expired、service 與 J。"
        "每次仍只改同一檔案的一個 marked block，先記 prediction，再做 baseline、candidate 與 withheld。"
        "受控區塊 marked block 的來源是 student_policy.py，作用是限制最小決策面，單位是 one block；"
        "判讀時，其他行與檔案必須保持不動，才能把差異歸因於本次實驗。"
    )


SLIDE_BUILDERS = [
    slide_1, slide_1a, slide_1b, slide_1c, slide_1d, slide_8b, slide_8c, slide_8d,
    slide_2, slide_3, slide_4, slide_5, slide_6, slide_7, slide_8, slide_8a,
]


def package_qa(path: Path) -> dict:
    forbidden = ["學生", "老師", "講師", "分鐘", "checksum", "ZIP test", "製作備註"]
    layout_targets = []
    notes = []
    creation_ids = []
    slide_backgrounds = []
    minimum_sizes = []
    text = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        slide_names = sorted(
            (name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda value: int(re.search(r"(\d+)", value).group(1)),
        )
        for slide_name in slide_names:
            root = ET.fromstring(archive.read(slide_name))
            slide_backgrounds.append(root.find(f"./{{{P_NS}}}cSld/{{{P_NS}}}bg") is not None)
            for node in root.iter():
                local = node.tag.rsplit("}", 1)[-1]
                if local == "t" and node.text:
                    text.append(node.text)
                if local in {"rPr", "defRPr", "endParaRPr"} and node.get("sz"):
                    minimum_sizes.append(int(node.get("sz")) / 100)
                if local == "creationId":
                    creation_ids.append(node.get("id") or node.get("val"))
            rel_name = f"ppt/slides/_rels/{Path(slide_name).name}.rels"
            rel_root = ET.fromstring(archive.read(rel_name))
            targets = [
                rel.get("Target")
                for rel in rel_root.findall(f"{{{PKG_REL_NS}}}Relationship")
                if rel.get("Type", "").endswith("/slideLayout")
            ]
            layout_targets.extend(targets)
        notes = [name for name in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)]
    joined = "".join(text)
    forbidden_hits = [term for term in forbidden if term in joined]
    if "你" in joined:
        forbidden_hits.append("你")
    report = {
        "output": str(path),
        "slide_count": len(layout_targets),
        "notes_count": len(notes),
        "layout_targets": sorted(set(layout_targets)),
        "all_layout2": bool(layout_targets) and all(target.endswith("slideLayout2.xml") for target in layout_targets),
        "slide_background_count": sum(slide_backgrounds),
        "creation_id_count": len(creation_ids),
        "creation_ids_unique": len(creation_ids) == len(set(creation_ids)),
        "minimum_authored_font_pt": min(minimum_sizes) if minimum_sizes else None,
        "forbidden_hits": forbidden_hits,
    }
    report["status"] = "PASS" if (
        report["slide_count"] == 16
        and report["notes_count"] == 16
        and report["all_layout2"]
        and report["slide_background_count"] == 0
        and report["creation_ids_unique"]
        and not report["forbidden_hits"]
        and (report["minimum_authored_font_pt"] or 0) >= 16
    ) else "FAIL"
    return report


def build() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)
    for builder in SLIDE_BUILDERS:
        builder(prs)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    overlay_template_parts(OUTPUT)
    emit_manifest(OUTPUT)
    report = package_qa(OUTPUT)
    QA_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    QA_OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if report["status"] != "PASS":
        raise RuntimeError(json.dumps(report, ensure_ascii=False))
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
