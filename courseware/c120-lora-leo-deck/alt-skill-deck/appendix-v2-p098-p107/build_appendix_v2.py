#!/usr/bin/env python3
"""Build the owned P098-P107 appendix V2 lane.

The lane is rebuilt from the native educate template with python-pptx.  It
uses only the second template layout, preserves the inherited footer/logo/
divider system, keeps authored content editable, and embeds formal notes.
The source appendix builder and its P098-P116 output remain untouched.
"""

from __future__ import annotations

import copy
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


OWNED = Path(__file__).resolve().parent
ALT_ROOT = OWNED.parent
LATEST_ROOT = ALT_ROOT / "latest"
LATEST_ROOT.mkdir(parents=True, exist_ok=True)
TEMPLATE = Path("/home/sat/pptx-wrap/assets/templates/educate.pptx")
SOURCE_MAP_SOURCE = ALT_ROOT / "appendix" / "source-map.json"
OUTPUT = LATEST_ROOT / "LoRaEnergySim-LEO-ALT-APPENDIX-V2-P098-P107-REVIEW.pptx"
SOURCE_MAP_OUTPUT = OWNED / "source-map.json"
MANIFEST_OUTPUT = OWNED / "manifest.json"
BUILD_REPORT = OWNED / "build-report.json"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"

NAVY = "25356B"
INK = "25324A"
MUTED = "65758B"
BLUE = "4967C2"
BLUE_PALE = "ECF1FE"
TEAL = "007374"
TEAL_PALE = "E8F7F4"
PURPLE = "660066"
PURPLE_PALE = "F4EDF7"
GOLD = "846826"
GOLD_PALE = "FDF9E2"
RED = "992A2A"
RED_PALE = "FDEFEF"
WHITE = "FFFFFF"
LINE = "D7DEEA"
SOFT = "F7F8FC"
GREEN = "2E6B52"

CLAIM_BOUNDARY = (
    "SIMULATED TEACHING DATA / NOT LIVE / NOT MEASURED / "
    "NOT CANONICAL-PARITY-VERIFIED"
)


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_fonts(run) -> None:
    """Apply the deck's Traditional-Chinese/Latin font contract."""
    run.font.name = "Times New Roman"
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("latin", "Times New Roman"), ("ea", "標楷體"), ("cs", "Times New Roman")):
        child = rpr.find(qn(f"a:{tag}"))
        if child is None:
            child = OxmlElement(f"a:{tag}")
            rpr.append(child)
        child.set("typeface", face)


def add_marked_runs(paragraph, text: str, size: float, color: str,
                    bold: bool = False, italic_all: bool = False) -> None:
    """Backtick spans are variables, field names, code, or formula tokens."""
    pieces = re.split(r"(`[^`]+`)", text)
    for piece in pieces:
        if not piece:
            continue
        marked = piece.startswith("`") and piece.endswith("`")
        value = piece[1:-1] if marked else piece
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = marked or italic_all
        run.font.color.rgb = rgb(color)
        set_run_fonts(run)


def write_text(shape, text: str, size: float = 24, color: str = INK,
               bold: bool = False, align=PP_ALIGN.LEFT,
               valign=MSO_ANCHOR.MIDDLE,
               margins=(0.10, 0.05, 0.10, 0.05),
               line_spacing: float = 1.0, italic_all: bool = False) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    for idx, line in enumerate(text.split("\n")):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        add_marked_runs(paragraph, line, size, color, bold, italic_all)


def add_text(slide, x: float, y: float, w: float, h: float, text: str,
             size: float = 24, color: str = INK, bold: bool = False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margins=(0.02, 0.01, 0.02, 0.01), line_spacing: float = 1.0,
             name: str = "Text", italic_all: bool = False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    write_text(shape, text, size, color, bold, align, valign, margins, line_spacing, italic_all)
    return shape


def add_box(slide, x: float, y: float, w: float, h: float, fill: str,
            line: str | None = LINE, radius: bool = True,
            name: str = "Box", line_width: float = 1.2):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_label_box(slide, x: float, y: float, w: float, h: float, text: str,
                  fill: str, line: str, size: float = 20, color: str = INK,
                  bold: bool = False, align=PP_ALIGN.CENTER,
                  name: str = "Label", margins=(0.10, 0.04, 0.10, 0.04)):
    shape = add_box(slide, x, y, w, h, fill, line, name=name)
    write_text(shape, text, size, color, bold, align, MSO_ANCHOR.MIDDLE,
               margins=margins, line_spacing=0.94)
    return shape


def add_chevron(slide, x: float, y: float, w: float = 0.34,
                h: float = 0.44, color: str = BLUE,
                name: str = "Mechanism connector"):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_connector(slide, x1: float, y1: float, x2: float, y2: float,
                  color: str = BLUE, width: float = 1.6,
                  name: str = "Mechanism line"):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.name = name
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_field_band(slide, field: str, chinese: str, source: str, purpose: str,
                   unit: str, interpretation: str, y: float = 5.60,
                   height: float = 0.78) -> None:
    band = add_box(slide, 0.82, y, 11.72, height, SOFT, BLUE,
                   name=f"Field contract {field}", line_width=1.0)
    write_text(
        band,
        f"`{field}`（{chinese}）｜來源：{source}｜作用：{purpose}\n"
        f"單位：{unit}｜判讀：{interpretation}",
        17.3, NAVY, False, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
        margins=(0.16, 0.03, 0.16, 0.03), line_spacing=0.88,
    )


def add_field_pair_band(slide, lines: list[str], y: float = 5.52,
                        height: float = 0.88, line: str = BLUE) -> None:
    band = add_box(slide, 0.82, y, 11.72, height, SOFT, line,
                   name="Field contract pair", line_width=1.0)
    write_text(band, "\n".join(lines), 17.0, NAVY, False, PP_ALIGN.LEFT,
               MSO_ANCHOR.MIDDLE, margins=(0.16, 0.03, 0.16, 0.03),
               line_spacing=0.88)


def add_status(slide, text: str, fill: str = GOLD_PALE, line: str = GOLD,
               color: str = NAVY):
    return add_label_box(slide, 9.32, 0.92, 3.02, 0.28, text,
                         fill, line, 16.0, color, True,
                         name="Evidence status", margins=(0.04, 0.01, 0.04, 0.01))


def delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def prepare_slide(prs: Presentation, title: str):
    """Create a slide from source slide 2 / slideLayout2.xml only."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    if title_shape is None:
        raise RuntimeError("slideLayout2 title placeholder missing")
    for placeholder in list(slide.placeholders):
        if (placeholder._element is title_shape._element
                or placeholder.placeholder_format.type == PP_PLACEHOLDER.TITLE):
            continue
        if placeholder.placeholder_format.type not in {
            PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER,
        }:
            delete_shape(placeholder)
    title_shape.left = Inches(0.718057)
    title_shape.top = Inches(0.204514)
    title_shape.width = Inches(10.34861)
    title_shape.height = Inches(0.525)
    title_shape.name = "Native layout2 title"
    write_text(title_shape, title, 28, NAVY, True, PP_ALIGN.LEFT,
               MSO_ANCHOR.MIDDLE, margins=(0.02, 0.0, 0.02, 0.0))
    return slide


def remove_all_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def overlay_template_parts(path: Path) -> None:
    """Restore native template-owned parts after python-pptx saves."""
    owned_prefixes = (
        "ppt/slideLayouts/", "ppt/slideMasters/", "ppt/theme/",
        "ppt/notesMasters/",
    )
    with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(path, "r") as old:
        data = {info.filename: old.read(info.filename) for info in old.infolist()}
        infos = {info.filename: copy.copy(info) for info in old.infolist()}
        for info in source.infolist():
            if info.filename.startswith(owned_prefixes):
                data[info.filename] = source.read(info.filename)
                infos[info.filename] = copy.copy(info)
    overlay = path.with_suffix(".overlay.pptx")
    with zipfile.ZipFile(overlay, "w", zipfile.ZIP_DEFLATED) as out:
        for name, payload in data.items():
            out.writestr(infos[name], payload)
    overlay.replace(path)


def _math_r(parent, text: str, size: int = 3000, italic: bool = False):
    run = ET.SubElement(parent, f"{{{M_NS}}}r")
    rpr = ET.SubElement(run, f"{{{A_NS}}}rPr", {
        "lang": "en-US", "altLang": "en-US", "sz": str(size), "dirty": "0",
    })
    if italic:
        rpr.set("i", "1")
    ET.SubElement(rpr, f"{{{A_NS}}}latin", {"typeface": "Cambria Math"})
    ET.SubElement(run, f"{{{M_NS}}}t").text = text
    return run


def _math_sub(parent, base: str, sub: str):
    node = ET.SubElement(parent, f"{{{M_NS}}}sSub")
    ET.SubElement(node, f"{{{M_NS}}}sSubPr")
    e1 = ET.SubElement(node, f"{{{M_NS}}}e")
    _math_r(e1, base, italic=True)
    e2 = ET.SubElement(node, f"{{{M_NS}}}e")
    _math_r(e2, sub)
    return node


def _formula_body() -> ET.Element:
    para = ET.Element(f"{{{M_NS}}}oMathPara")
    omath = ET.SubElement(para, f"{{{M_NS}}}oMath")
    _math_sub(omath, "E", "endpoint")
    _math_r(omath, " = ")
    nary = ET.SubElement(omath, f"{{{M_NS}}}nary")
    nary_pr = ET.SubElement(nary, f"{{{M_NS}}}naryPr")
    ET.SubElement(nary_pr, f"{{{M_NS}}}chr", {f"{{{M_NS}}}val": "∑"})
    ET.SubElement(nary_pr, f"{{{M_NS}}}limLoc", {f"{{{M_NS}}}val": "undOvr"})
    ET.SubElement(nary_pr, f"{{{M_NS}}}grow", {f"{{{M_NS}}}val": "1"})
    sub = ET.SubElement(nary, f"{{{M_NS}}}sub")
    _math_r(sub, "s", italic=True)
    _math_r(sub, " ∈ ")
    _math_r(sub, "S", italic=True)
    ET.SubElement(nary, f"{{{M_NS}}}sup")
    body = ET.SubElement(nary, f"{{{M_NS}}}e")
    _math_sub(body, "P", "s")
    _math_r(body, " ")
    _math_sub(body, "t", "s")
    return para


def add_native_math(path: Path, slide_number: int) -> None:
    """Add one native OMML object to the accounting slide."""
    ET.register_namespace("a", A_NS)
    ET.register_namespace("m", M_NS)
    ET.register_namespace("p", P_NS)
    ET.register_namespace("a14", A14_NS)
    with zipfile.ZipFile(path, "r") as archive:
        infos = [copy.copy(info) for info in archive.infolist()]
        data = {info.filename: archive.read(info.filename) for info in infos}
    name = f"ppt/slides/slide{slide_number}.xml"
    root = ET.fromstring(data[name])
    sp_tree = root.find(f"./{{{P_NS}}}cSld/{{{P_NS}}}spTree")
    if sp_tree is None:
        raise RuntimeError(f"formula slide {slide_number} has no shape tree")
    ids = [int(node.get("id")) for node in sp_tree.iter(f"{{{P_NS}}}cNvPr")
           if (node.get("id") or "").isdigit()]
    shape = ET.Element(f"{{{P_NS}}}sp")
    nv = ET.SubElement(shape, f"{{{P_NS}}}nvSpPr")
    next_id = max(ids, default=1) + 1
    ET.SubElement(nv, f"{{{P_NS}}}cNvPr", {
        "id": str(next_id), "name": "Native Office Math endpoint energy",
        "title": "endpoint energy formula", "descr": "editable Office Math",
    })
    c_nv_sp_pr = ET.SubElement(nv, f"{{{P_NS}}}cNvSpPr", {"hidden": "1"})
    ET.SubElement(c_nv_sp_pr, f"{{{A_NS}}}spLocks", {"noGrp": "1"})
    ET.SubElement(nv, f"{{{P_NS}}}nvPr")
    sp_pr = ET.SubElement(shape, f"{{{P_NS}}}spPr")
    xfrm = ET.SubElement(sp_pr, f"{{{A_NS}}}xfrm")
    ET.SubElement(xfrm, f"{{{A_NS}}}off", {"x": "0", "y": "0"})
    ET.SubElement(xfrm, f"{{{A_NS}}}ext", {"cx": "1", "cy": "1"})
    ET.SubElement(sp_pr, f"{{{A_NS}}}noFill")
    ln = ET.SubElement(sp_pr, f"{{{A_NS}}}ln")
    ET.SubElement(ln, f"{{{A_NS}}}noFill")
    tx_body = ET.SubElement(shape, f"{{{P_NS}}}txBody")
    ET.SubElement(tx_body, f"{{{A_NS}}}bodyPr", {"anchor": "ctr", "wrap": "none"})
    ET.SubElement(tx_body, f"{{{A_NS}}}lstStyle")
    para = ET.SubElement(tx_body, f"{{{A_NS}}}p")
    ET.SubElement(para, f"{{{A_NS}}}pPr", {"algn": "ctr"})
    wrapper = ET.SubElement(para, f"{{{A14_NS}}}m")
    wrapper.append(_formula_body())
    ET.SubElement(para, f"{{{A_NS}}}endParaRPr", {"lang": "en-US", "sz": "3000", "dirty": "0"})
    sp_tree.append(shape)
    data[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    temp = path.with_suffix(".math.pptx")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as out:
        for info in infos:
            out.writestr(info, data[info.filename])
    temp.replace(path)


def lead(slide, text: str, color: str = NAVY):
    return add_text(slide, 0.98, 1.03, 11.38, 0.42, text, 24, color, True,
                    PP_ALIGN.CENTER, name="Teaching proposition")


def node(slide, x: float, y: float, w: float, h: float, heading: str,
         body: str, fill: str, line: str, size: float = 19,
         title_size: float = 21, name: str = "Teaching node"):
    add_box(slide, x, y, w, h, fill, line, name=name)
    add_text(slide, x + 0.12, y + 0.12, w - 0.24, 0.34, heading,
             title_size, line, True, PP_ALIGN.CENTER, name=f"{name} heading")
    add_text(slide, x + 0.14, y + 0.53, w - 0.28, h - 0.63, body,
             size, INK, False, PP_ALIGN.CENTER, name=f"{name} body")


def note(slide, text: str) -> None:
    tf = slide.notes_slide.notes_text_frame
    tf.text = text
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            rpr = run._r.get_or_add_rPr()
            for tag, face in (("latin", "Times New Roman"), ("ea", "標楷體"), ("cs", "Times New Roman")):
                child = rpr.find(qn(f"a:{tag}"))
                if child is None:
                    child = OxmlElement(f"a:{tag}")
                    rpr.append(child)
                child.set("typeface", face)


def slide_p098(prs):
    slide = prepare_slide(prs, "P098｜TLE 來源與 scenario anchor")
    lead(slide, "來源識別、目標時間與 scenario identity 共同固定模型輸入")
    add_status(slide, "current source record｜待補", GOLD_PALE, GOLD)
    node(slide, 0.90, 1.62, 3.52, 2.52, "TLE source", "`tle_source_id`\noneweb-0314-archive-2026-08-08\n\n來源模式：bundled", BLUE_PALE, BLUE, 20, 22, "TLE source card")
    node(slide, 4.88, 1.62, 3.52, 2.52, "Scenario anchor", "`scenario_id`\nntpu-energy-decision-01\n\ncourse package identity", TEAL_PALE, TEAL, 20, 22, "Scenario anchor card")
    node(slide, 8.86, 1.62, 3.52, 2.52, "Target time", "`target_utc`\n2026-08-09T04:00:00Z\n\n模型推進的時間輸入", GOLD_PALE, GOLD, 20, 22, "Target time card")
    add_chevron(slide, 4.50, 2.60, 0.30, 0.40, BLUE, "Source to anchor connector")
    add_chevron(slide, 8.48, 2.60, 0.30, 0.40, TEAL, "Anchor to time connector")
    add_label_box(slide, 1.00, 4.48, 5.38, 0.66, "`epoch`｜元素基準時刻｜current source record：待補", SOFT, PURPLE, 19, NAVY, True, name="Epoch pending")
    add_label_box(slide, 6.94, 4.48, 5.38, 0.66, "`raw bytes`｜原始位元組｜current source record：待補", SOFT, PURPLE, 19, NAVY, True, name="Raw bytes pending")
    add_field_band(slide, "tle_source_id", "TLE 來源識別碼", "scenario package", "連結固定來源與後續模型推進", "字串識別碼", "identity 完整時建立 source provenance；epoch／raw bytes 狀態為待補")
    note(slide, "本頁建立 TLE source 與 scenario anchor 的關係。`tle_source_id`（TLE 來源識別碼）由 scenario package 提供，作用是連結固定來源與後續模型推進，值型態為字串識別碼。`scenario_id` 與 `target_utc` 共同界定本次模型輸入。當前 package 提供 source id、scenario id 與 target time；epoch 與 raw bytes 尚未進入 current source record，因此維持待補狀態。這一頁的證據層級是 source provenance，後續模型狀態必須沿相同 identity 生成。來源：source-map.json、ADR-004、C-120 SDD、current package scenario contract。")


def slide_p099(prs):
    slide = prepare_slide(prs, "P099｜SGP4 模型推進至目標時間")
    lead(slide, "平均元素與 target time 經模型推進後，形成可追溯的狀態向量")
    node(slide, 0.82, 1.68, 3.28, 2.38, "Input｜輸入", "`TLE / GP`\n來源：source record\n\n`target_time`\n來源：scenario clock", BLUE_PALE, BLUE, 19, 21, "SGP4 input")
    add_chevron(slide, 4.27, 2.64, 0.40, 0.48, BLUE, "Model connector")
    add_label_box(slide, 4.82, 1.98, 3.32, 1.76, "SGP4\n簡化一般擾動模型\n\n依 target time 推進", GOLD_PALE, GOLD, 22, NAVY, True, name="SGP4 model")
    add_chevron(slide, 8.31, 2.64, 0.40, 0.48, GOLD, "Derived state connector")
    node(slide, 8.84, 1.68, 3.70, 2.38, "Derived state｜推導狀態", "`position`（位置向量）\n單位：km\n\n`velocity`（速度向量）\n單位：km/s", TEAL_PALE, TEAL, 18.5, 21, "SGP4 output")
    add_label_box(slide, 1.10, 4.42, 5.20, 0.70, "model version + source identity + target time + units", SOFT, NAVY, 18.5, NAVY, True, name="Lineage lock")
    add_label_box(slide, 7.02, 4.42, 5.20, 0.70, "derived state｜模型推導資料分類", SOFT, TEAL, 20, TEAL, True, name="Derived class")
    add_field_pair_band(slide, [
        "`position`（位置向量）｜來源：SGP4 output｜作用：描述目標時間的 TEME 狀態｜單位：km｜判讀：保留 frame 與 time lineage",
        "`velocity`（速度向量）｜來源：SGP4 output｜作用：描述同一時刻的運動狀態｜單位：km/s｜判讀：與 position 使用同一 frame",
    ])
    note(slide, "本頁說明 SGP4 的資料鏈。`target_time`（目標時間）由 scenario clock 提供，作用是指定模型推進時點，值型態為時間戳。SGP4 讀取 TLE 或 GP 的來源資料，輸出 `position`（位置向量，單位 km）與 `velocity`（速度向量，單位 km/s）。兩個輸出均屬 derived model evidence；source identity、model version、target time 與 units 必須共同保留。current appendix 不引入 donor 數值，頁面只保留模型機制與資料分類。來源：source-map.json、ADR-004、C-120 SDD 的 scenario/model lineage。")


def slide_p100(prs):
    slide = prepare_slide(prs, "P100｜Reference frame 至觀測者座標")
    lead(slide, "同一狀態向量經 frame conversion 與 observer localization 後，形成方向欄位")
    frames = [
        ("TEME", "真赤道平春分點座標\nSGP4 state\nkm／km/s", BLUE_PALE, BLUE),
        ("ECEF", "地心地固座標\n旋轉地球 reference\nkm／km/s", GOLD_PALE, GOLD),
        ("ENU", "東北天座標\nobserver location\nkm／km/s", TEAL_PALE, TEAL),
        ("az / el / range", "方位角／仰角／斜距\nobserver-frame fields\ndeg／deg／km", PURPLE_PALE, PURPLE),
    ]
    x = 0.72
    for idx, (heading, body, fill, line) in enumerate(frames):
        node(slide, x, 1.72, 2.68 if idx < 3 else 3.28, 2.12, heading, body, fill, line, 17.5, 20, f"Frame {heading}")
        if idx < 3:
            add_chevron(slide, x + 2.80, 2.56, 0.28, 0.40, line, "Frame conversion connector")
        x += 3.00 if idx < 3 else 0
    add_label_box(slide, 0.96, 4.24, 3.56, 0.66, "地球旋轉映射", SOFT, GOLD, 19, GOLD, True, name="ECEF mechanism")
    add_label_box(slide, 4.90, 4.24, 3.56, 0.66, "觀測者位置局部化", SOFT, TEAL, 19, TEAL, True, name="ENU mechanism")
    add_label_box(slide, 8.84, 4.24, 3.56, 0.66, "方向與距離欄位形成", SOFT, PURPLE, 19, PURPLE, True, name="Observer mechanism")
    add_field_pair_band(slide, [
        "`TEME`（真赤道平春分點座標）｜來源：SGP4｜作用：保存軌道模型狀態｜單位：km、km/s｜判讀：reference frame",
        "`ECEF`（地心地固座標）／`ENU`（東北天座標）｜來源：frame conversion + observer location｜作用：連結旋轉地球與觀測者｜單位：km、km/s｜判讀：轉換鏈完整才進入 azimuth／elevation／range",
    ])
    note(slide, "本頁把 reference frame 與 observer-frame 欄位分開。`TEME`（真赤道平春分點座標）是 SGP4 的模型狀態，值型態包含 km 與 km/s。`ECEF`（地心地固座標）透過地球旋轉映射連接到地固 frame；`ENU`（東北天座標）再透過 observer location 形成觀測者局部座標。最後的 azimuth（方位角）與 elevation（仰角）以 degree 表示，range（斜距）以 km 表示。frame 或 unit mapping 不完整時，品質與 policy 比較保持未定義。來源：source-map.json、ADR-004、C-120 SDD；donor 僅提供概念位置。")


def slide_p101(prs):
    slide = prepare_slide(prs, "P101｜geometry 與 quality 的分層 gate")
    lead(slide, "幾何欄位、距離欄位與品質欄位進入不同判讀層級")
    stages = [
        ("G1", "elevation", "仰角\n可見性 context", BLUE_PALE, BLUE, "degree"),
        ("G2", "boresight angle", "離軸角\nbeam geometry context", GOLD_PALE, GOLD, "degree"),
        ("G3", "range", "斜距\npath／timing context", TEAL_PALE, TEAL, "km"),
        ("G4", "quality_band", "品質分類\npolicy observation", PURPLE_PALE, PURPLE, "ordinal"),
    ]
    x = 0.72
    for idx, (gate, field, body, fill, line, unit) in enumerate(stages):
        add_label_box(slide, x, 1.62, 0.72, 0.54, gate, fill, line, 19, line, True, name=f"Gate {gate}")
        add_label_box(slide, x + 0.82, 1.62, 2.10, 0.54, f"`{field}`", SOFT, line, 18, line, True, name=f"Gate field {field}")
        add_label_box(slide, x, 2.42, 2.72, 1.10, f"{body}\n單位：{unit}", fill, line, 19, NAVY, True, name=f"Gate body {field}")
        if idx < 3:
            add_chevron(slide, x + 2.82, 2.73, 0.24, 0.40, line, "Gate mechanism connector")
        x += 3.05
    add_label_box(slide, 1.06, 4.06, 5.28, 0.74, "visibility／geometry／quality 是不同 evidence layer", SOFT, NAVY, 20, NAVY, True, name="Gate distinction")
    add_label_box(slide, 6.98, 4.06, 5.28, 0.74, "service opportunity 需具備 current scenario 定義的必要 gate", SOFT, RED, 20, RED, True, name="Gate qualification")
    add_field_pair_band(slide, [
        "`elevation`（仰角）｜來源：geometry derivation｜作用：描述地平線可見性｜單位：degree｜判讀：可見性 context，不代表 delivery",
        "`quality_band`（品質帶）｜來源：runner observation｜作用：供 policy threshold 判讀｜單位：ordinal classification｜判讀：quality 不取代 service evidence",
    ])
    note(slide, "本頁建立四類 gate 的層級。`elevation`（仰角）來自 geometry derivation，單位 degree，描述地平線可見性。`boresight angle`（離軸角）描述波束中心偏離，同樣以 degree 表示。`range`（斜距）來自 observer frame，單位 km，提供 path 與 timing context。`quality_band`（品質帶）由 runner observation 產生，屬 ordinal classification，供 policy threshold 判讀。四類欄位各自提供 context；service opportunity 的合法性仍需依 current scenario 的 gate 定義。來源：source-map.json、ADR-004、C-120 SDD。")


def slide_p102(prs):
    slide = prepare_slide(prs, "P102｜quality／dB 作為 decision-time observation")
    lead(slide, "決策時的品質觀測提供 action context；service 與 energy 由後續 result fields 判定")
    add_status(slide, "dB source／timing｜待補", GOLD_PALE, GOLD)
    add_label_box(slide, 0.94, 1.58, 3.24, 0.72, "`quality_band`\nclosed／weak／usable／strong", BLUE_PALE, BLUE, 20, NAVY, True, name="Quality observation")
    add_chevron(slide, 4.38, 1.74, 0.34, 0.44, BLUE, "Observation connector")
    add_label_box(slide, 4.92, 1.58, 3.24, 0.72, "decision context\ncurrent／past observation", SOFT, NAVY, 20, NAVY, True, name="Decision context")
    add_chevron(slide, 8.36, 1.74, 0.34, 0.44, TEAL, "Policy connector")
    add_label_box(slide, 8.90, 1.58, 3.50, 0.72, "`threshold`\nmarked policy block", TEAL_PALE, TEAL, 20, NAVY, True, name="Policy threshold")
    add_label_box(slide, 0.98, 2.82, 3.52, 1.24, "`dB`（分貝）\nlogarithmic ratio\n來源：model field；待補", PURPLE_PALE, PURPLE, 21, NAVY, True, name="dB card")
    add_label_box(slide, 4.90, 2.82, 3.52, 1.24, "quality observation\nordinal classification\n單位：無物理 dB 單位", BLUE_PALE, BLUE, 21, NAVY, True, name="Quality band card")
    add_label_box(slide, 8.82, 2.82, 3.52, 1.24, "action branch\nWAIT／SLEEP／SEND family\nservice result 後續判讀", TEAL_PALE, TEAL, 21, NAVY, True, name="Action branch card")
    add_label_box(slide, 1.18, 4.45, 11.02, 0.64, "quality observation → policy branch；delivery、service 與 endpoint J 由各自 result fields 提供", SOFT, RED, 20, NAVY, True, name="Observation boundary")
    add_field_pair_band(slide, [
        "`quality_band`（品質帶）｜來源：runner observation｜作用：供 enter／hold／exit 條件讀取｜單位：ordinal、無物理單位｜判讀：描述 action context",
        "`dB`（分貝）／`threshold`（門檻）｜來源：model field／student_policy.py marked block｜作用：表示 logarithmic ratio／決定 branch｜單位：dB／contract type｜判讀：current source 或 timing 未定義時維持待補",
    ])
    note(slide, "本頁把 quality observation 與後續服務結果分層。`quality_band`（品質帶）由 runner observation 產生，採 closed、weak、usable、strong 的 ordinal classification，作用是供 enter、hold、exit 條件讀取，值型態沒有物理 dB 單位。`dB`（分貝）代表 logarithmic ratio；目前 appendix source record 未提供可核對的 dB source 與 observation timing，因此標示待補。`threshold`（門檻）由 student_policy.py 的 marked block 提供，作用是決定 action branch。品質觀測只提供 decision context，delivery、service 與 endpoint energy 由不同 result fields 判定。來源：source-map.json、current package quality contract、C-120 SDD。")


def slide_p103(prs):
    slide = prepare_slide(prs, "P103｜throughput、delivered bits 與 service")
    lead(slide, "傳輸過程、已送達工作量與服務判定分別描述不同結果")
    items = [
        ("throughput", "吞吐率\n指定 interval 的 delivered rate\nbit/s", BLUE_PALE, BLUE),
        ("delivered_bits", "已送達位元\npacket ledger 的交付工作量\nbit", TEAL_PALE, TEAL),
        ("service_pass", "服務通過\n整合 delivery、deadline、freshness\nBoolean", PURPLE_PALE, PURPLE),
    ]
    x = 0.82
    for idx, (field, body, fill, line) in enumerate(items):
        node(slide, x, 1.65, 3.45, 2.32, f"`{field}`", body, fill, line, 19, 21, f"Service field {field}")
        if idx < 2:
            add_chevron(slide, x + 3.56, 2.57, 0.30, 0.40, line, "Service mechanism connector")
        x += 4.05
    add_label_box(slide, 1.05, 4.30, 3.38, 0.66, "過程：rate／throughput", SOFT, BLUE, 19, BLUE, True, name="Rate role")
    add_label_box(slide, 4.98, 4.30, 3.38, 0.66, "輸出：delivered work", SOFT, TEAL, 19, TEAL, True, name="Delivered role")
    add_label_box(slide, 8.91, 4.30, 3.38, 0.66, "gate：service verdict", SOFT, PURPLE, 19, PURPLE, True, name="Service role")
    add_field_pair_band(slide, [
        "`throughput`（吞吐率）｜來源：result aggregation｜作用：描述指定 interval 的 delivered rate｜單位：bit/s｜判讀：不等同於完整 service",
        "`delivered_bits`（已送達位元）／`service_pass`（服務通過）｜來源：packet ledger／service evaluator｜作用：提供工作量與 gate｜單位：bit／Boolean｜判讀：energy efficiency 需在同一 qualified window 比較",
    ])
    note(slide, "本頁把 throughput、delivered bits 與 service 分開。`throughput`（吞吐率）由 result aggregation 產生，單位 bit/s，描述指定 interval 的 delivered rate。`delivered_bits`（已送達位元）由 packet ledger 產生，單位 bit，累積符合 delivery 定義的工作量。`service_pass`（服務通過）由 service evaluator 產生，值型態 Boolean，整合 delivery、deadline 與 freshness。判讀順序固定為：辨識過程與工作量，讀取 service gate，最後進入 endpoint energy efficiency 的比較。來源：source-map.json、C-120 SDD result contract、current package result schema。")


def slide_p104(prs):
    slide = prepare_slide(prs, "P104｜切換時機的兩類服務風險")
    lead(slide, "同一 frozen policy 在不同 trace 上，可能呈現延遲切換或往返切換")
    add_status(slide, "current frozen trace｜待補", GOLD_PALE, GOLD)
    add_label_box(slide, 0.92, 1.54, 5.54, 0.52, "型態 A｜delayed switch（延遲切換）", BLUE_PALE, BLUE, 21, BLUE, True, name="Delayed switch heading")
    add_label_box(slide, 6.86, 1.54, 5.54, 0.52, "型態 B｜ping-pong（往返切換）", PURPLE_PALE, PURPLE, 21, PURPLE, True, name="Ping pong heading")
    add_label_box(slide, 1.00, 2.34, 2.22, 1.08, "hold／threshold\n觸發較晚", SOFT, BLUE, 20, NAVY, True, name="Delayed cause")
    add_chevron(slide, 3.42, 2.65, 0.32, 0.42, BLUE, "Delayed mechanism connector")
    add_label_box(slide, 3.92, 2.34, 2.22, 1.08, "window／deadline\n機會縮短", SOFT, GOLD, 20, NAVY, True, name="Delayed context")
    add_label_box(slide, 6.94, 2.34, 2.22, 1.08, "enter／exit\n間隔過近", SOFT, PURPLE, 20, NAVY, True, name="Ping cause")
    add_chevron(slide, 9.36, 2.65, 0.32, 0.42, PURPLE, "Ping mechanism connector")
    add_label_box(slide, 9.86, 2.34, 2.22, 1.08, "MODE_CHANGE\n反覆出現", SOFT, RED, 20, NAVY, True, name="Ping context")
    add_label_box(slide, 1.00, 3.92, 5.54, 0.96, "觀察：transition timing、packet outcome、service continuity\n狀態：待補 current frozen trace", BLUE_PALE, BLUE, 19, NAVY, True, name="Delayed evidence")
    add_label_box(slide, 6.94, 3.92, 5.54, 0.96, "觀察：MODE_CHANGE、retry、service continuity、state duration\n狀態：待補 current frozen trace", PURPLE_PALE, PURPLE, 19, NAVY, True, name="Ping evidence")
    add_field_pair_band(slide, [
        "`delayed_switch`（延遲切換）｜來源：transition ledger｜作用：定位錯過 opportunity 的 event timing｜單位：event timing｜判讀：需連結 service／deadline",
        "`ping_pong`（往返切換）／`continuity`（服務連續性）｜來源：transition ledger／service trace｜作用：辨識反覆 transition 與服務結果｜單位：event count／Boolean 或 duration｜判讀：current trace 待補",
    ])
    note(slide, "本頁以 Lab B optional counterexample 的概念區分兩類切換風險。`delayed_switch`（延遲切換）由 transition ledger 提供，作用是定位 hold 或 threshold 使模式較晚進入的時機；判讀需連結 contact window、deadline 與 service evidence。`ping_pong`（往返切換）由反覆 MODE_CHANGE 事件描述，作用是辨識 enter／exit 間隔過近造成的狀態往返。`continuity`（服務連續性）由 service trace 產生，值型態為 Boolean 或 duration。current frozen trace 未凍結，因此事件數、能量與服務結果保留待補狀態；frozen policy 不變，後續 evidence 才能區分兩類機制。來源：source-map.json、Lab B field contract、C-120 SDD。")


def slide_p105(prs):
    slide = prepare_slide(prs, "P105｜Power × time 與 endpoint energy boundary")
    lead(slide, "Power 描述瞬時速率；state duration 將各 bucket 累積為 endpoint energy")
    add_box(slide, 0.92, 1.50, 5.70, 2.22, SOFT, PURPLE, name="Native formula frame", line_width=1.6)
    add_text(slide, 1.18, 1.82, 5.18, 0.60, "`E_endpoint = Σ P_s t_s`", 30, PURPLE, True, PP_ALIGN.CENTER, name="Editable formula transcription", italic_all=True)
    add_text(slide, 1.24, 2.62, 5.04, 0.34, "可編輯 Office Math｜Σ 對應 state buckets", 18, PURPLE, True, PP_ALIGN.CENTER, name="Formula contract label")
    add_label_box(slide, 7.00, 1.54, 2.50, 0.84, "`P(t)`\n功率｜W", BLUE_PALE, BLUE, 22, NAVY, True, name="Power field")
    add_label_box(slide, 9.82, 1.54, 2.50, 0.84, "`state duration`\n停留時間｜s", GOLD_PALE, GOLD, 20, NAVY, True, name="Duration field")
    add_connector(slide, 8.25, 2.44, 8.25, 3.08, BLUE, name="Power to ledger")
    add_connector(slide, 11.07, 2.44, 11.07, 3.08, GOLD, name="Duration to ledger")
    buckets = [("sleep", TEAL), ("awake_idle", BLUE), ("wake", GOLD), ("process", PURPLE), ("tx", RED), ("rx", NAVY)]
    x = 0.92
    for idx, (label, color) in enumerate(buckets):
        add_box(slide, x, 4.02, 1.72, 0.70, color, None, radius=False, name=f"Energy bucket {label}")
        add_text(slide, x + 0.05, 4.17, 1.62, 0.36, f"`{label}`", 18, WHITE, True, PP_ALIGN.CENTER, name=f"Energy bucket label {label}")
        x += 1.98
    add_label_box(slide, 9.04, 3.10, 3.48, 0.60, "endpoint boundary\nradio／processing assumptions", TEAL_PALE, TEAL, 18.5, NAVY, True, name="Endpoint boundary")
    add_field_pair_band(slide, [
        "`P(t)`（時間函數功率）／`state duration`（狀態停留時間）｜來源：energy profile／endpoint replay｜作用：形成各 bucket 的累積量｜單位：W／s｜判讀：時間與功率共同決定 J",
        "`endpoint_energy_j`（端點累積能量）｜來源：runner result｜作用：涵蓋 endpoint radio／processing course boundary｜單位：J｜判讀：system consumed J 與 canonical efficiency 由獨立 authority 提供",
    ])
    note(slide, "本頁以 native Office Math 表示 endpoint energy 的累積語義。`P(t)`（時間函數功率）由 energy profile 提供，單位 W；`state duration`（狀態停留時間）由 endpoint replay 產生，單位 s。公式 `E_endpoint = Σ P_s t_s` 將 sleep、awake_idle、wake、process、tx、rx 六個 bucket 的功率與時間累加為 `endpoint_energy_j`（端點累積能量），單位 J。這個 boundary 只涵蓋 endpoint radio 與 processing course assumptions；system consumed J 與 canonical efficiency 屬於獨立 system authority。公式由 native Office Math 與可編輯文字轉錄共同保存。來源：source-map.json、C-120 SDD energy semantics、current package result contract。")


def slide_p106(prs):
    slide = prepare_slide(prs, "P106｜公平 A/B 的固定條件與 consequential diff")
    lead(slide, "相同情境、seed、service window、units 與 predecessor 形成比較資格")
    locks = [
        ("scenario", "情境 identity\n固定 inputs／window", BLUE_PALE, BLUE),
        ("seed", "case seed\n固定 deterministic branch", GOLD_PALE, GOLD),
        ("predecessor", "前置結果\n保持 policy lineage", TEAL_PALE, TEAL),
        ("consequential diff", "result／replay\ntyped evidence change", PURPLE_PALE, PURPLE),
    ]
    x = 0.78
    for idx, (field, body, fill, line) in enumerate(locks):
        node(slide, x, 1.58, 2.78 if idx < 3 else 3.00, 2.18, f"`{field}`", body, fill, line, 18.5, 20, f"A/B lock {field}")
        if idx < 3:
            add_chevron(slide, x + 2.92, 2.45, 0.26, 0.38, line, "A/B fixed connector")
        x += 3.02 if idx < 3 else 0
    add_label_box(slide, 1.02, 4.15, 5.30, 0.76, "baseline：release policy + fixed input\ncandidate：one marked block change", BLUE_PALE, BLUE, 20, NAVY, True, name="Baseline candidate rule")
    add_label_box(slide, 7.00, 4.15, 5.30, 0.76, "identity mismatch 或 null diff\ncomparison gate：待補／failure", RED_PALE, RED, 20, NAVY, True, name="A/B failure gate")
    add_field_pair_band(slide, [
        "`scenario`（情境）／`seed`（隨機種子）｜來源：scenario package／runner case｜作用：固定 job、window、model inputs 與 deterministic branch｜單位：structured record／integer｜判讀：identity mismatch 停止比較",
        "`predecessor`（前置結果）／`consequential diff`（具因果影響的差異）｜來源：freeze receipt／result-replay comparison｜作用：維持 lineage、證明修改進入 runtime path｜單位：identity／typed events｜判讀：null diff 不形成 energy result",
    ])
    note(slide, "本頁定義公平 A/B 的固定條件。`scenario`（情境）由 scenario package 提供，固定 job、service window 與 model inputs；`seed`（隨機種子）由 runner case 提供，固定 deterministic branch；兩者的值型態分別為 structured record 與 integer。`predecessor`（前置結果）由 freeze receipt 提供，維持 policy lineage。`consequential diff`（具因果影響的差異）由 result／replay comparison 產生，內容是 typed events 或 state／packet／service／energy 的實質變化。baseline 與 candidate 只允許一個 marked policy block 產生差異；identity mismatch 或 null diff 形成 comparison gate failure。來源：source-map.json、ADR-004、C-120 SDD、package workflow contract。")


def slide_p107(prs):
    slide = prepare_slide(prs, "P107｜student_policy.py 至 evidence 的因果骨架")
    lead(slide, "唯一受控的 policy edit 經 action、state、packet／service 形成可反駁 evidence")
    add_label_box(slide, 0.84, 1.52, 2.54, 0.60, "`student_policy.py`\n唯一 marked edit surface", BLUE_PALE, BLUE, 19, NAVY, True, name="Policy source")
    add_chevron(slide, 3.58, 1.62, 0.36, 0.42, BLUE, "Policy action connector")
    add_label_box(slide, 4.10, 1.52, 2.54, 0.60, "`action`\nWAIT／SLEEP／SEND family", GOLD_PALE, GOLD, 19, NAVY, True, name="Policy action")
    add_chevron(slide, 6.84, 1.62, 0.36, 0.42, GOLD, "Action state connector")
    add_label_box(slide, 7.36, 1.52, 2.54, 0.60, "`radio_state`\nstate interval／duration", TEAL_PALE, TEAL, 19, NAVY, True, name="Radio state")
    add_chevron(slide, 10.10, 1.62, 0.36, 0.42, TEAL, "State packet connector")
    add_label_box(slide, 10.62, 1.52, 1.90, 0.60, "packet／service\nendpoint J", PURPLE_PALE, PURPLE, 18.5, NAVY, True, name="Evidence outcome")
    add_box(slide, 0.94, 2.50, 5.56, 2.05, SOFT, NAVY, name="Policy code excerpt", line_width=1.4)
    add_text(slide, 1.18, 2.78, 5.08, 1.48,
             "if observation.contact_open:\n    return action\n\nstate interval → packet event → service field",
             20, INK, False, PP_ALIGN.LEFT, name="Policy code")
    add_box(slide, 6.84, 2.50, 5.56, 2.05, TEAL_PALE, TEAL, name="Causal interpretation", line_width=1.4)
    add_text(slide, 7.14, 2.78, 4.96, 1.48,
             "中介 state evidence 將 code change\n連到 packet outcome 與 service verdict。\n\nendpoint energy 依同一 scope 判讀。",
             20, NAVY, False, PP_ALIGN.LEFT, name="Causal interpretation text")
    add_label_box(slide, 1.04, 4.84, 11.18, 0.42, "prediction 保留；中介 evidence 無差異時，claim 維持待補並返回 policy API 定義", RED_PALE, RED, 18.5, NAVY, True, name="Causal recovery")
    add_field_pair_band(slide, [
        "`student_policy.py`（受控策略檔）｜來源：package runner｜作用：唯一 marked edit surface｜單位：UTF-8/LF source｜判讀：policy identity 連結 permitted edit 與 artifact",
        "`action`（策略動作）／`radio_state`（無線電狀態）／`packet outcome`（封包結果）｜來源：policy callback／runner state machine／packet ledger｜作用：形成 action → state → service evidence｜單位：enum／enum／category＋timestamp｜判讀：中介 evidence 支持因果句",
    ])
    note(slide, "本頁收束 control-to-evidence causal spine。`student_policy.py`（受控策略檔）由 package runner 提供，作用是限制唯一 marked edit surface，值型態為 UTF-8/LF source；policy identity 連結 permitted edit 與 generated artifact。`action`（策略動作）由 policy callback 產生，值型態 enum；`radio_state`（無線電狀態）由 runner state machine 產生，描述 state interval 與 duration；`packet outcome`（封包結果）由 packet ledger 產生，連結 delivery、expiry 與 retry。三者共同形成 action → state → packet／service → endpoint energy 的可反駁因果句。prediction 保留；中介 evidence 無差異時，claim 維持待補並返回 policy API 定義。來源：source-map.json、ADR-004、C-120 SDD、current package student policy contract。")


SLIDE_BUILDERS = [
    slide_p098, slide_p099, slide_p100, slide_p101, slide_p102,
    slide_p103, slide_p104, slide_p105, slide_p106, slide_p107,
]


def source_map() -> list[dict]:
    data = json.loads(SOURCE_MAP_SOURCE.read_text(encoding="utf-8"))
    selected = data[:10]
    if [item["page_id"] for item in selected] != [f"P{n:03d}" for n in range(98, 108)]:
        raise RuntimeError("appendix source map does not start with P098-P107")
    return selected


def emit_source_outputs() -> None:
    selected = source_map()
    SOURCE_MAP_OUTPUT.write_text(json.dumps(selected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = {
        "deck": OUTPUT.stem,
        "page_ids": [item["page_id"] for item in selected],
        "source_map": str(SOURCE_MAP_SOURCE),
        "construction_route": "python-pptx from educate.pptx source slide 2 / slideLayout2.xml",
        "template": str(TEMPLATE),
        "template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        "fonts": {"cjk": "標楷體", "latin": "Times New Roman"},
        "title_pt": 28,
        "primary_text_pt": 24,
        "label_text_pt": "18-22",
        "background_author_fill": "unset",
        "current_evidence": "待補：TLE epoch/raw bytes, dB source/timing, and frozen counterexample trace are not in the current package evidence record",
        "donor_policy": "topic donor only; no donor bytes, screenshots, numbers, master, or stale semantics",
        "donor_sources": [item["donor_source"] for item in selected],
    }
    MANIFEST_OUTPUT.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def build() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    if not SOURCE_MAP_SOURCE.exists():
        raise FileNotFoundError(SOURCE_MAP_SOURCE)
    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)
    for builder in SLIDE_BUILDERS:
        builder(prs)
    prs.core_properties.title = "LoRaEnergySim + LEO Appendix V2 P098-P107"
    prs.core_properties.subject = "TLE, model lineage, geometry/quality gates, service, energy boundary, and causal evidence"
    prs.core_properties.author = "OpenAI Codex"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    overlay_template_parts(OUTPUT)
    add_native_math(OUTPUT, 8)
    emit_source_outputs()
    BUILD_REPORT.write_text(json.dumps({
        "status": "INITIAL_EDITABLE_REVIEW",
        "output": str(OUTPUT),
        "owned_workspace": str(OWNED),
        "template": str(TEMPLATE),
        "template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        "slide_count": len(SLIDE_BUILDERS),
        "page_ids": [f"P{n:03d}" for n in range(98, 108)],
        "layout_contract": "all authored slides use source slide 2 / slideLayout2.xml",
        "background_fill": "unset",
        "speaker_notes": len(SLIDE_BUILDERS),
        "native_office_math": {"slide": "P105", "count": 1},
        "current_evidence": "neutral待補 states used where current source/trace evidence is absent",
        "visual_rendering": "deferred: server has no LibreOffice executable",
        "powerpoint_reopen": "deferred: controller environment",
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "INITIAL_EDITABLE_REVIEW", "output": str(OUTPUT), "slides": len(SLIDE_BUILDERS)}, ensure_ascii=False))


if __name__ == "__main__":
    build()
