#!/usr/bin/env python3
"""Build the owned Part B V2 module P028--P045.

The module is authored from the exact educate template with python-pptx.  It
uses only slide layout 2, keeps the template master/background/footer intact,
and adds the two required formulas as native Office Math after the editable
DrawingML content has been created.
"""

from __future__ import annotations

import copy
import hashlib
import json
import os
import posixpath
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


REPO = Path("/home/u24/demo/leo-beam-sim")
OWNED = REPO / "courseware/c120-lora-leo-deck/alt-skill-deck/part-b-v2-p028-p045"
ALT_ROOT = OWNED.parent
LATEST_ROOT = ALT_ROOT / "latest"
WORK_ROOT = Path(os.environ.get("C120_PART_B_WORK_ROOT", "/tmp/c120-style-b028-p045"))
TEMPLATE = OWNED / "sources/educate.pptx"
OUTPUT = WORK_ROOT / "LoRaEnergySim-LEO-ALT-PART-B-V2-P028-P045-REVIEW.pptx"
PUBLISH_TARGET = LATEST_ROOT / "LoRaEnergySim-LEO-ALT-PART-B-V2-P028-P045-REVIEW.pptx"
BUILD_DIR = WORK_ROOT / "build"
SOURCE_DIR = OWNED / "sources"
QA_DIR = OWNED / "qa"
RENDER_DIR = WORK_ROOT / "renders"

EVIDENCE_DIR = ALT_ROOT / "current-evidence/course-20260811/.playwright-cli"
METRICS_IMAGE = EVIDENCE_DIR / "element-2026-08-11T03-06-45-152Z.png"
REPLAY_IMAGE = EVIDENCE_DIR / "element-2026-08-11T03-06-47-569Z.png"
LEDGER_IMAGE = EVIDENCE_DIR / "element-2026-08-11T03-06-50-920Z.png"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

SLIDE_CX = 12_192_000
SLIDE_CY = 6_858_000

NAVY = "25356B"
INK = "25324A"
MUTED = "65758B"
BLUE = "4967C2"
BLUE_PALE = "ECF1FE"
TEAL = "007374"
TEAL_PALE = "E8F7F4"
PURPLE = "660066"
PURPLE_PALE = "F4EDF7"
GOLD = "846826"
GOLD_PALE = "FDF9E2"
RED = "992A2A"
RED_PALE = "FDEFEF"
WHITE = "FFFFFF"
LINE = "D7DEEA"
SOFT = "F7F8FC"
GREEN = "2E6B52"

CLAIM_BOUNDARY = "SIMULATED TEACHING DATA｜NOT LIVE｜NOT MEASURED｜NOT CANONICAL-PARITY-VERIFIED"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_fonts(run) -> None:
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("latin", "Times New Roman"), ("ea", "標楷體"), ("cs", "Times New Roman")):
        child = rpr.find(qn(f"a:{tag}"))
        if child is None:
            child = OxmlElement(f"a:{tag}")
            rpr.append(child)
        child.set("typeface", face)


def add_marked_runs(paragraph, text: str, size: float, color: str,
                    bold: bool = False, italic_all: bool = False) -> None:
    """Backtick spans are variables, field names, code, or formula tokens."""
    pieces = re.split(r"(`[^`]+`)", text)
    for piece in pieces:
        if not piece:
            continue
        marked = piece.startswith("`") and piece.endswith("`")
        value = piece[1:-1] if marked else piece
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = marked or italic_all
        run.font.color.rgb = rgb(color)
        set_run_fonts(run)


def write_text(shape, text: str, size: float = 24, color: str = INK,
               bold: bool = False, align=PP_ALIGN.LEFT,
               valign=MSO_ANCHOR.MIDDLE, margins=(0.10, 0.05, 0.10, 0.05),
               line_spacing: float = 1.0, italic_all: bool = False) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    for idx, line in enumerate(text.split("\n")):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        add_marked_runs(paragraph, line, size, color, bold, italic_all)


def add_text(slide, x: float, y: float, w: float, h: float, text: str,
             size: float = 24, color: str = INK, bold: bool = False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margins=(0.02, 0.01, 0.02, 0.01), line_spacing: float = 1.0,
             name: str = "Text", italic_all: bool = False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    write_text(shape, text, size, color, bold, align, valign, margins, line_spacing, italic_all)
    return shape


def add_box(slide, x: float, y: float, w: float, h: float, fill: str,
            line: str | None = LINE, radius: bool = True,
            name: str = "Box", line_width: float = 1.2):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_label_box(slide, x: float, y: float, w: float, h: float, text: str,
                  fill: str, line: str, size: float = 20, color: str = INK,
                  bold: bool = False, align=PP_ALIGN.CENTER,
                  name: str = "Label", margins=(0.10, 0.04, 0.10, 0.04)):
    shape = add_box(slide, x, y, w, h, fill, line, name=name)
    write_text(shape, text, size, color, bold, align, MSO_ANCHOR.MIDDLE,
               margins=margins, line_spacing=0.94)
    return shape


def add_chevron(slide, x: float, y: float, w: float = 0.34, h: float = 0.44,
                color: str = BLUE, name: str = "Mechanism connector"):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_field_band(slide, field: str, chinese: str, source: str, purpose: str,
                   unit: str, interpretation: str, y: float = 5.64) -> None:
    band = add_box(slide, 0.82, y, 11.72, 0.72, SOFT, BLUE,
                   name=f"Field contract {field}", line_width=1.0)
    write_text(
        band,
        f"{chinese} `{field}`（欄位）｜來源：{source}｜作用：{purpose}\n"
        f"單位：{unit}｜判讀：{interpretation}",
        16.0, NAVY, False, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
        margins=(0.16, 0.02, 0.16, 0.02), line_spacing=0.78,
    )


def add_claim(slide, text: str = CLAIM_BOUNDARY, y: float = 6.40) -> None:
    add_text(slide, 0.92, y, 11.42, 0.26, text, 16, RED, True,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0),
             name="Claim boundary")


def add_picture(slide, path: Path, x: float, y: float, w: float, h: float,
                name: str):
    if not path.exists():
        raise FileNotFoundError(path)
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    picture.name = name
    return picture


def delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def prepare_slide(prs: Presentation, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    if title_shape is None:
        raise RuntimeError("slideLayout2 title placeholder missing")
    for placeholder in list(slide.placeholders):
        if placeholder._element is title_shape._element or placeholder.placeholder_format.type == PP_PLACEHOLDER.TITLE:
            continue
        if placeholder.placeholder_format.type not in {
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.SLIDE_NUMBER,
        }:
            delete_shape(placeholder)
    title_shape.left = Inches(0.718057)
    title_shape.top = Inches(0.204514)
    title_shape.width = Inches(10.34861)
    title_shape.height = Inches(0.525)
    title_shape.name = "Native layout2 title"
    write_text(title_shape, title, 28, NAVY, True, PP_ALIGN.LEFT,
               MSO_ANCHOR.MIDDLE, margins=(0.02, 0.0, 0.02, 0.0))
    return slide


def remove_all_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def overlay_template_parts(path: Path) -> None:
    """Restore template-owned inherited parts without dropping authored media."""
    owned_prefixes = (
        "ppt/slideLayouts/", "ppt/slideMasters/", "ppt/theme/", "ppt/notesMasters/",
    )
    with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(path, "r") as old:
        data = {info.filename: old.read(info.filename) for info in old.infolist()}
        infos = {info.filename: copy.copy(info) for info in old.infolist()}
        for info in source.infolist():
            if info.filename.startswith(owned_prefixes):
                data[info.filename] = source.read(info.filename)
                infos[info.filename] = copy.copy(info)
    overlay = path.with_suffix(".overlay.pptx")
    with zipfile.ZipFile(overlay, "w", zipfile.ZIP_DEFLATED) as out:
        for name, payload in data.items():
            out.writestr(infos[name], payload)
    overlay.replace(path)


def _math_r(parent, text: str, size: int = 3000, italic: bool = False):
    run = ET.SubElement(parent, f"{{{M_NS}}}r")
    rpr = ET.SubElement(run, f"{{{A_NS}}}rPr", {
        "lang": "en-US", "altLang": "en-US", "sz": str(size), "dirty": "0",
    })
    if italic:
        rpr.set("i", "1")
    ET.SubElement(rpr, f"{{{A_NS}}}latin", {"typeface": "Cambria Math"})
    ET.SubElement(run, f"{{{M_NS}}}t").text = text
    return run


def _math_sub(parent, base: str, sub: str):
    node = ET.SubElement(parent, f"{{{M_NS}}}sSub")
    ET.SubElement(node, f"{{{M_NS}}}sSubPr")
    e1 = ET.SubElement(node, f"{{{M_NS}}}e")
    _math_r(e1, base, italic=True)
    e2 = ET.SubElement(node, f"{{{M_NS}}}e")
    _math_r(e2, sub)
    return node


def _formula_body(kind: str):
    para = ET.Element(f"{{{M_NS}}}oMathPara")
    omath = ET.SubElement(para, f"{{{M_NS}}}oMath")
    if kind == "endpoint":
        _math_sub(omath, "E", "endpoint")
        _math_r(omath, " = ")
        nary = ET.SubElement(omath, f"{{{M_NS}}}nary")
        nary_pr = ET.SubElement(nary, f"{{{M_NS}}}naryPr")
        ET.SubElement(nary_pr, f"{{{M_NS}}}chr", {f"{{{M_NS}}}val": "∑"})
        ET.SubElement(nary_pr, f"{{{M_NS}}}limLoc", {f"{{{M_NS}}}val": "undOvr"})
        ET.SubElement(nary_pr, f"{{{M_NS}}}grow", {f"{{{M_NS}}}val": "1"})
        sub = ET.SubElement(nary, f"{{{M_NS}}}sub")
        _math_r(sub, "s", italic=True)
        _math_r(sub, " ∈ ")
        _math_r(sub, "S", italic=True)
        ET.SubElement(nary, f"{{{M_NS}}}sup")
        body = ET.SubElement(nary, f"{{{M_NS}}}e")
        _math_sub(body, "P", "s")
        _math_r(body, " ")
        _math_sub(body, "t", "s")
    elif kind == "efficiency":
        _math_sub(omath, "η", "E")
        _math_r(omath, " = ")
        frac = ET.SubElement(omath, f"{{{M_NS}}}f")
        ET.SubElement(ET.SubElement(frac, f"{{{M_NS}}}fPr"), f"{{{M_NS}}}type", {f"{{{M_NS}}}val": "bar"})
        num = ET.SubElement(frac, f"{{{M_NS}}}num")
        _math_sub(num, "D", "delivered")
        den = ET.SubElement(frac, f"{{{M_NS}}}den")
        _math_sub(den, "E", "endpoint")
    else:
        raise ValueError(kind)
    return para


def add_native_math(path: Path, formula_specs: dict[int, str]) -> None:
    """Insert native OMML shapes into the already-valid editable package."""
    ET.register_namespace("a", A_NS)
    ET.register_namespace("m", M_NS)
    ET.register_namespace("p", P_NS)
    ET.register_namespace("a14", A14_NS)
    with zipfile.ZipFile(path, "r") as archive:
        infos = [copy.copy(info) for info in archive.infolist()]
        data = {info.filename: archive.read(info.filename) for info in infos}
    for slide_number, kind in formula_specs.items():
        name = f"ppt/slides/slide{slide_number}.xml"
        root = ET.fromstring(data[name])
        sp_tree = root.find(f"./{{{P_NS}}}cSld/{{{P_NS}}}spTree")
        if sp_tree is None:
            raise RuntimeError(f"formula slide {slide_number} has no shape tree")
        ids = []
        for node in sp_tree.iter(f"{{{P_NS}}}cNvPr"):
            if node.get("id", "").isdigit():
                ids.append(int(node.get("id")))
        next_id = max(ids, default=1) + 1
        shape = ET.Element(f"{{{P_NS}}}sp")
        nv = ET.SubElement(shape, f"{{{P_NS}}}nvSpPr")
        ET.SubElement(nv, f"{{{P_NS}}}cNvPr", {
            "id": str(next_id), "name": f"Native Office Math {kind}",
            "title": kind, "descr": kind,
        })
        c_nv_sp_pr = ET.SubElement(nv, f"{{{P_NS}}}cNvSpPr", {"hidden": "1"})
        ET.SubElement(c_nv_sp_pr, f"{{{A_NS}}}spLocks", {"noGrp": "1"})
        ET.SubElement(nv, f"{{{P_NS}}}nvPr")
        sp_pr = ET.SubElement(shape, f"{{{P_NS}}}spPr")
        xfrm = ET.SubElement(sp_pr, f"{{{A_NS}}}xfrm")
        ET.SubElement(xfrm, f"{{{A_NS}}}off", {"x": "0", "y": "0"})
        ET.SubElement(xfrm, f"{{{A_NS}}}ext", {"cx": "1", "cy": "1"})
        ET.SubElement(sp_pr, f"{{{A_NS}}}noFill")
        ln = ET.SubElement(sp_pr, f"{{{A_NS}}}ln")
        ET.SubElement(ln, f"{{{A_NS}}}noFill")
        tx_body = ET.SubElement(shape, f"{{{P_NS}}}txBody")
        ET.SubElement(tx_body, f"{{{A_NS}}}bodyPr", {"anchor": "ctr", "wrap": "none"})
        ET.SubElement(tx_body, f"{{{A_NS}}}lstStyle")
        para = ET.SubElement(tx_body, f"{{{A_NS}}}p")
        ET.SubElement(para, f"{{{A_NS}}}pPr", {"algn": "ctr"})
        wrapper = ET.SubElement(para, f"{{{A14_NS}}}m")
        wrapper.append(_formula_body(kind))
        ET.SubElement(para, f"{{{A_NS}}}endParaRPr", {"lang": "en-US", "sz": "3000", "dirty": "0"})
        sp_tree.append(shape)
        data[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    temp = path.with_suffix(".math.pptx")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as out:
        for info in infos:
            out.writestr(info, data[info.filename])
    temp.replace(path)


def _add_title_lead(slide, text: str, color: str = NAVY):
    return add_text(slide, 0.98, 1.00, 11.38, 0.42, text, 24, color, True,
                    PP_ALIGN.CENTER, name="Teaching proposition")


def _node(slide, x, y, w, h, title, body, fill, line, size=20, title_size=21,
          name="Teaching node"):
    card = add_box(slide, x, y, w, h, fill, line, name=name)
    add_text(slide, x + 0.10, y + 0.12, w - 0.20, 0.38, title, title_size, line, True,
             PP_ALIGN.CENTER, name=f"{name} heading")
    add_text(slide, x + 0.14, y + 0.58, w - 0.28, h - 0.70, body, size, INK, False,
             PP_ALIGN.CENTER, name=f"{name} body", margins=(0.02, 0.01, 0.02, 0.01),
             line_spacing=0.92)
    return card


def slide_028(prs):
    slide = prepare_slide(prs, "P028｜Radio state 與停留時間")
    _add_title_lead(slide, "策略 action 是決策輸出；state interval 才留下時間與能量證據")
    states = [
        ("`SLEEP`｜低功耗休息", "窗口關閉或空檔\n`state interval` → sleep J", BLUE_PALE, BLUE),
        ("`WAKE`｜喚醒轉換", "送出前的延遲\ntransition → wake J", GOLD_PALE, GOLD),
        ("`PROCESS`｜處理", "封包處理區間\nprocess → process J", PURPLE_PALE, PURPLE),
        ("`TX`｜發射", "無線發射區間\nTX → tx J", TEAL_PALE, TEAL),
        ("`RX`｜接收", "回覆接收區間\nRX → rx J", RED_PALE, RED),
    ]
    x = 0.78
    for index, (heading, body, fill, line) in enumerate(states):
        _node(slide, x, 1.78, 2.15, 1.55, heading, body, fill, line, size=18.5, title_size=19,
              name=f"Radio state {index + 1}")
        if index < len(states) - 1:
            add_chevron(slide, x + 2.19, 2.32, 0.30, 0.38, line, "State mechanism connector")
        x += 2.48
    _node(slide, 3.70, 3.65, 5.95, 1.15, "`WAIT`｜清醒閒置", "policy 回傳 WAIT 時，radio 保持 awake idle；10 s clock step 會累積 awake_idle J。", GOLD_PALE, GOLD, size=20, title_size=21, name="WAIT state explanation")
    add_text(slide, 1.02, 5.08, 11.16, 0.34, "state interval × state power → energy bucket；action 名稱只標記一次決策。", 22, NAVY, True, PP_ALIGN.CENTER, name="State takeaway")
    add_field_band(slide, "radio_state", "無線狀態", "runner `events`", "標記每個停留區間", "state＋s", "把 interval 連到 `energy_breakdown_j`", y=5.62)


def slide_029(prs):
    slide = prepare_slide(prs, "P029｜Lab A 空檔策略的 service／energy gate")
    _add_title_lead(slide, "固定工作與服務窗口，只變更空檔 action 的因果假說")
    add_label_box(slide, 0.84, 1.58, 2.12, 0.78, "固定條件\nscenario／seed／traffic／window", SOFT, NAVY, 19, NAVY, True, name="Lab A fixed conditions")
    add_chevron(slide, 3.14, 1.78, 0.34, 0.38, NAVY)
    _node(slide, 3.58, 1.48, 2.56, 1.02, "Baseline", "`REST_DURING_GAP = SLEEP`\n低功耗休息＋後續 WAKE", BLUE_PALE, BLUE, size=18, title_size=20, name="Lab A baseline")
    add_chevron(slide, 6.36, 1.78, 0.34, 0.38, GOLD)
    _node(slide, 6.80, 1.48, 2.56, 1.02, "Candidate", "`REST_DURING_GAP = WAIT`\nawake idle＋部分 wake 避免", GOLD_PALE, GOLD, size=18, title_size=20, name="Lab A candidate")
    add_chevron(slide, 9.58, 1.78, 0.34, 0.38, TEAL)
    _node(slide, 10.02, 1.48, 2.48, 1.02, "Evidence", "state → packet → service → J", TEAL_PALE, TEAL, size=18, title_size=20, name="Lab A evidence path")
    add_label_box(slide, 0.98, 3.05, 3.68, 1.36, "SLEEP\nidle power ↓\nWAKE latency／energy ↑", BLUE_PALE, BLUE, 22, NAVY, True, name="SLEEP mechanism")
    add_label_box(slide, 4.82, 3.05, 3.68, 1.36, "WAIT\nawake-idle time ↑\n部分 WAKE transition ↓", GOLD_PALE, GOLD, 22, NAVY, True, name="WAIT mechanism")
    add_label_box(slide, 8.66, 3.05, 3.68, 1.36, "Service gate\n`delivered`／`deadline`／`freshness`\n通過後再讀 endpoint J", TEAL_PALE, TEAL, 20, NAVY, True, name="Lab A service gate")
    add_text(slide, 1.05, 4.84, 11.22, 0.48, "改前預測：WAIT 可能減少 WAKE，但 awake idle 可能增加；結果允許反駁方向。", 21, RED, True, PP_ALIGN.CENTER, name="Lab A prediction")
    add_field_band(slide, "service_pass", "服務布林狀態", "result `summary`", "表示 required packet、期限與 freshness 是否共同通過", "true／false", "J 的比較需保留服務結果", y=5.62)


def slide_030(prs):
    slide = prepare_slide(prs, "P030｜封包生命週期：送出與交付分開判定")
    _add_title_lead(slide, "一次 policy decision 只建立 transmission attempt；交付要走完整 packet outcome")
    phases = [
        ("產生", "`generated`\npacket 建立", BLUE_PALE, BLUE),
        ("排隊", "`queue`\n等待合法窗口", SOFT, NAVY),
        ("嘗試", "`attempt`\nPROCESS／TX／RX", GOLD_PALE, GOLD),
        ("失敗處理", "`collision` → `retry`\n可能再次嘗試", RED_PALE, RED),
        ("結果", "`delivered` 或 `expired`\n期限前後分流", TEAL_PALE, TEAL),
        ("服務判定", "`service_pass`\nrequired packet＋期限＋freshness", PURPLE_PALE, PURPLE),
    ]
    x = 0.72
    for index, (heading, body, fill, line) in enumerate(phases):
        _node(slide, x, 1.65, 1.92, 1.45, heading, body, fill, line, size=17.5, title_size=18.5, name=f"Packet phase {index + 1}")
        if index < len(phases) - 1:
            add_chevron(slide, x + 1.95, 2.15, 0.25, 0.34, line, "Packet mechanism connector")
        x += 2.08
    add_label_box(slide, 1.02, 3.55, 3.42, 1.20, "產生段\npacket identity、payload bit、deadline", BLUE_PALE, BLUE, 20, NAVY, True, name="Packet generation band")
    add_label_box(slide, 4.95, 3.55, 3.42, 1.20, "傳輸段\nattempt、collision、retry、PROCESS／TX／RX", GOLD_PALE, GOLD, 19, GOLD, True, name="Packet transmission band")
    add_label_box(slide, 8.88, 3.55, 3.42, 1.20, "結果段\ndelivered、expired、deadline、service verdict", TEAL_PALE, TEAL, 19, TEAL, True, name="Packet result band")
    add_text(slide, 1.00, 5.02, 11.30, 0.42, "SEND 次數是 action evidence；服務判定使用 packet ledger 的完整 outcome。", 21, NAVY, True, PP_ALIGN.CENTER, name="Packet takeaway")
    add_field_band(slide, "attempted_packets", "嘗試封包數", "result `summary`", "統計進入無線嘗試的封包數", "count", "與 `retransmissions`、`unique_delivered_packets`、`expired_packets` 一起判讀", y=5.62)


def slide_031(prs):
    slide = prepare_slide(prs, "P031｜Service gate 與 endpoint J 的比較順序")
    _add_title_lead(slide, "同一比較邊界 → service gate → endpoint energy；順序決定 claim 強度")
    add_label_box(slide, 0.92, 1.52, 3.45, 1.44, "比較邊界\nscenario／seed／traffic／window\nendpoint scope 一致", BLUE_PALE, BLUE, 20, NAVY, True, name="Comparison boundary gate")
    add_chevron(slide, 4.54, 2.04, 0.42, 0.48, BLUE)
    add_label_box(slide, 5.10, 1.52, 3.45, 1.44, "Service gate\n`delivered_bits`\n`deadline_pass`／`freshness_status`", GOLD_PALE, GOLD, 20, NAVY, True, name="Service gate")
    add_chevron(slide, 8.72, 2.04, 0.42, 0.48, GOLD)
    add_label_box(slide, 9.28, 1.52, 3.12, 1.44, "Endpoint energy\n`endpoint_energy_j`\n整段 state ledger 的總和", TEAL_PALE, TEAL, 20, NAVY, True, name="Endpoint energy gate")
    add_label_box(slide, 1.03, 3.45, 5.33, 1.20, "service_pass = true\nJ 可在相同工作邊界內比較\n效率比值另看 bit／J", TEAL_PALE, TEAL, 22, NAVY, True, name="Service pass outcome")
    add_label_box(slide, 6.96, 3.45, 5.33, 1.20, "service_pass = false\nJ 只描述 trade-off evidence\n較低 J 不自動升格為節能", RED_PALE, RED, 22, NAVY, True, name="Service fail outcome")
    add_text(slide, 1.08, 5.03, 11.16, 0.40, "比較句型：在相同 ______ 下，candidate 的 service ______；endpoint J 的差異解讀為 ______。", 20, NAVY, True, PP_ALIGN.CENTER, name="Comparison sentence")
    add_field_band(slide, "freshness_status", "新鮮度狀態", "result `summary`", "表示 required packet 是否在 freshness limit 內完成", "fresh／stale／expired／not-applicable", "與 deadline、delivery 一起決定 service gate", y=5.62)


def slide_032(prs):
    slide = prepare_slide(prs, "P032｜W 是瞬間功率，J 是整段累積")
    _add_title_lead(slide, "同一個功率值，停留時間不同就會留下不同的 endpoint energy")
    add_label_box(slide, 0.95, 1.55, 2.20, 2.86, "功率 `P`\nW\n某個 state 當下的高度", BLUE_PALE, BLUE, 24, NAVY, True, name="Power axis")
    add_box(slide, 3.45, 1.85, 3.55, 2.10, SOFT, LINE, name="Power area chart")
    add_box(slide, 3.80, 2.95, 0.82, 0.72, TEAL, None, radius=False, name="Low power bar")
    add_box(slide, 4.78, 2.32, 0.82, 1.35, GOLD, None, radius=False, name="Wake power bar")
    add_box(slide, 5.76, 2.15, 0.82, 1.52, PURPLE, None, radius=False, name="Process power bar")
    add_text(slide, 3.72, 3.72, 3.05, 0.22, "sleep｜wake｜process", 17.5, MUTED, False, PP_ALIGN.CENTER, name="Power labels")
    add_chevron(slide, 7.35, 2.58, 0.42, 0.48, NAVY)
    add_label_box(slide, 7.95, 1.55, 4.25, 2.86, "時間 `t`\ns\nstate interval 的長度\n\n功率 × 時間 = 能量\nW × s = J", GOLD_PALE, GOLD, 24, NAVY, True, name="Time area explanation")
    add_label_box(slide, 1.03, 4.82, 5.36, 0.55, "`awake_idle`、`sleep`、`wake`、`process`、`tx`、`rx` 各自累積 J", TEAL_PALE, TEAL, 20, NAVY, True, name="Energy buckets")
    add_label_box(slide, 6.94, 4.82, 5.36, 0.55, "固定 10 s clock step；收發流程另加入 transition energy", PURPLE_PALE, PURPLE, 20, NAVY, True, name="Clock mechanism")
    add_field_band(slide, "energy_breakdown_j", "能量分桶", "result artifact", "按 state／transition 保存 endpoint energy", "J per bucket", "將總 J 拆回可解釋的停留區間", y=5.62)


def slide_033(prs):
    slide = prepare_slide(prs, "P033｜Endpoint energy 公式的邊界")
    _add_title_lead(slide, "公式描述 endpoint state 的累積方式；artifact 提供這次 run 的數值")
    add_label_box(slide, 0.92, 1.52, 2.48, 2.20, "`P_s`\nstate power\nW\n來源：energy profile", BLUE_PALE, BLUE, 24, NAVY, True, name="Formula numerator input")
    add_label_box(slide, 9.92, 1.52, 2.48, 2.20, "`t_s`\n停留時間\ns\n來源：state interval", GOLD_PALE, GOLD, 24, NAVY, True, name="Formula time input")
    add_box(slide, 3.72, 1.64, 5.84, 1.88, SOFT, PURPLE, name="Native formula frame", line_width=1.6)
    add_text(slide, 3.92, 2.10, 5.44, 0.74, "`E_endpoint = Σ P_s t_s`", 30, PURPLE, True, PP_ALIGN.CENTER, name="Editable formula transcription", italic_all=True)
    add_text(slide, 4.05, 3.72, 5.20, 0.35, "可編輯 Office Math｜公式 slide 6", 18, PURPLE, True, PP_ALIGN.CENTER, name="Formula contract label")
    add_label_box(slide, 1.02, 4.36, 5.42, 0.88, "endpoint scope\nradio／processing 的 state 與 transition", TEAL_PALE, TEAL, 21, NAVY, True, name="Endpoint scope")
    add_label_box(slide, 6.88, 4.36, 5.42, 0.88, "scope boundary\nsatellite／gateway／whole-system wall-plug 不在此式", RED_PALE, RED, 20, NAVY, True, name="Scope boundary")
    add_field_band(slide, "endpoint_energy_j", "端點能量", "result `summary`", "提供同一 endpoint boundary 的總和", "J", "由 formula 對照 `energy_breakdown_j`", y=5.62)


def slide_034(prs):
    slide = prepare_slide(prs, "P034｜Energy efficiency 的分子與分母")
    _add_title_lead(slide, "bit/J 是同一 endpoint boundary 內的效率比值；service gate 仍獨立")
    add_label_box(slide, 0.96, 1.46, 3.05, 1.20, "`D_delivered`\n交付資料\n來源：result summary\n單位：bit", BLUE_PALE, BLUE, 22, NAVY, True, name="Delivered numerator")
    add_label_box(slide, 9.32, 1.46, 3.05, 1.20, "`E_endpoint`\n端點能量\n來源：result summary\n單位：J", TEAL_PALE, TEAL, 22, NAVY, True, name="Endpoint denominator")
    add_box(slide, 4.42, 1.42, 4.48, 1.40, SOFT, PURPLE, name="Native efficiency formula frame", line_width=1.6)
    add_text(slide, 4.57, 1.84, 4.18, 0.58, "`η_E = D_delivered / E_endpoint`", 28, PURPLE, True, PP_ALIGN.CENTER, name="Editable efficiency transcription", italic_all=True)
    add_text(slide, 4.56, 2.90, 4.20, 0.30, "bit ÷ J = bit/J｜可編輯 Office Math", 18, PURPLE, True, PP_ALIGN.CENTER, name="Efficiency unit label")
    add_label_box(slide, 1.00, 3.62, 5.38, 1.12, "高 bit/J 且 service_pass = false\n保留為效率觀察，不作服務成功判定", GOLD_PALE, GOLD, 21, NAVY, True, name="Efficiency caveat")
    add_label_box(slide, 6.92, 3.62, 5.38, 1.12, "service_pass、deadline_pass\n依 required packet 與期限另行判定", RED_PALE, RED, 21, NAVY, True, name="Efficiency gate")
    add_text(slide, 1.10, 5.07, 11.08, 0.40, "效率比值的分子、分母、單位與 service verdict 必須同時保留。", 21, NAVY, True, PP_ALIGN.CENTER, name="Efficiency takeaway")
    add_field_band(slide, "endpoint_energy_efficiency_bits_per_j", "端點能量效率", "result `summary`", "計算 delivered bit 與 endpoint J 的比值", "bit/J", "不取代 `service_pass` 或 `deadline_pass`", y=5.62)


def slide_035(prs):
    slide = prepare_slide(prs, "P035｜來源、模型、假設與結果分層")
    _add_title_lead(slide, "同一筆數字要沿來源、模型、固定假設與 result artifact 回溯")
    layers = [
        ("Source｜來源", "上游 repository／commit 是 provenance 參考", SOFT, NAVY, "reference"),
        ("Model｜模型", "coherent-course-simulated-adapter\n固定 runner 行為", BLUE_PALE, BLUE, "adapter"),
        ("Course｜課程假設", "scenario、contact window、traffic、endpoint profile", GOLD_PALE, GOLD, "fixed input"),
        ("Result｜執行結果", "JSON result＋同 run replay\nclaim boundary 隨 artifact 保存", TEAL_PALE, TEAL, "artifact"),
    ]
    y = 1.42
    for index, (heading, body, fill, line, connector) in enumerate(layers):
        add_label_box(slide, 1.05, y, 3.15, 0.92, heading, fill, line, 22, NAVY, True, name=f"Provenance layer {index + 1}")
        add_label_box(slide, 4.55, y, 5.22, 0.92, body, SOFT, line, 20, NAVY, True, name=f"Provenance layer explanation {index + 1}")
        add_label_box(slide, 10.08, y + 0.15, 2.18, 0.58, connector, fill, line, 18, line, True, name=f"Provenance connector {index + 1}")
        if index < len(layers) - 1:
            add_chevron(slide, 6.95, y + 1.02, 0.34, 0.28, line, "Provenance mechanism connector")
        y += 1.08
    add_text(slide, 1.02, 5.72, 11.20, 0.34, "`runner_provenance.upstream_execution = false`；資料類型維持 coherent simulated result。", 20, RED, True, PP_ALIGN.CENTER, name="Provenance claim")
    add_claim(slide)


def slide_036(prs):
    slide = prepare_slide(prs, "P036｜執行結果如何進入重播與工作簿")
    _add_title_lead(slide, "同一條 identity record 連結 result、replay 與 workbook 的比較資格")
    nodes = [
        ("RUN｜執行", "runner 以 scenario、case、policy 執行\n輸出 `result.json` 與同目錄 replay", BLUE_PALE, BLUE, "生成 result"),
        ("REPLAY｜事件重播", "依同一次 run 的 events\n顯示 action、state、queue、packet", TEAL_PALE, TEAL, "讀取 events"),
        ("WORKBOOK｜比較紀錄", "保存 prediction、baseline、candidate、freeze、withheld\n與 result／replay lineage", GOLD_PALE, GOLD, "寫入紀錄"),
    ]
    x = 0.90
    for index, (heading, body, fill, line, label) in enumerate(nodes):
        _node(slide, x, 1.55, 3.55, 1.68, heading, body, fill, line, size=19, title_size=21, name=f"Identity object {index + 1}")
        if index < len(nodes) - 1:
            add_chevron(slide, x + 3.68, 2.11, 0.38, 0.42, line, "Identity mechanism connector")
            add_text(slide, x + 3.42, 2.60, 0.92, 0.26, label, 16.5, line, True, PP_ALIGN.CENTER, name=f"Identity connector label {index + 1}")
        x += 4.20
    add_label_box(slide, 0.98, 3.78, 5.56, 1.08, "`scenario_id`（情境識別碼）\n來源：scenario／result｜單位：字串識別碼\n作用：指出固定情境，不是能量值", BLUE_PALE, BLUE, 19, NAVY, True, name="Scenario identity field")
    add_label_box(slide, 6.82, 3.78, 5.56, 1.08, "一致性欄位\n`run_id`、policy identity、units、provenance\n來源：result／replay／workbook", PURPLE_PALE, PURPLE, 19, NAVY, True, name="Identity consistency fields")
    add_text(slide, 1.05, 5.18, 11.18, 0.40, "identity mismatch 使匯入維持 fail closed；配對資料要回到 matching artifact。", 21, RED, True, PP_ALIGN.CENTER, name="Identity gate outcome")
    add_claim(slide, y=6.40)


def slide_037(prs):
    slide = prepare_slide(prs, "P037｜Baseline：固定條件下的對照執行")
    _add_title_lead(slide, "Baseline 是固定 scenario、seed 與原始 policy 的 control execution")
    columns = [
        ("固定輸入", "scenario\nseed\ntraffic\nwindow\nendpoint scope", BLUE_PALE, BLUE),
        ("唯一變動", "目前 lab 的\nmarked block\n一次一個 causal value", GOLD_PALE, GOLD),
        ("觀察輸出", "state duration\npacket outcome\nservice verdict\nendpoint J", TEAL_PALE, TEAL),
    ]
    x = 0.94
    for index, (heading, body, fill, line) in enumerate(columns):
        _node(slide, x, 1.54, 3.56, 2.28, heading, body, fill, line, size=22, title_size=23, name=f"Baseline rule {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.70, 2.43, 0.38, 0.44, line, "Baseline mechanism connector")
        x += 4.08
    add_label_box(slide, 1.02, 4.34, 5.30, 0.86, "Control condition\n兩次結果 identity 與 endpoint boundary 一致", SOFT, NAVY, 21, NAVY, True, name="Baseline control condition")
    add_label_box(slide, 6.98, 4.34, 5.30, 0.86, "Interpretation\n差異先回到 marked block，再讀 state／packet／service／J", PURPLE_PALE, PURPLE, 20, NAVY, True, name="Baseline interpretation")
    add_text(slide, 1.06, 5.45, 11.14, 0.36, "scenario、traffic 或 endpoint scope 改變時，A/B 不再是同一個 causal comparison。", 20, RED, True, PP_ALIGN.CENTER, name="Baseline warning")
    add_field_band(slide, "policy identity", "策略識別", "result `policy`", "指出此次執行使用的 policy surface", "識別碼", "與 predecessor、scenario、run_id 一起比對", y=5.86)


def slide_038(prs):
    slide = prepare_slide(prs, "P038｜LEO：changing-service-window trace")
    _add_title_lead(slide, "LEO 在本模組提供 changing-service-window trace，作為 endpoint 傳輸時機的輸入")
    segments = [
        ("closed", "窗口關閉\n只接受 SLEEP", "CBD5E1", INK, 1.20),
        ("contact-a", "窗口開啟\nquality 漸變", TEAL, WHITE, 2.25),
        ("closed", "窗口關閉\n服務機會消失", "CBD5E1", INK, 1.20),
        ("contact-b", "第二窗口\nquality 再變化", BLUE, WHITE, 2.25),
        ("closed", "窗口關閉\n回到休息", "CBD5E1", INK, 1.20),
    ]
    x = 0.93
    for index, (label, body, fill, color, width) in enumerate(segments):
        card = add_box(slide, x, 1.74, width, 1.30, fill, None, radius=False, name=f"LEO window segment {index + 1}")
        write_text(card, f"{label}\n{body}", 19, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margins=(0.04, 0.02, 0.04, 0.02), line_spacing=0.92)
        x += width + 0.08
    add_label_box(slide, 0.98, 3.55, 3.54, 1.20, "`contact_open`\n服務窗口布林狀態\nsource：fixed scenario trace", BLUE_PALE, BLUE, 21, NAVY, True, name="Contact open field")
    add_label_box(slide, 4.90, 3.55, 3.54, 1.20, "`quality_band`\n連線品質分級\nunit：0–3 category", GOLD_PALE, GOLD, 21, NAVY, True, name="Quality band field")
    add_label_box(slide, 8.82, 3.55, 3.54, 1.20, "`contact_remaining_s`\n窗口剩餘時間\nunit：s；影響 action 時機", TEAL_PALE, TEAL, 20, NAVY, True, name="Contact remaining field")
    add_text(slide, 1.07, 5.05, 11.16, 0.38, "window input 改變合法服務機會；policy 再由 observation 選擇 WAIT、SLEEP 或 SEND。", 20.5, NAVY, True, PP_ALIGN.CENTER, name="LEO takeaway")
    add_field_band(slide, "service_window", "服務窗口", "scenario quality trace", "定義合法傳輸區間與 quality context", "trace time／category", "用來解釋 action 是否具有服務機會", y=5.62)


def slide_039(prs):
    slide = prepare_slide(prs, "P039｜Lab A：等待空檔的 radio state")
    _add_title_lead(slide, "固定同一份工作，把 `REST_DURING_GAP` 的唯一 edit 交給 evidence 檢驗")
    add_label_box(slide, 0.92, 1.45, 2.40, 1.22, "Before edit\n`REST_DURING_GAP = SLEEP`\n`PACE_GAP_STEPS = 2`", BLUE_PALE, BLUE, 20, NAVY, True, name="Lab A before edit")
    add_chevron(slide, 3.54, 1.84, 0.38, 0.42, BLUE)
    add_label_box(slide, 4.08, 1.45, 2.40, 1.22, "Why edit\n測試空檔 state 的成本\nWAKE vs awake idle", GOLD_PALE, GOLD, 20, NAVY, True, name="Lab A why edit")
    add_chevron(slide, 6.70, 1.84, 0.38, 0.42, GOLD)
    add_label_box(slide, 7.24, 1.45, 2.40, 1.22, "Exact edit\n`REST_DURING_GAP = WAIT`\n只改 marked block", PURPLE_PALE, PURPLE, 20, NAVY, True, name="Lab A exact edit")
    add_chevron(slide, 9.86, 1.84, 0.38, 0.42, PURPLE)
    add_label_box(slide, 10.40, 1.45, 2.08, 1.22, "Evidence\nstate → packet\nservice → J", TEAL_PALE, TEAL, 20, NAVY, True, name="Lab A evidence chain")
    add_label_box(slide, 1.03, 3.18, 3.36, 1.30, "Prediction\nWAIT 可能少一次 WAKE\nawake idle 可能增加", GOLD_PALE, GOLD, 22, NAVY, True, name="Lab A prediction card")
    add_label_box(slide, 4.98, 3.18, 3.36, 1.30, "Service gate\nrequired packet、deadline、freshness\n先形成 verdict", TEAL_PALE, TEAL, 20, NAVY, True, name="Lab A service card")
    add_label_box(slide, 8.93, 3.18, 3.36, 1.30, "Claim boundary\nendpoint energy 只屬 radio／processing\nJ 不等於 system energy", RED_PALE, RED, 20, NAVY, True, name="Lab A energy scope card")
    add_text(slide, 1.08, 4.98, 11.08, 0.42, "Lab A sequence：baseline → exact edit → compact run → result／replay → causal explanation。", 20, NAVY, True, PP_ALIGN.CENTER, name="Lab A sequence")
    add_field_band(slide, "REST_DURING_GAP", "空檔休息動作", "student_policy.py `lab-a-pace-rest`", "指定 pacing gap 的 action", "action enum", "比較 SLEEP 與 WAIT 對 state ledger 的影響", y=5.62)


def slide_040(prs):
    slide = prepare_slide(prs, "P040｜A-00：Baseline decision 與 evidence")
    _add_title_lead(slide, "Release baseline 在 gap 回傳 SLEEP；P042 第一條 stdout path 提供 control evidence")
    add_label_box(slide, 0.92, 1.45, 4.10, 2.04, "Decision branch\n`contact_open = false` → `SLEEP`\n`steps_since_send < 2` → `REST_DURING_GAP`\nbaseline：`REST_DURING_GAP = SLEEP`", BLUE_PALE, BLUE, 20, NAVY, True, name="Baseline decision branch")
    add_label_box(slide, 5.34, 1.45, 2.08, 2.04, "Control\nA / baseline\nP042 第一條\nstdout `result_path`", SOFT, NAVY, 21, NAVY, True, name="Baseline control path")
    add_box(slide, 7.78, 1.45, 4.50, 2.04, "0E343A", "4E7F86", name="Current course metrics frame", line_width=1.0)
    add_picture(slide, METRICS_IMAGE, 7.94, 2.18, 4.18, 0.42, "Current course A metrics element evidence")
    add_text(slide, 7.96, 1.64, 4.12, 0.28, "current `/course` A metrics element capture", 17.5, WHITE, True, PP_ALIGN.CENTER, name="Metrics evidence label")
    add_text(slide, 7.98, 2.72, 4.08, 0.44, "同情境備用資料｜reference evidence\n實際值以 stdout `result_path` 為準", 18, "DDEFF0", False, PP_ALIGN.CENTER, name="Metrics provenance label")
    add_label_box(slide, 1.02, 3.86, 3.30, 1.02, "`endpoint_energy_j`\n6.92 J（參考）", TEAL_PALE, TEAL, 22, NAVY, True, name="Baseline energy metric")
    add_label_box(slide, 4.98, 3.86, 3.30, 1.02, "`delivered_bits`\n4,800 bit（參考）", BLUE_PALE, BLUE, 22, NAVY, True, name="Baseline delivered metric")
    add_label_box(slide, 8.94, 3.86, 3.30, 1.02, "`service_pass`\nfalse（參考）", RED_PALE, RED, 22, NAVY, True, name="Baseline service metric")
    add_text(slide, 1.06, 5.15, 11.16, 0.36, "Control 只建立原始 state／packet／service／J 對照；回答要留給 candidate evidence。", 20, NAVY, True, PP_ALIGN.CENTER, name="Baseline interpretation")
    add_field_band(slide, "result_path", "結果路徑", "runner stdout", "定位這次 `result.json` 與同 run replay", "file path", "不以投影片預填路徑替代 stdout", y=5.62)


def slide_041(prs):
    slide = prepare_slide(prs, "P041｜A-01：REST_DURING_GAP 的唯一 edit")
    _add_title_lead(slide, "只改 `student_policy.py` 的 `lab-a-pace-rest` marked block，原值與新值保持可對照")
    add_label_box(slide, 0.82, 1.42, 2.22, 3.05, "文件邊界\n`student_policy.py`\nmarked block：\n`lab-a-pace-rest`\n\n保留 scenario、runner、schema、generated JSON", SOFT, NAVY, 19, NAVY, True, name="Policy file boundary")
    code = add_box(slide, 3.32, 1.42, 4.25, 3.05, "F5F6FA", LINE, radius=False, name="Policy code anatomy")
    write_text(code, "# === LORA EDITABLE: lab-a-pace-rest ===\n`PACE_GAP_STEPS` = 2\n`REST_DURING_GAP` = `SLEEP`\n# === LORA END EDITABLE ===", 20, INK, False, PP_ALIGN.LEFT, MSO_ANCHOR.TOP, margins=(0.18, 0.16, 0.16, 0.12), line_spacing=1.06)
    add_chevron(slide, 7.78, 2.48, 0.38, 0.44, GOLD)
    add_label_box(slide, 8.36, 1.70, 3.94, 1.02, "Candidate edit\n`REST_DURING_GAP` = `WAIT`", GOLD_PALE, GOLD, 23, NAVY, True, name="Policy candidate edit")
    add_label_box(slide, 8.36, 2.98, 3.94, 1.02, "Mechanism\nWAIT → awake idle；可能少 WAKE\n再由 packet／service／J 驗證", TEAL_PALE, TEAL, 19, NAVY, True, name="Policy edit mechanism")
    add_label_box(slide, 3.34, 4.68, 4.24, 0.78, "POSIX\ncp student_policy.py student_policy.before-A-edit.py\n.venv/bin/python -m py_compile student_policy.py", BLUE_PALE, BLUE, 16.5, NAVY, True, name="Policy backup and compile POSIX", margins=(0.05, 0.03, 0.05, 0.03))
    add_label_box(slide, 8.36, 4.68, 3.94, 0.78, "Windows\ncopy /Y student_policy.py student_policy.before-A-edit.py\n.venv\\Scripts\\python.exe -m py_compile student_policy.py", BLUE_PALE, BLUE, 16.2, NAVY, True, name="Policy backup and compile Windows", margins=(0.05, 0.03, 0.05, 0.03))
    add_field_band(slide, "PACE_GAP_STEPS", "間隔步數", "student_policy.py marked block", "定義送出之間的固定 step 間隔", "step count", "本次保持 2；只改 `REST_DURING_GAP`", y=5.62)


def slide_042(prs):
    slide = prepare_slide(prs, "P042｜A-02：Lab A compact run／receipt")
    _add_title_lead(slide, "同一條 receipt ribbon 依序建立 baseline、candidate freeze 與 hidden evidence")
    stages = [
        ("01｜Baseline control", "不改 policy", "POSIX\n`bash course.sh run`\n`--lab A --case baseline`", "Windows\n`course.cmd run`\n`--lab A --case baseline`", BLUE_PALE, BLUE, "stdout `result_path`＋同 run replay"),
        ("02｜Candidate freeze", "完成 P041 edit", "POSIX\n`bash course.sh run`\n`--lab A --case candidate --freeze`", "Windows\n`course.cmd run`\n`--lab A --case candidate --freeze`", GOLD_PALE, GOLD, "stdout path＋lab-a-frozen checkpoint"),
        ("03｜Hidden check", "保持 frozen policy", "POSIX\n`bash course.sh run`\n`--lab A --case hidden`", "Windows\n`course.cmd run`\n`--lab A --case hidden`", TEAL_PALE, TEAL, "stdout path＋同 run replay"),
    ]
    x = 0.72
    for index, (heading, sub, posix, windows, fill, line, output) in enumerate(stages):
        add_box(slide, x, 1.38, 3.94, 4.12, fill, line, name=f"Run ribbon stage {index + 1}")
        add_text(slide, x + 0.14, 1.58, 3.66, 0.34, heading, 20, line, True, PP_ALIGN.CENTER, name=f"Run stage heading {index + 1}")
        add_text(slide, x + 0.20, 1.98, 3.54, 0.25, sub, 18.5, NAVY, True, PP_ALIGN.CENTER, name=f"Run stage condition {index + 1}")
        add_label_box(slide, x + 0.20, 2.40, 3.54, 0.94, posix, SOFT, line, 18, NAVY, True, name=f"Run POSIX command {index + 1}", margins=(0.05, 0.04, 0.05, 0.04))
        add_label_box(slide, x + 0.20, 3.50, 3.54, 0.94, windows, SOFT, line, 18, NAVY, True, name=f"Run Windows command {index + 1}", margins=(0.05, 0.04, 0.05, 0.04))
        add_label_box(slide, x + 0.20, 4.62, 3.54, 0.54, output, WHITE, line, 17.5, line, True, name=f"Run stage receipt {index + 1}")
        if index < 2:
            add_chevron(slide, x + 4.00, 3.18, 0.30, 0.40, line, "Run sequence connector")
        x += 4.20
    add_text(slide, 0.98, 5.72, 11.30, 0.32, "每次原樣保存 stdout `result_path`；replay 取同一 generated run 目錄；freeze 缺失時停止 hidden。", 18.5, RED, True, PP_ALIGN.CENTER, name="Run receipt rule")
    add_claim(slide, y=6.40)


def slide_043(prs):
    slide = prepare_slide(prs, "P043｜A-03：同一 boundary 下的比較")
    _add_title_lead(slide, "開啟 P042 的 baseline／candidate result path，先核對 identity，再讀 before／after")
    add_label_box(slide, 0.94, 1.44, 3.60, 0.62, "Baseline｜同情境備用資料參考", BLUE_PALE, BLUE, 19, NAVY, True, name="Baseline provenance")
    add_label_box(slide, 8.80, 1.44, 3.60, 0.62, "Candidate｜同情境備用資料參考", GOLD_PALE, GOLD, 19, NAVY, True, name="Candidate provenance")
    add_box(slide, 0.94, 2.20, 3.60, 2.24, SOFT, BLUE, name="Baseline result card")
    add_text(slide, 1.18, 2.42, 3.12, 0.34, "A / baseline", 23, BLUE, True, PP_ALIGN.CENTER, name="Baseline result heading")
    add_text(slide, 1.18, 2.96, 3.12, 1.10, "`endpoint_energy_j`：6.92 J\n`delivered_bits`：4,800 bit\n`service_pass`：false", 22, NAVY, False, PP_ALIGN.LEFT, name="Baseline result metrics")
    add_box(slide, 8.80, 2.20, 3.60, 2.24, SOFT, GOLD, name="Candidate result card")
    add_text(slide, 9.04, 2.42, 3.12, 0.34, "A / candidate", 23, GOLD, True, PP_ALIGN.CENTER, name="Candidate result heading")
    add_text(slide, 9.04, 2.96, 3.12, 1.10, "`endpoint_energy_j`：8.86 J\n`delivered_bits`：4,800 bit\n`service_pass`：false", 22, NAVY, False, PP_ALIGN.LEFT, name="Candidate result metrics")
    add_label_box(slide, 4.92, 2.35, 3.46, 1.94, "Identity gate\n`scenario_id`、seed、policy identity、units、provenance\n\n同一 boundary 才進入比較", PURPLE_PALE, PURPLE, 19.5, NAVY, True, name="A comparison identity gate")
    add_text(slide, 1.05, 4.82, 11.18, 0.48, "比較句：WAIT 讓 ______ state time 改變，packet service ______，endpoint J ______；結果稱為 ______。", 20, NAVY, True, PP_ALIGN.CENTER, name="A comparison sentence")
    add_field_band(slide, "policy identity", "策略識別", "P042 stdout result path", "確認 baseline／candidate 使用的 policy surface", "identity", "與 scenario、seed、units、provenance 同時核對", y=5.62)


def slide_044(prs):
    slide = prepare_slide(prs, "P044｜A-04：state ledger 連到 endpoint J")
    _add_title_lead(slide, "Candidate 的 events 與 energy breakdown 把 WAIT 的成本拆回 state bucket")
    buckets = [
        ("`awake_idle`", "清醒閒置", "6.00 J", GOLD_PALE, GOLD),
        ("`sleep`", "低功耗休息", "0.04 J", BLUE_PALE, BLUE),
        ("`wake`", "喚醒轉換", "0.02 J", RED_PALE, RED),
        ("`process`", "處理", "0.16 J", PURPLE_PALE, PURPLE),
        ("`tx`", "發射", "2.40 J", TEAL_PALE, TEAL),
        ("`rx`", "接收", "0.24 J", SOFT, NAVY),
    ]
    x = 0.84
    for index, (field, chinese, value, fill, line) in enumerate(buckets):
        add_label_box(slide, x, 1.44, 1.84, 1.12, f"{field}\n{chinese}\n{value}", fill, line, 18.5, NAVY, True, name=f"Energy bucket {index + 1}")
        if index < len(buckets) - 1:
            add_chevron(slide, x + 1.88, 1.82, 0.22, 0.30, line, "Energy bucket connector")
        x += 2.06
    add_box(slide, 0.88, 2.98, 6.28, 2.02, SOFT, PURPLE, name="State ledger reading order")
    add_text(slide, 1.12, 3.20, 5.80, 0.32, "讀取順序", 22, PURPLE, True, PP_ALIGN.CENTER, name="Ledger reading heading")
    add_text(slide, 1.22, 3.72, 5.60, 0.88, "1｜`STATE_INTERVAL`：指出 WAIT／SLEEP 的 duration\n2｜`WAKE`：指出 transition 是否出現\n3｜`energy_breakdown_j`：把 interval 對回 bucket", 20, NAVY, False, PP_ALIGN.LEFT, name="Ledger reading steps")
    add_box(slide, 7.52, 2.98, 4.72, 2.02, "0E343A", "4E7F86", name="Current course replay frame")
    add_picture(slide, REPLAY_IMAGE, 7.66, 3.44, 4.44, 0.92, "Current course A replay-frame element evidence")
    add_text(slide, 7.70, 3.14, 4.36, 0.25, "current `/course` A replay-frame element capture", 16.5, WHITE, True, PP_ALIGN.CENTER, name="Replay evidence label")
    add_text(slide, 7.72, 4.52, 4.32, 0.34, "同情境備用資料｜frame vocabulary reference\nCandidate 數值以 P042 stdout path 為準", 16.5, "DDEFF0", False, PP_ALIGN.CENTER, name="Replay provenance label")
    add_text(slide, 1.03, 5.20, 11.14, 0.36, "因果句：`REST_DURING_GAP = WAIT` 使 awake idle／wake bucket 改變，這才解釋 endpoint J。", 20.5, RED, True, PP_ALIGN.CENTER, name="Ledger causal sentence")
    add_field_band(slide, "STATE_INTERVAL", "狀態區間事件", "candidate result `events`", "提供 state、start／end time 與 event energy", "state＋s＋J", "沿 interval 判讀 WAIT 的成本", y=5.62)


def slide_045(prs):
    slide = prepare_slide(prs, "P045｜A-05：attempted 不等於 delivered")
    _add_title_lead(slide, "Packet outcome 連到 service verdict；SEND 次數單獨不足以證明服務完成")
    add_label_box(slide, 0.88, 1.42, 2.05, 1.02, "`attempted_packets`\n嘗試：2", BLUE_PALE, BLUE, 19, NAVY, True, name="Attempted packet metric")
    add_chevron(slide, 3.10, 1.73, 0.28, 0.34, BLUE)
    add_label_box(slide, 3.48, 1.42, 2.05, 1.02, "`retransmissions`\n重傳：1", GOLD_PALE, GOLD, 19, NAVY, True, name="Retransmission metric")
    add_chevron(slide, 5.70, 1.73, 0.28, 0.34, GOLD)
    add_label_box(slide, 6.08, 1.42, 2.05, 1.02, "`delivered_bits`\n交付：4,800 bit", TEAL_PALE, TEAL, 19, NAVY, True, name="Delivered bits metric")
    add_chevron(slide, 8.30, 1.73, 0.28, 0.34, TEAL)
    add_label_box(slide, 8.68, 1.42, 2.05, 1.02, "`expired_packets`\n逾期：3", RED_PALE, RED, 19, NAVY, True, name="Expired packet metric")
    add_chevron(slide, 10.90, 1.73, 0.28, 0.34, RED)
    add_label_box(slide, 11.28, 1.42, 1.20, 1.02, "gate\nFAIL", RED_PALE, RED, 19, RED, True, name="Packet service gate")
    add_box(slide, 0.90, 2.92, 6.34, 2.12, SOFT, RED, name="Packet service interpretation")
    add_text(slide, 1.14, 3.16, 5.86, 0.30, "Candidate reference summary", 22, RED, True, PP_ALIGN.CENTER, name="Packet summary heading")
    add_text(slide, 1.16, 3.70, 5.82, 0.96, "`deadline_pass`：false\n`service_pass`：false\n同情境備用資料；實際欄位取 P042 stdout `result_path`", 20, NAVY, False, PP_ALIGN.LEFT, name="Packet summary details")
    add_box(slide, 7.62, 2.92, 4.62, 2.12, "0E343A", "4E7F86", name="Current course ledger evidence")
    add_picture(slide, LEDGER_IMAGE, 7.76, 3.52, 4.34, 0.48, "Current course A ledger element evidence")
    add_text(slide, 7.78, 3.16, 4.30, 0.25, "current `/course` A execution-ledger element capture", 16.2, WHITE, True, PP_ALIGN.CENTER, name="Ledger evidence label")
    add_text(slide, 7.82, 4.32, 4.22, 0.36, "同情境備用資料｜baseline ledger reference\nCandidate packet outcome 以 result artifact 為準", 16.5, "DDEFF0", False, PP_ALIGN.CENTER, name="Ledger evidence provenance")
    add_text(slide, 1.04, 5.24, 11.16, 0.34, "因果句：action 改變 state，state 改變 packet outcome；service gate 再決定 J 的解讀層級。", 20, NAVY, True, PP_ALIGN.CENTER, name="Packet service causal sentence")
    add_field_band(slide, "deadline_pass", "期限通過狀態", "result `summary`", "表示 required packet 是否在 deadline 前交付", "true／false", "與 delivered、expired、freshness 共同形成 service verdict", y=5.62)


def _slide_028(prs):
    slide = prepare_slide(prs, "無線狀態與停留時間（Radio state）")
    _add_title_lead(slide, "策略動作 `action` 只記一次；狀態停留區間 `state interval` 才留下時間與能量證據")
    cards = [
        ("低功耗休息／清醒閒置", "`SLEEP`：窗口關閉／空檔，功率低\n`WAIT`：保持清醒，累積 `awake_idle`", BLUE_PALE, BLUE),
        ("喚醒轉換／封包處理", "`WAKE`：喚醒成本與延遲\n`PROCESS`：封包處理區間", GOLD_PALE, GOLD),
        ("發射／接收", "`TX`：發射區間；`RX`：接收區間\n停留或功率越高，累積 J 越多", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(cards):
        x = 0.95 + index * 3.94
        _node(slide, x, 1.62, 3.55, 2.42, heading, body, fill, line,
              size=18.5, title_size=21, name=f"Radio state group {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.66, 2.50, 0.30, 0.40, line, "State causal connector")
    add_label_box(slide, 1.02, 4.28, 5.34, 0.90,
                  "完整讀法：動作 `action` → 狀態區間 `state interval` → 能量分桶 `energy bucket`",
                  SOFT, NAVY, 18, NAVY, True, name="State causal chain")
    add_label_box(slide, 6.96, 4.28, 5.34, 0.90,
                  "同一個動作只標記一次；停留秒數 `duration` 才形成 J",
                  PURPLE_PALE, PURPLE, 18, NAVY, True, name="State interpretation")
    add_field_band(slide, "radio_state", "無線狀態", "runner events", "標記停留區間",
                   "state＋s", "對回 energy bucket")


def _slide_029(prs):
    slide = prepare_slide(prs, "空檔策略的服務與能量")
    _add_title_lead(slide, "固定工作與服務窗口，只改一個空檔策略值，檢驗它如何改變狀態、服務與 J")
    cards = [
        ("為何做", "同一份工作與窗口\n檢驗空檔策略是否改變服務與 J", BLUE_PALE, BLUE),
        ("機制：休息／閒置", "`SLEEP`：低功耗，可能付 `WAKE`\n`WAIT`：清醒閒置，可能少喚醒", GOLD_PALE, GOLD),
        ("判讀順序", "先讀交付、期限、新鮮度\n再讀端點 J；低 J 不等於成功", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(cards):
        x = 0.95 + index * 3.94
        _node(slide, x, 1.60, 3.55, 2.44, heading, body, fill, line,
              size=18.5, title_size=21, name=f"Lab A purpose card {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.66, 2.48, 0.30, 0.40, line, "Lab A causal connector")
    add_label_box(slide, 1.02, 4.12, 5.34, 1.32,
                  "固定不變：`student_policy.py` 的 `PACE_GAP_STEPS=2`、scenario／seed／traffic／window、Lab B/C",
                  BLUE_PALE, BLUE, 17.0, NAVY, True, name="Lab A baseline contract")
    add_label_box(slide, 6.96, 4.12, 5.34, 1.32,
                  "唯一變更：`lab-a-pace-rest`／`REST_DURING_GAP`\n"
                  "`SLEEP` → `WAIT`\n"
                  "成功：`status=OK`＋`result_path`\n"
                  "讀 `result.json`＋`endpoint-replay.json`\n"
                  "順序：服務→封包→狀態→J",
                  GOLD_PALE, GOLD, 16.0, NAVY, True, name="Lab A candidate contract")
    add_field_band(slide, "service_pass", "服務布林狀態", "result summary",
                   "required packet＋期限＋freshness",
                   "true／false", "先判 service，再讀 J")


def _slide_030(prs):
    slide = prepare_slide(prs, "封包生命週期：嘗試不等於交付")
    _add_title_lead(slide, "封包結果 `packet outcome` 要走完整生命週期；送出次數只是其中一段")
    phases = [
        ("建立／排隊", "建立 `generated`（已產生）封包\n`queue`（佇列）等待合法窗口", BLUE_PALE, BLUE),
        ("嘗試／重試", "進入 `attempt`（無線嘗試）\n碰撞 `collision` 可能造成 `retry`", GOLD_PALE, GOLD),
        ("交付／服務", "`delivered`（已交付）或 `expired`（過期）\n再由 `service_pass` 判定", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(phases):
        x = 0.95 + index * 3.94
        _node(slide, x, 1.62, 3.55, 2.42, heading, body, fill, line,
              size=18.5, title_size=21, name=f"Packet lifecycle stage {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.66, 2.50, 0.30, 0.40, line, "Packet lifecycle connector")
    add_label_box(slide, 1.02, 4.28, 5.34, 0.90,
                  "傳輸欄位：嘗試 `attempted`、碰撞 `collision`、重傳 `retransmissions`",
                  GOLD_PALE, GOLD, 17.5, NAVY, True, name="Packet transport fields")
    add_label_box(slide, 6.96, 4.28, 5.34, 0.90,
                  "結果欄位：`delivered_bits`、`expired`、`deadline`、`freshness`",
                  TEAL_PALE, TEAL, 17.5, NAVY, True, name="Packet outcome fields")
    add_field_band(slide, "attempted_packets", "嘗試封包數", "result summary",
                   "統計進入無線嘗試", "count",
                   "對照 retry、expired、delivered")


def _slide_031(prs):
    slide = prepare_slide(prs, "服務門檻先於端點能量")
    _add_title_lead(slide, "先固定同一比較邊界，再判定服務門檻 `service gate`，最後讀端點能量 J")
    cards = [
        ("1｜固定比較邊界", "情境 `scenario`、種子 `seed`、流量 `traffic`、窗口 `window`\n端點範圍 `endpoint scope` 必須一致", BLUE_PALE, BLUE),
        ("2｜服務門檻", "交付位元 `delivered_bits`、期限 `deadline_pass`\n新鮮度 `freshness_status` 共同形成 gate", GOLD_PALE, GOLD),
        ("3｜端點能量 J", "欄位 `endpoint_energy_j` 來自整段 state ledger\n單位 J；只在同一邊界內比較", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(cards):
        x = 0.95 + index * 3.94
        _node(slide, x, 1.62, 3.55, 2.42, heading, body, fill, line,
              size=18.5, title_size=21, name=f"Service ordering card {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.66, 2.48, 0.30, 0.40, line, "Service ordering connector")
    add_label_box(slide, 1.02, 4.28, 5.34, 0.90,
                  "服務通過 `service_pass=true`：才可在相同工作邊界內比較 J",
                  TEAL_PALE, TEAL, 17.5, NAVY, True, name="Service pass result")
    add_label_box(slide, 6.96, 4.28, 5.34, 0.90,
                  "服務未通過 `service_pass=false`：J 只描述能量／服務取捨",
                  RED_PALE, RED, 17.5, NAVY, True, name="Service fail result")
    add_field_band(slide, "freshness_status", "新鮮度狀態", "result summary",
                   "required packet 是否新鮮",
                   "fresh／stale／expired／N/A",
                   "與 deadline、delivery 決定 gate")


def _slide_032(prs):
    slide = prepare_slide(prs, "功率 W × 時間 s 形成能量 J")
    _add_title_lead(slide, "同一個功率值，停留時間不同就會留下不同的端點能量")
    _node(slide, 0.95, 1.62, 3.10, 2.48, "功率 `P`", "某個狀態當下的高度\n單位：W\n例如 `awake_idle` 或 `tx`", BLUE_PALE, BLUE,
          size=19, title_size=23, name="Power definition card")
    chart = add_box(slide, 4.35, 1.62, 4.18, 2.48, SOFT, LINE, name="Power interval chart")
    add_text(slide, 4.56, 1.80, 3.76, 0.38, "不同狀態的功率高度", 20, NAVY, True, PP_ALIGN.CENTER, name="Power chart heading")
    add_box(slide, 4.86, 3.03, 0.72, 0.62, TEAL, None, radius=False, name="Sleep power bar")
    add_box(slide, 5.78, 2.57, 0.72, 1.08, GOLD, None, radius=False, name="Wake power bar")
    add_box(slide, 6.70, 2.30, 0.72, 1.35, PURPLE, None, radius=False, name="Process power bar")
    add_box(slide, 7.62, 2.12, 0.72, 1.53, BLUE, None, radius=False, name="Tx power bar")
    add_text(slide, 4.58, 3.62, 3.72, 0.40, "休息 `sleep`｜喚醒 `wake`\n處理 `process`｜發射 `tx`", 16, MUTED, False, PP_ALIGN.CENTER, name="Power chart labels")
    add_chevron(slide, 8.78, 2.54, 0.30, 0.40, NAVY, "Power time connector")
    _node(slide, 9.25, 1.62, 3.10, 2.48, "時間 `t`", "狀態區間 `state interval` 的長度\n單位：s\n`P × t = J`", GOLD_PALE, GOLD,
          size=19, title_size=23, name="Time definition card")
    add_label_box(slide, 1.02, 4.34, 5.34, 0.86,
                  "能量分桶 `bucket` 分開累積：\nawake_idle／sleep／wake／process／tx／rx",
                  TEAL_PALE, TEAL, 17.5, NAVY, True, name="Energy bucket rule")
    add_label_box(slide, 6.96, 4.34, 5.34, 0.86,
                  "固定 10 s 時鐘步長 `clock step`；收發流程另加入轉換能量",
                  PURPLE_PALE, PURPLE, 17.5, NAVY, True, name="Clock energy rule")
    add_field_band(slide, "energy_breakdown_j", "能量分桶", "result artifact",
                   "按 state 保存 J", "J／bucket",
                   "拆回 state interval")


def _slide_033(prs):
    slide = prepare_slide(prs, "端點能量的累積公式")
    _add_title_lead(slide, "公式說明端點狀態如何累積；結果產物 `artifact` 提供這次執行的數值")
    _node(slide, 0.95, 1.62, 3.10, 2.42, "狀態功率 `P_s`", "某個狀態的功率高度\n單位：W\n來源：energy profile", BLUE_PALE, BLUE,
          size=18.5, title_size=23, name="Formula power input")
    frame = add_box(slide, 4.36, 1.62, 4.60, 2.42, SOFT, PURPLE, name="Endpoint formula frame", line_width=1.6)
    add_text(slide, 4.58, 2.10, 4.16, 0.70, "端點能量 `E_endpoint` = Σ\n`P_s` × `t_s`", 22, PURPLE, True, PP_ALIGN.CENTER, name="Editable endpoint formula", italic_all=True)
    add_text(slide, 4.62, 3.14, 4.08, 0.28, "分段累積｜單位：J", 17.5, PURPLE, True, PP_ALIGN.CENTER, name="Endpoint formula unit")
    _node(slide, 9.27, 1.62, 3.10, 2.42, "停留時間 `t_s`", "狀態區間 `state interval`\n單位：s\n來源：`STATE_INTERVAL` 事件", GOLD_PALE, GOLD,
          size=18.5, title_size=23, name="Formula time input")
    add_label_box(slide, 1.02, 4.28, 5.34, 0.88,
                  "端點範圍 `endpoint scope`：宣告的 radio／processing 狀態與轉換",
                  TEAL_PALE, TEAL, 17.5, NAVY, True, name="Endpoint scope boundary")
    add_label_box(slide, 6.96, 4.28, 5.34, 0.88,
                  "不包含衛星 `satellite`、閘道 `gateway`、整體插座 `wall-plug`",
                  RED_PALE, RED, 17.5, NAVY, True, name="Endpoint exclusion boundary")
    add_field_band(slide, "endpoint_energy_j", "端點能量", "result summary",
                   "同一 endpoint boundary 的總和", "J",
                   "對照 energy bucket")


def _slide_034(prs):
    slide = prepare_slide(prs, "端點能量效率的 bit/J 邊界")
    _add_title_lead(slide, "bit/J 是同一端點範圍內的效率比值；服務門檻仍獨立判定")
    _node(slide, 0.95, 1.62, 3.10, 2.30, "交付資料 `D_delivered`", "實際交付的資料量\n單位：bit\n來源：結果摘要", BLUE_PALE, BLUE,
          size=18.5, title_size=21, name="Delivered numerator card")
    frame = add_box(slide, 4.25, 1.62, 4.45, 2.30, SOFT, PURPLE, name="Efficiency formula frame", line_width=1.6)
    add_text(slide, 4.45, 2.04, 4.05, 0.72, "端點效率 `η_E` =\n`D_delivered` / `E_endpoint`", 20, PURPLE, True, PP_ALIGN.CENTER, name="Editable efficiency formula", italic_all=True)
    add_text(slide, 4.48, 3.18, 3.99, 0.26, "bit ÷ J = bit/J", 17.5, PURPLE, True, PP_ALIGN.CENTER, name="Efficiency unit")
    _node(slide, 8.83, 1.62, 3.10, 2.30, "端點能量 `E_endpoint`", "端點狀態的累積能量\n單位：J\n來源：結果摘要", TEAL_PALE, TEAL,
          size=18.5, title_size=21, name="Endpoint denominator card")
    add_label_box(slide, 1.02, 4.20, 5.34, 0.94,
                  "bit/J 高但 `service_pass=false`：\n只作效率觀察，不作服務成功判定",
                  GOLD_PALE, GOLD, 17.5, NAVY, True, name="Efficiency caveat card")
    add_label_box(slide, 6.96, 4.20, 5.34, 0.94,
                  "`service_pass`、`deadline_pass` 仍獨立判定：\n依必要封包與期限讀結果",
                  RED_PALE, RED, 17.5, NAVY, True, name="Efficiency service card")
    add_field_band(slide, "endpoint_energy_efficiency_bits_per_j", "端點能量效率",
                   "result summary", "delivered bit ÷ endpoint J", "bit/J",
                   "不取代 service gate")


def _slide_035(prs):
    slide = prepare_slide(prs, "來源、模型、課程假設與結果")
    _add_title_lead(slide, "同一筆數字要沿來源、模型、固定課程假設與結果產物 `artifact` 回溯")
    layers = [
        ("來源 Source", "可追溯來源 `provenance` 參考\n上游 repository／commit，不是本次量測", SOFT, NAVY),
        ("模型 Model", "課程模擬轉接器 `coherent-course-simulated-adapter`\n固定執行器 `runner` 的行為", BLUE_PALE, BLUE),
        ("課程假設 Course", "情境 `scenario`、窗口 `window`、流量 `traffic`\n端點設定 `endpoint profile` 固定", GOLD_PALE, GOLD),
        ("結果 Result", "JSON 結果＋同次事件重播 `replay`\n宣告的 claim boundary 隨產物保存", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(layers):
        x = 0.78 + index * 3.10
        _node(slide, x, 1.68, 2.78, 2.20, heading, body, fill, line,
              size=19.5, title_size=22, name=f"Provenance layer {index + 1}")
        if index < 3:
            add_chevron(slide, x + 2.88, 2.56, 0.22, 0.36, line, "Provenance connector")
    add_label_box(slide, 1.02, 4.34, 5.34, 0.82,
                  "上游執行旗標 `upstream_execution = false`：這不是上游即時執行",
                  RED_PALE, RED, 19, NAVY, True, name="Provenance execution boundary")
    add_label_box(slide, 6.96, 4.34, 5.34, 0.82,
                  "結果型態：一致的模擬結果；所有數值以本次產物 `artifact` 讀值",
                  TEAL_PALE, TEAL, 19.5, NAVY, True, name="Provenance result boundary")
    add_field_band(slide, "runner_provenance", "執行來源", "result artifact",
                   "保存 model／assumption／execution", "record",
                   "設定 claim ceiling")
    add_claim(slide)


def _slide_036(prs):
    slide = prepare_slide(prs, "執行、事件重播與工作簿的身分鏈")
    _add_title_lead(slide, "同一組身分欄位連結結果、重播與工作簿；不是只靠一個情境識別碼")
    nodes = [
        ("執行 RUN", "執行器 `runner` 讀情境、案例與策略\n輸出 `result.json` 與同目錄 `endpoint-replay.json`", BLUE_PALE, BLUE),
        ("事件重播 REPLAY", "讀同一次執行的 `events`\n顯示動作、狀態、佇列與封包，不重跑策略", TEAL_PALE, TEAL),
        ("比較紀錄 WORKBOOK", "保存預測、基準、候選、凍結與 withheld\n寫入結果／重播的 lineage（來源鏈）", GOLD_PALE, GOLD),
    ]
    for index, (heading, body, fill, line) in enumerate(nodes):
        x = 0.95 + index * 3.94
        _node(slide, x, 1.62, 3.55, 2.02, heading, body, fill, line,
              size=20, title_size=21, name=f"Identity object {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.66, 2.42, 0.30, 0.40, line, "Identity chain connector")
    add_label_box(slide, 1.02, 4.02, 5.34, 1.02,
                  "情境識別碼 `scenario_id`：指出固定情境\n是字串，不是能量值；來源：scenario／result",
                  BLUE_PALE, BLUE, 17.5, NAVY, True, name="Scenario identity field")
    add_label_box(slide, 6.96, 4.02, 5.34, 1.02,
                  "`run_id`＋策略身分＋`units`＋`provenance`\n任一不一致就拒絕匯入（fail closed）",
                  PURPLE_PALE, PURPLE, 17.5, NAVY, True, name="Identity mismatch field")
    add_claim(slide, y=6.40)


def _slide_037(prs):
    slide = prepare_slide(prs, "基準 Baseline：固定條件的對照執行")
    _add_title_lead(slide, "基準 `Baseline` 固定情境、種子與原始策略，才提供可比較的對照")
    columns = [
        ("固定輸入", "情境 `scenario`／種子 `seed`\n流量 `traffic`／窗口 `window`\n端點範圍 `endpoint scope`", BLUE_PALE, BLUE),
        ("唯一變動", "標記區塊 `marked block`\n一次只改一個因果值", GOLD_PALE, GOLD),
        ("觀察輸出", "狀態時間 `state duration`\n封包結果 `packet outcome`\n服務判定與端點 J", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(columns):
        x = 0.95 + index * 3.94
        _node(slide, x, 1.62, 3.55, 2.24, heading, body, fill, line,
              size=21, title_size=23, name=f"Baseline rule {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.66, 2.50, 0.30, 0.40, line, "Baseline causal connector")
    add_label_box(slide, 1.02, 4.32, 5.34, 0.82,
                  "公平條件：兩次結果的身分與端點範圍一致",
                  SOFT, NAVY, 20, NAVY, True, name="Baseline control condition")
    add_label_box(slide, 6.96, 4.32, 5.34, 0.82,
                  "解釋順序：先回到標記區塊，再讀狀態／封包／服務／J",
                  PURPLE_PALE, PURPLE, 19, NAVY, True, name="Baseline interpretation")
    add_field_band(slide, "policy identity", "策略識別", "result policy",
                   "指出 policy surface", "識別碼",
                   "與 scenario、run_id 比對")


def _slide_038(prs):
    slide = prepare_slide(prs, "LEO 的變動服務窗口軌跡")
    _add_title_lead(slide, "服務窗口輸入改變合法傳輸機會；策略再依觀察值選擇動作")
    segments = [
        ("窗口關閉 `closed`", "不接受傳輸", "CBD5E1", INK, 1.45),
        ("服務窗口 A `contact-a`", "品質逐步變化", TEAL, WHITE, 2.40),
        ("窗口關閉 `closed`", "服務機會消失", "CBD5E1", INK, 1.45),
        ("服務窗口 B `contact-b`", "品質再次變化", BLUE, WHITE, 2.40),
        ("窗口關閉 `closed`", "回到休息", "CBD5E1", INK, 1.45),
    ]
    x = 0.92
    for index, (label, body, fill, color, width) in enumerate(segments):
        card = add_box(slide, x, 1.72, width, 1.06, fill, None, radius=False, name=f"LEO window segment {index + 1}")
        write_text(card, f"{label}\n{body}", 19, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
                   margins=(0.03, 0.02, 0.03, 0.02), line_spacing=0.90)
        x += width + 0.08
    fields = [
        ("窗口是否開啟 `contact_open`", "布林狀態；來源：scenario trace\n作用：決定是否允許傳輸", BLUE_PALE, BLUE),
        ("連線品質 `quality_band`", "分類值 0–3；不是 dB\n作用：供策略判斷服務機會", GOLD_PALE, GOLD),
        ("窗口剩餘秒數 `contact_remaining_s`", "單位：s；來源：scenario trace\n作用：影響動作時機與期限", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(fields):
        _node(slide, 0.95 + index * 3.94, 3.18, 3.55, 1.60, heading, body, fill, line,
              size=19, title_size=20, name=f"LEO trace field {index + 1}")
    add_label_box(slide, 1.02, 4.96, 11.28, 0.52,
                  "窗口關閉時只接受低功耗休息 `SLEEP`；窗口開啟時，品質與剩餘時間共同改變動作的服務機會。",
                  SOFT, NAVY, 18.5, NAVY, True, name="LEO trace interpretation")
    add_field_band(slide, "service_window", "服務窗口", "scenario trace",
                   "定義合法傳輸區間", "time／category",
                   "解釋 action 的服務機會")


def _slide_039(prs):
    slide = prepare_slide(prs, "為何改 REST_DURING_GAP")
    _add_title_lead(slide, "先說明為何做，再定位檔案、標記區塊、命令與結果欄位")
    cards = [
        ("為何做 Why", "`SLEEP`↔`WAIT` 先改變\n`WAKE`／awake-idle／J；要驗證 state 與 service", BLUE_PALE, BLUE),
        ("檔案與區塊", "檔案：`student_policy.py`\n標記區塊：`lab-a-pace-rest`", GOLD_PALE, GOLD),
        ("改前 Before", "`REST_DURING_GAP = SLEEP`\n低功耗休息；可能出現 WAKE", BLUE_PALE, BLUE),
        ("改後 After", "`REST_DURING_GAP = WAIT`\n清醒閒置；可能少 WAKE、但多 awake-idle J", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(cards):
        x = 0.75 + index * 3.10
        _node(slide, x, 1.56, 2.78, 1.92, heading, body, fill, line,
              size=17.5, title_size=20, name=f"Lab A explanation card {index + 1}")
        if index < 3:
            add_chevron(slide, x + 2.86, 2.25, 0.20, 0.32, line, "Lab A explanation connector")
    add_label_box(slide, 1.02, 3.78, 5.34, 0.88,
                  "固定：`PACE_GAP_STEPS=2`、scenario／seed／traffic／window、runner／schema、Lab B/C",
                  SOFT, NAVY, 17.0, NAVY, True, name="Lab A fixed inputs")
    add_label_box(slide, 6.96, 3.78, 5.34, 0.88,
                  "成功：`status=OK`＋新 `result_path`；讀同目錄 `result.json`／`endpoint-replay.json`",
                  PURPLE_PALE, PURPLE, 17.0, NAVY, True, name="Lab A result fields")
    add_label_box(slide, 1.02, 4.86, 11.28, 0.62,
                  "完整解釋句：`SLEEP`→`WAIT` 先改變 state，再讀 packet／service，最後才解釋 endpoint J；方向可被結果反駁。",
                  RED_PALE, RED, 17.0, NAVY, True, name="Lab A falsifiable prediction")
    add_field_band(slide, "REST_DURING_GAP", "空檔策略欄位", "policy marked block",
                   "控制 gap state", "SLEEP／WAIT",
                   "先連 state，再讀 packet／service／J")


def _slide_040(prs):
    slide = prepare_slide(prs, "基準決策與對照證據")
    _add_title_lead(slide, "基準先證明原始策略在同一服務窗口如何決策，建立改前對照")
    cards = [
        ("決策分支", "`contact_open=false` → `SLEEP`\n`steps_since_send < 2` → `REST_DURING_GAP=SLEEP`", BLUE_PALE, BLUE),
        ("基準命令", "Linux／macOS 終端機（WSL 也使用 Linux 指令）\n`bash course.sh run --lab A --case baseline`", SOFT, NAVY),
        ("參考摘要", "`endpoint_energy_j=6.92 J`\n`delivered_bits=4,800 bit`；`service_pass=false`", TEAL_PALE, TEAL),
    ]
    for index, (heading, body, fill, line) in enumerate(cards):
        x = 0.95 + index * 3.94
        _node(slide, x, 1.56, 3.55, 2.48, heading, body, fill, line,
              size=17.5, title_size=21, name=f"Baseline evidence card {index + 1}")
        if index < 2:
            add_chevron(slide, x + 3.66, 2.50, 0.30, 0.40, line, "Baseline evidence connector")
    add_label_box(slide, 1.02, 4.30, 5.34, 0.90,
                  "固定：`student_policy.py`／`lab-a-pace-rest`；`PACE_GAP_STEPS=2`、`REST_DURING_GAP=SLEEP`；其他區塊不變",
                  BLUE_PALE, BLUE, 17.0, NAVY, True, name="Baseline why card")
    add_label_box(slide, 6.96, 4.30, 5.34, 0.90,
                  "讀取：`result.json` → `summary`／`events` → `endpoint-replay.json`；baseline 是 control，不回答省電。",
                  GOLD_PALE, GOLD, 17.0, NAVY, True, name="Baseline reference boundary")
    add_field_band(slide, "result_path", "結果路徑", "P042 stdout",
                   "定位 result.json＋同 run replay", "path",
                   "先核對 identity")


def _slide_041(prs):
    slide = prepare_slide(prs, "程式碼改前／改後")
    _add_title_lead(slide, "定位後只讀 5 行：淡化不變行，高亮唯一變更 `SLEEP` → `WAIT`")

    def _code_card(x, title, fill, line, changed_value, name):
        add_box(slide, x, 1.56, 5.20, 2.86, fill, line, name=name)
        add_text(slide, x + 0.14, 1.70, 4.92, 0.34, title, 22, line, True,
                 PP_ALIGN.CENTER, name=f"{name} heading")
        lines = [
            "# LORA EDITABLE: lab-a-pace-rest",
            "PACE_GAP_STEPS = 2",
            f"REST_DURING_GAP = {changed_value}",
            "# 其他 marked block 保持原樣",
            "# runner／scenario／schema／JSON 不變",
        ]
        for idx, code_line in enumerate(lines):
            color = line if idx == 2 else MUTED
            bold = idx == 2
            add_text(slide, x + 0.20, 2.14 + idx * 0.40, 4.80, 0.30, code_line,
                     17.5, color, bold, PP_ALIGN.LEFT, name=f"{name} code line {idx + 1}",
                     margins=(0.02, 0.01, 0.02, 0.01))

    _code_card(0.80, "改前 Before｜基準", BLUE_PALE, BLUE, "SLEEP", "Lab A before code")
    _code_card(6.20, "改後 After｜候選", TEAL_PALE, TEAL, "WAIT", "Lab A after code")
    add_chevron(slide, 5.62, 2.66, 0.36, 0.46, GOLD, "Lab A unique edit arrow")
    add_text(slide, 5.34, 2.24, 0.92, 0.28, "唯一變更", 16.5, GOLD, True, PP_ALIGN.CENTER,
             name="Lab A unique edit annotation")

    add_label_box(slide, 0.98, 4.62, 11.34, 0.92,
                  "只改 `student_policy.py` 的 `lab-a-pace-rest`：`REST_DURING_GAP`；`PACE_GAP_STEPS=2`、Lab B/C、runner、scenario、schema、JSON 不變。\n"
                  "原因：`WAIT` 先改 awake-idle state；P042 run 才看 `WAKE`、packet、service、J。",
                  SOFT, NAVY, 16.5, NAVY, True, name="Lab A edit mechanism and freeze",
                  margins=(0.16, 0.04, 0.16, 0.04))
    add_field_band(slide, "REST_DURING_GAP", "空檔策略欄位", "student_policy.py／lab-a-pace-rest",
                   "指定 gap action", "SLEEP／WAIT",
                   "先連 state，再用 A/B 結果判讀", y=5.62)


def _run_stage_card(slide, x, heading, sub, posix, windows, receipt, fill, line, index):
    add_box(slide, x, 1.62, 3.55, 3.90, fill, line, name=f"Run stage {index}")
    add_text(slide, x + 0.12, 1.78, 3.31, 0.34, heading, 18.5, line, True, PP_ALIGN.CENTER, name=f"Run stage heading {index}")
    add_text(slide, x + 0.16, 2.16, 3.23, 0.42, sub, 16.5, NAVY, True, PP_ALIGN.CENTER, name=f"Run stage condition {index}")
    add_text(slide, x + 0.18, 2.70, 3.19, 0.74, f"Linux／macOS：\n{posix}", 16.0, NAVY, False, PP_ALIGN.LEFT,
             margins=(0.04, 0.02, 0.04, 0.02), line_spacing=0.84, name=f"Run Linux exact command {index}")
    add_text(slide, x + 0.18, 3.52, 3.19, 0.74, f"Windows：\n{windows}", 16.0, NAVY, False, PP_ALIGN.LEFT,
             margins=(0.04, 0.02, 0.04, 0.02), line_spacing=0.84, name=f"Run Windows exact command {index}")
    add_label_box(slide, x + 0.16, 4.48, 3.23, 0.80, receipt, WHITE, line, 16.0, line, True,
                  name=f"Run stage receipt {index}", margins=(0.04, 0.02, 0.04, 0.02))


def _slide_042(prs):
    slide = prepare_slide(prs, "基準 → 候選 → 隱藏檢查")
    _add_title_lead(slide, "一條執行帶完成三步：建立基準、凍結候選、用同一策略做隱藏檢查")
    add_text(slide, 0.98, 1.38, 11.30, 0.22,
             "平台：Linux／macOS 終端機（WSL 也使用 Linux 指令）；Windows 命令提示字元。",
             16.0, NAVY, True, PP_ALIGN.CENTER, name="Run platform contract")
    stages = [
        ("01｜基準 Baseline", "不改 policy；A 區塊仍為 `SLEEP`", "bash course.sh run --lab A --case baseline", "course.cmd run --lab A --case baseline", "成功：`status=OK`＋`result_path`\n同目錄 `result.json`／`endpoint-replay.json`", BLUE_PALE, BLUE),
        ("02｜候選 Candidate", "完成 P041：A 區塊改 `WAIT`，再凍結", "bash course.sh run --lab A --case candidate --freeze", "course.cmd run --lab A --case candidate --freeze", "成功：`result_path`＋\n`lab-a-frozen.json`／`.py`", GOLD_PALE, GOLD),
        ("03｜隱藏 Hidden", "保持 frozen policy；不再改值", "bash course.sh run --lab A --case hidden", "course.cmd run --lab A --case hidden", "成功：`status=OK`＋`result_path`\n不建立第二個 freeze", TEAL_PALE, TEAL),
    ]
    for index, (heading, sub, posix, windows, receipt, fill, line) in enumerate(stages):
        _run_stage_card(slide, 0.74 + index * 4.18, heading, sub, posix, windows, receipt, fill, line, index + 1)
        if index < 2:
            add_chevron(slide, 4.40 + index * 4.18, 3.14, 0.28, 0.38, line, "Run ribbon connector")
    add_text(slide, 0.98, 5.64, 11.30, 0.42,
             "讀取順序：stdout `result_path` → 同一 generated 目錄 `result.json` → 配對 `endpoint-replay.json` → 服務／封包／狀態／J；freeze 缺失就停止 hidden。",
             16.0, RED, True, PP_ALIGN.CENTER, name="Run receipt rule")
    add_claim(slide, y=6.40)


def _slide_043(prs):
    slide = prepare_slide(prs, "同一邊界的改前／改後")
    _add_title_lead(slide, "三次 evidence 都未通過 service gate；先核對同一 boundary，再讀改前／改後／hidden")
    result_cards = [
        (
            "Baseline｜SLEEP",
            "`service`：false　`deadline`：false\n"
            "`delivered`：1/4　`retries`：1\n"
            "`collisions`：1　`expired`：3\n"
            "`wake`：2\n"
            "`endpoint`：6.92 J\n"
            "`bit/J`：693.641618",
            BLUE_PALE, BLUE, "Lab A baseline evidence",
        ),
        (
            "Candidate｜WAIT",
            "`service`：false　`deadline`：false\n"
            "`delivered`：1/4　`retries`：1\n"
            "`collisions`：1　`expired`：3\n"
            "`wake`：1\n"
            "`endpoint`：8.86 J\n"
            "`bit/J`：541.760722",
            GOLD_PALE, GOLD, "Lab A candidate evidence",
        ),
        (
            "Hidden frozen｜WAIT",
            "`service`：false　`deadline`：false\n"
            "`delivered`：0/3　`retries`：0\n"
            "`collisions`：0　`expired`：3\n"
            "`wake`：0\n"
            "`endpoint`：3.61 J\n"
            "`bit/J`：0",
            RED_PALE, RED, "Lab A hidden evidence",
        ),
    ]
    for index, (heading, body, fill, line, name) in enumerate(result_cards):
        _node(slide, 0.66 + index * 4.02, 1.52, 3.72, 3.58, heading, body, fill, line,
              size=17.0, title_size=20.0, name=name)
    add_label_box(slide, 0.82, 5.22, 5.74, 0.40,
                  "WAIT：service 不變（false→false），endpoint `+1.94 J`",
                  GOLD_PALE, GOLD, 16.0, NAVY, True, name="Lab A candidate comparison")
    add_label_box(slide, 6.72, 5.22, 5.74, 0.40,
                  "Hidden 低 J 來自 `0/3` 未交付；不可視為節能成功",
                  RED_PALE, RED, 16.0, NAVY, True, name="Lab A hidden interpretation")
    add_field_band(slide, "policy identity", "策略識別", "P042 result path",
                   "確認三次 run 的 policy 與 frozen identity", "identity",
                   "先核對 scenario、seed、units、provenance")


def _slide_044(prs):
    slide = prepare_slide(prs, "狀態帳本回到端點能量")
    _add_title_lead(slide, "用事件與能量分桶，把 `WAIT` 的成本拆回可檢查的狀態區間")
    _node(slide, 0.95, 1.62, 5.48, 2.42, "帳本讀取順序 Ledger",
          "1｜狀態區間 `STATE_INTERVAL`：找 WAIT／SLEEP 停留秒數\n2｜喚醒 `WAKE`：確認轉換是否出現\n3｜能量分桶 `energy_breakdown_j`：對回每個 bucket",
          SOFT, PURPLE, size=20, title_size=22, name="State ledger order card")
    _node(slide, 6.90, 1.62, 5.48, 2.42, "能量分桶 `energy_breakdown_j`",
          "醒著閒置 `awake_idle` 6.00 J｜休息 `sleep` 0.04 J\n喚醒 `wake` 0.02 J｜處理 `process` 0.16 J\n發射 `tx` 2.40 J｜接收 `rx` 0.24 J",
          TEAL_PALE, TEAL, size=19, title_size=22, name="Energy bucket result card")
    add_label_box(slide, 1.02, 4.34, 5.34, 0.88,
                  "來源：候選 `result_path` 的 `events`\n保留 state、起訖秒數、事件能量與 bucket J",
                  BLUE_PALE, BLUE, 17.0, NAVY, True, name="Ledger result fields")
    add_label_box(slide, 6.96, 4.34, 5.34, 0.88,
                  "讀取：`result_path` → `STATE_INTERVAL`／`WAKE` → `energy_breakdown_j`\n哪個 interval 改變，就解釋哪個 J bucket",
                  GOLD_PALE, GOLD, 16.5, NAVY, True, name="Ledger artifact boundary")
    add_field_band(slide, "STATE_INTERVAL", "狀態區間事件", "candidate events",
                   "提供 state、time、event energy", "state＋s＋J",
                   "沿 interval 判讀 WAIT 成本")


def _slide_045(prs):
    slide = prepare_slide(prs, "封包結果回到服務判定")
    _add_title_lead(slide, "WAIT 沒改善服務，卻多耗 1.94 J；hidden 的低 J 是未送達，不是節能成功")
    _node(
        slide, 0.84, 1.52, 5.52, 2.72, "SLEEP → WAIT｜state／packet",
        "`wake`：2 → 1　（少 1 次）\n"
        "`retries`：1 → 1　`collisions`：1 → 1\n"
        "`expired`：3 → 3　`delivered`：1/4 → 1/4\n"
        "`service`：false → false　`deadline`：false → false",
        BLUE_PALE, BLUE, size=18.0, title_size=21, name="Lab A state packet interpretation",
    )
    _node(
        slide, 6.72, 1.52, 4.66, 2.72, "Endpoint trade-off｜J／bit/J",
        "`endpoint`：6.92 → 8.86 J\n"
        "差異：`+1.94 J`\n"
        "`bit/J`：693.641618 → 541.760722\n\n"
        "WAIT 沒改善 service，且端點能量更高、效率更低。",
        GOLD_PALE, GOLD, size=18.0, title_size=21, name="Lab A endpoint tradeoff",
    )
    add_label_box(
        slide, 0.84, 4.48, 10.54, 0.86,
        "Hidden frozen WAIT｜`delivered`：0/3　`retries`：0　`collisions`：0　`expired`：3　`wake`：0　`endpoint`：3.61 J　`bit/J`：0",
        RED_PALE, RED, 16.5, NAVY, True, name="Lab A hidden frozen result",
        margins=(0.08, 0.03, 0.08, 0.03),
    )
    add_label_box(
        slide, 0.84, 5.08, 10.54, 0.40,
        "Hidden 低 J 的原因是完全未送達；`service=false`／`deadline=false`，不能把它稱為節能成功。",
        RED_PALE, RED, 16.0, NAVY, True, name="Lab A hidden no success",
    )
    add_field_band(slide, "deadline_pass", "期限通過狀態", "result summary",
                   "required packet 是否準時交付", "true／false",
                   "先讀 delivered／expired／freshness，再形成 service gate")


SLIDE_BUILDERS = [
    _slide_028, _slide_029, _slide_030, _slide_031, _slide_032, _slide_033,
    _slide_034, _slide_035, _slide_036, _slide_037, _slide_038, _slide_039,
    _slide_040, _slide_041, _slide_042, _slide_043, _slide_044, _slide_045,
]


PAGE_META = [
    {"page": 28, "title": "Radio state 與停留時間", "source": "part-b-visible-content P028; Part B field audit", "notes": "P028。Radio state 與停留時間構成 endpoint energy 的狀態積分基礎。策略 action 是一次決策輸出；runner events 以 STATE_INTERVAL 保存 state、起訖時間與 event energy，energy_breakdown_j 再按 sleep、awake_idle、wake、process、tx、rx 分桶。SLEEP 是低功耗休息，WAIT 是清醒閒置，WAKE、PROCESS、TX、RX 是送出流程中的 transition 或 radio state。判讀時把 action 與完整 state ledger 分開。"},
    {"page": 29, "title": "Lab A 空檔策略的 service／energy gate", "source": "part-b-visible-content P029; experiment-operation-contract Lab A", "notes": "P029。Lab A 固定同一份工作、服務窗口、traffic、seed 與 endpoint boundary，只比較空檔的 SLEEP 與 WAIT。SLEEP 降低 idle power，但後續可能支付 WAKE latency 與 wake energy；WAIT 保持 awake idle，可能減少部分 wake transition，卻增加 awake-idle 停留。service_pass 先由 delivered、deadline 與 freshness 形成，再讀整段 endpoint_energy_j。改前預測允許兩個方向，結果用 baseline、candidate 與 withheld evidence 來檢驗。"},
    {"page": 30, "title": "封包生命週期：送出與交付分開判定", "source": "part-b-visible-content P030; engine packet lifecycle", "notes": "P030。封包從 generated 進入 queue，policy action 讓它進入 attempt；collision 可能產生 retry，之後才有 delivered 或 expired。attempted_packets、retransmissions、unique_delivered_packets 與 expired_packets 來自 result summary，單位是 count；delivered_bits 來自交付 payload，單位是 bit。service_pass 需要完整 packet outcome、deadline 與 freshness，SEND 次數不取代服務判定。"},
    {"page": 31, "title": "Service gate 與 endpoint J 的比較順序", "source": "part-b-visible-content P031; field-explanation-contract", "notes": "P031。公平比較先固定 scenario、seed、traffic、window 與 endpoint scope，再讀 delivered_bits、deadline_pass、freshness_status 與 service_pass。service_pass 是 result summary 的布林狀態，表示 required packet、期限、最小交付量、逾期數與 freshness 條件共同通過。只有在相同邊界下，endpoint_energy_j 才能被放進同一個 trade-off 句子；較低 J 不自動成為服務成功。"},
    {"page": 32, "title": "W 是瞬間功率，J 是整段累積", "source": "part-b-visible-content P032; engine energy buckets", "notes": "P032。W 是某個 state 當下的功率，J 是功率在時間上的累積。WAIT 與 SLEEP 使用固定 10 s clock step；PROCESS、TX、RX 與 WAKE 依 transition profile 另行累加。energy_breakdown_j 以 J 保存每個 bucket，能把 total endpoint energy 拆回 state interval 與 transition。讀值時同時保留功率、時間與 bucket，避免以單一峰值代替整段累積。"},
    {"page": 33, "title": "Endpoint energy 公式的邊界", "source": "part-b-visible-content P033; native Office Math requirement", "notes": "P033。Office Math 公式 E_endpoint = Σ P_s t_s 表示 endpoint state 的功率與停留時間如何累加。P_s 是 state power，單位 W；t_s 是 state interval 的停留時間，單位 s；E_endpoint 是 result summary 的端點能量，單位 J。endpoint scope 只涵蓋宣告的 radio 與 processing 假設，satellite、gateway 與 whole-system wall-plug 不進入這個分母。"},
    {"page": 34, "title": "Energy efficiency 的分子與分母", "source": "part-b-visible-content P034; native Office Math requirement", "notes": "P034。Office Math 公式 η_E = D_delivered / E_endpoint 把交付資料與端點能量放在同一個 endpoint boundary。D_delivered 來自 result summary 的 delivered_bits，單位 bit；E_endpoint 來自同一筆 result 的 endpoint_energy_j，單位 J；η_E 的單位是 bit/J。service_pass 與 deadline_pass 仍是獨立 verdict，因此 bit/J 的升高不會自動改寫服務結果。"},
    {"page": 35, "title": "來源、模型、假設與結果分層", "source": "part-b-visible-content P035; ADR-004 claim ceiling", "notes": "P035。Source 是 provenance 參考；model 是 coherent-course-simulated-adapter；course layer 固定 scenario、contact window、traffic 與 endpoint profile；result layer 保存這次 run 的 JSON 與 replay。runner_provenance.upstream_execution 固定為 false，claim boundary 也隨 artifact 保存。這些資料可追溯 input、policy 與 output 的關聯，證據類型維持 coherent simulated result。"},
    {"page": 36, "title": "執行結果如何進入重播與工作簿", "source": "part-b-visible-content P036; field-explanation-contract P036", "notes": "P036。RUN 是 runner 以指定 scenario、case 與 policy 產生 result.json 與同目錄 endpoint-replay.json。REPLAY 依同一次 run 的 events 顯示 action、radio state、queue 與 packet 時序，不重新執行 policy。WORKBOOK 是保存 prediction、baseline、candidate、freeze、withheld 與 result／replay lineage 的比較紀錄。scenario_id 是情境識別碼，run_id、policy identity、units 與 provenance 共同形成一致性核對；任一 mismatch 使匯入維持 fail closed。"},
    {"page": 37, "title": "Baseline：固定條件下的對照執行", "source": "part-b-visible-content P037; donor insertion map topic-level redraw", "notes": "P037。Baseline 使用固定 scenario、seed、traffic、window 與 endpoint scope，並保留 release policy 作為對照執行。Candidate 只改目前 lab 的一個 marked block，觀察 state duration、packet outcome、service verdict 與 endpoint J。policy identity 來自 result 的 policy 欄位，單位是識別碼；它與 scenario、seed、units、provenance 一起建立公平比較。"},
    {"page": 38, "title": "LEO：changing-service-window trace", "source": "part-b-visible-content P038; v2 window authority", "notes": "P038。LEO 在本模組只提供 changing-service-window trace。contact_open 是服務窗口布林狀態，quality_band 是 0 到 3 的品質分類，contact_remaining_s 是窗口剩餘時間，三者都由 fixed scenario trace 進入 policy observation。窗口關閉時 runner 只接受 SLEEP；窗口開啟時，quality 與剩餘時間共同改變 action 的服務機會。"},
    {"page": 39, "title": "Lab A：等待空檔的 radio state", "source": "part-b-visible-content P039; part-b-field-audit P039", "notes": "P039。Lab A 的固定條件包含 scenario、case workload、service window、seed、endpoint energy model 與 PACE_GAP_STEPS = 2。baseline 的 REST_DURING_GAP = SLEEP；candidate 的唯一 edit 是 REST_DURING_GAP = WAIT。這個 edit 先改變空檔的 radio state，再由 state interval 影響 packet timing、service gate 與 endpoint J。改前先保存 WAIT 可能減少 WAKE、也可能增加 awake idle 的可反駁預測。"},
    {"page": 40, "title": "A-00：Baseline decision 與 evidence", "source": "part-b-visible-content P040; current /course A metrics element capture", "notes": "P040。Baseline decision 在 contact_open 為 false 時回傳 SLEEP；steps_since_send 小於 PACE_GAP_STEPS 時回傳 REST_DURING_GAP，因此 release baseline 的 gap action 仍是 SLEEP。A / baseline 的 control evidence 使用 P042 第一條 stdout result_path，參考 summary 是 endpoint_energy_j 6.92 J、delivered_bits 4,800 bit、service_pass false。current /course metrics element capture 標示同情境備用資料；實際教學判讀仍以 stdout path 的 result.json 與同 run replay 為準。"},
    {"page": 41, "title": "A-01：REST_DURING_GAP 的唯一 edit", "source": "part-b-visible-content P041; student_policy.py marked block", "notes": "P041。可編輯檔案是 student_policy.py，標記區段是 lab-a-pace-rest。PACE_GAP_STEPS = 2 保持原值；REST_DURING_GAP 從 SLEEP 改為 WAIT，其他 marked block、runner、scenario、schema 與 generated JSON 保持原狀。備份命令先保存原始 policy，py_compile 只確認語法與編譯回傳碼。機制預測是 WAIT 先改變 awake-idle state，再可能改變 wake、packet、service 與 endpoint J。"},
    {"page": 42, "title": "A-02：Lab A compact run／receipt", "source": "part-b-visible-content P042; venv platform contract", "notes": "P042。第一條命令以 baseline 建立 control；完成 REST_DURING_GAP = WAIT 後，第二條命令以 candidate --freeze 建立 frozen policy 與 checkpoint；第三條命令保持 frozen policy 執行 hidden。POSIX 使用 bash course.sh，Windows Command Prompt 使用 course.cmd。每次成功 stdout 都提供新的 result_path，配對的 endpoint-replay.json 位於同一 generated run 目錄，candidate 另保存 lab-a-frozen checkpoint。"},
    {"page": 43, "title": "A-03：同一 boundary 下的比較", "source": "part-b-visible-content P043; current result contract", "notes": "P043。baseline 與 candidate 都從 P042 stdout result_path 開啟。比較前核對 scenario_id、seed、policy identity、units、provenance 與 endpoint scope，確定兩筆結果仍屬同一 boundary。參考 summary 為 baseline 6.92 J／4,800 bit、candidate 8.86 J／4,800 bit，兩者 service_pass 都是 false；因此解釋必須回到 state 與 packet evidence，而不是只報 J。"},
    {"page": 44, "title": "A-04：state ledger 連到 endpoint J", "source": "part-b-visible-content P044; current /course A replay-frame element capture", "notes": "P044。Candidate result 的 events 先找 STATE_INTERVAL 與 WAKE，再對照 energy_breakdown_j。參考 bucket 是 awake_idle 6.00 J、sleep 0.04 J、wake 0.02 J、process 0.16 J、tx 2.40 J、rx 0.24 J；這些值要以該次 stdout result_path 的 artifact 為準。current /course replay-frame element capture 是同情境備用資料，用來對齊 state、queue、contact 與累積 J 的畫面 vocabulary，不冒充 candidate 的新 run。"},
    {"page": 45, "title": "A-05：attempted 不等於 delivered", "source": "part-b-visible-content P045; current /course A ledger element capture", "notes": "P045。Candidate packet ledger 的參考摘要是 attempted_packets 2、retransmissions 1、delivered_bits 4,800 bit、expired_packets 3、deadline_pass false、service_pass false。current /course execution-ledger element capture 顯示同情境備用資料的 baseline row，candidate 的數值仍由 P042 stdout result_path 的 result.json 取得。因果解釋要指出 action 先改變 state，再改變 packet outcome，最後由 service gate 決定 endpoint J 的解讀層級。"},
]

# Keep the narration authoritative and directly readable after the visual
# rewrite.  The older metadata above is retained as source provenance, while
# these overrides make the exported notes use the same Chinese-first grammar
# as the repaired pages.
_PAGE_META_OVERRIDES = {
    28: ("無線狀態與停留時間（Radio state）", "這一張先建立狀態 vocabulary。低功耗休息 `SLEEP` 的功率較低；清醒閒置 `WAIT` 會累積醒著閒置；喚醒 `WAKE` 與封包處理 `PROCESS` 是轉換或處理區間；發射 `TX` 與接收 `RX` 是收發區間。每個狀態由 runner events 保存停留秒數與事件能量，之後才能把動作 `action` 連到 state interval、energy bucket 與端點 J。動作名稱只記一次，不能代替完整狀態帳本。"),
    29: ("實驗 A：空檔策略的服務與能量", "這一張說明為何做實驗 A。固定同一份工作、scenario、seed、traffic、service window、endpoint scope、runner、schema 與 Lab B/C 區塊，只把 `student_policy.py` 的 `lab-a-pace-rest` 中 `REST_DURING_GAP` 由 `SLEEP` 改成 `WAIT`。`SLEEP` 降低空檔功率但可能支付 `WAKE`；`WAIT` 保持清醒閒置，可能少喚醒卻增加 awake-idle J。每次成功要看 stdout 的 `status=OK` 與新 `result_path`，再讀同一 generated run 目錄的 `result.json` 與 `endpoint-replay.json`，順序是服務、封包、狀態、J。"),
    30: ("封包生命週期：嘗試不等於交付", "封包先建立 `generated`，再進入 `queue` 等待合法窗口；策略動作造成 `attempt`，碰撞 `collision` 可能產生 `retry`；最後才有 `delivered` 或 `expired`。嘗試封包數、重傳次數、交付位元與過期封包各有來源與單位，服務通過 `service_pass` 還要共同檢查必要封包、期限與新鮮度。送出次數只能說明嘗試，不是交付或服務成功。"),
    31: ("服務門檻先於端點能量", "公平比較先固定 scenario、seed、traffic、window 與 endpoint scope，再看交付位元 `delivered_bits`、期限 `deadline_pass`、新鮮度 `freshness_status` 與服務通過 `service_pass`。服務門檻先回答工作是否在期限內完成；只有在相同邊界且服務條件清楚時，端點能量 `endpoint_energy_j` 才能用來描述取捨。較低 J 不會自動變成服務成功。"),
    32: ("功率 W × 時間 s 形成能量 J", "功率 P 是狀態當下的高度，單位 W；時間 t 是狀態區間的長度，單位 s；兩者相乘才形成能量 J。每個 bucket 分開保存 awake_idle、sleep、wake、process、tx、rx，WAIT 與 SLEEP 使用固定 10 s 時鐘步長，收發流程另有轉換能量。讀值時保留功率、停留時間與 bucket，避免用單一峰值代替整段累積。"),
    33: ("端點能量的累積公式", "端點能量 `E_endpoint` 的可編輯 Office Math 公式是各狀態功率 `P_s` 與停留時間 `t_s` 的分段累積。`P_s` 單位為 W，來自 energy profile；`t_s` 單位為 s，來自 `STATE_INTERVAL`；結果摘要的 `endpoint_energy_j` 單位為 J。這個 endpoint scope 只涵蓋宣告的 radio 與 processing 狀態及轉換，不包含衛星、閘道或整體插座功率。"),
    34: ("端點能量效率的 bit/J 邊界", "端點效率 `η_E` 的分子是交付資料 `D_delivered`，單位 bit；分母是端點能量 `E_endpoint`，單位 J；結果單位 bit/J。兩者必須來自同一份 result artifact 與同一 endpoint boundary。`service_pass` 與 `deadline_pass` 仍是獨立判定，因此 bit/J 變高不能覆寫服務失敗。"),
    35: ("來源、模型、課程假設與結果", "來源 `Source` 是可追溯的上游 repository／commit 參考，不等於本次量測；模型 `Model` 是固定 runner 行為的課程模擬轉接器；課程假設 `Course` 固定 scenario、window、traffic 與 endpoint profile；結果 `Result` 保存本次 JSON 與同次 replay。`upstream_execution=false` 表示不是上游即時執行，所有數值回到產物與宣告的 claim boundary。"),
    36: ("執行、事件重播與工作簿的身分鏈", "執行 `RUN` 讀 scenario、case 與 policy，輸出 `result.json` 與同目錄 `endpoint-replay.json`。事件重播 `REPLAY` 只讀同一次 run 的 events，顯示 action、radio state、queue 與 packet，不重新執行 policy。比較紀錄 `WORKBOOK` 保存 prediction、baseline、candidate、freeze、withheld 與 result／replay lineage。匯入前核對 `scenario_id`、`run_id`、policy identity、units 與 provenance；任一 mismatch 都維持 fail closed。"),
    37: ("基準 Baseline：固定條件的對照執行", "基準 `Baseline` 使用固定 scenario、seed、traffic、window 與 endpoint scope，保留原始 `student_policy.py` 作為 control。唯一變動只能是目前實驗的 marked block，輸出再觀察 state duration、packet outcome、service verdict 與 endpoint J。兩次結果的身分、單位與端點範圍必須一致，差異才可回到該標記區塊解釋。"),
    38: ("LEO 的變動服務窗口軌跡", "LEO 在本模組只提供變動服務窗口軌跡。`contact_open` 是窗口是否開啟的布林狀態；`quality_band` 是 0–3 的品質分類，不是 dB；`contact_remaining_s` 是剩餘秒數。三者由固定 scenario trace 進入 policy observation；窗口關閉時只接受 `SLEEP`，窗口開啟時品質與剩餘時間共同改變動作的服務機會。"),
    39: ("實驗 A：為何改 REST_DURING_GAP", "實驗 A 的目的，是在固定工作與服務窗口下檢驗一個可反駁機制。檔案是 `student_policy.py`，標記區塊是 `lab-a-pace-rest`；基準為 `REST_DURING_GAP=SLEEP`，候選只改成 `WAIT`。固定 `PACE_GAP_STEPS=2`、scenario、seed、traffic、window、runner、schema 與 Lab B/C 區塊。成功訊號是 stdout `status=OK` 與新 `result_path`；同一目錄有 `result.json` 與 `endpoint-replay.json`。匯入或檢視時依服務、封包、狀態、J 順序讀取。"),
    40: ("A-00：基準決策與對照證據", "基準的 `student_policy.py`／`lab-a-pace-rest` 保持 `PACE_GAP_STEPS=2` 與 `REST_DURING_GAP=SLEEP`。窗口未開時 `contact_open=false` 回傳 `SLEEP`；送出間隔不足時回傳空檔策略。Linux／macOS 終端機或 WSL 執行 `bash course.sh run --lab A --case baseline`；成功 stdout 印出 `status=OK` 與新的 `result_path`，同目錄產生 `result.json` 與 `endpoint-replay.json`。先讀 summary／events，再讀 replay。6.92 J、4,800 bit、service false 是參考摘要，實際數值以該 path 為準；基準建立 before 對照，不回答是否省電。"),
    41: ("A-01：只改 REST_DURING_GAP", "這一張的唯一 edit 是 `student_policy.py` 的 `lab-a-pace-rest`：把 `REST_DURING_GAP = SLEEP` 改成 `REST_DURING_GAP = WAIT`。`PACE_GAP_STEPS=2`、其他 marked blocks、runner、scenario、schema 與 generated JSON 都保持不變。Linux／macOS 終端機或 WSL 先執行 `cp student_policy.py student_policy.before-A-edit.py`，再執行 `.venv/bin/python -m py_compile student_policy.py`；Windows 使用 `copy /Y` 與 `.venv\\Scripts\\python.exe -m py_compile`。成功是無語法錯誤且回傳碼 0；這一張只做編輯檢查，下一張才取得 run 的 result_path。預測是 WAIT 可能減少 WAKE 但增加 awake-idle，方向必須由後續 state、packet、service、J 證據決定。"),
    42: ("A-02：基準 → 候選 → 隱藏檢查", "這是實驗 A 唯一的執行與 receipt 頁。第一步不改 policy 執行 `bash course.sh run --lab A --case baseline` 或 Windows `course.cmd`，成功要有 `status=OK` 與新 `result_path`。第二步完成 P041 的 WAIT edit，執行 candidate `--freeze`；成功除新 path 外，還要有 `artifacts/checkpoints/lab-a-frozen.json` 與 `.py`。第三步保持凍結策略執行 hidden，不再編輯；成功只接受新的 path，不建立第二個 freeze。每個 path 讀同一 generated 目錄的 result 與 replay，順序是服務、封包、狀態、J；缺少 freeze 就停止 hidden。"),
    43: ("同一邊界的改前／改後", "三次 evidence 都從 P042 stdout 的 result_path 讀取；比較前核對 scenario_id、seed、policy identity、units、provenance 與 endpoint scope。baseline SLEEP：service=false、deadline=false、delivered=1/4、retries=1、collisions=1、expired=3、wake=2、endpoint=6.92 J、bit/J=693.641618。candidate WAIT：service=false、deadline=false、delivered=1/4、retries=1、collisions=1、expired=3、wake=1、endpoint=8.86 J、bit/J=541.760722。hidden frozen WAIT：service=false、deadline=false、delivered=0/3、retries=0、collisions=0、expired=3、wake=0、endpoint=3.61 J、bit/J=0。WAIT 沒改善服務，且 endpoint 多耗 1.94 J；hidden 的低 J 來自完全未送達，不能視為節能成功。"),
    44: ("A-04：狀態帳本回到端點能量", "使用 candidate 的 result_path，先在 events 找 `STATE_INTERVAL` 的 WAIT／SLEEP 停留秒數，再找 `WAKE` 轉換，最後對照 `energy_breakdown_j`。參考 bucket 是 awake_idle 6.00 J、sleep 0.04 J、wake 0.02 J、process 0.16 J、tx 2.40 J、rx 0.24 J；實際數值以本次 artifact 為準。解釋句要指出哪個狀態區間增加或減少、哪個 bucket 因而改變，才能把 WAIT 的機制連回 endpoint J。"),
    45: ("封包結果回到服務判定", "先把 baseline SLEEP 與 candidate WAIT 對齊：delivered 皆為 1/4、retries=1、collisions=1、expired=3、deadline=false、service=false；WAIT 只讓 wake 由 2 降到 1，endpoint 卻由 6.92 J 升到 8.86 J（+1.94 J），bit/J 由 693.641618 降到 541.760722。hidden frozen WAIT 是 delivered=0/3、retries=0、collisions=0、expired=3、wake=0、endpoint=3.61 J、bit/J=0，低 J 的原因是完全未送達，不是節能成功。使用各自 P042 stdout result_path 的 result.json 與配對 endpoint-replay.json；服務門檻先於端點 J，嘗試或較低 J 都不能覆寫 service=false。"),
}
for _meta in PAGE_META:
    if _meta["page"] in _PAGE_META_OVERRIDES:
        _meta["title"], _meta["notes"] = _PAGE_META_OVERRIDES[_meta["page"]]


def emit_sources() -> None:
    SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    template_copy = SOURCE_DIR / "educate.pptx"
    if TEMPLATE.resolve() != template_copy.resolve():
        shutil.copy2(TEMPLATE, template_copy)
    manifest = {
        "schema": "c120-part-b-v2-p028-p045-source-v1",
        "template": str(TEMPLATE),
        "template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        "layout": "ppt/slideLayouts/slideLayout2.xml",
        "font_contract": {"cjk": "標楷體", "latin": "Times New Roman"},
        "claim_boundary": CLAIM_BOUNDARY,
        "evidence_assets": {
            "metrics": str(METRICS_IMAGE),
            "replay_frame": str(REPLAY_IMAGE),
            "ledger": str(LEDGER_IMAGE),
            "label": "current /course element captures; same-scenario fallback labels preserved",
        },
        "slides": PAGE_META,
    }
    (OWNED / "slides.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    notes_md = ["# Part B V2 speaker notes (P028–P045)", ""]
    source_md = ["# Part B V2 owned source (P028–P045)", "", "Source authority: teaching-rewrite/part-b-visible-content.md and part-b/slides.json.", "", "The builder is the owning implementation source; authority files remain read-only.", ""]
    for item in PAGE_META:
        notes_md.extend([f"## P{item['page']:03d}｜{item['title']}", "", item["notes"], ""])
        source_md.extend([f"## P{item['page']:03d}｜{item['title']}", "", f"Source: {item['source']}", "", ""])
    (SOURCE_DIR / "speaker-notes.md").write_text("\n".join(notes_md), encoding="utf-8")
    (SOURCE_DIR / "module-source.md").write_text("\n".join(source_md), encoding="utf-8")


def _ordered_slide_parts(data: dict[str, bytes]) -> list[str]:
    pres = ET.fromstring(data["ppt/presentation.xml"])
    rels = ET.fromstring(data["ppt/_rels/presentation.xml.rels"])
    targets = {rel.get("Id"): rel.get("Target") for rel in rels.findall(f"{{{PKG_REL_NS}}}Relationship")}
    parts = []
    for node in pres.findall(f"./{{{P_NS}}}sldIdLst/{{{P_NS}}}sldId"):
        target = targets.get(node.get(f"{{{OFFICE_REL_NS}}}id"))
        if target:
            parts.append(posixpath.normpath(posixpath.join("ppt", target)))
    return parts


def _rel_source_part(rel_name: str) -> str:
    if rel_name == "_rels/.rels":
        return ""
    rel_dir = posixpath.dirname(rel_name)
    source_dir = posixpath.dirname(rel_dir)
    base = posixpath.basename(rel_name)[:-5]
    return posixpath.join(source_dir, base)


def _resolved_target(rel_name: str, target: str) -> str:
    source_part = _rel_source_part(rel_name)
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _text_nodes(root) -> list[str]:
    return [node.text or "" for node in root.iter(f"{{{A_NS}}}t")]


def _placeholder_type(shape) -> str | None:
    ph = shape.find(f"./{{{P_NS}}}nvSpPr/{{{P_NS}}}nvPr/{{{P_NS}}}ph")
    return ph.get("type") if ph is not None else None


def package_qa(path: Path) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    rel_errors: list[str] = []
    xml_errors: list[str] = []
    low_text: list[dict] = []
    forbidden_hits: list[dict] = []
    slide_reports: list[dict] = []
    creation_ids: list[str] = []
    c_nv_ids: dict[str, list[str]] = {}
    formula_slides: list[dict] = []
    authored_font_sizes: list[float] = []
    authored_font_attributes: dict[str, dict[str, int]] = {"latin": {}, "ea": {}, "cs": {}}
    font_contract_errors: list[dict] = []
    notes_body_placeholders = 0
    expected_fonts = {"latin": "Times New Roman", "ea": "標楷體", "cs": "Times New Roman"}
    forbidden = [
        "學生", "老師", "講師", "你", "請", "分鐘", "checksum", "SHA", "ZIP test",
        "製作備註", "舊截圖", "截圖", "講稿提示", "版面", "Fit", "本頁", "這一頁",
        "能回答", "能指出", "能說出", "可以回答", "可教", "先辨認 radio state，才談省電",
        "不代表結果變成",
    ]
    with zipfile.ZipFile(path, "r") as deck:
        names = set(deck.namelist())
        if deck.testzip() is not None:
            errors.append("zip_integrity_failed")
        data = {name: deck.read(name) for name in names}
        for name in sorted(names):
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    ET.fromstring(data[name])
                except ET.ParseError as exc:
                    xml_errors.append(f"{name}: {exc}")
        if xml_errors:
            errors.append(f"recursive_xml_parse_errors={len(xml_errors)}")

        for rel_name in sorted(name for name in names if name.endswith(".rels")):
            try:
                root = ET.fromstring(data[rel_name])
            except ET.ParseError:
                continue
            for rel in root.findall(f"{{{PKG_REL_NS}}}Relationship"):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target", "")
                resolved = _resolved_target(rel_name, target)
                if resolved not in names:
                    rel_errors.append(f"{rel_name}: {target} -> {resolved}")
        if rel_errors:
            errors.append(f"relationship_target_errors={len(rel_errors)}")

        parts = _ordered_slide_parts(data)
        notes = sorted((name for name in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)), key=lambda item: int(re.search(r"(\d+)", item).group(1)))
        if len(parts) != 18:
            errors.append(f"slide_count={len(parts)} expected=18")
        if len(notes) != 18:
            errors.append(f"notes_count={len(notes)} expected=18")

        for index, part in enumerate(parts, start=1):
            root = ET.fromstring(data[part])
            rel_name = f"ppt/slides/_rels/{Path(part).name}.rels"
            rel_root = ET.fromstring(data[rel_name])
            layout_targets = []
            for rel in rel_root.findall(f"{{{PKG_REL_NS}}}Relationship"):
                if rel.get("Type", "").endswith("/slideLayout"):
                    layout_targets.append(posixpath.normpath(posixpath.join(posixpath.dirname(part), rel.get("Target", ""))))
            if layout_targets != ["ppt/slideLayouts/slideLayout2.xml"]:
                errors.append(f"slide_{index}_layout={layout_targets}")
            if root.find(f"./{{{P_NS}}}cSld/{{{P_NS}}}bg") is not None:
                errors.append(f"slide_{index}_authored_background")
            shape_tree = root.find(f"./{{{P_NS}}}cSld/{{{P_NS}}}spTree")
            ids = []
            if shape_tree is not None:
                for shape in list(shape_tree):
                    c_nv_pr = shape.find(f"./{{{P_NS}}}nvSpPr/{{{P_NS}}}cNvPr")
                    if c_nv_pr is not None:
                        value = c_nv_pr.get("id")
                        if value:
                            ids.append(value)
                        if shape.tag == f"{{{P_NS}}}sp" and _placeholder_type(shape) not in {"date", "footer", "sldNum"}:
                            for rpr in shape.findall(f".//{{{A_NS}}}rPr"):
                                if rpr.get("sz"):
                                    authored_font_sizes.append(int(rpr.get("sz")) / 100)
                            for run in shape.findall(f".//{{{A_NS}}}r"):
                                rpr = run.find(f"{{{A_NS}}}rPr")
                                if rpr is None:
                                    continue
                                for attr, expected in expected_fonts.items():
                                    font = rpr.find(f"{{{A_NS}}}{attr}")
                                    if font is None:
                                        font_contract_errors.append({"slide": index, "attribute": attr, "value": None})
                                        continue
                                    value = font.get("typeface")
                                    authored_font_attributes[attr][value] = authored_font_attributes[attr].get(value, 0) + 1
                                    if value != expected:
                                        font_contract_errors.append({"slide": index, "attribute": attr, "value": value})
            c_nv_ids[part] = ids
            if len(ids) != len(set(ids)):
                errors.append(f"slide_{index}_duplicate_cNvPr_ids")
            for creation in root.iter():
                if creation.tag.rsplit("}", 1)[-1] == "creationId":
                    value = creation.get("id") or creation.get("val")
                    if value:
                        creation_ids.append(value)
            text = "".join(_text_nodes(root))
            for term in forbidden:
                if term.lower() in text.lower():
                    forbidden_hits.append({"location": part, "term": term})
            omath = root.findall(f".//{{{M_NS}}}oMath")
            if omath:
                formula_slides.append({"slide": index, "omath": len(omath)})
            slide_reports.append({
                "slide": index,
                "part": part,
                "layout_targets": layout_targets,
                "authored_text_count": len(_text_nodes(root)),
                "omath_count": len(omath),
                "has_slide_background": root.find(f"./{{{P_NS}}}cSld/{{{P_NS}}}bg") is not None,
            })
        for note_name in notes:
            root = ET.fromstring(data[note_name])
            note_text = "".join(_text_nodes(root))
            if any(_placeholder_type(shape) == "body" for shape in root.findall(f".//{{{P_NS}}}sp")):
                notes_body_placeholders += 1
            if not note_text.strip():
                errors.append(f"empty_notes={note_name}")
            for term in forbidden:
                if term.lower() in note_text.lower():
                    forbidden_hits.append({"location": note_name, "term": term})

        if len(formula_slides) != 2 or [(item["slide"], item["omath"]) for item in formula_slides] != [(6, 1), (7, 1)]:
            errors.append(f"omml_contract={formula_slides}")
        if any(item["omath_count"] > 0 and "pic" in str(item) for item in slide_reports):
            warnings.append("formula_picture_check_requires_full_shape_scan")

    if forbidden_hits:
        errors.append(f"forbidden_language_hits={len(forbidden_hits)}")
    if creation_ids and len(creation_ids) != len(set(creation_ids)):
        errors.append("duplicate_creation_ids")
    if authored_font_sizes and min(authored_font_sizes) < 16:
        errors.append(f"font_floor={min(authored_font_sizes)}")
    if font_contract_errors:
        errors.append(f"font_contract_errors={len(font_contract_errors)}")
    if notes_body_placeholders != len(notes):
        errors.append(f"notes_body_placeholders={notes_body_placeholders} expected={len(notes)}")
    template_preserved = True
    with zipfile.ZipFile(TEMPLATE, "r") as template, zipfile.ZipFile(path, "r") as deck:
        for name in template.namelist():
            if name.startswith(("ppt/slideLayouts/", "ppt/slideMasters/", "ppt/theme/", "ppt/notesMasters/")):
                if name not in deck.namelist() or hashlib.sha256(template.read(name)).hexdigest() != hashlib.sha256(deck.read(name)).hexdigest():
                    template_preserved = False
            if name.startswith("ppt/media/") and name in deck.namelist() and hashlib.sha256(template.read(name)).hexdigest() != hashlib.sha256(deck.read(name)).hexdigest():
                template_preserved = False
        if not template_preserved:
            errors.append("template_owned_parts_changed")
    report = {
        "schema": "c120-part-b-v2-p028-p045-qa-v1",
        "status": "PASS" if not errors else "FAIL",
        "output": str(path),
        "slide_count": len(parts),
        "notes_count": len(notes),
        "layout_targets": sorted({target for slide in slide_reports for target in slide["layout_targets"]}),
        "all_slides_layout2": all(slide["layout_targets"] == ["ppt/slideLayouts/slideLayout2.xml"] for slide in slide_reports),
        "slide_level_background_count": sum(1 for slide in slide_reports if slide["has_slide_background"]),
        "creation_id_count": len(creation_ids),
        "creation_ids_unique": len(creation_ids) == len(set(creation_ids)),
        "authored_font_floor_pt": min(authored_font_sizes) if authored_font_sizes else None,
        "authored_font_attributes": authored_font_attributes,
        "font_contract_errors": font_contract_errors,
        "notes_body_placeholders": notes_body_placeholders,
        "formula_slides": formula_slides,
        "template_owned_parts_preserved": template_preserved,
        "recursive_xml_parse_errors": xml_errors,
        "relationship_target_errors": rel_errors,
        "forbidden_language_hits": forbidden_hits,
        "slide_reports": slide_reports,
        "warnings": warnings,
        "limitations": [
            "LibreOffice executable is absent in the server environment; render and PowerPoint reopen are deferred to controller environment.",
            "Current evidence images are clean element captures labelled as same-scenario fallback; no SHA-like identifier is embedded.",
        ],
    }
    QA_DIR.mkdir(parents=True, exist_ok=True)
    (QA_DIR / "qa_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def write_visual_qa() -> dict:
    renderer = shutil.which("soffice") or shutil.which("libreoffice")
    result = {
        "schema": "c120-part-b-v2-p028-p045-visual-qa-v1",
        "status": "DEFERRED" if not renderer else "PENDING_RENDER",
        "renderer": renderer,
        "rendered_slides": 0,
        "inspection": "deferred to controller environment" if not renderer else "renderer available; run original-size conversion and inspect every slide",
        "fix_and_rerender_cycle": "deferred" if not renderer else "required",
        "evidence": [str(METRICS_IMAGE), str(REPLAY_IMAGE), str(LEDGER_IMAGE)],
    }
    (QA_DIR / "visual-qa.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def build() -> tuple[dict, dict]:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    for directory in (BUILD_DIR, SOURCE_DIR, QA_DIR, RENDER_DIR):
        directory.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)
    for index, builder in enumerate(SLIDE_BUILDERS):
        builder(prs)
        # Notes are authored with the slide so each exported page has a
        # directly readable narration in the native notes relationship.
        prs.slides[-1].notes_slide.notes_text_frame.text = PAGE_META[index]["notes"]
    working = BUILD_DIR / "python-pptx-working.pptx"
    prs.save(str(working))
    shutil.copy2(working, OUTPUT)
    overlay_template_parts(OUTPUT)
    # Early editable export is now present at the assigned root-level path.
    print(json.dumps({"early_export": str(OUTPUT), "slides": len(SLIDE_BUILDERS)}, ensure_ascii=False))
    add_native_math(OUTPUT, {6: "endpoint", 7: "efficiency"})
    shutil.copy2(OUTPUT, BUILD_DIR / "post-math.pptx")
    report = package_qa(OUTPUT)
    visual = write_visual_qa()
    emit_sources()
    readme = {
        "schema": "c120-part-b-v2-p028-p045-readme-v1",
        "output": str(PUBLISH_TARGET),
        "work_output": str(OUTPUT),
        "slides": "P028–P045 (18 slides)",
        "template": str(TEMPLATE),
        "template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        "layout": "slideLayout2.xml only",
        "qa": str(QA_DIR / "qa_report.json"),
        "visual_qa": visual,
        "renderer_unknowns": ["original-size render and PowerPoint reopen remain pending until controller QA"],
        "no_commit_push": True,
    }
    (OWNED / "README.md").write_text(json.dumps(readme, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report, visual


if __name__ == "__main__":
    qa, visual = build()
    print(json.dumps({"output": str(OUTPUT), "qa_status": qa["status"], "slides": qa["slide_count"], "notes": qa["notes_count"], "visual_status": visual["status"]}, ensure_ascii=False, indent=2))
    raise SystemExit(0 if qa["status"] == "PASS" else 2)
