#!/usr/bin/env python3
"""Build the owned P019–P027b section from the educate layout-2 shell.

The source slide manifest is read from the authoritative Part A slides.json.
This builder deliberately uses python-pptx against the exact template and
restores template-owned parts after save so every authored page inherits the
same background, logo, divider, footer, and slide-number system.
"""

from __future__ import annotations

import copy
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ALT_ROOT = ROOT.parent
REPO_ROOT = ALT_ROOT.parent.parent.parent
TEMPLATE = Path("/home/u24/pptx-wrap/assets/templates/educate.pptx")
SOURCE_JSON = ALT_ROOT / "part-a" / "slides.json"
LATEST_ROOT = ALT_ROOT / "latest"
OUTPUT = LATEST_ROOT / "LoRaEnergySim-LEO-ALT-PART-A-V2-P019-P027B-REVIEW.pptx"
SOURCE_SNAPSHOT = ROOT / "source-snapshot.json"

NAVY = "25356B"
INK = "25324A"
MUTED = "65758B"
BLUE = "4967C2"
BLUE_PALE = "ECF1FE"
TEAL = "007374"
TEAL_PALE = "E8F7F4"
PURPLE = "660066"
PURPLE_PALE = "F4EDF7"
GOLD = "846826"
GOLD_PALE = "FDF9E2"
RED = "992A2A"
RED_PALE = "FDEFEF"
WHITE = "FFFFFF"
LINE = "D7DEEA"
COOL = "F7F8FC"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_fonts(run) -> None:
    """Use the contract fonts while allowing PowerPoint script fallback."""
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("latin", "Times New Roman"), ("ea", "標楷體"), ("cs", "Times New Roman")):
        child = rpr.find(qn(f"a:{tag}"))
        if child is None:
            child = OxmlElement(f"a:{tag}")
            rpr.append(child)
        child.set("typeface", face)


def add_marked_runs(paragraph, text: str, size: float, color: str,
                    bold: bool = False, italic: bool = False) -> None:
    """Backtick spans are editable field names, variables, or code tokens."""
    pieces = re.split(r"(`[^`]+`)", text)
    for piece in pieces:
        if not piece:
            continue
        marked = piece.startswith("`") and piece.endswith("`")
        value = piece[1:-1] if marked else piece
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic or marked
        run.font.color.rgb = rgb(color)
        set_run_fonts(run)


def write_text(shape, text: str, size: float = 24, color: str = INK,
               bold: bool = False, align=PP_ALIGN.LEFT,
               valign=MSO_ANCHOR.MIDDLE, margins=(0.08, 0.04, 0.08, 0.04),
               line_spacing: float = 1.0, italic: bool = False) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    for index, line in enumerate(text.split("\n")):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        add_marked_runs(paragraph, line, size, color, bold, italic)


def add_text(slide, x, y, w, h, text, size=24, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margins=(0.02, 0.01, 0.02, 0.01), line_spacing=1.0,
             name="Text", italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    write_text(shape, text, size, color, bold, align, valign, margins,
               line_spacing, italic)
    return shape


def add_box(slide, x, y, w, h, fill, line=LINE, radius=True,
            name="Box", line_width=1.1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_chevron(slide, x, y, w=0.42, h=0.52, color=BLUE, name="Mechanism connector"):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_label_box(slide, x, y, w, h, text, fill, line, size=22,
                  color=INK, bold=False, align=PP_ALIGN.CENTER,
                  name="Label", italic=False):
    shape = add_box(slide, x, y, w, h, fill, line, name=name)
    write_text(shape, text, size, color, bold, align, MSO_ANCHOR.MIDDLE,
               margins=(0.12, 0.04, 0.12, 0.04), line_spacing=0.95,
               italic=italic)
    return shape


def add_card(slide, x, y, w, h, title, body, fill=BLUE_PALE, line=BLUE,
             title_color=NAVY, body_color=INK, title_size=22, body_size=19,
             name="Teaching card", title_italic=False, title_h=None):
    add_box(slide, x, y, w, h, fill, line, name=name)
    # The first pass used a fixed 0.35-inch title frame.  A wrapped heading
    # then occupied the body frame and made the card appear to have duplicated
    # or colliding text.  Reserve title height from the actual heading shape;
    # the caller can override it for a deliberately compact card.
    if title_h is None:
        title_lines = max(1, title.count("\n") + 1)
        # Approximate the extra wrap caused by narrow cards.  This is kept
        # conservative so that a heading never falls into the body region.
        title_len = len(title.replace("`", ""))
        if w <= 2.7 and title_len > 15:
            title_lines += 1
        elif w <= 3.6 and title_len > 24:
            title_lines += 1
        title_h = 0.16 + 0.30 * min(title_lines, 3)
    add_text(slide, x + 0.16, y + 0.10, w - 0.32, title_h, title, title_size,
             title_color, True, PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             name=f"{name} heading", italic=title_italic,
             line_spacing=0.90)
    body_y = y + title_h + 0.12
    body_h = max(0.20, h - title_h - 0.20)
    add_text(slide, x + 0.16, body_y, w - 0.32, body_h, body, body_size,
             body_color, False, PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             name=f"{name} explanation", line_spacing=0.90)


def add_code_box(slide, x, y, w, h, code, fill="F6F7FB", line=BLUE,
                 size=19, name="Editable code"):
    add_box(slide, x, y, w, h, fill, line, name=name, line_width=1.0)
    add_text(slide, x + 0.16, y + 0.10, w - 0.32, h - 0.20, code, size,
             NAVY, False, PP_ALIGN.LEFT, MSO_ANCHOR.TOP,
             margins=(0.01, 0.01, 0.01, 0.01), line_spacing=0.96,
             name=f"{name} text")


def add_field_band(slide, field: str, chinese: str, source: str,
                   purpose: str, unit: str, interpretation: str,
                   color=BLUE, y=5.42, h=0.90):
    band = add_box(slide, 0.88, y, 11.55, h, COOL, color,
                   name=f"Field contract {field}", line_width=1.0)
    text = (
        f"`{field}`（{chinese}）｜來源：{source}｜作用：{purpose}\n"
        f"單位：{unit}｜判讀：{interpretation}"
    )
    write_text(band, text, 18, NAVY, False, PP_ALIGN.LEFT,
               MSO_ANCHOR.MIDDLE, margins=(0.18, 0.03, 0.18, 0.03),
               line_spacing=0.90)


def delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def remove_all_slides(prs: Presentation) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        slide_ids.remove(slide_id)


def prepare_slide(prs: Presentation, title: str):
    """Create only from the second source layout and retain footer placeholders."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    if title_shape is None:
        raise RuntimeError("slideLayout2 title placeholder missing")
    for placeholder in list(slide.placeholders):
        if placeholder._element is title_shape._element:
            continue
        typ = placeholder.placeholder_format.type
        top = placeholder.top / 914400
        if typ in {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}:
            continue
        if top >= 6.55:
            continue
        delete_shape(placeholder)
    title_shape.left = Inches(0.718057)
    title_shape.top = Inches(0.204514)
    title_shape.width = Inches(10.34861)
    title_shape.height = Inches(0.525)
    title_shape.name = "Native layout2 title"
    write_text(title_shape, title, 28, NAVY, True, PP_ALIGN.LEFT,
               MSO_ANCHOR.MIDDLE, margins=(0.02, 0.0, 0.02, 0.0))
    return slide


def overlay_template_parts(path: Path) -> None:
    """Restore template-owned parts after python-pptx saves the authored slides."""
    owned_prefixes = (
        "ppt/slideLayouts/",
        "ppt/slideMasters/",
        "ppt/theme/",
        "ppt/media/",
        "ppt/notesMasters/",
    )
    with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(path, "r") as old:
        data = {info.filename: old.read(info.filename) for info in old.infolist()}
        infos = {info.filename: copy.copy(info) for info in old.infolist()}
        for info in source.infolist():
            if info.filename.startswith(owned_prefixes):
                data[info.filename] = source.read(info.filename)
                infos[info.filename] = copy.copy(info)
    overlay = path.with_suffix(".overlay.pptx")
    with zipfile.ZipFile(overlay, "w", zipfile.ZIP_DEFLATED) as out:
        for name, payload in data.items():
            out.writestr(infos[name], payload)
    overlay.replace(path)


def load_authority() -> list[dict]:
    data = json.loads(SOURCE_JSON.read_text(encoding="utf-8"))
    wanted = {
        "P019", "P020", "P021a", "P021b", "P021c", "P022", "P023",
        "P024a", "P024b", "P025a", "P025b", "P026a", "P026b", "P027a", "P027b",
    }
    slides = [slide for slide in data["slides"] if slide["id"] in wanted]
    if [slide["id"] for slide in slides] != [
        "P019", "P020", "P021a", "P021b", "P021c", "P022", "P023",
        "P024a", "P024b", "P025a", "P025b", "P026a", "P026b", "P027a", "P027b",
    ]:
        raise RuntimeError("authoritative source does not contain the ordered P019–P027b lane")
    return slides


def write_source_snapshot(authority: list[dict]) -> None:
    payload = {
        "deck": "LoRaEnergySim-LEO-ALT-PART-A-V2-P019-P027B-REVIEW",
        "owned_lane": "part-a-v2-p019-p027b",
        "source": str(SOURCE_JSON),
        "template": str(TEMPLATE),
        "template_shell": "source slide 2 / slideLayout2.xml only",
        "typography": {"title_pt": 28, "body_pt": 24, "minimum_body_pt": 16,
                        "cjk": "標楷體", "latin": "Times New Roman"},
        "slides": authority,
        "package_sources": [
            "/home/sat/lora-energy-lab-reference/README.zh-TW.md",
            "/home/u24/lora-energy-lab/student_policy.py",
            "/home/sat/lora-energy-lab-reference/run_lab.py",
            "/home/sat/lora-energy-lab-reference/lora_energy_lab/engine.py",
            "/home/sat/lora-energy-lab-reference/lora_energy_lab/replay.py",
            "/home/sat/lora-energy-lab-reference/scenarios/ntpu-energy-decision-01.json",
        ],
    }
    SOURCE_SNAPSHOT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def slide_019(prs: Presentation):
    slide = prepare_slide(prs, "P019｜WSL package path")
    add_text(slide, 0.95, 1.05, 11.25, 0.42,
             "release asset 的讀取位置與 Linux package root 必須分屬清楚的路徑。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P019 proposition")
    xs = [0.88, 3.70, 6.55, 9.42]
    labels = [
        ("Windows Downloads", "來源：Windows filesystem\n作用：讀取 release asset\n單位：folder path\n判讀：transfer source", BLUE_PALE, BLUE),
        ("`/mnt/c/.../Downloads`", "來源：WSL mount\n作用：讀取 Windows file\n單位：filesystem path\n判讀：transfer source only", TEAL_PALE, TEAL),
        ("`~/c120-course`", "來源：Ubuntu home\n作用：保存 release／root\n單位：directory\n判讀：Linux `.venv` boundary", GOLD_PALE, GOLD),
        ("`lora-energy-lab/`", "來源：extracted release\n作用：執行 setup\n單位：directory\n判讀：setup runs here", PURPLE_PALE, PURPLE),
    ]
    for index, (title, body, fill, line) in enumerate(labels):
        add_card(slide, xs[index], 1.72, 2.32, 1.58, title, body, fill, line,
                 title_size=17 if index in {1, 2, 3} else 19,
                 body_size=17, name=f"P019 path {index + 1}")
        if index < 3:
            add_chevron(slide, xs[index] + 2.45, 2.26, 0.34, 0.43,
                        (TEAL, GOLD, PURPLE)[index], f"P019 mechanism {index + 1}")
    add_code_box(
        slide, 0.90, 3.68, 7.10, 1.28,
        "mkdir -p ~/c120-course\n"
        "cp /mnt/c/Users/<YourName>/Downloads/lora-energy-lab-v1.zip ~/c120-course/\n"
        "cd ~/c120-course/lora-energy-lab && pwd",
        fill="F6F7FB", line=BLUE, size=18, name="P019 path commands")
    add_card(
        slide, 8.28, 3.68, 4.10, 1.28,
        "`cp`（複製命令）",
        "來源：POSIX shell｜作用：copy release files\n單位：command｜判讀：把檔案帶入 Linux staging path",
        fill=RED_PALE, line=RED, title_size=20, body_size=17,
        name="P019 boundary")
    add_field_band(slide, "pwd", "目前目錄命令", "POSIX shell",
                   "確認 package root", "path string",
                   "輸出位於 Linux home 下的 `lora-energy-lab/`")
    slide.notes_slide.notes_text_frame.text = (
        "WSL 會讀取 Windows Downloads 中的 release asset，但 Linux 的 setup 與虛擬環境必須在 Linux home 完成。"
        "命令中的 `/mnt/c/.../Downloads` 是 WSL mount 讀取位置，`cp` 將兩個 release files 複製到 `~/c120-course`，"
        "最後進入 `lora-energy-lab/` 並以 `pwd` 核對 package root。這個路徑分離保持 `.venv/bin/python`、lock 與 shell script 屬於同一份 Linux package。"
    )


def slide_020(prs: Presentation):
    slide = prepare_slide(prs, "P020｜POSIX／WSL setup 與 verify")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "exact interpreter 進入 setup；verify 只在環境與契約一致時產生 READY。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P020 proposition")
    add_card(slide, 0.92, 1.62, 5.75, 2.10,
             "入口 A｜`setup.sh`（POSIX 設定啟動器）",
             "`PYTHON_BIN=python3.11 bash setup.sh`\n"
             "`bash course.sh verify`\n\n"
             "來源：package root｜作用：建立 `.venv`／安裝 lock\n"
             "單位：executable file｜判讀：準備 POSIX runner。",
             BLUE_PALE, BLUE, title_size=18, body_size=17, name="P020 python path")
    add_card(slide, 6.75, 1.62, 5.65, 2.10,
             "入口 B｜`course.sh verify`（POSIX 驗證命令）",
             "`PYTHON_BIN=\"$(uv python find 3.11)\" bash setup.sh`\n"
             "`bash course.sh verify`\n"
             "來源：package launcher｜作用：核對 interpreter／contract\n"
             "單位：command｜判讀：寫入 READY receipt。",
             TEAL_PALE, TEAL, title_size=18, body_size=17, name="P020 uv path")
    add_chevron(slide, 5.98, 4.12, 0.50, 0.44, GOLD, "P020 setup connector")
    add_label_box(slide, 1.24, 4.02, 4.20, 0.72, "`setup.sh`\n建立 `.venv/bin/python`",
                  GOLD_PALE, GOLD, 21, NAVY, True, name="P020 setup stage")
    add_label_box(slide, 7.00, 4.02, 5.02, 0.72, "`course.sh verify`\n預期 `READY`（可進入 case 的狀態）+ Python 3.11.x",
                  PURPLE_PALE, PURPLE, 20, NAVY, True, name="P020 verify stage")
    add_field_band(slide, "PYTHON_BIN", "Python 執行器路徑變數", "shell environment",
                   "選擇 exact 3.11 interpreter", "path string",
                   "提供給 setup；成功後由 package-local `.venv/bin/python` 執行 verify",
                   color=TEAL)
    slide.notes_slide.notes_text_frame.text = (
        "POSIX 與 WSL 的 setup 需要一個明確的 Python 3.11.x interpreter。"
        "`PYTHON_BIN` 來自 shell environment，作用是把 `python3.11` 或 `uv python find 3.11` 的 path 交給 `setup.sh`；"
        "setup 建立 package-local `.venv`、安裝 lock dependencies，`course.sh verify` 再核對 exact interpreter 與 course contract。"
        "預期狀態是 `READY`、Python 3.11.x 與 `.venv/bin/python`；此 gate 尚未產生任何 lab case result。"
    )


def slide_021a(prs: Presentation):
    slide = prepare_slide(prs, "P021a｜查看 READY receipt")
    add_text(slide, 0.95, 1.05, 11.30, 0.42,
             "三種 shell 讀取同一份 JSON receipt；檔案來源與 case result 保持分離。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P021a proposition")
    cols = [
        (0.94, "POSIX／WSL", "sed -n '1,80p'\nartifacts/verify-receipt.json", BLUE_PALE, BLUE),
        (4.40, "PowerShell", "Get-Content .\\artifacts\\verify-receipt.json", TEAL_PALE, TEAL),
        (7.86, "Windows CMD", "type artifacts\\verify-receipt.json", GOLD_PALE, GOLD),
    ]
    for x, title, command, fill, line in cols:
        add_card(slide, x, 1.68, 3.10, 2.05, title, command,
                 fill, line, title_size=21, body_size=19, name=f"P021a {title}")
    add_label_box(slide, 1.20, 4.20, 10.85, 0.70,
                  "一個檔案：artifacts/verify-receipt.json｜作用：讀取 setup evidence｜預期：欄位可見｜失敗：保留檔案來源",
                  PURPLE_PALE, PURPLE, 22, NAVY, True, name="P021a receipt ribbon")
    add_field_band(slide, "verify-receipt.json", "驗證收據檔", "setup／verify command",
                   "記錄 environment 與 contract gate", "JSON artifact",
                   "解讀 setup evidence；不代替 case result", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "`verify-receipt.json` 是 setup 與 verify command 寫入的驗證收據檔。"
        "POSIX 使用 `sed`，PowerShell 使用 `Get-Content`，Windows CMD 使用 `type`；三個命令都只讀取同一份 artifact。"
        "收據的欄位用來描述環境與契約 gate，case result 仍由後續 run command 另行產生，兩種 provenance 不混用。"
    )


def slide_021b(prs: Presentation):
    slide = prepare_slide(prs, "P021b｜READY receipt：狀態與識別欄位")
    add_text(slide, 0.95, 1.05, 11.30, 0.42,
             "receipt 的 identity 欄位把環境、情境與 policy API 綁在同一個 gate。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P021b proposition")
    rows = [
        ("`status`（狀態）", "來源：receipt｜作用：gate result｜單位：enum｜判讀：READY 才能進入 case", BLUE_PALE, BLUE),
        ("`python_version`（Python 版本）", "來源：receipt｜作用：記錄 interpreter｜單位：version｜判讀：必須是 3.11.x", TEAL_PALE, TEAL),
        ("`scenario_id`（情境識別碼）", "來源：scenario contract／receipt｜作用：綁定 fixed case｜單位：string identifier｜判讀：它是情境，不是能量結果", GOLD_PALE, GOLD),
        ("`policy_api_version`（策略介面版本）", "來源：package contract｜作用：綁定 observation／action｜單位：string identifier｜判讀：避免 API 漂移", PURPLE_PALE, PURPLE),
    ]
    y = 1.62
    for index, (title, body, fill, line) in enumerate(rows):
        add_box(slide, 0.96, y, 4.10, 0.82, fill, line, name=f"P021b field {index + 1}")
        add_text(slide, 1.16, y + 0.12, 3.70, 0.30, title, 19, NAVY, True,
                 name=f"P021b field {index + 1} heading")
        add_text(slide, 5.28, y + 0.07, 7.00, 0.70, body, 18, INK, False,
                 valign=MSO_ANCHOR.MIDDLE, name=f"P021b field {index + 1} explanation")
        y += 0.93
    add_field_band(slide, "status", "狀態", "verify receipt",
                   "表示 gate 是否可進入 case", "enum",
                   "READY 只描述環境與契約一致；result identity 仍需另行產生", color=BLUE)
    slide.notes_slide.notes_text_frame.text = (
        "READY receipt 的第一組欄位描述狀態與 identity。"
        "`status` 必須是 READY，`python_version` 必須是 3.11.x；`scenario_id` 將固定情境和 case 綁定，"
        "`policy_api_version` 將允許的 observation 與 action 契約綁定。這些欄位表示 setup gate 通過，並不表示能量結果已經存在。"
    )


def slide_021c(prs: Presentation):
    slide = prepare_slide(prs, "P021c｜READY receipt：引擎與主張欄位")
    add_text(slide, 0.95, 1.05, 11.30, 0.42,
             "producer mode 與 claim boundary 決定結果的證據範圍；READY 仍只是 setup gate。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P021c proposition")
    add_card(slide, 0.94, 1.65, 3.72, 2.40,
             "`engine_mode`\n執行引擎模式",
             "來源：runner metadata\n作用：標示 producer\n單位：enum\n判讀：`coherent-course-simulated-adapter`",
             BLUE_PALE, BLUE, title_size=19, body_size=19, name="P021c engine")
    add_card(slide, 4.82, 1.65, 3.72, 2.40,
             "`upstream_execution`\n上游程式執行狀態",
             "來源：runner metadata\n作用：標示上游是否執行\n單位：Boolean\n判讀：course adapter 為 `false`",
             TEAL_PALE, TEAL, title_size=19, body_size=19, name="P021c upstream")
    add_card(slide, 8.70, 1.65, 3.72, 2.40,
             "`claim_boundary`\n主張邊界",
             "來源：course contract\n作用：標示 evidence class\n單位：fixed string\n判讀：simulated／非 live／非 measured／非 canonical parity",
             GOLD_PALE, GOLD, title_size=19, body_size=18, name="P021c claim")
    add_label_box(slide, 1.44, 4.48, 10.50, 0.62,
                  "READY：環境與契約一致｜case result：另有 run identity 與 result.json",
                  PURPLE_PALE, PURPLE, 21, NAVY, True, name="P021c gate distinction")
    add_field_band(slide, "engine_mode", "執行引擎模式", "runner metadata",
                   "分類 producer", "enum",
                   "`coherent-course-simulated-adapter`；不延伸為上游程式執行", color=TEAL)
    slide.notes_slide.notes_text_frame.text = (
        "`engine_mode`、`upstream_execution` 與 `claim_boundary` 共同描述 producer 和 evidence scope。"
        "engine mode 是 `coherent-course-simulated-adapter`，上游程式執行狀態為 false，claim boundary 固定標示 simulated teaching data、非 live、非 measured 與非 canonical parity。"
        "因此 READY 表示環境與契約一致；case result 仍需用自己的 run identity 保存，兩者不互相代替。"
    )


def slide_022(prs: Presentation):
    slide = prepare_slide(prs, "P022｜Setup gate recovery 與 same-scenario fallback")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "gate 的兩條合法路徑：修復指定邊界，或保留同情境回復資料的來源分類。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P022 proposition")
    add_label_box(slide, 0.98, 1.62, 3.18, 0.66, "環境 gate", BLUE_PALE, BLUE,
                  22, NAVY, True, name="P022 gate")
    add_chevron(slide, 4.36, 1.76, 0.46, 0.40, BLUE, "P022 gate connector")
    add_card(slide, 4.96, 1.43, 3.32, 2.45,
             "READY 通過",
             "source：local setup\nidentity：current package\n作用：進入 exact case\n判讀：另行產生 result／replay",
             TEAL_PALE, TEAL, title_size=22, body_size=20, name="P022 ready branch")
    add_card(slide, 8.72, 1.43, 3.32, 2.45,
             "READY 受阻",
             "boundary：Python／root／lock／policy\n修復：只處理指出的 gate\n替代：使用 exact case pair\n判讀：保留 fallback source",
             GOLD_PALE, GOLD, title_size=22, body_size=19, name="P022 blocked branch")
    add_text(slide, 4.30, 4.20, 7.40, 0.30,
             "同情境回復資料：result.json + endpoint-replay.json",
             21, PURPLE, True, PP_ALIGN.CENTER, name="P022 fallback pair title")
    add_label_box(slide, 1.10, 4.65, 5.24, 0.65,
                  "`fallback_artifacts/<case>/result.json`\n同一 scenario identity",
                  PURPLE_PALE, PURPLE, 19, NAVY, True, name="P022 fallback result")
    add_label_box(slide, 6.96, 4.65, 5.24, 0.65,
                  "`endpoint-replay.json`\n同一 run pairing",
                  PURPLE_PALE, PURPLE, 19, NAVY, True, name="P022 fallback replay")
    add_field_band(slide, "artifact_source", "資料來源分類", "result metadata",
                   "標示 local policy run 或 fallback", "enum",
                   "來源模式維持原值；同情境回復資料不改寫成本機執行", color=GOLD)
    slide.notes_slide.notes_text_frame.text = (
        "setup 或 verify 受阻時，artifact source 必須保持可見。"
        "可修復的問題只在 Python minor、package root、lock 或 policy boundary 內處理；若 local setup 仍無法建立 READY，"
        "同一個 lab／case 的 `result.json` 與 `endpoint-replay.json` 以 pair 方式提供 recovery。"
        "`artifact_source` 是 result metadata 的資料來源分類，fallback 保持 `same-scenario-fallback`，不改寫成本機 policy execution claim。"
    )


def slide_023(prs: Presentation):
    slide = prepare_slide(prs, "P023｜Bounded edit surface：student_policy.py")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "可編輯面集中在一個檔案的三個 marked blocks；其餘 package inputs／outputs 保持原樣。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P023 proposition")
    add_code_box(
        slide, 0.96, 1.58, 7.08, 3.26,
        "# === LORA EDITABLE: lab-a-pace-rest ===\n"
        "PACE_GAP_STEPS = 2\n"
        "REST_DURING_GAP = SLEEP\n"
        "# === LORA END EDITABLE: lab-a-pace-rest ===\n\n"
        "# === LORA EDITABLE: lab-b-enter-exit-hold ===\n"
        "ENTER_QUALITY = 2\nEXIT_QUALITY = 1\nSTABLE_STEPS = 2",
        fill="F6F7FB", line=BLUE, size=18, name="P023 policy markers")
    add_card(slide, 8.36, 1.58, 4.02, 1.06,
             "Lab C block",
             "`BATCH_SIZE = 3`\n`URGENT_MARGIN_S = 20`",
             GOLD_PALE, GOLD, title_size=21, body_size=19, name="P023 lab c")
    add_card(slide, 8.36, 2.83, 4.02, 2.01,
             "讀取範圍",
             "`scenario`：固定 input\n`schemas`：契約檢查\n`runner`：產生 events\n`generated JSON`：run output",
             RED_PALE, RED, title_size=21, body_size=19, name="P023 read-only")
    add_label_box(slide, 1.24, 5.04, 10.86, 0.50,
                  "一次只開啟目前 lab 的 marked block；policy identity 與 result lineage 因此保持可解釋",
                  TEAL_PALE, TEAL, 20, NAVY, True, name="P023 bounded rule")
    add_field_band(slide, "marked block", "標記區段", "file markers",
                   "界定一個 lab 的 edit surface", "source range",
                   "每次只保留一個 active region；其他 package boundary 維持 read-only", color=BLUE)
    slide.notes_slide.notes_text_frame.text = (
        "`student_policy.py` 是 package root 內唯一的策略檔。"
        "目前檔案的三個 markers 是 `lab-a-pace-rest`、`lab-b-enter-exit-hold` 與 `lab-c-batch-urgent`；每個 lab 只在自己的區段變更常數。"
        "scenario、schemas、runner 與 generated JSON 分別是固定輸入、契約、執行器與輸出；保留這些 boundary，policy identity 和結果 lineage 才能被比較。"
    )


def slide_024a(prs: Presentation):
    slide = prepare_slide(prs, "P024a｜Observation 是 decision-time input")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "每一次決策只讀目前與已允許的 observation，回傳一個 legal action。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P024a proposition")
    add_card(slide, 0.94, 1.60, 3.25, 2.20,
             "目前 observation",
             "`contact_open`（窗口開啟狀態）｜來源：scenario｜作用：send guard｜單位：Boolean｜判讀：false → SLEEP\n"
             "`quality_band`（品質分帶）｜來源：current trace｜作用：present quality｜單位：ordinal band｜判讀：threshold input\n"
             "`stable_steps`（穩定步數）｜來源：trace accumulator｜作用：hold gate｜單位：steps｜判讀：consecutive stable count",
             BLUE_PALE, BLUE, title_size=21, body_size=17, name="P024a inputs")
    add_chevron(slide, 4.42, 2.35, 0.52, 0.46, TEAL, "P024a input connector")
    add_code_box(slide, 5.12, 1.60, 3.15, 2.20,
                 "def choose_action(observation):\n"
                 "    if not observation.contact_open:\n"
                 "        return SLEEP",
                 fill=TEAL_PALE, line=TEAL, size=19, name="P024a function")
    add_chevron(slide, 8.56, 2.35, 0.52, 0.46, GOLD, "P024a output connector")
    add_card(slide, 9.24, 1.60, 3.18, 2.20,
             "legal action",
             "`SLEEP`：窗口關閉的安全分支\n其餘 action 由下一個條件決定",
             GOLD_PALE, GOLD, title_size=22, body_size=20, name="P024a output")
    add_label_box(slide, 1.18, 4.18, 10.92, 0.70,
                  "policy scope 外：future quality／future energy／result summary",
                  RED_PALE, RED, 21, RED, True, name="P024a scope boundary")
    add_field_band(slide, "choose_action(observation)", "策略函式", "policy API",
                   "讀 current observation 並回傳一個合法 action", "function contract",
                   "不讀 future 或 result summary；輸出交給 runner 產生事件", color=TEAL)
    slide.notes_slide.notes_text_frame.text = (
        "`choose_action(observation)` 是 policy API 的函式契約。"
        "`contact_open` 來自當下 scenario observation，`quality_band` 來自目前 trace，`stable_steps` 來自連續觀察；"
        "窗口關閉時，函式回傳 `SLEEP`，讓 runner 進入合法的低功耗分支。future quality、future energy 與 result summary 不屬於 decision-time input，"
        "因此 action 是對當下資訊的控制輸出，而不是對結果的回看。"
    )


def slide_024b(prs: Presentation):
    slide = prepare_slide(prs, "P024b｜Observation 欄位與合法 action")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "queue、deadline、前一動作與窗口剩餘時間補齊 decision-time 的可觀察邊界。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P024b proposition")
    add_card(slide, 0.94, 1.55, 5.28, 2.45,
             "輸入欄位｜來源／單位／作用",
             "`steps_since_send`（距上次送出步數）｜runner state｜steps｜pace input\n"
             "`queue_size`（佇列數量）｜endpoint queue｜packet count｜batch pressure\n"
             "`urgent_due_in_s`（緊急期限剩餘秒數）｜scenario deadline｜s｜urgency input\n"
             "`previous_action`（前一動作）｜runner state｜action enum｜continuity\n"
             "`contact_remaining_s`（窗口剩餘秒數）｜contact trace｜s｜send timing",
             BLUE_PALE, BLUE, title_size=20, body_size=17, name="P024b fields")
    add_card(slide, 6.54, 1.55, 5.88, 2.45,
             "合法輸出",
             "`WAIT`｜清醒等待\n`SLEEP`｜低功耗休息\n`SEND_ONE`｜單筆傳輸\n`SEND_URGENT`｜緊急傳輸\n`FLUSH_BATCH`｜批次送出",
             TEAL_PALE, TEAL, title_size=22, body_size=21, name="P024b actions")
    add_code_box(slide, 1.20, 4.28, 10.90, 0.74,
                 "if observation.steps_since_send < PACE_GAP_STEPS:\n    return REST_DURING_GAP",
                 fill=GOLD_PALE, line=GOLD, size=20, name="P024b pacing branch")
    add_field_band(slide, "queue_size", "佇列數量", "endpoint queue",
                   "提供 batching／urgency input", "packet count",
                   "決定送出壓力；合法輸出仍限於 policy API enum", color=BLUE)
    slide.notes_slide.notes_text_frame.text = (
        "`steps_since_send`、`queue_size`、`urgent_due_in_s`、`previous_action` 與 `contact_remaining_s` 都是 runner 在決策時提供的欄位。"
        "它們的單位分別是 steps、packet count、秒、action enum 與秒；作用是控制節奏、批次、期限、轉移與窗口時機。"
        "函式只能回傳 `WAIT`、`SLEEP`、`SEND_ONE`、`SEND_URGENT` 或 `FLUSH_BATCH`，因此條件分支的結果仍被限制在合法 action 邊界。"
    )


def slide_025a(prs: Presentation):
    slide = prepare_slide(prs, "P025a｜Lab A／B 的受控常數")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "常數是可追蹤的控制點；差異比較固定為一次一個值，沿 state 與 service 回溯。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P025a proposition")
    add_card(slide, 0.96, 1.54, 5.62, 1.72,
             "Lab A｜pace／rest",
             "`PACE_GAP_STEPS = 2`（節奏間隔步數）｜來源：A block｜單位：steps｜作用：minimum gap\n"
             "`REST_DURING_GAP = SLEEP`（空檔休息動作）｜來源：A block｜單位：action enum｜判讀：awake idle／wake\n"
             "比較 `SLEEP` ↔ `WAIT`，觀察 service／J",
             BLUE_PALE, BLUE, title_size=21, body_size=17, name="P025a lab a")
    add_card(slide, 6.76, 1.54, 5.62, 1.72,
             "Lab B｜enter／exit／hold",
             "`ENTER_QUALITY = 2`（進入品質門檻）｜來源：B block｜單位：quality band｜作用：enter\n"
             "`EXIT_QUALITY = 1`（退出品質門檻）｜來源：B block｜單位：quality band｜作用：exit\n"
             "`STABLE_STEPS = 2`（穩定步數門檻）｜來源：B block｜單位：steps｜判讀：hold duration",
             TEAL_PALE, TEAL, title_size=21, body_size=17, name="P025a lab b")
    add_label_box(slide, 1.18, 3.70, 5.05, 0.78,
                  "A candidate：`REST_DURING_GAP = WAIT`\n改變空檔的 state action",
                  GOLD_PALE, GOLD, 19, NAVY, True, name="P025a candidate a")
    add_label_box(slide, 7.10, 3.70, 5.05, 0.78,
                  "B candidate：`STABLE_STEPS = 1`\n改變 send-ready 進入時機",
                  PURPLE_PALE, PURPLE, 19, NAVY, True, name="P025a candidate b")
    add_text(slide, 1.18, 4.85, 10.92, 0.34,
             "hysteresis：enter／exit／hold 的門檻與持續條件分開，降低來回轉移的機會。",
             20, NAVY, True, PP_ALIGN.CENTER, name="P025a hysteresis")
    add_field_band(slide, "hysteresis", "遲滯控制", "enter／exit／hold constants",
                   "分開進入、退出與 hold 條件", "rule",
                   "降低 ping-pong transition；仍需用 packet／service／energy evidence 判讀", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "Lab A 使用 `PACE_GAP_STEPS` 與 `REST_DURING_GAP` 描述送出間隔和空檔 action；封裝 baseline 是 2 與 `SLEEP`。"
        "Lab B 使用 `ENTER_QUALITY = 2`、`EXIT_QUALITY = 1` 與 `STABLE_STEPS = 2` 描述 send-ready 的進入、退出與連續穩定條件。"
        "候選變更分別是把 A 的休息 action 改為 `WAIT`，或把 B 的 hold 改為 1；每次只動一個 marked constant，並以 state、packet、service 與 endpoint energy 比較。"
    )


def slide_025b(prs: Presentation):
    slide = prepare_slide(prs, "P025b｜Lab C 的佇列與緊急常數")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "Lab C 把 queue aggregation 與 deadline protection 連接到同一個 action surface。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P025b proposition")
    add_card(slide, 0.96, 1.56, 4.18, 2.42,
             "`BATCH_SIZE = 3`\n批次大小",
             "來源：Lab C marked block\n作用：決定一次 flush 的 packet count\n單位：packet count\n判讀：影響 queue wait、process／TX／RX 與 delivery",
             BLUE_PALE, BLUE, title_size=21, body_size=19, name="P025b batch")
    add_card(slide, 5.46, 1.56, 4.18, 2.42,
             "`URGENT_MARGIN_S = 20`\n緊急餘裕秒數",
             "來源：Lab C marked block\n作用：期限前觸發 urgent action\n單位：s\n判讀：影響 deadline protection 與送出時機",
             GOLD_PALE, GOLD, title_size=20, body_size=19, name="P025b urgent")
    add_card(slide, 9.96, 1.56, 2.42, 2.42,
             "`FLUSH_BATCH`\n批次送出",
             "來源：policy API\n作用：將 queue 成組傳輸\n單位：action enum\n判讀：連接 queue 到 packet path",
             TEAL_PALE, TEAL, title_size=18, body_size=17, name="P025b flush")
    add_label_box(slide, 1.25, 4.38, 3.28, 0.66, "queue：● ● ● → FLUSH_BATCH",
                  PURPLE_PALE, PURPLE, 19, NAVY, True, name="P025b queue visual")
    add_chevron(slide, 4.78, 4.50, 0.45, 0.38, PURPLE, "P025b queue connector")
    add_label_box(slide, 5.42, 4.38, 3.42, 0.66, "urgent_due_in_s ≤ margin → SEND_URGENT",
                  RED_PALE, RED, 18, RED, True, name="P025b urgent visual")
    add_label_box(slide, 9.22, 4.38, 2.98, 0.66, "candidate／revision：各改一個常數",
                  GOLD_PALE, GOLD, 18, NAVY, True, name="P025b one change")
    add_field_band(slide, "URGENT_MARGIN_S", "緊急餘裕秒數", "Lab C marked block",
                   "提供期限分支的比較值", "s",
                   "值變更後觀察 urgent action、delivery、deadline 與 endpoint J", color=GOLD)
    slide.notes_slide.notes_text_frame.text = (
        "Lab C 的 `BATCH_SIZE` 來自 marked block，單位是 packet count，作用是決定 queue 何時成組送出；"
        "`URGENT_MARGIN_S` 的單位是秒，作用是讓期限逼近時啟動 `SEND_URGENT`。"
        "`FLUSH_BATCH` 是 policy API 的合法 action，將 queue 連接到 process、TX、RX 與 delivery events。"
        "candidate 與 revision 各自只改一個 marked constant，withheld case 保持 frozen policy，用來讀取條件式的 queue、service 與 energy 後果。"
    )


def slide_026a(prs: Presentation):
    slide = prepare_slide(prs, "P026a｜讀取條件與 action 分支")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "branch order 把目前 observation 轉成合法 action；每一段都有清楚的機制語意。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P026a proposition")
    add_code_box(slide, 0.96, 1.52, 6.30, 3.60,
                 "if not observation.contact_open:\n"
                 "    return SLEEP\n\n"
                 "if observation.urgent_pending and observation.urgent_due_in_s <= URGENT_MARGIN_S:\n"
                 "    return SEND_URGENT\n\n"
                 "if observation.steps_since_send < PACE_GAP_STEPS:\n"
                 "    return REST_DURING_GAP",
                 fill="F6F7FB", line=BLUE, size=17, name="P026a policy branches")
    add_card(slide, 7.62, 1.52, 4.80, 0.98,
             "窗口狀態",
             "`contact_open = false` → `SLEEP`\n保護 send legality",
             BLUE_PALE, BLUE, title_size=20, body_size=18, name="P026a contact")
    add_card(slide, 7.62, 2.70, 4.80, 0.98,
             "緊急期限",
             "`urgent_due_in_s ≤ URGENT_MARGIN_S` → `SEND_URGENT`\n保護 deadline",
             RED_PALE, RED, title_size=20, body_size=17, name="P026a urgent")
    add_card(slide, 7.62, 3.88, 4.80, 1.24,
             "節奏間隔",
             "`steps_since_send < PACE_GAP_STEPS` → `REST_DURING_GAP`\nA block 決定 `WAIT` 或 `SLEEP`",
             GOLD_PALE, GOLD, title_size=20, body_size=18, name="P026a pace")
    add_field_band(slide, "urgent_pending", "緊急封包待處理", "queue／deadline observation",
                   "顯示 urgent packet 是否存在", "Boolean",
                   "啟動 deadline branch；action 後仍需讀 packet／service evidence", color=RED)
    slide.notes_slide.notes_text_frame.text = (
        "目前 policy 的 branch order 先檢查窗口，再檢查緊急期限，最後檢查送出節奏。"
        "窗口關閉時，`contact_open` 為 false，回傳 `SLEEP` 以維持 send legality；urgent packet 逼近 `URGENT_MARGIN_S` 時，回傳 `SEND_URGENT`；"
        "距離上次送出仍小於 `PACE_GAP_STEPS` 時，回傳 A block 指定的 `REST_DURING_GAP`。"
        "`urgent_pending` 是 queue／deadline observation 的 Boolean，作用是啟動期限分支；條件本身是機制入口，結果仍要從事件 ledger 判讀。"
    )


def slide_026b(prs: Presentation):
    slide = prepare_slide(prs, "P026b｜Action 後果的事件讀法")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "action 的後果要沿事件 ledger 讀取：state interval、packet outcome、service 與 endpoint J。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P026b proposition")
    add_card(slide, 0.94, 1.55, 3.40, 1.54,
             "`WAIT`｜清醒等待",
             "`AWAKE_IDLE` interval\n來源：runner ledger\n判讀：保持反應，累積時間積分能量",
             TEAL_PALE, TEAL, title_size=21, body_size=19, name="P026b wait")
    add_card(slide, 0.94, 3.35, 3.40, 1.54,
             "`SLEEP`｜低功耗休息",
             "`SLEEP` interval + `WAKE` transition\n來源：runner ledger\n判讀：休息功率、喚醒延遲與喚醒能量",
             BLUE_PALE, BLUE, title_size=21, body_size=18, name="P026b sleep")
    add_chevron(slide, 4.62, 2.96, 0.50, 0.46, PURPLE, "P026b event connector")
    add_label_box(slide, 5.40, 1.86, 3.00, 0.62,
                  "`SEND_*` → PROCESS → TX → RX",
                  GOLD_PALE, GOLD, 20, NAVY, True, name="P026b send states")
    add_label_box(slide, 5.40, 2.76, 3.00, 0.62,
                  "attempt → retry → delivery／expiry",
                  PURPLE_PALE, PURPLE, 20, NAVY, True, name="P026b packet states")
    add_chevron(slide, 8.74, 2.96, 0.50, 0.46, PURPLE, "P026b result connector")
    add_card(slide, 9.46, 1.55, 2.94, 3.34,
             "結果層",
             "`service result`\npass／fail + reason\n\n`endpoint_energy_j`\n來源：endpoint model\n單位：J\n\n較低 J 且 service fail：energy／service trade-off",
             PURPLE_PALE, PURPLE, title_size=21, body_size=18, name="P026b outcome")
    add_field_band(slide, "endpoint_energy_j", "端點能量", "endpoint radio／processing model",
                   "累積狀態與轉換能量", "J",
                   "只屬 endpoint evidence layer；與 LEO／system／canonical energy 分開", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "`WAIT` 會產生 `AWAKE_IDLE` interval，`SLEEP` 會產生 `SLEEP` interval，後續再次啟動 action 時可能出現 `WAKE` transition。"
        "`SEND_ONE`、`SEND_URGENT` 與 `FLUSH_BATCH` 會沿 process、TX、RX 形成 packet attempt、retry、delivery 或 expiry。"
        "`endpoint_energy_j` 來自 endpoint radio／processing model，單位是 J；它必須和 service result 一起判讀。"
        "較低 J 與 service fail 同時出現時，證據分類是 energy／service trade-off；它不會被改稱為 LEO、system 或 canonical energy。"
    )


def slide_027a(prs: Presentation):
    slide = prepare_slide(prs, "P027a｜Python indentation 與 return")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "Python 的縮排決定分支歸屬；`return` 把一個合法 action 送回 runner。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P027a proposition")
    add_code_box(slide, 0.96, 1.55, 5.72, 2.64,
                 "def choose_action(observation):\n"
                 "    if not observation.contact_open:\n"
                 "        return SLEEP\n"
                 "    return WAIT",
                 fill=BLUE_PALE, line=BLUE, size=21, name="P027a skeleton")
    add_card(slide, 7.02, 1.55, 5.34, 1.08,
             "`indentation`（縮排）",
             "來源：Python syntax｜作用：將 statement 放入 branch｜單位：code structure｜判讀：決定 branch membership",
             TEAL_PALE, TEAL, title_size=21, body_size=18, name="P027a indentation")
    add_card(slide, 7.02, 2.86, 5.34, 1.08,
             "`return`（回傳）",
             "來源：policy function｜作用：送出一個 legal action｜單位：action enum｜判讀：結束 decision step",
             GOLD_PALE, GOLD, title_size=21, body_size=18, name="P027a return")
    add_label_box(slide, 1.20, 4.58, 10.90, 0.62,
                  "`.venv/bin/python -m py_compile student_policy.py`\nsyntax gate：修復 active marked block",
                  PURPLE_PALE, PURPLE, 19, NAVY, True, name="P027a compile")
    add_field_band(slide, "return", "回傳", "policy function",
                   "將一個合法 action 交給 runner", "action enum",
                   "若縮排或名稱錯誤，保留錯誤並只修 active marked block", color=GOLD)
    slide.notes_slide.notes_text_frame.text = (
        "Python 的 `indentation` 是 code structure，作用是把 statements 放入正確的條件分支；"
        "`return` 來自 policy function，作用是將一個 action enum 送回 runner，結束這一個 decision step。"
        "目前 package 的最小骨架在窗口關閉時回傳 `SLEEP`，窗口開啟後回傳 `WAIT`。"
        "compile command 只檢查語法；若出現錯誤，修復目前 active marked block，scenario、runner 與 schema 保持原狀。"
    )


def slide_027b(prs: Presentation):
    slide = prepare_slide(prs, "P027b｜Lab A baseline run 與結果入口")
    add_text(slide, 0.95, 1.05, 11.35, 0.42,
             "packaged policy + fixed scenario 建立 baseline control artifact，結果路徑直接交給後續觀察。",
             24, NAVY, True, PP_ALIGN.CENTER, name="P027b proposition")
    add_card(slide, 0.94, 1.52, 3.20, 2.20,
             "baseline",
             "來源：packaged policy + fixed scenario／case\n作用：建立 control artifact\n單位：case label\n判讀：成為後續比較 anchor",
             BLUE_PALE, BLUE, title_size=23, body_size=19, name="P027b baseline")
    add_chevron(slide, 4.38, 2.34, 0.46, 0.42, TEAL, "P027b command connector")
    add_code_box(slide, 5.05, 1.52, 3.22, 2.20,
                 "bash course.sh run --lab A --case baseline\n\n"
                 "course.cmd run --lab A --case baseline",
                 fill=TEAL_PALE, line=TEAL, size=19, name="P027b commands")
    add_chevron(slide, 8.54, 2.34, 0.46, 0.42, GOLD, "P027b output connector")
    add_card(slide, 9.20, 1.52, 3.20, 2.20,
             "stdout result",
             "`status`：OK（artifact status）\n`result_path`：指向 `result.json`\n`run_id`：本次 identity\nclaim boundary：course-simulated",
             GOLD_PALE, GOLD, title_size=21, body_size=19, name="P027b stdout")
    add_label_box(slide, 1.20, 4.08, 4.60, 0.90,
                  "`result_path`（結果檔路徑）\n`artifacts/<run_id>/result.json`｜command JSON output",
                  PURPLE_PALE, PURPLE, 18, NAVY, True, name="P027b result path")
    add_chevron(slide, 6.00, 4.30, 0.46, 0.42, PURPLE, "P027b replay connector")
    add_label_box(slide, 6.72, 4.08, 5.36, 0.90,
                  "`endpoint-replay.json`（端點重播檔）\n來源：same run directory｜作用：replay action／state／queue／packet timing｜單位：JSON artifact",
                  PURPLE_PALE, PURPLE, 17, NAVY, True, name="P027b replay path")
    add_field_band(slide, "result_path", "結果檔路徑", "command JSON output",
                   "定位 `result.json`", "path string",
                   "與同一 run directory 的 endpoint replay 配對；作為觀察入口", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "Lab A baseline 使用 packaged policy、fixed scenario 與 baseline case 建立 control artifact。"
        "POSIX 使用 `bash course.sh run --lab A --case baseline`，Windows 使用 `course.cmd run --lab A --case baseline`；"
        "runner stdout 會輸出 `status: OK`、`run_id` 與 `result_path`，其中 OK 只描述 artifact status。"
        "`result_path` 指向 `result.json`，同一個 run directory 內的 `endpoint-replay.json` 提供 action、state、queue 與 packet timing。"
        "保存 prediction 與這兩個配對入口後，下一段即可依 state、packet、service 與 endpoint energy 讀取 baseline，並開始 Part B 的受控比較。"
    )


# Beginner-comprehension pass: keep the approved donor geometry, but introduce
# each visible token only after a plain-language explanation.  The detailed
# Do/Why/Mechanism/Expect/Interpret narration remains in speaker notes.
def _beginner_page(prs: Presentation, title: str, takeaway: str,
                   cards: list[tuple], causal: str, field: tuple,
                   notes: str, causal_style: tuple = (GOLD_PALE, GOLD, NAVY)):
    slide = prepare_slide(prs, title)
    add_takeaway(slide, takeaway, f"{title.split('｜', 1)[0]} beginner takeaway")
    count = len(cards)
    if count == 2:
        xs, width = [0.92, 6.86], 5.54
    elif count == 3:
        xs, width = [0.92, 4.77, 8.62], 3.54
    elif count == 4:
        xs, width = [0.92, 3.97, 7.02, 10.07], 2.74
    else:
        raise ValueError(f"unsupported beginner card count: {count}")
    for i, item in enumerate(cards):
        heading, body, fill, line = item[:4]
        heading_size = item[4] if len(item) > 4 else (20 if count == 4 else 22)
        body_size = item[5] if len(item) > 5 else (18 if count == 4 else 20)
        add_visual_card(slide, xs[i], 1.82, width, 2.62 if count < 4 else 2.58,
                        heading, body, fill, line,
                        f"{title.split('｜', 1)[0]} beginner card {i + 1}",
                        heading_size=heading_size, body_size=body_size)
        if i < count - 1 and count == 3:
            add_chevron(slide, xs[i] + width + 0.14, 2.86, 0.36, 0.48,
                        (TEAL, GOLD)[i] if i < 2 else BLUE,
                        f"{title.split('｜', 1)[0]} beginner connector {i + 1}")
    fill, line, color = causal_style
    add_causal_strip(slide, causal,
                     f"{title.split('｜', 1)[0]} beginner causal strip",
                     fill=fill, line=line, color=color)
    add_field_band(slide, *field)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def slide_019(prs: Presentation):
    return _beginner_page(
        prs, "把下載資料搬進 Linux 課程資料夾",
        "WSL 是可執行 Linux 命令的工作區；先找到檔案，再進入課程資料夾",
        [
            ("Windows 檔案區", "保存下載檔：`Downloads`", BLUE_PALE, BLUE, 21, 20),
            ("WSL 讀取位置", "只從 Windows 讀檔\n`/mnt/c/.../Downloads`", TEAL_PALE, TEAL, 20, 18),
            ("Linux 課程資料夾", "先集中檔案\n`~/c120-course`", GOLD_PALE, GOLD, 20, 19),
            ("程式套件根目錄", "後續安裝都在這裡\n`lora-energy-lab/`", PURPLE_PALE, PURPLE, 19, 18),
        ],
        "先複製，再確認位置；不同資料夾不混用",
        ("pwd", "目前目錄命令", "Linux／macOS 終端機（WSL 也使用 Linux 指令）", "確認目前位置", "路徑文字", "輸出應位於 Linux home 的課程資料夾", BLUE),
        "WSL 會讀取 Windows Downloads 中的 release asset，但 Linux 的 setup 與虛擬環境必須在 Linux home 完成。命令中的 `/mnt/c/.../Downloads` 是 WSL mount 讀取位置，`cp` 將兩個 release files 複製到 `~/c120-course`，最後進入 `lora-energy-lab/` 並以 `pwd` 核對 package root。這個路徑分離保持 `.venv/bin/python`、lock 與 shell script 屬於同一份 Linux package。",
    )


def slide_020(prs: Presentation):
    return _beginner_page(
        prs, "P020｜POSIX／WSL setup 與 verify",
        "先找到 Python 3.11，再確認環境真的可以使用",
        [
            ("建立虛擬環境", "建立隔離的 Python 空間\n`setup.sh`", BLUE_PALE, BLUE, 21, 19),
            ("確認環境", "核對版本與必要檔案\n顯示可用狀態", TEAL_PALE, TEAL, 21, 19),
        ],
        "指定 Python 3.11 → 建立隔離空間 → 環境通過",
        ("PYTHON_BIN", "Python 執行器路徑變數", "命令工具設定", "指定要使用的版本", "路徑文字", "供建立環境的命令讀取", TEAL),
        "POSIX 與 WSL 的 setup 需要一個明確的 Python 3.11.x interpreter。`PYTHON_BIN` 來自 shell environment，作用是把 `python3.11` 或 `uv python find 3.11` 的 path 交給 `setup.sh`；setup 建立 package-local `.venv`、安裝 lock dependencies，`course.sh verify` 再核對 exact interpreter 與 course contract。預期狀態是 `READY`、Python 3.11.x 與 `.venv/bin/python`；此 gate 尚未產生任何 lab case result。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_021a(prs: Presentation):
    return _beginner_page(
        prs, "P021a｜查看 READY receipt",
        "同一張驗證收據，在三種命令工具中都只讀不改",
        [
            ("Linux 命令工具", "讀取收據內容\n不修改檔案", BLUE_PALE, BLUE, 20, 19),
            ("Windows PowerShell", "讀取同一份收據\n不修改檔案", TEAL_PALE, TEAL, 19, 19),
            ("Windows 命令提示字元", "讀取同一份收據\n不修改檔案", GOLD_PALE, GOLD, 18, 18),
        ],
        "三種工具都讀同一份檔案；來源不變",
        ("verify-receipt.json", "驗證收據檔", "環境檢查命令", "記錄環境與契約關卡", "JSON 檔案", "只解讀環境證據，不代替案例結果", PURPLE),
        "`verify-receipt.json` 是 setup 與 verify command 寫入的驗證收據檔。POSIX 使用 `sed`，PowerShell 使用 `Get-Content`，Windows CMD 使用 `type`；三個命令都只讀取同一份 artifact。收據的欄位用來描述環境與契約 gate，case result 仍由後續 run command 另行產生，兩種 provenance 不混用。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_021b(prs: Presentation):
    return _beginner_page(
        prs, "P021b｜READY receipt：狀態與識別欄位",
        "收據先說明環境是否可進入案例，再把案例綁到固定情境",
        [
            ("關卡狀態", "通過後才能進入案例", BLUE_PALE, BLUE, 21, 20),
            ("Python 版本", "必須是 3.11.x", TEAL_PALE, TEAL, 20, 20),
            ("固定情境", "把案例綁在同一情境\n`scenario_id`", GOLD_PALE, GOLD, 20, 18),
            ("允許範圍", "限定可讀資料與可用動作", PURPLE_PALE, PURPLE, 19, 18),
        ],
        "先確認環境，再確認案例身份；兩者都不是能量結果",
        ("status", "狀態", "驗證收據", "表示關卡是否可進入案例", "狀態值", "通過只表示環境與契約一致", BLUE),
        "READY receipt 的第一組欄位描述狀態與 identity。`status` 必須是 READY，`python_version` 必須是 3.11.x；`scenario_id` 將固定情境和 case 綁定，`policy_api_version` 將允許的 observation 與 action 契約綁定。這些欄位表示 setup gate 通過，並不表示能量結果已經存在。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_021c(prs: Presentation):
    return _beginner_page(
        prs, "P021c｜READY receipt：引擎與主張欄位",
        "收據還要說明資料從哪裡來，以及可以主張到哪裡",
        [
            ("執行模式", "課程使用模擬資料\n`engine_mode`", BLUE_PALE, BLUE, 21, 18),
            ("上游是否執行", "這次沒有執行上游程式\n`false`", TEAL_PALE, TEAL, 19, 18),
            ("主張邊界", "模擬、非即時、非實測", GOLD_PALE, GOLD, 21, 19),
        ],
        "環境通過不等於上游實測；案例結果仍另有身份",
        ("engine_mode", "執行模式", "執行器紀錄", "分類資料產生方式", "模式值", "模擬資料不延伸成上游程式執行", TEAL),
        "`engine_mode`、`upstream_execution` 與 `claim_boundary` 共同描述 producer 和 evidence scope。engine mode 是 `coherent-course-simulated-adapter`，上游程式執行狀態為 false，claim boundary 固定標示 simulated teaching data、非 live、非 measured 與非 canonical parity。因此 READY 表示環境與契約一致；case result 仍需用自己的 run identity 保存，兩者不互相代替。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_022(prs: Presentation):
    return _beginner_page(
        prs, "P022｜Setup gate recovery 與 same-scenario fallback",
        "環境卡住時，不假裝已經執行；保留同一案例的資料來源",
        [
            ("正常路徑", "環境通過後執行案例", BLUE_PALE, BLUE, 22, 20),
            ("修復路徑", "只修 Python、資料夾與套件鎖定", GOLD_PALE, GOLD, 20, 18),
            ("回復路徑", "同一情境的結果與重播成對保存\n`fallback`", PURPLE_PALE, PURPLE, 20, 18),
        ],
        "修復或回復都保留原始來源，不改寫執行事實",
        ("artifact_source", "資料來源分類", "結果紀錄", "標示本機執行或回復資料", "來源值", "來源模式保持原值，不改寫成本機執行", GOLD),
        "setup 或 verify 受阻時，artifact source 必須保持可見。可修復的問題只在 Python minor、package root、lock 或 policy boundary 內處理；若 local setup 仍無法建立 READY，同一個 lab／case 的 `result.json` 與 `endpoint-replay.json` 以 pair 方式提供 recovery。`artifact_source` 是 result metadata 的資料來源分類，fallback 保持 `same-scenario-fallback`，不改寫成本機 policy execution claim。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_023(prs: Presentation):
    return _beginner_page(
        prs, "P023｜Bounded edit surface：student_policy.py",
        "只改策略檔裡標好的區段；固定輸入與輸出不動",
        [
            ("檔案", "唯一可編輯的策略檔\n`student_policy.py`", BLUE_PALE, BLUE, 21, 18),
            ("區段", "A、B、C 三個標記區", GOLD_PALE, GOLD, 21, 19),
            ("固定輸入", "情境與規格保持不變", TEAL_PALE, TEAL, 20, 19),
            ("結果", "執行器產生結果與重播", PURPLE_PALE, PURPLE, 20, 19),
        ],
        "一次只開啟目前 lab 的區段；其他內容維持原樣",
        ("marked block", "標記區段", "檔案標記", "界定一個實驗的可編輯範圍", "程式碼區段", "其他套件邊界維持唯讀", BLUE),
        "`student_policy.py` 是 package root 內唯一的策略檔。目前檔案的三個 markers 是 `lab-a-pace-rest`、`lab-b-enter-exit-hold` 與 `lab-c-batch-urgent`；每個 lab 只在自己的區段變更常數。scenario、schemas、runner 與 generated JSON 分別是固定輸入、契約、執行器與輸出；保留這些 boundary，policy identity 和結果 lineage 才能被比較。",
        (TEAL_PALE, TEAL, NAVY),
    )


def slide_024a(prs: Presentation):
    return _beginner_page(
        prs, "P024a｜Observation 是 decision-time input",
        "每次判斷只看眼前資料，再回傳一個合法動作",
        [
            ("目前資料", "窗口、品質、穩定步數", BLUE_PALE, BLUE, 22, 20),
            ("判斷函式", "讀目前資料\n回傳一個動作", TEAL_PALE, TEAL, 21, 20),
            ("合法動作", "窗口關閉時休息\n`SLEEP`", GOLD_PALE, GOLD, 21, 19),
        ],
        "不讀未來品質、未來能量或結果摘要",
        ("choose_action(observation)", "策略函式", "策略介面", "讀目前觀察並回傳一個動作", "函式契約", "輸出交給執行器產生事件", TEAL),
        "`choose_action(observation)` 是 policy API 的函式契約。`contact_open` 來自當下 scenario observation，`quality_band` 來自目前 trace，`stable_steps` 來自連續觀察；窗口關閉時，函式回傳 `SLEEP`，讓 runner 進入合法的低功耗分支。future quality、future energy 與 result summary 不屬於 decision-time input，因此 action 是對當下資訊的控制輸出，而不是對結果的回看。",
        (RED_PALE, RED, RED),
    )


def slide_024b(prs: Presentation):
    return _beginner_page(
        prs, "P024b｜Observation 欄位與合法 action",
        "判斷時還要知道等待多久、手上有多少資料、期限剩多少",
        [
            ("節奏與前一步", "等待多久、上一個動作", BLUE_PALE, BLUE, 20, 19),
            ("佇列與期限", "有多少資料、還剩多少秒", GOLD_PALE, GOLD, 20, 19),
            ("輸出種類", "等待、休息、單筆、緊急、批次", TEAL_PALE, TEAL, 20, 18),
        ],
        "窗口剩餘時間影響送出時機；輸出種類仍固定",
        ("queue_size", "佇列數量", "端點資料佇列", "提供批次與期限判斷", "封包數量", "決定送出壓力，但不能超出允許動作", BLUE),
        "`steps_since_send`、`queue_size`、`urgent_due_in_s`、`previous_action` 與 `contact_remaining_s` 都是 runner 在決策時提供的欄位。它們的單位分別是 steps、packet count、秒、action enum 與秒；作用是控制節奏、批次、期限、轉移與窗口時機。函式只能回傳 `WAIT`、`SLEEP`、`SEND_ONE`、`SEND_URGENT` 或 `FLUSH_BATCH`，因此條件分支的結果仍被限制在合法 action 邊界。",
        (GOLD_PALE, GOLD, NAVY),
    )


def slide_025a(prs: Presentation):
    return _beginner_page(
        prs, "P025a｜Lab A／B 的受控常數",
        "先說原本規則，再一次只改一個數值",
        [
            ("A：送出間隔", "兩次送出至少隔 2 步", BLUE_PALE, BLUE, 20, 19),
            ("A：空檔動作", "空檔採低功耗休息\n`SLEEP`", GOLD_PALE, GOLD, 20, 18),
            ("B：進入與退出", "進入門檻 2；退出門檻 1", TEAL_PALE, TEAL, 19, 18),
            ("B：穩定條件", "連續穩定 2 步才保持", PURPLE_PALE, PURPLE, 19, 18),
        ],
        "候選只改一個值：A 改清醒等待，B 改穩定步數",
        ("hysteresis", "遲滯控制", "進入、退出與保持規則", "分開三種條件", "控制規則", "降低來回切換，再用事件與能量判讀", PURPLE),
        "Lab A 使用 `PACE_GAP_STEPS` 與 `REST_DURING_GAP` 描述送出間隔和空檔 action；封裝 baseline 是 2 與 `SLEEP`。Lab B 使用 `ENTER_QUALITY = 2`、`EXIT_QUALITY = 1` 與 `STABLE_STEPS = 2` 描述 send-ready 的進入、退出與連續穩定條件。候選變更分別是把 A 的休息 action 改為 `WAIT`，或把 B 的 hold 改為 1；每次只動一個 marked constant，並以 state、packet、service 與 endpoint energy 比較。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_025b(prs: Presentation):
    return _beginner_page(
        prs, "P025b｜Lab C 的佇列與緊急常數",
        "批次規則處理一般資料；期限逼近時優先處理緊急資料",
        [
            ("批次大小", "三筆一起送出\n`BATCH_SIZE`", BLUE_PALE, BLUE, 21, 18),
            ("期限餘裕", "剩 20 秒時進入緊急分支", GOLD_PALE, GOLD, 20, 18),
            ("動作結果", "批次送出／緊急送出", TEAL_PALE, TEAL, 21, 19),
        ],
        "每次只改一個值，再看佇列、送達與端點能量",
        ("URGENT_MARGIN_S", "緊急餘裕秒數", "Lab C 標記區", "提供期限分支的比較值", "秒", "改值後觀察緊急動作、送達與期限", GOLD),
        "Lab C 的 `BATCH_SIZE` 來自 marked block，單位是 packet count，作用是決定 queue 何時成組送出；`URGENT_MARGIN_S` 的單位是秒，作用是讓期限逼近時啟動 `SEND_URGENT`。`FLUSH_BATCH` 是 policy API 的合法 action，將 queue 連接到 process、TX、RX 與 delivery events。candidate 與 revision 各自只改一個 marked constant，withheld case 保持 frozen policy，用來讀取條件式的 queue、service 與 energy 後果。",
        (GOLD_PALE, GOLD, NAVY),
    )


def slide_026a(prs: Presentation):
    return _beginner_page(
        prs, "P026a｜讀取條件與 action 分支",
        "先判斷窗口，再判斷期限，最後判斷節奏",
        [
            ("01｜窗口關閉", "不能送出，先休息\n`SLEEP`", BLUE_PALE, BLUE, 20, 19),
            ("02｜期限逼近", "優先送出緊急資料", RED_PALE, RED, 20, 19),
            ("03｜間隔未到", "保持低功耗或清醒等待", GOLD_PALE, GOLD, 20, 19),
        ],
        "每段都回傳合法動作；順序本身就是機制",
        ("urgent_pending", "緊急封包待處理", "佇列與期限觀察", "顯示緊急資料是否存在", "真假值", "啟動期限分支，再讀事件證據", RED),
        "目前 policy 的 branch order 先檢查窗口，再檢查緊急期限，最後檢查送出節奏。窗口關閉時，`contact_open` 為 false，回傳 `SLEEP` 以維持 send legality；urgent packet 逼近 `URGENT_MARGIN_S` 時，回傳 `SEND_URGENT`；距離上次送出仍小於 `PACE_GAP_STEPS` 時，回傳 A block 指定的 `REST_DURING_GAP`。`urgent_pending` 是 queue／deadline observation 的 Boolean，作用是啟動期限分支；條件本身是機制入口，結果仍要從事件 ledger 判讀。",
        (RED_PALE, RED, RED),
    )


def slide_026b(prs: Presentation):
    return _beginner_page(
        prs, "P026b｜Action 後果的事件讀法",
        "動作留下狀態與封包事件，最後才看服務和端點能量",
        [
            ("清醒等待", "清醒閒置時間增加", TEAL_PALE, TEAL, 21, 19),
            ("低功耗休息", "休息後要喚醒", BLUE_PALE, BLUE, 21, 19),
            ("送出", "處理、發射、接收，再看是否送達", GOLD_PALE, GOLD, 20, 18),
            ("結果", "服務成敗與能量一起判讀\n`endpoint_energy_j`", PURPLE_PALE, PURPLE, 20, 17),
        ],
        "較低端點能量但服務失敗，是能量與服務的取捨",
        ("endpoint_energy_j", "端點能量", "端點無線與處理模型", "累積狀態與轉換能量", "焦耳", "只屬端點證據，不等於系統能量", PURPLE),
        "`WAIT` 會產生 `AWAKE_IDLE` interval，`SLEEP` 會產生 `SLEEP` interval，後續再次啟動 action 時可能出現 `WAKE` transition。`SEND_ONE`、`SEND_URGENT` 與 `FLUSH_BATCH` 會沿 process、TX、RX 形成 packet attempt、retry、delivery 或 expiry。`endpoint_energy_j` 來自 endpoint radio／processing model，單位是 J；它必須和 service result 一起判讀。較低 J 與 service fail 同時出現時，證據分類是 energy／service trade-off；它不會被改稱為 LEO、system 或 canonical energy。",
        (RED_PALE, RED, RED),
    )


def slide_027a(prs: Presentation):
    return _beginner_page(
        prs, "P027a｜Python indentation 與 return",
        "縮排決定哪個條件生效；回傳把動作交回執行器",
        [
            ("縮排", "把程式放進正確分支", BLUE_PALE, BLUE, 22, 20),
            ("回傳", "交回一個合法動作", GOLD_PALE, GOLD, 22, 20),
            ("檢查語法", "先檢查、不改內容\n`py_compile`", PURPLE_PALE, PURPLE, 20, 18),
        ],
        "錯誤時只修目前區段；情境與執行器保持原樣",
        ("return", "回傳", "策略函式", "把一個合法動作交給執行器", "動作值", "縮排或名稱錯誤時保留錯誤並只修目前區段", GOLD),
        "Python 的 `indentation` 是 code structure，作用是把 statements 放入正確的條件分支；`return` 來自 policy function，作用是將一個 action enum 送回 runner，結束這一個 decision step。目前 package 的最小骨架在窗口關閉時回傳 `SLEEP`，窗口開啟後回傳 `WAIT`。compile command 只檢查語法；若出現錯誤，修復目前 active marked block，scenario、runner 與 schema 保持原狀。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_027b(prs: Presentation):
    return _beginner_page(
        prs, "P027b｜Lab A baseline run 與結果入口",
        "先建立基準，再用同一案例比較改動後果",
        [
            ("基準案例", "固定策略＋固定情境", BLUE_PALE, BLUE, 22, 20),
            ("執行命令", "讀取固定案例並產生結果\n`course.sh`", TEAL_PALE, TEAL, 20, 18),
            ("結果檔案", "保存結果與重播位置\n`result.json`", GOLD_PALE, GOLD, 21, 18),
        ],
        "先保存基準結果，下一頁才有可比較的起點",
        ("result_path", "結果檔路徑", "命令輸出", "定位結果檔案", "路徑文字", "與同一案例的重播資料配對，作為觀察入口", PURPLE),
        "Lab A baseline 使用 packaged policy、fixed scenario 與 baseline case 建立 control artifact。POSIX 使用 `bash course.sh run --lab A --case baseline`，Windows 使用 `course.cmd run --lab A --case baseline`；runner stdout 會輸出 `status: OK`、`run_id` 與 `result_path`，其中 OK 只描述 artifact status。`result_path` 指向 `result.json`，同一個 run directory 內的 `endpoint-replay.json` 提供 action、state、queue 與 packet timing。保存 prediction 與這兩個配對入口後，下一段即可依 state、packet、service 與 endpoint energy 讀取 baseline，並開始 Part B 的受控比較。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


# ---------------------------------------------------------------------------
# V2 donor-style redraw
# ---------------------------------------------------------------------------
#
# The first implementation used paragraph-shaped cards: every field's source,
# purpose, unit, and interpretation was placed in the visual board.  That is
# useful for a source manifest, but it is not the approved eight-page classroom
# style.  The accepted donor uses one short takeaway, a small number of large
# cards, one causal strip, and the detailed field contract only at the bottom.
# These replacement builders keep the authoritative notes and field contracts
# while making the visible teaching move legible at a glance.


def add_takeaway(slide, text: str, name: str):
    return add_text(slide, 1.02, 1.05, 11.30, 0.52, text, 24, NAVY, True,
                    PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
                    margins=(0.02, 0.0, 0.02, 0.0), name=name,
                    line_spacing=0.96)


def add_visual_card(slide, x: float, y: float, w: float, h: float,
                    heading: str, body: str, fill: str, line: str,
                    name: str, heading_size: float = 23,
                    body_size: float = 20):
    """Donor-style card: short heading plus at most two short body lines."""
    add_box(slide, x, y, w, h, fill, line, name=name)
    add_text(slide, x + 0.16, y + 0.20, w - 0.32, 0.58, heading,
             heading_size, line, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
             margins=(0.03, 0.01, 0.03, 0.01), line_spacing=0.92,
             name=f"{name} heading")
    add_text(slide, x + 0.18, y + 0.98, w - 0.36, h - 1.18, body,
             body_size, INK, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
             margins=(0.05, 0.01, 0.05, 0.01), line_spacing=0.94,
             name=f"{name} body")


def add_causal_strip(slide, text: str, name: str, fill=GOLD_PALE, line=GOLD,
                     color=NAVY):
    return add_label_box(slide, 1.24, 4.62, 10.88, 0.58, text, fill, line,
                         20, color, True, name=name)


def slide_019(prs: Presentation):
    slide = prepare_slide(prs, "P019｜WSL package path")
    add_takeaway(slide, "release 來源與 Linux package root 必須分開", "P019 takeaway")
    xs = [0.92, 3.98, 7.04, 10.10]
    cards = [
        ("Windows\nDownloads", "release asset", BLUE_PALE, BLUE),
        ("`/mnt/c/.../Downloads`", "WSL 讀取位置", TEAL_PALE, TEAL),
        ("`~/c120-course`", "Linux staging path", GOLD_PALE, GOLD),
        ("`lora-energy-lab/`", "setup + `.venv` root", PURPLE_PALE, PURPLE),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, xs[i], 1.82, 2.42, 2.52, heading, body,
                        fill, line, f"P019 path card {i + 1}",
                        heading_size=20 if i != 1 else 18, body_size=21)
        if i < 3:
            add_chevron(slide, xs[i] + 2.56, 2.82, 0.36, 0.48,
                        (TEAL, GOLD, PURPLE)[i], f"P019 path connector {i + 1}")
    add_causal_strip(slide, "`cp` 複製 release → `pwd` 核對 package root",
                     "P019 causal strip")
    add_field_band(slide, "pwd", "目前目錄命令", "POSIX shell",
                   "確認 package root", "path string",
                   "輸出位於 Linux home 下的 `lora-energy-lab/`", color=BLUE)
    slide.notes_slide.notes_text_frame.text = (
        "WSL 會讀取 Windows Downloads 中的 release asset，但 Linux 的 setup 與虛擬環境必須在 Linux home 完成。"
        "命令中的 `/mnt/c/.../Downloads` 是 WSL mount 讀取位置，`cp` 將兩個 release files 複製到 `~/c120-course`，"
        "最後進入 `lora-energy-lab/` 並以 `pwd` 核對 package root。這個路徑分離保持 `.venv/bin/python`、lock 與 shell script 屬於同一份 Linux package。"
    )


def slide_020(prs: Presentation):
    slide = prepare_slide(prs, "P020｜POSIX／WSL setup 與 verify")
    add_takeaway(slide, "先固定 Python 3.11，再建立 READY receipt", "P020 takeaway")
    add_visual_card(slide, 0.92, 1.82, 5.54, 2.62,
                    "`setup.sh`", "建立 `.venv`\n安裝 lock dependencies",
                    BLUE_PALE, BLUE, "P020 setup card", 24, 21)
    add_visual_card(slide, 6.86, 1.82, 5.54, 2.62,
                    "`course.sh verify`", "核對 interpreter / contract\n輸出 `READY` receipt",
                    TEAL_PALE, TEAL, "P020 verify card", 22, 20)
    add_causal_strip(slide, "`PYTHON_BIN=python3.11` → `.venv/bin/python` → `READY`",
                     "P020 causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "PYTHON_BIN", "Python 執行器路徑變數", "shell environment",
                   "選擇 exact 3.11 interpreter", "path string",
                   "提供給 setup；成功後由 package-local `.venv/bin/python` 執行 verify",
                   color=TEAL)
    slide.notes_slide.notes_text_frame.text = (
        "POSIX 與 WSL 的 setup 需要一個明確的 Python 3.11.x interpreter。"
        "`PYTHON_BIN` 來自 shell environment，作用是把 `python3.11` 或 `uv python find 3.11` 的 path 交給 `setup.sh`；"
        "setup 建立 package-local `.venv`、安裝 lock dependencies，`course.sh verify` 再核對 exact interpreter 與 course contract。"
        "預期狀態是 `READY`、Python 3.11.x 與 `.venv/bin/python`；此 gate 尚未產生任何 lab case result。"
    )


def slide_021a(prs: Presentation):
    slide = prepare_slide(prs, "P021a｜查看 READY receipt")
    add_takeaway(slide, "同一份 receipt，可用三種 shell 讀取", "P021a takeaway")
    cards = [
        ("POSIX／WSL", "`sed -n '1,80p'`\n`artifacts/verify-receipt.json`", BLUE_PALE, BLUE),
        ("PowerShell", "`Get-Content`\n`verify-receipt.json`", TEAL_PALE, TEAL),
        ("Windows CMD", "`type`\n`verify-receipt.json`", GOLD_PALE, GOLD),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P021a shell card {i + 1}",
                        heading_size=21, body_size=18)
    add_causal_strip(slide, "一份 `verify-receipt.json`；讀取 provenance 不變",
                     "P021a causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "verify-receipt.json", "驗證收據檔", "setup／verify command",
                   "記錄 environment 與 contract gate", "JSON artifact",
                   "解讀 setup evidence；不代替 case result", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "`verify-receipt.json` 是 setup 與 verify command 寫入的驗證收據檔。"
        "POSIX 使用 `sed`，PowerShell 使用 `Get-Content`，Windows CMD 使用 `type`；三個命令都只讀取同一份 artifact。"
        "收據的欄位用來描述環境與契約 gate，case result 仍由後續 run command 另行產生，兩種 provenance 不混用。"
    )


def slide_021b(prs: Presentation):
    slide = prepare_slide(prs, "P021b｜READY receipt：狀態與識別欄位")
    add_takeaway(slide, "READY 只證明 identity 一致，不代表 case result", "P021b takeaway")
    rows = [
        ("`status`", "READY gate", BLUE_PALE, BLUE),
        ("`python_version`", "3.11.x", TEAL_PALE, TEAL),
        ("`scenario_id`", "fixed case identity", GOLD_PALE, GOLD),
        ("`policy_api_version`", "observation / action contract", PURPLE_PALE, PURPLE),
    ]
    for i, (heading, body, fill, line) in enumerate(rows):
        add_visual_card(slide, 0.92 + i * 3.05, 1.82, 2.74, 2.58,
                        heading, body, fill, line, f"P021b field card {i + 1}",
                        heading_size=19 if i in {1, 3} else 21, body_size=18)
    add_causal_strip(slide, "READY：環境與契約一致｜result identity 另行產生",
                     "P021b causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "status", "狀態", "verify receipt",
                   "表示 gate 是否可進入 case", "enum",
                   "READY 只描述環境與契約一致；result identity 仍需另行產生", color=BLUE)
    slide.notes_slide.notes_text_frame.text = (
        "READY receipt 的第一組欄位描述狀態與 identity。"
        "`status` 必須是 READY，`python_version` 必須是 3.11.x；`scenario_id` 將固定情境和 case 綁定，"
        "`policy_api_version` 將允許的 observation 與 action 契約綁定。這些欄位表示 setup gate 通過，並不表示能量結果已經存在。"
    )


def slide_021c(prs: Presentation):
    slide = prepare_slide(prs, "P021c｜READY receipt：引擎與主張欄位")
    add_takeaway(slide, "producer mode 決定 evidence boundary", "P021c takeaway")
    cards = [
        ("`engine_mode`", "simulated adapter", BLUE_PALE, BLUE),
        ("`upstream_execution`", "`false`", TEAL_PALE, TEAL),
        ("`claim_boundary`", "simulated / non-live", GOLD_PALE, GOLD),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P021c claim card {i + 1}",
                        heading_size=19, body_size=20)
    add_causal_strip(slide, "READY 是 setup gate；`result.json` 另有 run identity",
                     "P021c causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "engine_mode", "執行引擎模式", "runner metadata",
                   "分類 producer", "enum",
                   "`coherent-course-simulated-adapter`；不延伸為上游程式執行", color=TEAL)
    slide.notes_slide.notes_text_frame.text = (
        "`engine_mode`、`upstream_execution` 與 `claim_boundary` 共同描述 producer 和 evidence scope。"
        "engine mode 是 `coherent-course-simulated-adapter`，上游程式執行狀態為 false，claim boundary 固定標示 simulated teaching data、非 live、非 measured 與非 canonical parity。"
        "因此 READY 表示環境與契約一致；case result 仍需用自己的 run identity 保存，兩者不互相代替。"
    )


def slide_022(prs: Presentation):
    slide = prepare_slide(prs, "P022｜Setup gate recovery 與 same-scenario fallback")
    add_takeaway(slide, "gate 受阻時，保留資料來源分類", "P022 takeaway")
    cards = [
        ("READY", "exact case run\n產生 result / replay", BLUE_PALE, BLUE),
        ("READY 受阻", "repair Python / root / lock", GOLD_PALE, GOLD),
        ("fallback", "same scenario\nresult + replay pair", PURPLE_PALE, PURPLE),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P022 recovery card {i + 1}",
                        heading_size=22, body_size=19)
    add_causal_strip(slide, "`result.json` + `endpoint-replay.json` 保留同一 scenario",
                     "P022 causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "artifact_source", "資料來源分類", "result metadata",
                   "標示 local policy run 或 fallback", "enum",
                   "來源模式維持原值；同情境回復資料不改寫成本機執行", color=GOLD)
    slide.notes_slide.notes_text_frame.text = (
        "setup 或 verify 受阻時，artifact source 必須保持可見。"
        "可修復的問題只在 Python minor、package root、lock 或 policy boundary 內處理；若 local setup 仍無法建立 READY，"
        "同一個 lab／case 的 `result.json` 與 `endpoint-replay.json` 以 pair 方式提供 recovery。"
        "`artifact_source` 是 result metadata 的資料來源分類，fallback 保持 `same-scenario-fallback`，不改寫成本機 policy execution claim。"
    )


def slide_023(prs: Presentation):
    slide = prepare_slide(prs, "P023｜Bounded edit surface：student_policy.py")
    add_takeaway(slide, "只編輯 `student_policy.py` 的 marked block", "P023 takeaway")
    cards = [
        ("01｜檔案", "`student_policy.py`", BLUE_PALE, BLUE),
        ("02｜區段", "A / B / C marked block", GOLD_PALE, GOLD),
        ("03｜固定", "scenario + schemas", TEAL_PALE, TEAL),
        ("04｜輸出", "runner → result / replay", PURPLE_PALE, PURPLE),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.05, 1.82, 2.74, 2.58,
                        heading, body, fill, line, f"P023 edit card {i + 1}",
                        heading_size=19, body_size=18)
    add_causal_strip(slide, "一次只開啟目前 lab 的 marked block；其他 boundary 保持原樣",
                     "P023 causal strip", fill=TEAL_PALE, line=TEAL, color=NAVY)
    add_field_band(slide, "marked block", "標記區段", "file markers",
                   "界定一個 lab 的 edit surface", "source range",
                   "每次只保留一個 active region；其他 package boundary 維持 read-only", color=BLUE)
    slide.notes_slide.notes_text_frame.text = (
        "`student_policy.py` 是 package root 內唯一的策略檔。"
        "目前檔案的三個 markers 是 `lab-a-pace-rest`、`lab-b-enter-exit-hold` 與 `lab-c-batch-urgent`；每個 lab 只在自己的區段變更常數。"
        "scenario、schemas、runner 與 generated JSON 分別是固定輸入、契約、執行器與輸出；保留這些 boundary，policy identity 和結果 lineage 才能被比較。"
    )


def slide_024a(prs: Presentation):
    slide = prepare_slide(prs, "P024a｜Observation 是 decision-time input")
    add_takeaway(slide, "決策只讀當下 observation，回傳一個 legal action", "P024a takeaway")
    cards = [
        ("目前 observation", "`contact_open`\n`quality_band` / `stable_steps`", BLUE_PALE, BLUE),
        ("`choose_action`", "讀 current input\n回傳一個 action", TEAL_PALE, TEAL),
        ("legal action", "窗口關閉\n→ `SLEEP`", GOLD_PALE, GOLD),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P024a observation card {i + 1}",
                        heading_size=20, body_size=20)
        if i < 2:
            add_chevron(slide, 4.60 + i * 3.85, 2.86, 0.36, 0.48,
                        (TEAL, GOLD)[i], f"P024a connector {i + 1}")
    add_causal_strip(slide, "future quality／future energy／result summary 不屬於 input",
                     "P024a causal strip", fill=RED_PALE, line=RED, color=RED)
    add_field_band(slide, "choose_action(observation)", "策略函式", "policy API",
                   "讀 current observation 並回傳一個合法 action", "function contract",
                   "不讀 future 或 result summary；輸出交給 runner 產生事件", color=TEAL)
    slide.notes_slide.notes_text_frame.text = (
        "`choose_action(observation)` 是 policy API 的函式契約。"
        "`contact_open` 來自當下 scenario observation，`quality_band` 來自目前 trace，`stable_steps` 來自連續觀察；"
        "窗口關閉時，函式回傳 `SLEEP`，讓 runner 進入合法的低功耗分支。future quality、future energy 與 result summary 不屬於 decision-time input，"
        "因此 action 是對當下資訊的控制輸出，而不是對結果的回看。"
    )


def slide_024b(prs: Presentation):
    slide = prepare_slide(prs, "P024b｜Observation 欄位與合法 action")
    add_takeaway(slide, "queue、deadline、前一動作補齊 decision-time 邊界", "P024b takeaway")
    cards = [
        ("runner state", "`steps_since_send`\n`previous_action`", BLUE_PALE, BLUE),
        ("queue + deadline", "`queue_size`\n`urgent_due_in_s`", GOLD_PALE, GOLD),
        ("合法輸出", "`WAIT` / `SLEEP`\n`SEND_*` / `FLUSH_BATCH`", TEAL_PALE, TEAL),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P024b field card {i + 1}",
                        heading_size=21, body_size=19)
    add_causal_strip(slide, "`contact_remaining_s` 影響 send timing；輸出仍限於 API enum",
                     "P024b causal strip", fill=GOLD_PALE, line=GOLD)
    add_field_band(slide, "queue_size", "佇列數量", "endpoint queue",
                   "提供 batching／urgency input", "packet count",
                   "決定送出壓力；合法輸出仍限於 policy API enum", color=BLUE)
    slide.notes_slide.notes_text_frame.text = (
        "`steps_since_send`、`queue_size`、`urgent_due_in_s`、`previous_action` 與 `contact_remaining_s` 都是 runner 在決策時提供的欄位。"
        "它們的單位分別是 steps、packet count、秒、action enum 與秒；作用是控制節奏、批次、期限、轉移與窗口時機。"
        "函式只能回傳 `WAIT`、`SLEEP`、`SEND_ONE`、`SEND_URGENT` 或 `FLUSH_BATCH`，因此條件分支的結果仍被限制在合法 action 邊界。"
    )


def slide_025a(prs: Presentation):
    slide = prepare_slide(prs, "P025a｜Lab A／B 的受控常數")
    add_takeaway(slide, "一次只改一個常數，才能比較 A／B 後果", "P025a takeaway")
    cards = [
        ("Lab A｜pace", "`PACE_GAP_STEPS = 2`", BLUE_PALE, BLUE),
        ("Lab A｜rest", "`REST_DURING_GAP = SLEEP`", GOLD_PALE, GOLD),
        ("Lab B｜enter／exit", "`ENTER = 2` / `EXIT = 1`", TEAL_PALE, TEAL),
        ("Lab B｜hold", "`STABLE_STEPS = 2`", PURPLE_PALE, PURPLE),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.05, 1.82, 2.74, 2.58,
                        heading, body, fill, line, f"P025a constant card {i + 1}",
                        heading_size=18 if i in {2, 3} else 20, body_size=18)
    add_causal_strip(slide, "candidate：A → `WAIT`｜B → `STABLE_STEPS = 1`",
                     "P025a causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "hysteresis", "遲滯控制", "enter／exit／hold constants",
                   "分開進入、退出與 hold 條件", "rule",
                   "降低 ping-pong transition；仍需用 packet／service／energy evidence 判讀", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "Lab A 使用 `PACE_GAP_STEPS` 與 `REST_DURING_GAP` 描述送出間隔和空檔 action；封裝 baseline 是 2 與 `SLEEP`。"
        "Lab B 使用 `ENTER_QUALITY = 2`、`EXIT_QUALITY = 1` 與 `STABLE_STEPS = 2` 描述 send-ready 的進入、退出與連續穩定條件。"
        "候選變更分別是把 A 的休息 action 改為 `WAIT`，或把 B 的 hold 改為 1；每次只動一個 marked constant，並以 state、packet、service 與 endpoint energy 比較。"
    )


def slide_025b(prs: Presentation):
    slide = prepare_slide(prs, "P025b｜Lab C 的佇列與緊急常數")
    add_takeaway(slide, "Lab C 把 queue batching 與 deadline protection 接到 action", "P025b takeaway")
    cards = [
        ("`BATCH_SIZE`", "`= 3`\nqueue → flush", BLUE_PALE, BLUE),
        ("`URGENT_MARGIN_S`", "`= 20`\ndeadline → urgent", GOLD_PALE, GOLD),
        ("action surface", "`FLUSH_BATCH`\n`SEND_URGENT`", TEAL_PALE, TEAL),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P025b constant card {i + 1}",
                        heading_size=20 if i != 1 else 18, body_size=20)
    add_causal_strip(slide, "candidate／revision：每次只改一個 marked constant",
                     "P025b causal strip", fill=GOLD_PALE, line=GOLD)
    add_field_band(slide, "URGENT_MARGIN_S", "緊急餘裕秒數", "Lab C marked block",
                   "提供期限分支的比較值", "s",
                   "值變更後觀察 urgent action、delivery、deadline 與 endpoint J", color=GOLD)
    slide.notes_slide.notes_text_frame.text = (
        "Lab C 的 `BATCH_SIZE` 來自 marked block，單位是 packet count，作用是決定 queue 何時成組送出；"
        "`URGENT_MARGIN_S` 的單位是秒，作用是讓期限逼近時啟動 `SEND_URGENT`。"
        "`FLUSH_BATCH` 是 policy API 的合法 action，將 queue 連接到 process、TX、RX 與 delivery events。"
        "candidate 與 revision 各自只改一個 marked constant，withheld case 保持 frozen policy，用來讀取條件式的 queue、service 與 energy 後果。"
    )


def slide_026a(prs: Presentation):
    slide = prepare_slide(prs, "P026a｜讀取條件與 action 分支")
    add_takeaway(slide, "branch order 把 observation 轉成 legal action", "P026a takeaway")
    cards = [
        ("01｜窗口", "`contact_open = false`\n→ `SLEEP`", BLUE_PALE, BLUE),
        ("02｜期限", "urgent due ≤ margin\n→ `SEND_URGENT`", RED_PALE, RED),
        ("03｜節奏", "steps < gap\n→ `WAIT` / `SLEEP`", GOLD_PALE, GOLD),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P026a branch card {i + 1}",
                        heading_size=20, body_size=19)
        if i < 2:
            add_chevron(slide, 4.60 + i * 3.85, 2.86, 0.36, 0.48,
                        (RED, GOLD)[i], f"P026a connector {i + 1}")
    add_causal_strip(slide, "先窗口，再期限，再節奏；每段回傳合法 action",
                     "P026a causal strip", fill=RED_PALE, line=RED, color=RED)
    add_field_band(slide, "urgent_pending", "緊急封包待處理", "queue／deadline observation",
                   "顯示 urgent packet 是否存在", "Boolean",
                   "啟動 deadline branch；action 後仍需讀 packet／service evidence", color=RED)
    slide.notes_slide.notes_text_frame.text = (
        "目前 policy 的 branch order 先檢查窗口，再檢查緊急期限，最後檢查送出節奏。"
        "窗口關閉時，`contact_open` 為 false，回傳 `SLEEP` 以維持 send legality；urgent packet 逼近 `URGENT_MARGIN_S` 時，回傳 `SEND_URGENT`；"
        "距離上次送出仍小於 `PACE_GAP_STEPS` 時，回傳 A block 指定的 `REST_DURING_GAP`。"
        "`urgent_pending` 是 queue／deadline observation 的 Boolean，作用是啟動期限分支；條件本身是機制入口，結果仍要從事件 ledger 判讀。"
    )


def slide_026b(prs: Presentation):
    slide = prepare_slide(prs, "P026b｜Action 後果的事件讀法")
    add_takeaway(slide, "action 後果沿 state → packet → service → endpoint J 讀取", "P026b takeaway")
    cards = [
        ("`WAIT`", "`AWAKE_IDLE`\n清醒閒置", TEAL_PALE, TEAL),
        ("`SLEEP`", "`SLEEP` + `WAKE`\n休息與喚醒", BLUE_PALE, BLUE),
        ("`SEND_*`", "PROCESS → TX → RX\npacket outcome", GOLD_PALE, GOLD),
        ("結果層", "service result\n`endpoint_energy_j`", PURPLE_PALE, PURPLE),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.05, 1.82, 2.74, 2.58,
                        heading, body, fill, line, f"P026b event card {i + 1}",
                        heading_size=19 if i in {2, 3} else 21, body_size=18)
    add_causal_strip(slide, "較低 J 且 service fail：energy／service trade-off",
                     "P026b causal strip", fill=RED_PALE, line=RED, color=RED)
    add_field_band(slide, "endpoint_energy_j", "端點能量", "endpoint radio／processing model",
                   "累積狀態與轉換能量", "J",
                   "只屬 endpoint evidence layer；與 LEO／system／canonical energy 分開", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "`WAIT` 會產生 `AWAKE_IDLE` interval，`SLEEP` 會產生 `SLEEP` interval，後續再次啟動 action 時可能出現 `WAKE` transition。"
        "`SEND_ONE`、`SEND_URGENT` 與 `FLUSH_BATCH` 會沿 process、TX、RX 形成 packet attempt、retry、delivery 或 expiry。"
        "`endpoint_energy_j` 來自 endpoint radio／processing model，單位是 J；它必須和 service result 一起判讀。"
        "較低 J 與 service fail 同時出現時，證據分類是 energy／service trade-off；它不會被改稱為 LEO、system 或 canonical energy。"
    )


def slide_027a(prs: Presentation):
    slide = prepare_slide(prs, "P027a｜Python indentation 與 return")
    add_takeaway(slide, "縮排決定 branch；`return` 把 action 交給 runner", "P027a takeaway")
    cards = [
        ("`indentation`", "`if ...:`\n`return SLEEP`", BLUE_PALE, BLUE),
        ("`return`", "一個 legal action\n回到 runner", GOLD_PALE, GOLD),
        ("syntax gate", "`py_compile`\n只檢查語法", PURPLE_PALE, PURPLE),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P027a code card {i + 1}",
                        heading_size=20, body_size=20)
    add_causal_strip(slide, "錯誤時只修 active marked block；scenario／runner 保持原樣",
                     "P027a causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "return", "回傳", "policy function",
                   "將一個合法 action 交給 runner", "action enum",
                   "若縮排或名稱錯誤，保留錯誤並只修 active marked block", color=GOLD)
    slide.notes_slide.notes_text_frame.text = (
        "Python 的 `indentation` 是 code structure，作用是把 statements 放入正確的條件分支；"
        "`return` 來自 policy function，作用是將一個 action enum 送回 runner，結束這一個 decision step。"
        "目前 package 的最小骨架在窗口關閉時回傳 `SLEEP`，窗口開啟後回傳 `WAIT`。"
        "compile command 只檢查語法；若出現錯誤，修復目前 active marked block，scenario、runner 與 schema 保持原狀。"
    )


def slide_027b(prs: Presentation):
    slide = prepare_slide(prs, "P027b｜Lab A baseline run 與結果入口")
    add_takeaway(slide, "baseline 先建立 control artifact，再進入比較", "P027b takeaway")
    cards = [
        ("baseline", "packaged policy\n+ fixed scenario", BLUE_PALE, BLUE),
        ("run command", "`course.sh run --lab A`\n`--case baseline`", TEAL_PALE, TEAL),
        ("result + replay", "`result.json`\n`endpoint-replay.json`", GOLD_PALE, GOLD),
    ]
    for i, (heading, body, fill, line) in enumerate(cards):
        add_visual_card(slide, 0.92 + i * 3.85, 1.82, 3.54, 2.62,
                        heading, body, fill, line, f"P027b baseline card {i + 1}",
                        heading_size=21 if i != 1 else 20, body_size=19)
        if i < 2:
            add_chevron(slide, 4.60 + i * 3.85, 2.86, 0.36, 0.48,
                        (TEAL, GOLD)[i], f"P027b connector {i + 1}")
    add_causal_strip(slide, "`result.json` + `endpoint-replay.json` → Part B observation",
                     "P027b causal strip", fill=PURPLE_PALE, line=PURPLE)
    add_field_band(slide, "result_path", "結果檔路徑", "command JSON output",
                   "定位 `result.json`", "path string",
                   "與同一 run directory 的 endpoint replay 配對；作為觀察入口", color=PURPLE)
    slide.notes_slide.notes_text_frame.text = (
        "Lab A baseline 使用 packaged policy、fixed scenario 與 baseline case 建立 control artifact。"
        "POSIX 使用 `bash course.sh run --lab A --case baseline`，Windows 使用 `course.cmd run --lab A --case baseline`；"
        "runner stdout 會輸出 `status: OK`、`run_id` 與 `result_path`，其中 OK 只描述 artifact status。"
        "`result_path` 指向 `result.json`，同一個 run directory 內的 `endpoint-replay.json` 提供 action、state、queue 與 packet timing。"
        "保存 prediction 與這兩個配對入口後，下一段即可依 state、packet、service 與 endpoint energy 讀取 baseline，並開始 Part B 的受控比較。"
    )


# ---------------------------------------------------------------------------
# Final beginner / lab-acceptance pass
# ---------------------------------------------------------------------------
#
# Keep the approved eight-page donor geometry while making the policy teaching
# sequence explicit.  This block is intentionally last: the earlier redraw
# functions are retained as provenance, but the names below are the builders
# used by build().  The executable student_policy.py is never edited here.


def slide_019(prs: Presentation):
    return _beginner_page(
        prs, "把下載資料搬進 Linux 課程資料夾",
        "這一步只做路徑準備：把 release 帶進 Linux，再確認程式根目錄",
        [
            ("① 這個工作區", "WSL＝Windows 內的 Linux\n讀 `/mnt/c/.../Downloads`", BLUE_PALE, BLUE, 18, 17),
            ("② 程式不改", "`student_policy.py` 不動\n只複製 release 副本", TEAL_PALE, TEAL, 19, 17),
            ("③ 進入根目錄", "下一步集中到 Linux\n`lora-energy-lab/`", GOLD_PALE, GOLD, 19, 18),
            ("④ 核對結果", "執行 `pwd` 只讀位置\n路徑不對就先停止", PURPLE_PALE, PURPLE, 19, 17),
        ],
        "先把下載副本放進 Linux，再確認根目錄；路徑錯會讓後續設定讀錯檔",
        ("pwd", "目前目錄命令", "Linux／macOS 終端機（WSL 也使用 Linux 指令）", "確認目前位置", "路徑文字", "輸出應位於 Linux home 的課程資料夾", BLUE),
        "命令先讀取 Windows Downloads，再把 release 副本寫到 Linux staging；它不修改 Windows 原檔、`student_policy.py` 或任何 marked block。下一步進入 `lora-energy-lab/`，用 `pwd` 讀取目前位置；要檢查的 website／artifact field 是 package root 與來源路徑，路徑不一致時不要進入 setup。",
    )


def slide_020(prs: Presentation):
    return _beginner_page(
        prs, "先找到或安裝 Python 3.11",
        "這一步只找符合版本的 Python；找不到或版本不符就停止",
        [
            ("Linux／WSL 先檢查", "`command -v python3.11`\n`python3.11 --version`\n找不到：`sudo apt install -y curl`\n再 `uv python install 3.11`", BLUE_PALE, BLUE, 16, 16),
            ("macOS 先檢查", "`command -v python3.11`\n`python3.11 --version`\n沒有時：`brew install python@3.11`", TEAL_PALE, TEAL, 16, 16),
            ("Windows PowerShell", "`py -3.11 --version`\n找不到：`winget install --id=astral-sh.uv -e`\n再 `uv python install 3.11`", GOLD_PALE, GOLD, 16, 16),
            ("Windows Command Prompt", "`py -3.11 --version`\n只讀版本結果\n沒有 3.11：回 PowerShell 安裝", PURPLE_PALE, PURPLE, 16, 16),
        ],
        "先檢查 → 缺少就安裝 uv／Python 3.11 → 再檢查；版本不對就停在這一頁",
        ("PYTHON_BIN", "Python 3.11 執行器", "README.zh-TW.md", "選擇符合版本的 interpreter", "path", "結果必須是 Python 3.11.x；不符合就停止，不建 venv", TEAL),
        "本頁目的只是找到可用的 Python 3.11.x，不建立 venv，也不修改 `/home/u24/lora-energy-lab/student_policy.py`。Linux／WSL 先執行 `command -v python3.11` 與 `python3.11 --version`；若找不到，依 README 執行 `sudo apt update`、`sudo apt install -y curl`、`curl -LsSf https://astral.sh/uv/install.sh | sh`、`export PATH=\"$HOME/.local/bin:$PATH\"`、`uv python install 3.11`、`PYTHON_BIN=\"$(uv python find 3.11)\"`，最後以 `\"$PYTHON_BIN\" --version` 重查。macOS 先同樣檢查；沒有時可執行 `brew install python@3.11`、`PYTHON_BIN=\"$(brew --prefix python@3.11)/bin/python3.11\"`，或依 README 改走 uv。Windows PowerShell 先執行 `py -3.11 --version`，再以 `$python311 = (py -3.11 -c \"import sys; print(sys.executable)\").Trim()` 取得路徑；若找不到，執行 `winget install --id=astral-sh.uv -e`、`uv python install 3.11`、`$python311 = (uv python find 3.11).Trim()`、`& $python311 --version`。Command Prompt 只以 `py -3.11 --version` 檢查；缺少時回到 PowerShell 安裝，不在 CMD 混用 PowerShell 語法。預期輸出一律是 `Python 3.11.x`；若沒有，停止於本頁，不建立 `.venv`。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_021a(prs: Presentation):
    return _beginner_page(
        prs, "確認 venv 可用並建立 .venv",
        "這一步先測 venv module，再建立隔離資料夾；失敗就停，不安裝套件",
        [
            ("Linux／WSL", "`\"$PYTHON_BIN\" -m venv --help`\n`\"$PYTHON_BIN\" -m venv .venv`\n成功：出現 `.venv/`", BLUE_PALE, BLUE, 16, 16),
            ("macOS", "`\"$PYTHON_BIN\" -m venv --help`\n`\"$PYTHON_BIN\" -m venv .venv`\n成功：出現 `.venv/`", TEAL_PALE, TEAL, 17, 16),
            ("Windows PowerShell", "`& $python311 -m venv --help`\n`& $python311 -m venv .venv`\n成功：出現 `.venv\\`", GOLD_PALE, GOLD, 16, 16),
            ("Windows Command Prompt", "`py -3.11 -m venv --help`\n`py -3.11 -m venv .venv`\n成功：出現 `.venv\\`", PURPLE_PALE, PURPLE, 16, 16),
        ],
        "先測 `venv` module → 建立 `.venv`；出現 `venv`／`ensurepip` 錯誤就停，不執行 pip install",
        (".venv", "隔離 Python 環境", "Python 3.11 venv module", "建立 package-local 環境", "directory", "看到 `.venv/` 或 `.venv\\` 才進入啟用；失敗回到 Python 版本頁", BLUE),
        "本頁只做一件事：確認 Python 3.11 內建的 `venv` module 可用，並在 package root 建立 `.venv`；不啟用、不安裝 requirements、不修改 policy。Linux／WSL 與 macOS 執行 `\"$PYTHON_BIN\" -m venv --help`，再執行 `\"$PYTHON_BIN\" -m venv .venv`；PowerShell 執行 `& $python311 -m venv --help` 與 `& $python311 -m venv .venv`；Command Prompt 執行 `py -3.11 -m venv --help` 與 `py -3.11 -m venv .venv`。成功判斷是 package root 出現 `.venv` 目錄；若回報缺少 `venv` 或 `ensurepip`，停止、不使用 `pip install venv`，刪除未完成的 package-local `.venv` 後回到上一頁改用 uv-managed Python 再建立。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_021b(prs: Presentation):
    return _beginner_page(
        prs, "啟用 .venv 並確認 Python 路徑",
        "這一步讓 `python` 指向隔離環境；版本或路徑不對就停",
        [
            ("Linux／WSL", "`source .venv/bin/activate`\n`python --version`\n`python -c 'import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)'", BLUE_PALE, BLUE, 16, 16),
            ("macOS", "`source .venv/bin/activate`\n`python --version`\n`python -c 'import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)'", TEAL_PALE, TEAL, 17, 16),
            ("Windows PowerShell", "`.\\.venv\\Scripts\\Activate.ps1`\n`python --version`\n`python -c \"import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)\"`", GOLD_PALE, GOLD, 16, 16),
            ("Windows Command Prompt", "`call .venv\\Scripts\\activate.bat`\n`python --version`\n`python -c \"import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)\"`", PURPLE_PALE, PURPLE, 16, 16),
        ],
        "啟用 → `python` 指向 `.venv` → 版本與 `sys.prefix` 通過；不通過就停在啟用頁",
        ("sys.prefix", "隔離環境判斷", "Python runtime", "確認目前 shell 已啟用 `.venv`", "path", "`sys.prefix` 不等於 `sys.base_prefix`，且執行檔位於 `.venv`；否則不安裝", TEAL),
        "本頁只啟用剛建立的 `.venv` 並確認 interpreter；不安裝 requirements、不執行 verify、不修改 policy。Linux／WSL 與 macOS 使用 `source .venv/bin/activate`；PowerShell 使用 `.\\.venv\\Scripts\\Activate.ps1`；Command Prompt 使用 `call .venv\\Scripts\\activate.bat`。各 shell 接著執行 `python --version` 與 `python -c 'import sys; assert sys.prefix != sys.base_prefix; print(sys.executable)'`（Windows 使用雙引號版本）。成功輸出為 Python 3.11.x，且 executable path 在 `.venv` 內、`sys.prefix` 不等於 base prefix。若 PowerShell 的 `Activate.ps1` 被 execution policy 擋住，不放寬全系統 policy；停止並改在同一 package 的 Command Prompt 執行 `call .venv\\Scripts\\activate.bat`。版本或路徑不正確時不要進入 pip install。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_021c(prs: Presentation):
    return _beginner_page(
        prs, "在 venv 安裝鎖定套件並驗證 READY",
        "四種 shell 都先安裝鎖定檔再驗證；看到 READY 才進案例，失敗就停",
        [
            ("Linux／WSL", "`python -m pip install`\n`--require-hashes --no-deps`\n`-r requirements-lock.txt`\n`bash course.sh verify`", BLUE_PALE, BLUE, 16, 16),
            ("macOS", "`python -m pip install`\n`--require-hashes --no-deps`\n`-r requirements-lock.txt`\n`bash course.sh verify`", TEAL_PALE, TEAL, 17, 16),
            ("Windows PowerShell", "`python -m pip install`\n`--require-hashes --no-deps`\n`-r .\\requirements-lock.txt`\n`.\\course.cmd verify`", GOLD_PALE, GOLD, 16, 16),
            ("Windows Command Prompt", "`python -m pip install`\n`--require-hashes --no-deps`\n`-r requirements-lock.txt`\n`course.cmd verify`", PURPLE_PALE, PURPLE, 16, 16),
        ],
        "locked requirements → verify → `status: READY`；沒有 READY 就停止，不執行 lab",
        ("verify-receipt.json", "驗證收據檔", "`course.sh`／`course.cmd verify`", "記錄 Python、scenario 與 policy contract gate", "JSON artifact", "看到 `status: READY` 才能進案例；它不是 `result.json`", PURPLE),
        "本頁只在已啟用的 `.venv` 內安裝 package 的 exact locked file `requirements-lock.txt`，不是自行猜測的 `requirements.txt`；不修改 `/home/u24/lora-energy-lab/student_policy.py`、marked blocks、runner 或 scenario。Linux／WSL 執行 `python -m pip install --disable-pip-version-check --require-hashes --no-deps -r requirements-lock.txt`，接著 `bash course.sh verify`；macOS 使用相同兩個命令。PowerShell 執行相同 pip 參數但使用 `-r .\\requirements-lock.txt`，接著 ` .\\course.cmd verify`；Command Prompt 使用 `-r requirements-lock.txt`，接著 `course.cmd verify`。預期產生 `artifacts/verify-receipt.json`，其中 `status` 是 `READY`、Python 是 3.11.x，並保留 scenario identity、policy API 與 claim boundary。requirements install 或 verify 任何一步失敗就停，不進入 lab；此收據不是案例的 `result.json`，案例結果要等後續 run。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_022(prs: Presentation):
    return _beginner_page(
        prs, "卡住時修復或保留來源",
        "這一步處理卡關：修得好就重跑，修不好也不冒充本機執行",
        [
            ("① 正常", "READY 後執行案例\n產生新結果", BLUE_PALE, BLUE, 20, 18),
            ("② 修復什麼", "只修 Python、root、lock\n不改 policy block", GOLD_PALE, GOLD, 19, 17),
            ("③ 回復什麼", "同情境結果成對保存\n`result.json`／`endpoint-replay.json`", PURPLE_PALE, PURPLE, 18, 16),
            ("④ 看什麼", "讀 `artifact_source`\n判斷本機或回復", TEAL_PALE, TEAL, 19, 17),
        ],
        "修復不改策略；回復不冒充本機執行",
        ("artifact_source", "資料來源分類", "結果紀錄", "標示本機執行或回復資料", "來源值", "來源模式保持原值，不改寫成本機執行", GOLD),
        "setup 或 verify 受阻時，只修復 Python minor、package root、lock 或 policy boundary；不能把回復資料說成本機執行。下一步依 `artifact_source` 判斷是否可進 run；讀取同一 lab／case 的 `result.json` 與 `endpoint-replay.json` pair，再按 service → packet → state → endpoint J → efficiency 判讀。本頁不改 student_policy.py 或任何 marked block。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_023(prs: Presentation):
    return _beginner_page(
        prs, "只可修改三個標記區塊",
        "這一步畫出可編輯邊界：三個實驗區可改，其他程式只讀",
        [
            ("① 唯一檔案", "只在這裡做實驗\n`student_policy.py`", BLUE_PALE, BLUE, 20, 18),
            ("② 可改區段", "A／B／C 三個標記\n`marked block`", GOLD_PALE, GOLD, 20, 18),
            ("③ 保持不變", "中文註解、函式、情境\nrunner、schemas 不動", TEAL_PALE, TEAL, 18, 17),
            ("④ 下一步／證據", "只開目前 lab\n結果仍看 result／replay", PURPLE_PALE, PURPLE, 18, 17),
        ],
        "只開目前實驗區；其他檔案與輸出契約保持不變",
        ("marked block", "標記區段", "檔案標記", "界定一個實驗的可編輯範圍", "程式碼區段", "其他套件邊界維持唯讀", BLUE),
        "實際 `/home/u24/lora-energy-lab/student_policy.py` 現在包含中文註解；本頁只鏡像說明，不能把投影片文字貼回程式。UTF-8/LF bytes 參與 policy identity，outside-block 的註解、函式、runner、scenario、schemas 與輸出契約都不改；可編輯面只有 `lab-a-pace-rest`、`lab-b-enter-exit-hold`、`lab-c-batch-urgent`。下一步只改當前 block 的一個 value，再看成對 result／replay 證據。",
        (TEAL_PALE, TEAL, NAVY),
    )


def slide_024a(prs: Presentation):
    return _beginner_page(
        prs, "程式每次判斷會看什麼",
        "這一步先說清楚輸入：每個 step 只看目前／過去資料",
        [
            ("① 讀目前資料", "窗口、急件、間隔、品質\n都是當下 snapshot", BLUE_PALE, BLUE, 19, 17),
            ("② 不讀未來", "不看結果摘要\n不回看 energy", RED_PALE, RED, 20, 18),
            ("③ 交給函式", "把目前資料交給 `choose_action()`\n再回一個合法動作", GOLD_PALE, GOLD, 18, 17),
            ("④ 下一步／證據", "回一個合法動作\nrunner 再記 event／result", TEAL_PALE, TEAL, 18, 17),
        ],
        "當下觀察 → 合法動作；這一頁沒有程式值變更",
        ("choose_action", "策略函式", "student_policy.py", "依當下觀察回傳一個合法動作", "函式契約", "只讀目前／過去欄位，不讀未來結果", TEAL),
        "`observation` 是 runner 每個固定 step 提供的唯讀快照，實際使用窗口、急件、期限、距上次傳送步數、傳送模式、品質、穩定步數與佇列數。此頁是中文說明，不可貼回 `/home/u24/lora-energy-lab/student_policy.py`；沒有 file／block／value 變更。下一步由 `choose_action(observation)` 回傳合法 action，再從 result／replay 讀 state／event、packet／delivery、service 與 endpoint J。",
        (RED_PALE, RED, RED),
    )


def slide_024b(prs: Presentation):
    return _beginner_page(
        prs, "判斷順序如何產生動作",
        "這一步把合法動作順序翻成白話；鏡像可讀，不能貼回程式",
        [
            ("① 窗口關閉", "不能送出，先休息\n`SLEEP`", BLUE_PALE, BLUE, 20, 18),
            ("② 期限逼近", "急件先送出\n`SEND_URGENT`", RED_PALE, RED, 20, 18),
            ("③ 節奏／品質", "間隔未到休息；品質未穩等待\n`REST_DURING_GAP`／`WAIT`", GOLD_PALE, GOLD, 18, 16),
            ("④ 佇列／保底", "夠量批次，否則單筆；無事等待\n`FLUSH_BATCH`／`SEND_ONE`／`WAIT`", TEAL_PALE, TEAL, 17, 16),
        ],
        "先看窗口，再看期限、間隔、品質與佇列；順序會改變動作",
        ("action", "合法動作", "policy API", "把條件轉成 runner 可執行的動作", "enum", "下一步從 state／packet／service／endpoint J 讀後果", TEAL),
        "這一頁是 `choose_action()` 的 read-only 中文註解鏡像，不可貼回程式。實際順序是窗口關閉回傳 `SLEEP`；urgent deadline 回傳 `SEND_URGENT`；pacing gap 回傳 `REST_DURING_GAP`；品質 hold 未滿足回傳 `WAIT`；quality ready 且 queue 足量回傳 `FLUSH_BATCH`，否則 `SEND_ONE`；所有條件都不要求傳送時回傳 `WAIT`。沒有 file／block／value 變更；下一步 run 後按 service → packet／delivery → state／event → endpoint J → efficiency 判讀。",
        (RED_PALE, RED, RED),
    )


def slide_025a(prs: Presentation):
    return _beginner_page(
        prs, "Lab A：空檔休息與等待",
        "Lab A 的目的：比較空檔休息與清醒等待，找服務／能量取捨",
        [
            ("① 檔案／區段", "`student_policy.py`\n`lab-a-pace-rest`", BLUE_PALE, BLUE, 18, 17),
            ("② Before（唯一變更）", "空檔休息\n`SLEEP`", GOLD_PALE, GOLD, 17, 22),
            ("③ After（唯一變更）", "清醒等待\n`WAIT`", TEAL_PALE, TEAL, 17, 22),
            ("④ 固定／觀察", "gap=2；A/B/C、runner 不動\n看 service／state／J", PURPLE_PALE, PURPLE, 18, 16),
        ],
        "只改空檔動作 → 改變休息／清醒／喚醒 → 影響服務與端點焦耳",
        ("REST_DURING_GAP", "空檔動作", "Lab A marked block", "控制送出間隔中的狀態", "action", "比較 SLEEP 與 WAIT 對狀態、服務、端點 J 的影響", GOLD),
        "Lab A 的 exact file 是 `/home/u24/lora-energy-lab/student_policy.py`，exact block 是 `lab-a-pace-rest`；mirror 顯示 Before `SLEEP`、After `WAIT`，只改一個 value。`PACE_GAP_STEPS = 2`、Lab B/C、runner、scenario、schemas、中文註解與函式都不變。機制是 action 改變 sleep／awake-idle／wake 分布。下一步 baseline → candidate with freeze → hidden frozen policy；每次保留 `result.json` + `endpoint-replay.json`，網站先看 service，再看 packet／state／endpoint J／efficiency。較低 J 但 service fail 仍是 trade-off。",
        (GOLD_PALE, GOLD, NAVY),
    )


def slide_025b(prs: Presentation):
    return _beginner_page(
        prs, "Lab C：批次與期限餘裕",
        "Lab C 的目的：比較批次門檻與期限餘裕，找傳送時機代價",
        [
            ("① 檔案／區段", "`student_policy.py`\n`lab-c-batch-urgent`", BLUE_PALE, BLUE, 18, 17),
            ("② Before（唯一變更）", "期限餘裕\n`20 秒`", GOLD_PALE, GOLD, 17, 22),
            ("③ Candidate／Revision", "候選 `5 秒`\n修訂 `30 秒`", RED_PALE, RED, 18, 21),
            ("④ 固定／觀察", "`BATCH_SIZE=3`；A/B、runner 不動\n看 service／deadline／delivery／J", PURPLE_PALE, PURPLE, 16, 16),
        ],
        "只改期限餘裕 → 改變緊急分支提早量 → 影響期限與送達",
        ("URGENT_MARGIN_S", "緊急餘裕秒數", "Lab C marked block", "提供期限分支的比較值", "秒", "觀察緊急動作、送達、期限與端點 J", RED),
        "Lab C 的 exact file 是 `/home/u24/lora-energy-lab/student_policy.py`，exact block 是 `lab-c-batch-urgent`；Before 20、candidate 5、revision 30，只有這個 value 變更。`BATCH_SIZE = 3`、frozen A/B、runner、scenario、schemas、中文註解與函式均不變。機制是 threshold 改變 `SEND_URGENT` 提前量。下一步 baseline → candidate → revision with freeze → withheld surprise；保存 result／replay pair，先看 service，再看 deadline／delivery／state／endpoint J／efficiency。",
        (RED_PALE, RED, RED),
    )


def slide_026a(prs: Presentation):
    return _beginner_page(
        prs, "Lab B：穩定觀察如何影響動作",
        "這一步用 Lab B 示範品質條件：穩定次數會改變 action 何時出現",
        [
            ("① 檔案／區段", "`student_policy.py`\n`lab-b-enter-exit-hold`", BLUE_PALE, BLUE, 18, 17),
            ("② Before（唯一變更）", "穩定 `2 步`\n才保持", GOLD_PALE, GOLD, 17, 22),
            ("③ After（唯一變更）", "穩定 `1 步`\n即可", TEAL_PALE, TEAL, 17, 22),
            ("④ 固定／觀察", "ENTER=2、EXIT=1；A/C、runner 不動\n看 transition／service／J", PURPLE_PALE, PURPLE, 16, 16),
        ],
        "只改穩定步數 → 改變可送模式出現時間 → 影響送達與服務",
        ("STABLE_STEPS", "穩定步數", "Lab B marked block", "決定可送模式何時成立", "steps", "比較 transition、送達、服務與端點 J", TEAL),
        "Lab B 的 exact file 是 `/home/u24/lora-energy-lab/student_policy.py`，exact block 是 `lab-b-enter-exit-hold`；Before `STABLE_STEPS = 2`，After 只改 1。`ENTER_QUALITY = 2`、`EXIT_QUALITY = 1`、frozen A/C、runner、scenario、schemas、中文註解與函式均不變。機制是 stable observations 改變 quality_ready，因而改變 send-ready transition timing。下一步 Trace A baseline → candidate with freeze → withheld Trace B；保留 result／replay pair，先看 service，再看 packet／state／endpoint J／efficiency。",
        (TEAL_PALE, TEAL, NAVY),
    )


def slide_026b(prs: Presentation):
    return _beginner_page(
        prs, "動作之後要讀哪些事件",
        "這一步把動作接到證據：先看到狀態事件，再判斷服務與成本",
        [
            ("① WAIT 後", "清醒閒置時間增加\n改看 awake-idle", TEAL_PALE, TEAL, 20, 18),
            ("② SLEEP 後", "休息再喚醒\n改看 sleep／wake", BLUE_PALE, BLUE, 20, 18),
            ("③ SEND 後", "處理、發射、接收\n再看 packet 是否送達", GOLD_PALE, GOLD, 19, 17),
            ("④ 結果", "服務成敗與端點焦耳\n`endpoint_energy_j`", PURPLE_PALE, PURPLE, 19, 17),
        ],
        "動作 → 狀態／封包事件 → 服務結果 → 端點焦耳；不能倒過來",
        ("endpoint_energy_j", "端點能量", "端點無線與處理模型", "累積狀態與轉換能量", "焦耳", "只屬端點證據，不等於系統能量", PURPLE),
        "`WAIT` 會留下 awake-idle，`SLEEP` 會留下 sleep 並在恢復時產生 wake；send action 沿 process、TX、RX 形成 attempt／retry／delivery／expiry。下一步依 service → packet → state → endpoint J → efficiency 讀取。`endpoint_energy_j` 只屬 endpoint layer，與 LEO／system／canonical energy 分開；較低 J 且 service fail 只能說明 trade-off。",
        (RED_PALE, RED, RED),
    )


def slide_027a(prs: Presentation):
    return _beginner_page(
        prs, "縮排與回傳決定程式走哪條路",
        "這一步只讀一小段程式：看縮排如何選分支，再把動作交回執行器",
        [
            ("① 先看意思", "窗口關閉就不送\n`if ...: return SLEEP`", BLUE_PALE, BLUE, 19, 17),
            ("② 縮排做什麼", "把條件放進正確分支\n只修 active block", TEAL_PALE, TEAL, 19, 17),
            ("③ 先檢查", "執行 `py_compile`\n只檢查語法、不改檔", GOLD_PALE, GOLD, 18, 17),
            ("④ 下一步／看什麼", "通過後跑 baseline\n讀 result／replay pair", PURPLE_PALE, PURPLE, 18, 17),
        ],
        "縮排 → 分支 → 回傳 → runner 事件；語法檢查不改內容",
        ("return", "回傳", "policy function", "把一個合法動作交給執行器", "動作值", "語法通過後再看 result／replay 的實際事件", GOLD),
        "目前 `/home/u24/lora-energy-lab/student_policy.py` 的中文註解已說明每個 branch。這一頁只顯示小段可讀鏡像：窗口關閉時由 `if ...` 進入 `return SLEEP`；縮排決定 statement 屬於哪個條件。`.venv/bin/python -m py_compile student_policy.py` 只檢查語法，不改檔；若錯，只修目前 active marked block，中文註解、函式、runner、scenario 與 schemas 不動。下一步才執行 P027b baseline，讀成對 result／replay。",
        (PURPLE_PALE, PURPLE, NAVY),
    )


def slide_027b(prs: Presentation):
    return _beginner_page(
        prs, "先建立基準結果再比較",
        "這一步先保存沒有改動的基準，下一步才比較候選值",
        [
            ("① 固定輸入", "packaged policy＋固定情境\n不讀未來結果", BLUE_PALE, BLUE, 20, 18),
            ("② 執行命令", "只跑 baseline\n`bash course.sh run --lab A --case baseline`", TEAL_PALE, TEAL, 17, 16),
            ("③ 產生結果", "兩個 JSON\nresult.json／endpoint-replay.json\n原始程式不回寫", GOLD_PALE, GOLD, 18, 17),
            ("④ 下一步／看什麼", "確認 `status: OK`、路徑\n先 service，再成本", PURPLE_PALE, PURPLE, 17, 16),
        ],
        "先保存基準，再凍結規則，最後比較候選；結果不能回頭改 block",
        ("result_path", "結果檔路徑", "命令輸出", "定位 result.json", "路徑文字", "網站只匯入 result.json；endpoint replay 只做配對，runner 不是網站", PURPLE),
        "POSIX exact command 是 `bash course.sh run --lab A --case baseline`；Windows 是 `course.cmd run --lab A --case baseline`。它讀 packaged `student_policy.py` 與 fixed scenario，只在 run directory 寫 `result.json`、`endpoint-replay.json`，不修改 policy source、marked blocks、runner、scenario 或 schemas。下一步確認 `status: OK`、run identity、`result_path`，網站先看 service，再看 packet／state／endpoint J／efficiency；Lab A／B／C 後續依各自 baseline → candidate → freeze → withheld 順序比較。",
        (GOLD_PALE, GOLD, NAVY),
    )


def build() -> None:
    authority = load_authority()
    write_source_snapshot(authority)
    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)
    builders = [
        slide_019, slide_020, slide_021a, slide_021b, slide_021c,
        slide_022, slide_023, slide_024a, slide_024b, slide_025a,
        slide_025b, slide_026a, slide_026b, slide_027a, slide_027b,
    ]
    for builder in builders:
        builder(prs)
    prs.core_properties.title = "LoRaEnergySim + LEO｜Part A P019–P027b"
    prs.core_properties.subject = "WSL setup、READY receipt、bounded policy、Lab A baseline handoff"
    prs.core_properties.author = "C-120 deck lane P019–P027b"
    prs.core_properties.comments = "Editable section built from educate slideLayout2.xml"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    overlay_template_parts(OUTPUT)
    print(json.dumps({
        "output": str(OUTPUT),
        "slides": len(prs.slides),
        "template": str(TEMPLATE),
        "template_shell": "slideLayout2.xml",
        "source_snapshot": str(SOURCE_SNAPSHOT),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
