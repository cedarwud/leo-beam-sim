#!/usr/bin/env python3
"""Geometry guard for the template-safe authored area."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT.parent / "latest" / "LoRaEnergySim-LEO-ALT-PART-A-V2-P019-P027B-REVIEW.pptx"
OUT = ROOT / "qa" / "geometry-qa.json"

SAFE = {"left": 0.718057, "top": 1.05, "right": 13.004168, "bottom": 6.627083}
LOGO_ZONES = [
    (10.799630, 0.0, 13.277154, 0.353933),
    (11.166667, 0.0, 13.333334, 0.309524),
]


def rect(shape):
    return {
        "left": shape.left / 914400,
        "top": shape.top / 914400,
        "right": (shape.left + shape.width) / 914400,
        "bottom": (shape.top + shape.height) / 914400,
    }


def intersects(a, b):
    return a["left"] < b[2] and a["right"] > b[0] and a["top"] < b[3] and a["bottom"] > b[1]


def qa():
    prs = Presentation(str(PPTX))
    failures = []
    slides = []
    for index, slide in enumerate(prs.slides, 1):
        authored = [shape for shape in slide.shapes if not shape.is_placeholder and shape.name != "Native layout2 title"]
        out_of_bounds = []
        logo_collisions = []
        text_count = 0
        visual_count = 0
        for shape in authored:
            box = rect(shape)
            if box["left"] < SAFE["left"] - 0.001 or box["top"] < SAFE["top"] - 0.001 or box["right"] > SAFE["right"] + 0.001 or box["bottom"] > SAFE["bottom"] + 0.001:
                out_of_bounds.append({"name": shape.name, "rect": box})
            if getattr(shape, "has_text_frame", False) and getattr(shape, "text", "").strip():
                text_count += 1
            else:
                visual_count += 1
            for zone in LOGO_ZONES:
                if intersects(box, zone):
                    logo_collisions.append({"name": shape.name, "rect": box})
        if out_of_bounds:
            failures.append(f"slide {index}: authored shape outside safe bounds")
        if logo_collisions:
            failures.append(f"slide {index}: authored shape collides with protected logo zone")
        if text_count == 0 or visual_count == 0:
            failures.append(f"slide {index}: missing text or dominant visual element")
        slides.append({
            "slide": index,
            "title": slide.shapes.title.text if slide.shapes.title else "",
            "authored_shapes": len(authored),
            "text_shapes": text_count,
            "visual_shapes": visual_count,
            "out_of_bounds": out_of_bounds,
            "logo_collisions": logo_collisions,
        })
    report = {
        "safe_bounds": SAFE,
        "protected_logo_zones": LOGO_ZONES,
        "slide_count": len(prs.slides),
        "slides": slides,
        "all_checks_pass": not failures,
        "failures": failures,
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(qa())
