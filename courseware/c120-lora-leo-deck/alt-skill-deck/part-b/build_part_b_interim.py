#!/usr/bin/env python3
"""Build the C-120 Part B interim deck using only educate slideLayout2.

The authoritative Part B content/drawing helpers remain read-only in the
courseware root.  This wrapper owns the alternative-skill output directory and
overrides only the template/output route plus slide construction shell.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET

from pptx import Presentation


OWNED = Path(__file__).resolve().parent
REPO = OWNED.parents[3]
SOURCE_BUILDER = REPO / "courseware/c120-lora-leo-deck/build-direct-teaching-part-b.py"
TEMPLATE = Path("/home/u24/pptx-wrap/assets/templates/educate.pptx")
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def load_source_module():
    spec = importlib.util.spec_from_file_location("c120_part_b_source_builder", SOURCE_BUILDER)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load source builder: {SOURCE_BUILDER}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def configure(module) -> None:
    module.ROOT = REPO / "courseware/c120-lora-leo-deck"
    module.TEMPLATE = TEMPLATE
    module.PROJECT = OWNED / "build"
    module.SOURCE_COPY = module.PROJECT / "sources/educate.pptx"
    # The PPTX is the one explicitly shared review handoff at alt-skill-deck
    # root; every script, render, report, and source copy remains under part-b.
    module.EXPORT = OWNED.parent / "LoRaEnergySim-LEO-ALT-PART-B-REVIEW.pptx"
    module.RENDER_DIR = module.PROJECT / "renders"
    module.RENDER_SOURCE = module.PROJECT / "validation/.render-source.pptx"
    module.SOURCE_SCRIPT = module.ROOT / "teaching-rewrite/part-b-visible-content.md"
    module.EQUATION_TOOL = module.ROOT / "insert-native-equations.py"


def _remove_named_shapes(slide, names: set[str]) -> None:
    """Remove only the authored shapes targeted by a local source-page repair."""
    for shape in list(slide.shapes):
        if shape.name in names:
            shape._element.getparent().remove(shape._element)


def repair_source_page_layout(module, slide, page) -> None:
    """Repair two local source-page wraps without touching the read-only builder.

    Part B page P030 (source slide 3; merged logical slide 36) uses two lower
    cards whose heading/body boxes are too close for the native template
    renderer.  Recompose only those authored card shapes with explicit
    vertical separation.  Part B page P036 (source slide 9; merged logical
    slide 42) keeps the identity card content but inserts deliberate line
    breaks so ``provenance`` remains a complete word.
    """
    if page.number == 30 and page.kind == "packet":
        _remove_named_shapes(
            slide,
            {
                "Packet ledger",
                "Packet ledger heading",
                "Packet ledger body",
                "Packet service",
                "Packet service heading",
                "Packet service body",
            },
        )
        # The stage row ends at y=2.92.  These taller cards start lower and
        # reserve separate heading/body bands, leaving the inherited footer
        # and claim boundary untouched.
        module.box(slide, 0.92, 3.24, 5.34, 1.58, fill=module.PALE_BLUE, line=module.BLUE, name="Packet ledger")
        module.text_box(
            slide, 1.10, 3.38, 4.98, 0.28, "PACKET LEDGER｜封包紀錄", 18,
            bold=True, color=module.BLUE, align=module.PP_ALIGN.CENTER,
            valign=module.MSO_ANCHOR.MIDDLE, name="Packet ledger heading",
            italic_code=False, margins=(0.01, 0.0, 0.01, 0.0),
        )
        module.text_box(
            slide, 1.12, 3.82, 4.94, 0.76,
            "attempted（嘗試次數）≠ delivered（交付 bit）\ncollision／retry／expired = packet outcome",
            18, color=module.INK, align=module.PP_ALIGN.CENTER,
            valign=module.MSO_ANCHOR.MIDDLE, name="Packet ledger body",
            margins=(0.01, 0.0, 0.01, 0.0),
        )
        module.box(slide, 6.66, 3.24, 5.46, 1.58, fill=module.PALE_GOLD, line=module.GOLD, name="Packet service")
        module.text_box(
            slide, 6.84, 3.38, 5.10, 0.28, "SERVICE VERDICT｜服務判定", 18,
            bold=True, color=module.GOLD, align=module.PP_ALIGN.CENTER,
            valign=module.MSO_ANCHOR.MIDDLE, name="Packet service heading",
            italic_code=False, margins=(0.01, 0.0, 0.01, 0.0),
        )
        module.text_box(
            slide, 6.88, 3.82, 5.02, 0.76,
            "deadline（期限）／freshness（新鮮度）／required packet\n共同定義 service_pass（服務布林狀態）",
            18, color=module.INK, align=module.PP_ALIGN.CENTER,
            valign=module.MSO_ANCHOR.MIDDLE, name="Packet service body",
            margins=(0.01, 0.0, 0.01, 0.0),
        )
    elif page.number == 36 and page.kind == "identity":
        for shape in slide.shapes:
            if shape.name == "Identity replay body":
                module.write_text(
                    shape.text_frame,
                    "REPLAY：同 run events 的\nstate／queue／packet\nWORKBOOK：baseline／candidate／\nfreeze／withheld\n一致性：run_id、policy identity、\nunits、provenance",
                    18,
                    color=module.INK,
                    align=module.PP_ALIGN.CENTER,
                    valign=module.MSO_ANCHOR.MIDDLE,
                    italic_code=True,
                    margins=(0.01, 0.0, 0.01, 0.0),
                )
                break


def build_layout2_only(module):
    """Run the existing content/drawing pipeline with layout2 for every slide."""
    module.validate_pages()
    module.ensure_dirs()
    prs = Presentation(str(module.TEMPLATE))
    module.remove_all_slides(prs)
    # python-pptx indexes the source slide 2 layout as 1.  No other layout is
    # ever used; the template master/background/footer stay inherited.
    source_layout2 = prs.slide_layouts[1]
    for page in module.PAGES:
        slide = prs.slides.add_slide(source_layout2)
        module.compose(slide, page)
        repair_source_page_layout(module, slide, page)
    working = module.PROJECT / "validation/python-pptx-working.pptx"
    prs.save(str(working))
    shutil.copy2(working, module.EXPORT)
    shutil.copy2(module.EXPORT, module.RENDER_SOURCE)
    module.overlay_template_parts(module.EXPORT)
    equations = module.insert_equations()
    add_render_safe_formula_transcriptions(module)
    rewrite_notes(module)
    report = module.qa_package()
    return report, {"equations": equations}


def formal_note(page) -> str:
    """Turn the structured page contract into directly readable narration."""
    def sentence(value: str) -> str:
        return str(value).strip().rstrip("。")

    return (
        f"P{page.number:03d}。{sentence(page.title)}。"
        f"目的：{sentence(page.purpose)}。"
        f"作用機制：{sentence(page.mechanism)}。"
        f"觀察與判讀：{sentence(page.expected)}。"
        f"操作受阻時的復原：{sentence(page.recovery)}。"
    )


def rewrite_notes(module) -> None:
    """Embed complete formal notes while preserving notes/master relationships."""
    with zipfile.ZipFile(module.EXPORT, "r") as archive:
        infos = [info for info in archive.infolist()]
        data = {info.filename: archive.read(info.filename) for info in infos}
    notes_names = sorted(
        (name for name in data if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")),
        key=lambda name: int(re.search(r"notesSlide(\d+)\.xml$", name).group(1)),
    )
    if len(notes_names) != len(module.PAGES):
        raise RuntimeError(f"notes count {len(notes_names)} does not match pages {len(module.PAGES)}")
    for page, name in zip(module.PAGES, notes_names):
        root = ET.fromstring(data[name])
        body = None
        for shape in root.findall(f".//{{{P_NS}}}sp"):
            ph = shape.find(f"./{{{P_NS}}}nvSpPr/{{{P_NS}}}nvPr/{{{P_NS}}}ph")
            if ph is not None and ph.get("type") == "body":
                body = shape
                break
        if body is None:
            raise RuntimeError(f"notes body placeholder missing: {name}")
        tx_body = body.find(f"{{{P_NS}}}txBody")
        if tx_body is None:
            raise RuntimeError(f"notes text body missing: {name}")
        paragraphs = tx_body.findall(f"{{{A_NS}}}p")
        if not paragraphs:
            paragraphs = [ET.SubElement(tx_body, f"{{{A_NS}}}p")]
        paragraph = paragraphs[0]
        for child in list(paragraph):
            paragraph.remove(child)
        run = ET.SubElement(paragraph, f"{{{A_NS}}}r")
        text_node = ET.SubElement(run, f"{{{A_NS}}}t")
        text_node.text = formal_note(page)
        data[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    temp = module.EXPORT.with_suffix(".notes.pptx")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as archive:
        for info in infos:
            archive.writestr(info, data[info.filename])
    temp.replace(module.EXPORT)


def add_render_safe_formula_transcriptions(module) -> None:
    """Keep native OMML while making LibreOffice's rendered view legible.

    LibreOffice currently mishandles PowerPoint's ``a14:m`` native choice.
    The native shape remains in the package with a one-unit geometry footprint
    for that renderer; an editable DrawingML text transcription carries the
    same formula visibly.  No picture or raster fallback is introduced.
    """
    formulas = {6: ("E_endpoint = Σ P_s t_s", "E_endpoint"), 7: ("η_E = D_delivered / E_endpoint", "eta_E")}
    with zipfile.ZipFile(module.EXPORT, "r") as archive:
        infos = [info for info in archive.infolist()]
        data = {info.filename: archive.read(info.filename) for info in infos}
    ET.register_namespace("a", A_NS)
    ET.register_namespace("p", P_NS)
    ET.register_namespace("m", M_NS)
    for logical_slide, (formula, key) in formulas.items():
        name = f"ppt/slides/slide{logical_slide}.xml"
        root = ET.fromstring(data[name])
        sp_tree = root.find(f"./{{{P_NS}}}cSld/{{{P_NS}}}spTree")
        if sp_tree is None:
            raise RuntimeError(f"formula slide has no shape tree: {logical_slide}")
        native = None
        used_ids = []
        for shape in sp_tree.findall(f"{{{P_NS}}}sp"):
            c_nv_pr = shape.find(f"./{{{P_NS}}}nvSpPr/{{{P_NS}}}cNvPr")
            if c_nv_pr is None:
                continue
            value = c_nv_pr.get("id")
            if value and value.isdigit():
                used_ids.append(int(value))
            if c_nv_pr.get("name", "").startswith("Native Office Math"):
                native = shape
        if native is None:
            raise RuntimeError(f"native formula shape missing: {logical_slide}")
        native_nv = native.find(f"./{{{P_NS}}}nvSpPr/{{{P_NS}}}cNvSpPr")
        if native_nv is None:
            raise RuntimeError(f"native formula locks missing: {logical_slide}")
        native_nv.set("hidden", "1")
        xfrm = native.find(f"./{{{P_NS}}}spPr/{{{A_NS}}}xfrm")
        if xfrm is None:
            raise RuntimeError(f"native formula geometry missing: {logical_slide}")
        original_xfrm_children = [ET.fromstring(ET.tostring(child, encoding="utf-8")) for child in list(xfrm)]
        native_off = xfrm.find(f"{{{A_NS}}}off")
        native_ext = xfrm.find(f"{{{A_NS}}}ext")
        if native_off is None or native_ext is None:
            raise RuntimeError(f"native formula transform incomplete: {logical_slide}")
        native_off.set("x", "0")
        native_off.set("y", "0")
        native_ext.set("cx", "1")
        native_ext.set("cy", "1")
        new_id = str(max(used_ids, default=0) + 1)
        visual = ET.Element(f"{{{P_NS}}}sp")
        nv_sp_pr = ET.SubElement(visual, f"{{{P_NS}}}nvSpPr")
        ET.SubElement(nv_sp_pr, f"{{{P_NS}}}cNvPr", {"id": new_id, "name": f"Editable formula transcription {key}"})
        ET.SubElement(nv_sp_pr, f"{{{P_NS}}}cNvSpPr", {"txBox": "1"})
        ET.SubElement(nv_sp_pr, f"{{{P_NS}}}nvPr")
        sp_pr = ET.SubElement(visual, f"{{{P_NS}}}spPr")
        sp_xfrm = ET.SubElement(sp_pr, f"{{{A_NS}}}xfrm")
        for child in original_xfrm_children:
            sp_xfrm.append(child)
        ET.SubElement(sp_pr, f"{{{A_NS}}}prstGeom", {"prst": "rect"})
        ET.SubElement(sp_pr, f"{{{A_NS}}}noFill")
        line_node = ET.SubElement(sp_pr, f"{{{A_NS}}}ln")
        ET.SubElement(line_node, f"{{{A_NS}}}noFill")
        tx_body = ET.SubElement(visual, f"{{{P_NS}}}txBody")
        ET.SubElement(tx_body, f"{{{A_NS}}}bodyPr", {"anchor": "ctr", "wrap": "none"})
        ET.SubElement(tx_body, f"{{{A_NS}}}lstStyle")
        para = ET.SubElement(tx_body, f"{{{A_NS}}}p")
        ET.SubElement(para, f"{{{A_NS}}}pPr", {"algn": "ctr"})
        run = ET.SubElement(para, f"{{{A_NS}}}r")
        ET.SubElement(run, f"{{{A_NS}}}rPr", {"lang": "en-US", "sz": "3000", "i": "1"})
        ET.SubElement(run, f"{{{A_NS}}}t").text = formula
        sp_tree.append(visual)
        data[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    temp = module.EXPORT.with_suffix(".formula-visible.pptx")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as archive:
        for info in infos:
            archive.writestr(info, data[info.filename])
    temp.replace(module.EXPORT)


def render_and_contact_local(module) -> dict:
    """Render with an isolated LibreOffice profile and save owned evidence."""
    render_dir = module.RENDER_DIR
    render_dir.mkdir(parents=True, exist_ok=True)
    for old in render_dir.glob("slide-*.png"):
        old.unlink()
    contact_dir = render_dir / "contact-sheets"
    if contact_dir.exists():
        shutil.rmtree(contact_dir)
    contact_dir.mkdir(parents=True, exist_ok=True)

    profile = Path("/tmp/c120-alt-part-b-libreoffice")
    profile_uri = f"file://{profile}"
    env = os.environ.copy()
    env["SAL_USE_VCLPLUGIN"] = "svp"
    command = [
        "libreoffice", "--headless", f"-env:UserInstallation={profile_uri}",
        "--convert-to", "pdf", "--outdir", str(render_dir), str(module.EXPORT),
    ]
    result = subprocess.run(command, cwd=str(REPO), env=env, text=True, capture_output=True, timeout=180, check=False)
    produced = render_dir / f"{module.EXPORT.stem}.pdf"
    final_pdf = render_dir / "direct-teaching-part-b-final.pdf"
    if produced.exists():
        if final_pdf.exists():
            final_pdf.unlink()
        produced.replace(final_pdf)
    pdf = render_dir / "direct-teaching-part-b.pdf"
    if final_pdf.exists():
        shutil.copy2(final_pdf, pdf)

    images = []
    if pdf.exists():
        subprocess.run(["pdftoppm", "-png", "-r", "144", str(pdf), str(render_dir / "slide")], cwd=str(REPO), check=True, timeout=180)
        images = sorted(render_dir.glob("slide-*.png"), key=lambda p: int(re.search(r"(\d+)$", p.stem).group(1)))
        from PIL import Image, ImageDraw
        for start in range(0, len(images), 9):
            subset = images[start:start + 9]
            tw, th = 360, 203
            sheet = Image.new("RGB", (tw * 3, (th + 24) * 3), "white")
            draw = ImageDraw.Draw(sheet)
            for j, path in enumerate(subset):
                with Image.open(path) as image:
                    image = image.convert("RGB")
                    image.thumbnail((tw, th))
                    x = (j % 3) * tw
                    y = (j // 3) * (th + 24)
                    sheet.paste(image, (x, y))
                    draw.text((x + 6, y + th + 4), f"P{start + j + 28:03d}", fill="black")
            sheet.save(contact_dir / f"contact-{start + 28:03d}-{start + len(subset) + 27:03d}.jpg", quality=90)
    visual = {
        "pdf": str(pdf),
        "chosen_render_source": str(final_pdf),
        "slides_rendered": len(images),
        "conversion": {"returncode": result.returncode, "stdout": result.stdout[-1000:], "stderr": result.stderr[-1000:]},
        "contact_sheets": [str(path) for path in sorted(contact_dir.glob("contact-*.jpg"))],
        "page_by_page_human_inspection": "PENDING",
        "fix_and_rerender_pass": "PENDING",
    }
    (module.PROJECT / "validation/visual-qa.json").write_text(json.dumps(visual, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return visual


def write_machine_manifest(module, equation_data: dict) -> None:
    donor = {
        "model_bridge": "donor-analysis/donor-insertion-map.md:e2 14,16/17,19-21",
        "fair_baseline": "donor-analysis/donor-insertion-map.md:e2 20,36-37",
        "lab_b_branch": "donor-analysis/donor-insertion-map.md:e2 11-12",
        "provenance": "donor-analysis/donor-insertion-map.md; topic-level redraw only",
        "disposition": "adopt-after-rewrite; no donor bytes, screenshots, KPIs, formulas, or notes copied",
    }
    rows = []
    for index, page in enumerate(module.PAGES, start=1):
        note = formal_note(page)
        rows.append({
            "order": index,
            "source_page": f"P{page.number:03d}",
            "title": page.title,
            "body": {
                "evidence": page.evidence,
                "purpose": page.purpose,
                "mechanism": page.mechanism,
                "expected": page.expected,
                "recovery": page.recovery,
                "context": page.context,
                "command": page.command,
                "visual": page.visual,
            },
            "layout": page.kind,
            "notes": note,
            "formulas": (["E_endpoint = Σ P_s t_s"] if page.kind == "formula-energy" else
                         ["η_E = D_delivered / E_endpoint"] if page.kind == "formula-eff" else []),
            "donor-source": donor if page.number in {32, 33, 34, 37, 51, 60} else {},
        })
    manifest = {
        "schema": "c120-alt-skill-part-b-slides-v1",
        "template": str(TEMPLATE),
        "template_sha256": module.hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        "layout_contract": {
            "source_slide": 2,
            "source_layout": "ppt/slideLayouts/slideLayout2.xml",
            "all_slides_use_layout2": True,
            "background": "inherited/unset",
            "footer_logo_page_number": "inherited",
        },
        "typography": {
            "title_pt": 28,
            "body_baseline_pt": 24,
            "minimum_pt": 16,
            "cjk": "標楷體/DFKai-SB",
            "latin": "Times New Roman",
            "formula": "native Office Math / OMML",
        },
        "claim_boundary": module.COMMON_BOUNDARY,
        "slides": rows,
        "equations": equation_data.get("equations", []),
    }
    (OWNED / "slides.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    notes_md = ["# Part B embedded speaker notes", ""]
    for page in module.PAGES:
        notes_md.extend([f"## P{page.number:03d} — {page.title}", "", formal_note(page), ""])
    (OWNED / "speaker-notes.md").write_text("\n".join(notes_md), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-render", action="store_true")
    args = parser.parse_args(argv)
    module = load_source_module()
    configure(module)
    report, equation_data = build_layout2_only(module)
    module.write_sources(equation_data.get("equations", []))
    write_machine_manifest(module, equation_data)
    if args.skip_render:
        visual = {"slides_rendered": 0, "skip": True}
        (module.PROJECT / "validation/visual-qa.json").write_text(json.dumps(visual, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    else:
        visual = render_and_contact_local(module)
    module.write_readme(report, visual, equation_data)
    print(json.dumps({
        "output": str(module.EXPORT),
        "slides": len(module.PAGES),
        "qa_status": report["status"],
        "omml_count": report["omml_count"],
        "rendered": visual.get("slides_rendered", 0),
        "manifest": str(OWNED / "slides.json"),
    }, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
