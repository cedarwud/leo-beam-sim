#!/usr/bin/env python3
"""Geometry sanity checks used when a renderer is unavailable."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT.parent / "LoRaEnergySim-LEO-ALT-OPENING-V2-REVIEW.pptx"
REPORT = ROOT / "qa" / "geometry-qa.json"
EMU = 914400
FOOTER_TOP = 6.6858
SAFE = {"left": 0.70, "right": 13.04, "top": 1.02, "bottom": 6.66}


def box(shape):
    return {
        "left": shape.left / EMU,
        "top": shape.top / EMU,
        "right": (shape.left + shape.width) / EMU,
        "bottom": (shape.top + shape.height) / EMU,
    }


def contains(a, b, eps=1e-5):
    return (
        a["left"] - eps <= b["left"] and a["top"] - eps <= b["top"]
        and a["right"] + eps >= b["right"] and a["bottom"] + eps >= b["bottom"]
    )


def overlaps(a, b, eps=1e-5):
    return not (
        a["right"] <= b["left"] + eps or b["right"] <= a["left"] + eps
        or a["bottom"] <= b["top"] + eps or b["bottom"] <= a["top"] + eps
    )


def main():
    prs = Presentation(str(OUTPUT))
    report = {"output": str(OUTPUT), "slides": [], "failures": []}
    for slide_index, slide in enumerate(prs.slides, start=1):
        authored = [shape for shape in slide.shapes if not shape.is_placeholder]
        shapes = []
        for shape in authored:
            b = box(shape)
            entry = {"name": shape.name, **b}
            shapes.append(entry)
            if b["left"] < SAFE["left"] or b["right"] > SAFE["right"]:
                report["failures"].append(
                    f"slide {slide_index} {shape.name}: horizontal bounds {b}"
                )
            if b["top"] < SAFE["top"] or b["bottom"] > FOOTER_TOP:
                report["failures"].append(
                    f"slide {slide_index} {shape.name}: vertical/footer bounds {b}"
                )
        suspicious = []
        for index, first in enumerate(shapes):
            for second in shapes[index + 1:]:
                a = {key: first[key] for key in ("left", "top", "right", "bottom")}
                b = {key: second[key] for key in ("left", "top", "right", "bottom")}
                if overlaps(a, b) and not contains(a, b) and not contains(b, a):
                    suspicious.append((first["name"], second["name"]))
        report["slides"].append({
            "slide": slide_index,
            "authored_shape_count": len(authored),
            "suspicious_noncontainment_overlaps": suspicious,
        })
    report["status"] = "PASS" if not report["failures"] else "FAIL"
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
