#!/usr/bin/env python3
"""Extract reproducible OOXML/text/notes/media/layout evidence from donor decks.

This script is intentionally read-only with respect to the donor files.  All
outputs are written beneath the audit directory supplied with ``--output``.
"""

from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import json
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


DECKS = [
    (
        "e2",
        Path("/home/u24/papers/beamshift/e2.pptx"),
        "content-and-notes-donor",
    ),
    (
        "delivery-variant",
        Path(
            "/home/u24/papers/beamshift/.scratch/teaching-course/v3-ee/delivery/"
            "satellite-energy-course-combined-handoff-v1.pptx"
        ),
        "visual-and-provenance-comparison-only",
    ),
]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean_newlines(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n")


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", clean_newlines(value))
    value = value.casefold()
    return re.sub(r"\s+", " ", value).strip()


def enum_name(value: Any) -> str:
    return getattr(value, "name", str(value))


def shape_text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return clean_newlines(shape.text)


def shape_bbox(shape: Any) -> dict[str, int | float | None]:
    result: dict[str, int | float | None] = {}
    for field in ("left", "top", "width", "height", "rotation"):
        try:
            value = getattr(shape, field)
        except (AttributeError, ValueError):
            result[field] = None
        else:
            result[field] = int(value) if field != "rotation" else float(value)
    return result


def shape_record(shape: Any) -> dict[str, Any]:
    text = shape_text(shape)
    record: dict[str, Any] = {
        "shape_id": int(getattr(shape, "shape_id", -1)),
        "name": str(getattr(shape, "name", "")),
        "shape_type": enum_name(getattr(shape, "shape_type", "")),
        "bbox_emu": shape_bbox(shape),
        "text": text,
        "text_sha256": sha256_bytes(text.encode("utf-8")) if text else None,
        "has_text_frame": bool(getattr(shape, "has_text_frame", False)),
        "is_placeholder": bool(getattr(shape, "is_placeholder", False)),
    }
    if record["is_placeholder"]:
        try:
            record["placeholder_type"] = enum_name(shape.placeholder_format.type)
            record["placeholder_idx"] = int(shape.placeholder_format.idx)
        except (AttributeError, ValueError):
            record["placeholder_type"] = None
            record["placeholder_idx"] = None
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.TABLE:
        try:
            record["table_rows"] = len(shape.table.rows)
            record["table_columns"] = len(shape.table.columns)
        except (AttributeError, ValueError):
            record["table_rows"] = None
            record["table_columns"] = None
    return record


def notes_text(slide: Any) -> str:
    """Return the actual notes body, excluding image/number placeholders."""
    try:
        notes = slide.notes_slide
    except Exception:
        return ""
    chunks: list[str] = []
    for shape in notes.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if getattr(shape, "is_placeholder", False):
            try:
                if enum_name(shape.placeholder_format.type) != "BODY":
                    continue
            except (AttributeError, ValueError):
                continue
        text = clean_newlines(shape.text).strip()
        if text and text not in {"Click to add notes", "Click to add text"}:
            chunks.append(text)
    return "\n".join(chunks)


def media_records(slide: Any) -> list[dict[str, Any]]:
    media: list[dict[str, Any]] = []
    for rel in slide.part.rels.values():
        if rel.reltype not in {RT.MEDIA, RT.IMAGE}:
            continue
        part = rel.target_part
        blob = getattr(part, "blob", b"")
        media.append(
            {
                "relationship_id": rel.rId,
                "partname": str(part.partname),
                "content_type": str(part.content_type),
                "bytes": len(blob),
                "sha256": sha256_bytes(blob),
            }
        )
    media.sort(key=lambda item: (item["partname"], item["relationship_id"]))
    return media


def slide_xml_names(index: int) -> tuple[str, str]:
    slide_name = f"ppt/slides/slide{index}.xml"
    rel_name = f"ppt/slides/_rels/slide{index}.xml.rels"
    return slide_name, rel_name


def deck_record(deck_id: str, path: Path, role: str) -> dict[str, Any]:
    data = path.read_bytes()
    source_sha = sha256_bytes(data)
    prs = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as package:
        for index, slide in enumerate(prs.slides, 1):
            slide_name, rel_name = slide_xml_names(index)
            slide_xml = package.read(slide_name)
            rel_xml = package.read(rel_name) if rel_name in package.namelist() else b""
            all_shapes = [shape_record(shape) for shape in slide.shapes]
            title_shape = slide.shapes.title
            title = shape_text(title_shape).strip() if title_shape is not None else ""
            title_id = int(title_shape.shape_id) if title_shape is not None else None
            body_chunks: list[str] = []
            for shape in slide.shapes:
                if int(getattr(shape, "shape_id", -1)) == title_id:
                    continue
                text = shape_text(shape).strip()
                if text:
                    body_chunks.append(text)
            body = "\n".join(body_chunks)
            visible = "\n".join(part for part in (title, body) if part)
            note = notes_text(slide)
            normalized = normalize_text(visible)
            media = media_records(slide)
            layout = slide.slide_layout
            layout_part = layout.part
            layout_xml_name = str(layout_part.partname).lstrip("/")
            master_name = str(layout_part.slide_master.part.partname).lstrip("/")
            slides.append(
                {
                    "deck_id": deck_id,
                    "source_path": str(path),
                    "source_role": role,
                    "source_sha256": source_sha,
                    "source_bytes": len(data),
                    "slide_no": index,
                    "slide_xml_name": slide_name,
                    "slide_xml_sha256": sha256_bytes(slide_xml),
                    "slide_rels_sha256": sha256_bytes(rel_xml) if rel_xml else None,
                    "title": title,
                    "body": body,
                    "notes": note,
                    "notes_nonempty": bool(note),
                    "visible_text": visible,
                    "visible_text_sha256": sha256_bytes(visible.encode("utf-8")),
                    "normalized_text": normalized,
                    "normalized_text_sha256": sha256_bytes(normalized.encode("utf-8")),
                    "notes_sha256": sha256_bytes(note.encode("utf-8")) if note else None,
                    "layout": {
                        "name": str(layout.name),
                        "type": enum_name(getattr(layout, "type", "")),
                        "layout_part": layout_xml_name,
                        "master_part": master_name,
                        "shape_count": len(layout.shapes),
                    },
                    "shape_count": len(all_shapes),
                    "text_shape_count": sum(1 for item in all_shapes if item["text"]),
                    "shape_type_counts": {
                        shape_type: sum(1 for item in all_shapes if item["shape_type"] == shape_type)
                        for shape_type in sorted({item["shape_type"] for item in all_shapes})
                    },
                    "shapes": all_shapes,
                    "media": media,
                    "media_count": len(media),
                    "media_types": sorted({item["content_type"] for item in media}),
                    "media_sha256": sorted({item["sha256"] for item in media}),
                }
            )
    return {
        "deck_id": deck_id,
        "source_path": str(path),
        "source_role": role,
        "source_sha256": source_sha,
        "source_bytes": len(data),
        "slide_count": len(slides),
        "slide_width_emu": int(prs.slide_width),
        "slide_height_emu": int(prs.slide_height),
        "slide_width_inches": round(prs.slide_width / 914400, 4),
        "slide_height_inches": round(prs.slide_height / 914400, 4),
        "notes_nonempty_count": sum(1 for slide in slides if slide["notes_nonempty"]),
        "slide_xml_hashes_unique": len({slide["slide_xml_sha256"] for slide in slides}),
        "slides": slides,
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def csv_text(value: str) -> str:
    """Keep flattened text readable without changing the JSON lossless record."""
    return "\n".join(line.rstrip() for line in value.splitlines())


def pairwise_candidates(slides: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for left_index, left in enumerate(slides):
        for right in slides[left_index + 1 :]:
            if left["deck_id"] == right["deck_id"] and left["slide_no"] == right["slide_no"]:
                continue
            left_text = left["normalized_text"]
            right_text = right["normalized_text"]
            similarity = difflib.SequenceMatcher(None, left_text, right_text).ratio()
            exact = left["visible_text_sha256"] == right["visible_text_sha256"]
            normalized_exact = left["normalized_text_sha256"] == right["normalized_text_sha256"]
            title_similarity = difflib.SequenceMatcher(
                None, normalize_text(left["title"]), normalize_text(right["title"])
            ).ratio()
            if exact or normalized_exact or similarity >= 0.86 or title_similarity >= 0.92:
                rows.append(
                    {
                        "left_deck": left["deck_id"],
                        "left_slide": left["slide_no"],
                        "right_deck": right["deck_id"],
                        "right_slide": right["slide_no"],
                        "exact_visible_text": exact,
                        "exact_normalized_text": normalized_exact,
                        "text_similarity": round(similarity, 4),
                        "title_similarity": round(title_similarity, 4),
                        "left_title": left["title"],
                        "right_title": right["title"],
                        "left_xml_sha256": left["slide_xml_sha256"],
                        "right_xml_sha256": right["slide_xml_sha256"],
                    }
                )
    rows.sort(
        key=lambda row: (
            not row["exact_visible_text"],
            not row["exact_normalized_text"],
            -row["text_similarity"],
            row["left_deck"],
            row["left_slide"],
        )
    )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    decks: list[dict[str, Any]] = []
    for deck_id, path, role in DECKS:
        if not path.is_file():
            decks.append(
                {
                    "deck_id": deck_id,
                    "source_path": str(path),
                    "source_role": role,
                    "status": "MISSING",
                }
            )
        else:
            decks.append(deck_record(deck_id, path, role))
    slides = [slide for deck in decks if "slides" in deck for slide in deck["slides"]]
    inventory = {
        "schema": "beamshift-donor-audit-v2/slide-inventory-1",
        "generated_by": str(Path(__file__).resolve()),
        "decks": decks,
    }
    (args.output / "slide-inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    inventory_rows: list[dict[str, Any]] = []
    for slide in slides:
        inventory_rows.append(
            {
                "deck_id": slide["deck_id"],
                "source_path": slide["source_path"],
                "source_sha256": slide["source_sha256"],
                "source_bytes": slide["source_bytes"],
                "slide_no": slide["slide_no"],
                "slide_xml_sha256": slide["slide_xml_sha256"],
                "title": csv_text(slide["title"]),
                "body": csv_text(slide["body"]),
                "notes": csv_text(slide["notes"]),
                "notes_nonempty": slide["notes_nonempty"],
                "visible_text_sha256": slide["visible_text_sha256"],
                "normalized_text_sha256": slide["normalized_text_sha256"],
                "notes_sha256": slide["notes_sha256"],
                "layout_name": slide["layout"]["name"],
                "layout_type": slide["layout"]["type"],
                "layout_part": slide["layout"]["layout_part"],
                "master_part": slide["layout"]["master_part"],
                "shape_count": slide["shape_count"],
                "text_shape_count": slide["text_shape_count"],
                "shape_type_counts": json.dumps(slide["shape_type_counts"], ensure_ascii=False),
                "media_count": slide["media_count"],
                "media_types": ";".join(slide["media_types"]),
                "media_sha256": ";".join(slide["media_sha256"]),
            }
        )
    write_csv(
        args.output / "slide-inventory.csv",
        inventory_rows,
        list(inventory_rows[0].keys()) if inventory_rows else ["deck_id"],
    )
    write_csv(
        args.output / "dedup-candidates.csv",
        pairwise_candidates(slides),
        [
            "left_deck",
            "left_slide",
            "right_deck",
            "right_slide",
            "exact_visible_text",
            "exact_normalized_text",
            "text_similarity",
            "title_similarity",
            "left_title",
            "right_title",
            "left_xml_sha256",
            "right_xml_sha256",
        ],
    )


if __name__ == "__main__":
    main()
